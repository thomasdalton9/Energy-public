"""
Weather master workbook: population-weighted heating and cooling degree days per region, one 'Dashboard - <Region>' page for
every region workbook found in "output/Data and Chart Outputs/" (<region>_degree_days.xlsx, written by weather/degree_days.py).
Same layout and chart code as the other masters (south_america/SOUTH_AMERICA_MASTER.py): dashboards, a data tab per chart,
raw tabs, Sources.

  Dashboard - Europe   cumulative HDD since 1 Oct (Oct-Sep season chart: 5-year min-max band, 5-year average, previous and
                       current season, 1991-2020 normal) for the Europe total (EU27 + UK) and the main gas countries, monthly
                       CDD vs the 1991-2020 normal (Spain, Italy, Greece, Turkey), monthly HDD minus the normal.
  <chart> data         the table each Dashboard chart plots
  <region> <sheet> raw Cities / Countries / Monthly / Normals month of the region workbook (Daily, Anomaly and Coverage stay in
                       the region workbook)
  Sources              where each dataset comes from

TO ADD A REGION: add it to weather/regions.json with its cities/countries csv (see the docstring of weather/degree_days.py), run
degree_days.py --region <region> (a copy of .github/workflows/weather_degree_days_europe.yml), and rebuild this master: the page
appears by itself. The season start of the cumulative-HDD chart is the region's season_start_month (10 for northern regions, 4
for southern ones). Other existing NASA POWER degree-day stores (Texas, MACRO_DRIVERS cities) can be added as regions the same way
by writing their cities in the csv format.

Usage: python3 weather/WEATHER_MASTER.py [--out "output/Data and Chart Outputs/Master Outputs/weather_master.xlsx"]
"""
import argparse
import os
import sys
from types import SimpleNamespace

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
sys.path.insert(0, HERE)
import add_charts  # noqa: E402
import degree_days  # noqa: E402
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR
NASA_URL = "https://power.larc.nasa.gov/"
NASA_PUB = ("NASA POWER daily T2M (MERRA-2 reanalysis), population-weighted city temperatures; city and country populations from "
            "official statistics (Eurostat Urban Audit / GISCO LAU / national statistics offices, listed on the Cities sheet of "
            "the region workbook)")
RAW_SHEETS = ("Cities", "Countries", "Monthly", "Normals month")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "Master Outputs", "weather_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()

    regions = degree_days.load_regions()
    found = [(r, c) for r, c in regions.items() if os.path.exists(os.path.join(args.data_dir, f"{r}_degree_days.xlsx"))]
    if not found:
        raise SystemExit("no <region>_degree_days.xlsx found")
    wb = Workbook()
    names = [f"Dashboard - {c['title']}" for _, c in found]
    wb.active.title = names[0]
    for n in names[1:]:
        wb.create_sheet(n)
    used = set(names) | {"Sources"}
    sources, counts = [], {}
    for (region, cfg), dash in zip(found, [wb[n] for n in names]):
        fname = f"{region}_degree_days.xlsx"
        src = (cfg.get("source_line") or NASA_PUB, NASA_URL)
        ns = SimpleNamespace(SOURCES={fname: (src[0], src[1])}, MASTER_SPECS={}, HYDRO_DATASETS=[], HYDRO_EXTRA={}, EMBER=set(),
                             OPERATORS={}, DASHBOARD_ONLY={}, GAS_BCFD=False)
        charts, rows, missing = sam.collect(wb, [(region[:3].upper(), cfg["title"], fname, RAW_SHEETS, f"{cfg['title']} degree days")],
                                            args.data_dir, used, sources, cfg=ns)
        sam.draw_dashboard(dash, f"{cfg['title']} - population-weighted degree days (NASA POWER), base 18 C", charts, rows, missing)
        counts[cfg["title"]] = len(charts)
        if missing:
            print("missing:", missing)

    s = wb.create_sheet("Sources")
    s.append(["Region", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in s[1]:
        cell.font = Font(bold=True)
    for row in sources:
        s.append(list(row))
        if row[4]:
            s.cell(row=s.max_row, column=5).hyperlink = row[4]
    for col, w in (("A", 18), ("B", 26), ("C", 34), ("D", 90), ("E", 40), ("F", 160)):
        s.column_dimensions[col].width = w
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: charts per dashboard {counts}; tabs {len(wb.sheetnames)}")


if __name__ == "__main__":
    main()
