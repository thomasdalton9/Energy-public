"""
BOLIVIA_CNDC_API_INSPECT.py found the real archivo_url pattern for CNDC's
monthly statistics XLSX files (e.g. gen_dia_0826.xlsx, iny_dia_0826.xlsx,
ret_dia_0826.xlsx under
https://www.cndc.bo/wp-content/uploads/mem/estadisticas/mensual/{YYYY}/{MM}/)
but only tried wrong guessed URLs before the real one appeared in the API
response, so it never actually inspected a file's structure.

This downloads one real month of each (Aug 2026) and dumps sheet names,
headers and a few sample rows - needed to know whether "Inyecciones de
Energia en Nodos del STI" / "Retiros de Energia en Nodos del STI" break
down by plant/node (which would let us map to a region) or something else.
"""
import io
import sys

import requests

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
}
TIMEOUT = (10, 60)

FILES = {
    "gen_dia_0826.xlsx": "https://www.cndc.bo/wp-content/uploads/mem/estadisticas/mensual/2026/08/gen_dia_0826.xlsx",
    "iny_dia_0826.xlsx": "https://www.cndc.bo/wp-content/uploads/mem/estadisticas/mensual/2026/08/iny_dia_0826.xlsx",
    "ret_dia_0826.xlsx": "https://www.cndc.bo/wp-content/uploads/mem/estadisticas/mensual/2026/08/ret_dia_0826.xlsx",
}


def inspect(label, url):
    print(f"\n{'=' * 70}\n{label}: {url}\n{'=' * 70}", file=sys.stderr)
    r = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    print(f"  status={r.status_code} bytes={len(r.content)}", file=sys.stderr)
    if r.status_code != 200:
        print(r.text[:500], file=sys.stderr)
        return
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    print(f"  sheets: {wb.sheetnames}", file=sys.stderr)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n  --- sheet '{sheet_name}': {ws.max_row} rows x {ws.max_column} cols ---", file=sys.stderr)
        for row in ws.iter_rows(min_row=1, max_row=min(15, ws.max_row), values_only=True):
            print(f"    {row}", file=sys.stderr)


def main():
    for label, url in FILES.items():
        try:
            inspect(label, url)
        except Exception as e:
            print(f"  ERROR: {type(e).__name__}: {e}", file=sys.stderr)


if __name__ == "__main__":
    main()
