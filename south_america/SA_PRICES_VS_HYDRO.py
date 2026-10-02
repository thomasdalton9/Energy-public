"""
Wholesale power price vs hydro storage, monthly from 2021, for the hydro-driven South American markets - built only
from this repo's workbooks (no downloads): prices from south_america_power_prices_daily.xlsx ('USD daily'), storage
from each country's hydro workbook ('Daily'). One sheet per country holding the monthly table and two native charts
stacked on the same dates: price (US$/MWh) on top, storage (% full, 0-100) underneath - two panels rather than one
chart with two y-axes. Charts are drawn here in the same save as the data (add_charts.py's registry entry is None).

  Brazil    ONS CMO SE/CO subsystem            vs  ONS stored energy, SE/CO subsystem (% of maximum)
  Colombia  XM Precio de Bolsa                 vs  XM useful stored energy (% of capacity)
  Peru      COES marginal cost, Santa Rosa 220  vs  COES reservoirs and lagoons, useful volume (% of capacity)
  Uruguay   ADME spot sancionado               vs  Rincon del Bonete level (% of its 70-80 m operating range)
  Chile     CNE Precio Medio de Mercado (monthly) vs DGA total of 8 generation reservoirs (% of capacity)

Usage: python3 SA_PRICES_VS_HYDRO.py [--out PATH]
"""
import argparse
import os
import sys

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import xlsx_charts  # noqa: E402
import xlsx_notes  # noqa: E402

DIR = os.path.join("output", "Data and Chart Outputs")
OUT = os.path.join(DIR, "south_america_prices_vs_hydro.xlsx")
START = "2021-01-01"
MARKETS = [  # country, price column (USD daily), price label, hydro workbook, storage column, storage label
    ("Brazil", "Brazil SE/CO (CMO)", "ONS CMO, SE/CO", "brazil_hydro_reservoirs.xlsx", "SE_CO_pct",
     "ONS stored energy, SE/CO (% of max)"),
    ("Colombia", "Colombia (Precio de Bolsa)", "XM Precio de Bolsa", "colombia_hydro_reservoirs.xlsx", "Storage_pct",
     "XM useful stored energy (% of capacity)"),
    ("Peru", "Peru (CMg Santa Rosa 220 kV)", "COES marginal cost, Santa Rosa 220 kV", "peru_hydro_reservoirs.xlsx",
     "Total_pct", "COES reservoirs and lagoons, useful volume (%)"),
    ("Uruguay", "Uruguay (spot sancionado)", "ADME spot sancionado", "uruguay_hydro_reservoirs.xlsx",
     "Bonete_pct_of_range", "Rincon del Bonete level (% of 70-80 m range)"),
    ("Chile", "Chile (CNE PMM, monthly value)", "CNE Precio Medio de Mercado", "chile_hydro_reservoirs.xlsx",
     "Total_pct", "DGA generation reservoirs, total (% of capacity)"),
]


def monthly(path, sheet, col):
    d = pd.read_excel(os.path.join(DIR, path), sheet_name=sheet)
    d["date"] = pd.to_datetime(d["date"].astype(str), errors="coerce")
    s = pd.to_numeric(d.set_index("date")[col], errors="coerce").dropna()
    return s.resample("MS").mean()


def build():
    out, notes = {}, []
    for country, pcol, plabel, hfile, hcol, hlabel in MARKETS:
        try:
            price = monthly("south_america_power_prices_daily.xlsx", "USD daily", pcol)
            storage = monthly(hfile, "Daily", hcol)
        except Exception as e:  # noqa: BLE001 - a missing workbook / column skips that country only
            print(f"  {country}: skipped ({type(e).__name__}: {e})", flush=True)
            continue
        df = pd.DataFrame({f"Price, {plabel} (US$/MWh)": price, f"Storage, {hlabel}": storage})
        df = df[df.index >= START].dropna(how="all").round(2)
        df.index.name = "date"
        out[country] = df
        notes.append(f"{country}: price = {plabel} (monthly mean of daily US$/MWh); storage = {hlabel} (monthly mean). "
                     f"{df.index.min():%b-%y} to {df.index.max():%b-%y}.")
    return out, notes


def add_charts(path, frames):
    """Two charts per country sheet, built in one load/save (re-reading charts loses their formatting)."""
    wb = load_workbook(path)
    for country, df in frames.items():
        name = f"Chart - {country}"
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name, 1)
        price = df.iloc[:, [0]].dropna()
        storage = df.iloc[:, [1]].dropna()
        lo, hi = min(price.index.min(), storage.index.min()), max(price.index.max(), storage.index.max())
        span = pd.date_range(lo, hi, freq="MS")
        price, storage = price.reindex(span), storage.reindex(span)   # same months on both panels
        xlsx_charts.write_table(ws, price, start_col=1)
        xlsx_charts.write_table(ws, storage, start_col=4)
        c1 = xlsx_charts.build_chart(ws, price, 1, f"{country}: wholesale power price", "US$/MWh", "line",
                                     start_col=1, width=26, height=9)
        c2 = xlsx_charts.build_chart(ws, storage, 1, f"{country}: hydro storage", "% full", "line",
                                     start_col=4, width=26, height=9)
        c2.y_axis.scaling.min, c2.y_axis.scaling.max, c2.y_axis.majorUnit = 0, 100, 20
        ws.add_chart(c1, "H2")
        ws.add_chart(c2, "H21")
    xlsx_charts.save_atomic(wb, path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=OUT)
    args = ap.parse_args()
    frames, notes = build()
    if not frames:
        sys.exit("no country could be built")
    lines = ["SOUTH AMERICA: WHOLESALE POWER PRICE VS HYDRO STORAGE (monthly)",
             "Built only from this repo's workbooks (south_america_power_prices_daily.xlsx and each country's hydro "
             "workbook) - no downloads, no estimates. Monthly means. Each 'Chart - <country>' sheet stacks two charts "
             "on the same months: price (US$/MWh) above, storage (% full, 0-100) below.",
             "Chile's price is CNE's Precio Medio de Mercado: an average of contract prices over a 4-month window, "
             "published with a lag - not a spot price, so it does not move inversely with storage the way the spot "
             "prices of Brazil, Colombia, Peru and Uruguay do. Chile's spot marginal costs (Coordinador) need a key.", "",
             "MARKETS"] + notes + ["", "UPDATES",
             "Rebuilt on the 1st and 15th after the price and hydro pulls (south_america_prices_vs_hydro.yml)."]
    xlsx_notes.write_workbook(args.out, {c: df for c, df in frames.items()}, lines, ["MARKETS", "UPDATES"])
    add_charts(args.out, frames)
    for c, df in frames.items():
        print(f"  {c}: {len(df)} months, latest {df.dropna(how='all').index.max():%b-%y}; "
              f"corr(price, storage) = {df.iloc[:, 0].corr(df.iloc[:, 1]):.2f}", flush=True)
    print(f"saved {args.out}", flush=True)


if __name__ == "__main__":
    main()
