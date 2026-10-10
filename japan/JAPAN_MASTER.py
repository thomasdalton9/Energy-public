"""
Japan master workbook: every Japan dataset this repo pulls in one file, with Dashboard front pages carrying all their
charts - the same layout and chart code as the other masters (south_america/SOUTH_AMERICA_MASTER.py).

  Dashboard - Power      power generation by source (average GW, the nine OCCTO-linked TSO areas), area demand, solar and wind
                         output control, JEPX day-ahead spot prices by area (JPY per kWh and US$ per MWh, converted at the
                         Federal Reserve H.10 yen rate of each day), installed capacity by type (METI/ANRE statistics)
  Dashboard - Gas        LNG imports (Ministry of Finance, Mt per month), power-company LNG stock (ANRE weekly; METI month-end
                         as an Oct-Sep water-year chart), LNG receipts and consumption at power stations (METI)
  Dashboard - Fuels      crude oil, coal and LPG imports (Ministry of Finance)
  Dashboard - Long-term  annual summary and history from 2000 (fundamentals.py)
  <chart> data           the table each Dashboard chart plots
  <dataset> raw          the full data sheet(s) from each source workbook
  Sources                where each dataset comes from, units and notes, and the sources checked but not built

Power generation is the general transmission and distribution operators' own "area supply and demand results" (nine areas;
Okinawa's island grid is not reachable from GitHub Actions and is not in it). Prices are JEPX's own published results. Ember
appears only on the annual Long-term page.

Gas: Japan imports LNG only (no pipeline gas); the Ministry of Finance monthly release gives weights (Mt, not converted to a
gas volume) without an origin split, METI/ANRE give the power sector's LNG stock and use. No city-gas/LNG sales by sector or
LNG imports by origin were reachable (see the Sources tab); nothing is estimated.

Reads (doesn't refetch) the workbooks the scheduled pulls write to "output/Data and Chart Outputs/". A missing input
is listed on its Dashboard and skipped. To add a dataset: append it to the lists below plus a SOURCES entry.

Usage: python3 japan/JAPAN_MASTER.py [--out "output/Data and Chart Outputs/Master Outputs/japan_master.xlsx"]
"""
import argparse
import os
import sys
from types import SimpleNamespace

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
sys.path.insert(0, os.path.join(ROOT, "japan"))
import add_charts  # noqa: E402
import fundamentals  # noqa: E402
import japan_charts  # noqa: E402
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR

GEN = "japan_power_generation_daily.xlsx"
JEPX = "japan_jepx_spot_prices.xlsx"
FX = "japan_fx_usd_daily.xlsx"
MOF = "japan_mof_energy_imports_monthly.xlsx"
METI = "japan_meti_electric_power_stats.xlsx"
LNGSTOCK = "japan_lng_stock_weekly.xlsx"

# (dashboard, [(code, country, workbook, raw sheets, short name, charts to show (None = all))])
POWER = [("JP", "Japan", GEN, ("Daily", "Areas", "Coverage"), "power generation", None),
         ("JP", "Japan", JEPX, ("Daily",), "JEPX spot prices", None),
         ("JP", "Japan", METI, ("Capacity",), "METI statistics", {"Capacity"})]
GAS_SETS = [("JP", "Japan", MOF, ("Monthly", "Releases"), "energy imports", {"LNG imports"}),
            ("JP", "Japan", LNGSTOCK, ("Weekly", "Month-end reference"), "LNG stock", None),
            ("JP", "Japan", METI, ("Fuel",), "METI statistics", {"LNG stock", "LNG use"})]
FUELS = [("JP", "Japan", MOF, (), "energy imports", {"Coal and LPG imports", "Crude oil imports"})]
FED_URL = "https://www.federalreserve.gov/releases/h10/hist/dat00_ja.htm"

SOURCES = {
    GEN: ("General transmission and distribution operators' area supply and demand results (エリア需給実績データ), monthly "
          "CSVs: Hokkaido Electric Power Network, Tohoku Electric Power Network, TEPCO Power Grid, Chubu Electric Power Grid, "
          "Hokuriku Electric Power Transmission & Distribution, Kansai Transmission and Distribution, Energia Chugoku, "
          "Shikoku T&D, Kyushu T&D (OCCTO-linked grid, Okinawa not included)",
          "https://www.tepco.co.jp/forecast/html/area_jukyu-j.html"),
    JEPX: ("Japan Electric Power Exchange (JEPX), spot (day-ahead) market trading information",
           "https://www.jepx.jp/electricpower/market-data/spot/"),
    FX: ("Board of Governors of the Federal Reserve System, H.10 Foreign Exchange Rates: noon buying rates in New York, "
         "Japanese yen per US dollar (the series FRED republishes as DEXJPUS); used to convert JEPX prices to US$", FED_URL),
    MOF: ("Ministry of Finance (Japan Customs), Trade Statistics of Japan, monthly press release: major import commodities",
          "https://www.customs.go.jp/toukei/shinbun/happyou.htm"),
    METI: ("Agency for Natural Resources and Energy (METI), electric power statistics: power stations and maximum output of "
           "electricity business operators (table 1) and thermal fuel receipts, consumption and stocks (table 4)",
           "https://www.enecho.meti.go.jp/statistics/electric_power/ep002/results.html"),
    LNGSTOCK: ("Agency for Natural Resources and Energy (METI), weekly LNG stock of the large power companies "
               "(発電用LNGの在庫状況)",
               "https://www.enecho.meti.go.jp/category/electricity_and_gas/electricity_measures/pdf/denryoku_LNG_stock.pdf"),
}

MASTER_SPECS = {}
HYDRO_DATASETS = []
HYDRO_EXTRA = {}
EMBER = set()
OPERATORS = {}
DASHBOARD_ONLY = {}

# Sources checked from GitHub Actions on 10 Oct 2026 (discovery_archive/japan/JAPAN_SOURCES_PROBE*.py) and NOT built.
NOT_AVAILABLE = [
    ("Japan", "Okinawa power generation by source", "Okinawa Electric Power Company (okinawa-epco.co.jp / okiden.co.jp)",
     "Not reachable", "https://www.okinawa-epco.co.jp/",
     "okinawa-epco.co.jp does not resolve and okiden.co.jp answers HTTP 403 from GitHub Actions. Okinawa is an island grid of "
     "about 1 GW (around 1% of Japan) and is left out of the area sum; nothing is added for it."),
    ("Japan", "Power generation by source before April 2024", "the nine operators' older supply-demand files (HEPCO sup_dem_results_*q.csv, Kyushu area_jyukyu_jisseki_*Q.csv, "
     "Shikoku jukyuYYYY.xlsx, TEPCO yearly files, Tohoku juyo_YYYY_tohoku.csv)", "Reachable, not built",
     "https://www.hepco.co.jp/network/con_service/public_document/supply_demand_results/index.html",
     "The common monthly eria_jukyu_YYYYMM_NN.csv format exists from Oct 2023 (Tokyo, Kyushu) to Apr 2024 (Hokkaido, Chubu, Kansai), so the "
     "raw power series starts in April 2024. Earlier files use a different layout per operator (and some only list demand); they were not "
     "parsed. The annual history back to 2000 is on the Long-term page (Ember yearly, labelled)."),
    ("Japan", "Chubu and Kansai monthly by-source files", "Chubu Electric Power Grid (getFilesInfo.php annual zips eria_jukyu_YYYY.zip), "
     "Kansai T&D (area-performance/filelist.json)", "Built (script-listed)",
     "https://powergrid.chuden.co.jp/denkiyoho/eriajukyu_data/",
     "Both operators' pages build their download lists with JavaScript; the master's pull reads the same JSON lists. Chubu's monthly "
     "'keito' zips hold only total demand and generation (not by source) and are not used. A month a feed lacks (Chubu Jul and Sep 2025 and "
     "Aug 2026 were not in its annual zips on 10 Oct 2026; Tohoku's September file was not yet posted) is a gap in the nine-area total: "
     "months with fewer than 80% of days are left out of the charts, nothing is filled."),
    ("Japan", "TEPCO Power Grid area supply-demand page", "tepco.co.jp/forecast/html/area_jukyu-j.html", "Page 403; CSV reachable",
     "https://www.tepco.co.jp/forecast/html/area_jukyu-j.html",
     "The HTML page answers HTTP 403 from GitHub Actions but the monthly CSV files (forecast/html/images/eria_jukyu_YYYYMM_03.csv) "
     "download, and are what the master uses."),
    ("Japan", "OCCTO system information service (supply-demand actuals by area and source)", "occtonet.occto.or.jp",
     "Not reachable", "https://www.occtonet.occto.or.jp/",
     "occtonet.occto.or.jp does not resolve from GitHub Actions; occto.or.jp itself loads but has no data links. The same "
     "area-by-source figures are published by the nine operators and are used instead."),
    ("Japan", "LNG imports by origin country, city-gas and LNG sales by sector", "MOF customs trade statistics database (e-Stat / "
     "customs.go.jp/toukei/srch), METI gas statistics", "Not built",
     "https://www.customs.go.jp/toukei/srch/indexe.htm",
     "The customs commodity-by-country tables are a JavaScript query form (the CSV downloads are totals by country only) and "
     "e-Stat's API needs a registered application key. METI's gas statistics pages were not found as downloadable tables. "
     "The monthly press release gives LNG, crude oil, LPG and coal weights without an origin split, and that is what is built."),
    ("Japan", "Spot LNG price (JKM) and Asian LNG benchmarks", "S&P Global Platts / CME", "Paid only",
     "https://www.spglobal.com/commodityinsights/en/market-insights/topics/lng",
     "No free official source for JKM was found; not scraped."),
    ("Japan", "Weekly LNG stock before the first run", "ANRE weekly PDF", "No history at the source",
     "https://www.enecho.meti.go.jp/category/electricity_and_gas/electricity_measures/",
     "The PDF shows only the current season; the history starts at the repo's first run. The METI month-end series "
     "(table 4, all business operators, from April 2021) is the long record."),
    ("Japan", "METI/ANRE 'Energy White Paper' tables", "ANRE", "Not built",
     "https://www.enecho.meti.go.jp/about/whitepaper/",
     "Annual PDFs; the annual power and gas history is already on the Long-term page from Ember and the Energy Institute."),
]


def collect(wb, datasets, data_dir, used, sources):
    """sam.collect per dataset so each can show a chosen subset of its charts (names = spec names; None = all)."""
    charts, rows, missing = [], [], []
    for code, country, fname, raw, short, only in datasets:
        cfg = SimpleNamespace(**{k: getattr(sys.modules[__name__], k) for k in (
            "SOURCES", "MASTER_SPECS", "HYDRO_DATASETS", "HYDRO_EXTRA", "EMBER", "OPERATORS")},
            DASHBOARD_ONLY={fname: only} if only else {}, GAS_BCFD=True)
        before = set(wb.sheetnames)
        c, r, m = sam.collect(wb, [(code, country, fname, raw, short)], data_dir, used, sources, cfg=cfg)
        for ws in wb.worksheets:   # 'JP Data raw 2' -> 'JP industrial output Data raw': say which dataset it is
            if ws.title not in before and ws.title.endswith(" raw") or (ws.title not in before and " raw " in ws.title):
                sheet = ws.title.split(" ", 1)[1].rsplit(" raw", 1)[0]
                used.discard(ws.title)
                ws.title = sam.sheet_name(f"JP {short} {sheet} raw", used)
        charts += c
        rows += r
        missing += m
    return charts, rows, missing


def fx_for(dates, fx):
    """Rate for each date = the last published rate on or before it (no other filling); returns (rate, rate date)."""
    left = pd.DataFrame({"d": pd.to_datetime(list(dates))}).sort_values("d")
    right = fx.rename("rate").rename_axis("rate_date").reset_index()
    m = pd.merge_asof(left, right, left_on="d", right_on="rate_date", direction="backward").set_index("d")
    return m["rate"], m["rate_date"]




USD_AREAS = ["System", "Hokkaido", "Tokyo", "Kansai", "Kyushu"]


def jepx_usd_chart(wb, used, data_dir, sources):
    """JEPX daily average prices (JPY per kWh) in US$ per MWh: price x 1000 / (yen per US$), with the Fed H.10 rate of the
    delivery date or, where the Fed published none that day (weekends, US holidays), the last rate published before it;
    monthly mean of the daily US$ prices. The data tab carries the monthly mean rate beside the prices."""
    jpath, fpath = os.path.join(data_dir, JEPX), os.path.join(data_dir, FX)
    if not (os.path.exists(jpath) and os.path.exists(fpath)):
        return [], [], [f"Japan JEPX prices in US$ ({JEPX} or {FX} missing)"]
    try:
        fx = add_charts.by_date(add_charts.read(fpath, "Data"), "date")["JPY_per_USD"].dropna()
        d = add_charts.by_date(add_charts.read(jpath, "Daily"), "date")
    except Exception as e:  # noqa: BLE001
        return [], [], [f"Japan JEPX prices in US$ ({type(e).__name__}: {e})"]
    cols = {a: f"{a}_JPY_per_kWh" for a in USD_AREAS if f"{a}_JPY_per_kWh" in d.columns}
    d = d[list(cols.values())].dropna(how="all")
    rate, rdate = fx_for(d.index, fx)
    rate = rate.reindex(d.index)
    usd = d.mul(1000.0).div(rate, axis=0)
    usd = usd[rate.notna().values]
    last = usd.index.max()
    m = usd.resample("MS").mean()
    if last < last + pd.offsets.MonthEnd(0):   # month in progress
        m = m[m.index < last.to_period("M").to_timestamp()]
    m = m.rename(columns={v: k for k, v in cols.items()}).round(2)
    m = m[m.index >= "2021-04-01"].dropna(how="all")
    title = "Japan JEPX day-ahead spot price in US$ (monthly average)"
    df, n_bars = xlsx_charts.prepare(m, ())
    ws = wb.create_sheet(sam.sheet_name("JP spot prices US$ data", used))
    xlsx_charts.write_table(ws, df, "%Y-%m")
    c0 = len(df.columns) + 3
    ws.cell(row=1, column=c0, value="Mean yen per US$ (Fed H.10) in the month")
    mrate = rate[rate.notna()].resample("MS").mean()
    for i, ts in enumerate(df.index, start=2):
        ws.cell(row=i, column=c0, value=float(mrate.get(ts, float("nan"))))
    ws.cell(row=1, column=c0).font = Font(bold=True)
    ws.column_dimensions[ws.cell(row=1, column=c0).column_letter].width = 30
    src = ("JEPX day-ahead spot prices, converted at the Federal Reserve H.10 yen per US$ rate", FED_URL)
    chart = xlsx_charts.build_chart(ws, df, n_bars, title, "US$ per MWh (monthly average)", "line", "%Y-%m",
                                    width=sam.CHART_W, height=sam.CHART_H, gridlines=False,
                                    inner=xlsx_charts.DASHBOARD_INNER)
    sam.write_frame(wb.create_sheet(sam.sheet_name("JP FX Data raw", used)), add_charts.read(fpath, "Data"))
    sources.append(("Japan", "yen per US dollar", FX, *SOURCES[FX], sam.notes_text(fpath)))
    data_tab = ws.title
    ws = wb.create_sheet(sam.sheet_name("Conversion factors", used))
    for line in ("Currency conversion of the JEPX prices (the only conversion applied)",
                 "Prices: JEPX day-ahead (spot) market, simple average of the day's 48 half-hourly prices, yen per kWh, as published; "
                 "the yen series are on the 'JP JEPX spot prices Daily raw' tab and charted in JPY per kWh on the same dashboard.",
                 "US$ per MWh = (yen per kWh x 1,000) / (yen per US$), per day, then averaged over the month.",
                 "Exchange rate: Board of Governors of the Federal Reserve System, H.10 Foreign Exchange Rates, Historical Rates for the "
                 "Japanese Yen (noon buying rates in New York; the series FRED republishes as DEXJPUS): " + FED_URL,
                 "Rate used for each day: that day's rate or, where the Fed published none (weekends, US holidays), the last rate published "
                 "before it. Nothing else is used to fill days.",
                 f"Rate series: {fx.index.min():%d %b %Y} to {fx.index.max():%d %b %Y}, latest {fx.iloc[-1]:.2f} yen per US$.",
                 "Workbook: japan_fx_usd_daily.xlsx (japan/JAPAN_FX_USD.py, 1st and 15th); an ECB cross-rate column in it is validation only."):
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 180
    return [(chart, src)], [("Japan", title, df.index.max().strftime("%b/%y"), data_tab, *src)], []


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "Master Outputs", "japan_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()

    wb = Workbook()
    dash_power = wb.active
    dash_power.title = "Dashboard - Power"
    names = ["Dashboard - Gas", "Dashboard - Fuels"]
    dashes = [wb.create_sheet(n) for n in names]
    used = {"Dashboard - Power", *names, "Dashboard - Long-term", "Sources"}
    sources = []

    sections = [(dash_power, "Japan - power generation (TSO area data, average GW), demand, JEPX prices and capacity", POWER),
                (dashes[0], "Japan - LNG: imports (Ministry of Finance) and power-company stocks and use (METI/ANRE)", GAS_SETS),
                (dashes[1], "Japan - crude oil, coal and LPG imports (Ministry of Finance)", FUELS)]
    counts = []
    for dash, heading, datasets in sections:
        charts, rows, missing = collect(wb, datasets, args.data_dir, used, sources)
        if datasets is POWER:
            pc, pr, pm = jepx_usd_chart(wb, used, args.data_dir, sources)
            at = next((i for i, r in enumerate(rows) if "JEPX" in r[1]), len(rows)) + 1   # right after the yen price chart
            charts[at:at] = pc
            rows[at:at] = pr
            missing += pm
        sam.draw_dashboard(dash, heading, charts, rows, missing)
        counts.append(len(charts))
        if missing:
            print(f"missing ({heading}):", missing)

    fundamentals.add_long_term_dashboard(wb, used, sam.sheet_name, "Japan", args.data_dir, sam.CHART_W, sam.CHART_H,
                                         rows_per_chart=sam.ROWS_PER_CHART, degree_days=("CDD_18", "HDD_18"),
                                         index=sum(s.startswith("Dashboard") for s in wb.sheetnames))

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    seen = set()
    for s in sources:
        if s[2] in seen:   # a workbook read for two dashboards is listed once
            continue
        seen.add(s[2])
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    if NOT_AVAILABLE:
        src.append([])
        src.append(["Checked, NOT built", "Dataset", "Publisher / page", "Result from GitHub Actions", "Link", "Notes"])
        for cell in src[src.max_row]:
            cell.font = Font(bold=True)
        for row in NOT_AVAILABLE:
            src.append(list(row))
            if row[4]:
                src.cell(row=src.max_row, column=5).hyperlink = row[4]
    for col, w in (("A", 18), ("B", 22), ("C", 44), ("D", 60), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: charts per dashboard {dict(zip(['power', *names], counts))}; tabs {wb.sheetnames}")


if __name__ == "__main__":
    main()
