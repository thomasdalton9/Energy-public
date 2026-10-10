"""
Pull two METI / Agency for Natural Resources and Energy (ANRE) electric power statistics tables and keep them as a
growing archive (output/Data and Chart Outputs/japan_meti_electric_power_stats.xlsx):

  Capacity  Table 1 "number of power stations and maximum output of electric business operators", the '合計' (total) row
            of each monthly sheet, kW by type: general hydro, pumped hydro, thermal by fuel (coal, LNG, oil, LPG, other
            gas, bituminous mixture, other), nuclear, wind, solar, geothermal, battery, other; biomass and waste are
            memo lines inside thermal and are not added again.
  Fuel      Table 4 "thermal power fuel results (receipts, consumption, month-end stock)": LNG receipts, consumption and
            month-end stock in tonnes, all electric business operators.

Source: https://www.enecho.meti.go.jp/statistics/electric_power/ep002/results.html - one workbook per fiscal year and table
(xls/<FY>/1-1-<FY>.xlsx, 4-<FY>.xlsx; the current year's file is published as ...n.xlsx), one sheet per month. The statistics
cover registered electricity business operators; they are not a census of every solar panel or self-generation set.
Reachable from GitHub Actions (probe 10 Oct 2026). History from fiscal 2021 (April 2021).

Incremental: the committed workbook is the history store; the fiscal-year workbook(s) from the latest stored month's year
are re-read each run (METI revises recent months), earlier years are kept as stored.

    python3 japan/JAPAN_METI_ELECTRIC.py --out "output/Data and Chart Outputs/japan_meti_electric_power_stats.xlsx"
"""
import argparse
import io
import os
import re
import sys
import time
import unicodedata

import openpyxl
import pandas as pd
import requests

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # repo root, for xlsx_notes
import xlsx_notes  # noqa: E402

BASE = "https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"}
FIRST_FY = 2021
DEFAULT_OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "output",
                           "Data and Chart Outputs", "japan_meti_electric_power_stats.xlsx")
# (category, sub-category) after NFKC -> column
CAP = {
    ("水力発電所", "一般"): "Hydro_kW", ("水力発電所", "揚水式"): "Pumped hydro_kW",
    ("火力発電", "石炭"): "Coal_kW", ("火力発電", "LNG"): "LNG_kW", ("火力発電", "石油"): "Oil_kW",
    ("火力発電", "LPG"): "LPG_kW", ("火力発電", "その他ガス"): "Other gas_kW", ("火力発電", "歴青質混合物"): "Bituminous mixture_kW",
    ("火力発電", "その他"): "Other thermal_kW", ("原子力発電所", ""): "Nuclear_kW",
    ("新エネルギー等発電所", "風力"): "Wind_kW", ("新エネルギー等発電所", "太陽光"): "Solar_kW",
    ("新エネルギー等発電所", "地熱"): "Geothermal_kW", ("新エネルギー等発電所", "蓄電池"): "Battery_kW",
    ("その他", ""): "Other_kW", ("計", ""): "Total_kW",
}


def nf(x):
    return unicodedata.normalize("NFKC", str(x)).strip() if x is not None else ""


def get_wb(fy, table):
    for suffix in ("n", ""):
        url = f"{BASE}{fy}/{table}-{fy}{suffix}.xlsx"
        for i in range(3):
            try:
                r = requests.get(url, headers=UA, timeout=(10, 120))
            except requests.RequestException:
                time.sleep(3)
                continue
            if r.status_code == 404:
                break
            if r.status_code == 200:
                return openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True), url
            time.sleep(3)
    return None, None


def month_of(sheet):
    m = re.match(r"^(\d{4})\.(\d{1,2})", sheet)
    return pd.Timestamp(int(m.group(1)), int(m.group(2)), 1) if m else None


def parse_capacity(ws):
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows[:12]) if r and nf(r[0]) == "時間軸コード")
    cat, sub, item = rows[hdr - 3], rows[hdr - 2], rows[hdr - 1]
    tot = next(r for r in rows[hdr + 1:] if r and nf(r[2]) == "合計")
    out = {}
    c0, s0 = "", ""
    for c in range(10, len(rows[hdr])):
        c0 = nf(cat[c]) or c0
        if nf(item[c]) != "最大出力":
            continue
        key = (c0, nf(sub[c]))
        if "再掲" in key[1] or key[1] == "小計":
            continue
        col = CAP.get(key) or CAP.get((key[0], ""))
        if col is None:
            col = f"{key[0]}/{key[1]}_kW"
        v = tot[c]
        if isinstance(v, (int, float)):
            out[col] = out.get(col, 0.0) + float(v)
    return out


def parse_fuel(ws):
    for r in ws.iter_rows(values_only=True):
        if r and any("LNG" in nf(c) for c in r[:3] if c):
            nums = [float(c) for c in r[2:] if isinstance(c, (int, float))]
            if len(nums) >= 4:       # receipts, consumption, heat value (kJ/kg), month-end stock
                return {"LNG_receipts_t": nums[0], "LNG_consumption_t": nums[1], "LNG_stock_t": nums[3]}
            if len(nums) == 3:       # receipts, consumption, month-end stock
                return {"LNG_receipts_t": nums[0], "LNG_consumption_t": nums[1], "LNG_stock_t": nums[2]}
    return {}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()
    old_cap = old_fuel = None
    if os.path.exists(args.out):
        try:
            old_cap = pd.read_excel(args.out, sheet_name="Capacity", index_col=0, parse_dates=True)
            old_fuel = pd.read_excel(args.out, sheet_name="Fuel", index_col=0, parse_dates=True)
        except Exception as e:  # noqa: BLE001
            print(f"  stored workbook unreadable ({type(e).__name__}); starting again", file=sys.stderr)
            old_cap = old_fuel = None
    today = pd.Timestamp.today()
    cur_fy = today.year if today.month >= 4 else today.year - 1
    first = FIRST_FY
    if old_cap is not None and len(old_cap):
        last = old_cap.index.max()
        first = max(FIRST_FY, (last.year if last.month >= 4 else last.year - 1) - (1 if today.month in (4, 5, 6) else 0))
    cap, fuel = {}, {}
    if old_cap is not None:
        cap = {k: v.dropna().to_dict() for k, v in old_cap.iterrows()}
        fuel = {k: v.dropna().to_dict() for k, v in old_fuel.iterrows()}
    for fy in range(first, cur_fy + 1):
        for table, store, parse in (("1-1", cap, parse_capacity), ("4", fuel, parse_fuel)):
            wb, url = get_wb(fy, table)
            if wb is None:
                print(f"  FY{fy} table {table}: not available", file=sys.stderr)
                continue
            n = 0
            for ws in wb.worksheets:
                month = month_of(ws.title)
                if month is None:
                    continue
                try:
                    vals = parse(ws)
                except Exception as e:  # noqa: BLE001
                    print(f"  {url} {ws.title}: {type(e).__name__}: {e}", file=sys.stderr)
                    continue
                if vals:
                    store[month] = vals
                    n += 1
            print(f"  FY{fy} table {table}: {n} months from {url}")
    if not cap and not fuel:
        raise SystemExit("no data")
    capdf = pd.DataFrame.from_dict(cap, orient="index").sort_index()
    capdf.index.name = "month"
    fueldf = pd.DataFrame.from_dict(fuel, orient="index").sort_index()
    fueldf.index.name = "month"
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_notes.write_workbook(args.out, {"Capacity": capdf, "Fuel": fueldf}, NOTES, TITLES)
    c = capdf.iloc[-1]
    print(f"Saved capacity {len(capdf)} months ({capdf.index.min():%Y-%m}..{capdf.index.max():%Y-%m}), total "
          f"{c.get('Total_kW', float('nan')) / 1e6:.1f} GW; fuel {len(fueldf)} months; latest LNG stock "
          f"{fueldf['LNG_stock_t'].iloc[-1] / 1e6:.2f} Mt -> {args.out}")


NOTES = [
    "UNITS",
    "Capacity: *_kW = maximum output in kW of registered electricity business operators' power stations at month end, the "
    "'total' (合計) row of METI table 1; Total_kW is METI's own total. Biomass and waste are memo lines inside thermal and are "
    "not added again. Fuel: LNG_*_t = tonnes of LNG received, burned and in stock at month end at those operators' thermal "
    "power stations (METI table 4).",
    "",
    "SOURCE",
    "Agency for Natural Resources and Energy (METI), electric power statistics, tables 1 and 4: "
    "https://www.enecho.meti.go.jp/statistics/electric_power/ep002/results.html .",
    "",
    "UPDATES",
    "Incremental: stored months are kept; the latest fiscal-year workbook(s) are re-read each run.",
]
TITLES = {"UNITS", "SOURCE", "UPDATES"}

if __name__ == "__main__":
    main()
