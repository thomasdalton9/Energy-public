"""
Japan + South Korea + Taiwan master workbook: one file with every JKT dataset this repo pulls, and Dashboard front
pages carrying all their charts - the same layout as the South & Central America, North America and Australia/NZ
masters (south_america/SOUTH_AMERICA_MASTER.py, whose table/chart/dashboard code this reuses).

  Dashboard                  - gas (added as the pulls come online)
  Dashboard - Power          - Japan generation by fuel (sum of the 10 regional TSOs) and JEPX spot prices; Korea and
                               Taiwan generation by fuel and installed capacity per country
  <CC> <chart> data          - the table each Dashboard chart plots
  <CC> <dataset> raw         - the full data sheet(s) from each source workbook
  Sources                    - where each dataset comes from, units and notes

Prefers raw sources (TSOs, market operators, ministries). Ember is used only for countries with no raw feed yet and is
labelled as such. Reads (doesn't refetch) the workbooks the scheduled pulls write to "output/Data and Chart Outputs/".
A missing input is listed on the Dashboard and skipped rather than stopping the rest.

Usage: python3 asia/JKT_MASTER.py [--out "output/Data and Chart Outputs/japan_korea_taiwan_master.xlsx"]
"""
import argparse
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR

# (country code, country, workbook, raw sheet / tuple of raw sheets / "*", short dataset name)
DATASETS = [
]
RAW_POWER_DATASETS = [
    ("JP", "Japan", "japan_power_generation_daily.xlsx", ("Daily", "By area"), "power"),
    ("JP", "Japan", "japan_power_prices_daily.xlsx", "Daily", "prices"),
    ("KR/TW", "Korea / Taiwan", "korea_taiwan_power_by_type.xlsx", "*", "power (Ember)"),
]
CAPACITY_DATASETS = [
    ("KR", "South Korea", "south_korea_power_capacity.xlsx", "Monthly", "capacity"),
    ("TW", "Taiwan", "taiwan_power_capacity.xlsx", "Monthly", "capacity"),
]
HYDRO_DATASETS = []
HYDRO_EXTRA = {}
DASHBOARD_ONLY = {}
MASTER_SPECS = {}
EMBER = {"korea_taiwan_power_by_type.xlsx", "south_korea_power_capacity.xlsx", "taiwan_power_capacity.xlsx"}
OPERATORS = {"South Korea": "KPX / MOTIE", "Taiwan": "Taipower / MOEA"}

SOURCES = {
    "japan_power_generation_daily.xlsx": ("The 10 regional TSOs' area supply-demand actuals (eria_jukyu CSVs: Hokkaido, "
                                          "Tohoku, TEPCO PG, Chubu, Hokuriku, Kansai, Chugoku, Shikoku, Kyushu, Okinawa)",
                                          "https://www.occto.or.jp/en/"),
    "japan_power_prices_daily.xlsx": ("JEPX day-ahead spot market", "https://www.jepx.jp/electricpower/market-data/spot/"),
    "korea_taiwan_power_by_type.xlsx": ("Ember monthly electricity data (no raw KPX / Taipower feed yet)",
                                        "https://ember-energy.org/data/monthly-electricity-data/"),
    "south_korea_power_capacity.xlsx": ("Ember yearly electricity data (no raw capacity feed yet)",
                                        "https://ember-energy.org/data/yearly-electricity-data/"),
    "taiwan_power_capacity.xlsx": ("Ember yearly electricity data (no raw capacity feed yet)",
                                   "https://ember-energy.org/data/yearly-electricity-data/"),
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "japan_korea_taiwan_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()
    cfg = sys.modules[__name__]

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    dash2 = wb.create_sheet("Dashboard - Power")
    used = {"Dashboard", "Dashboard - Power", "Sources"}
    sources = []

    gas = sam.collect(wb, DATASETS, args.data_dir, used, sources, cfg=cfg)
    parts = [sam.collect(wb, RAW_POWER_DATASETS, args.data_dir, used, sources, cfg=cfg),
             sam.collect(wb, CAPACITY_DATASETS, args.data_dir, used, sources, cfg=cfg)]
    power = tuple(sum((p[i] for p in parts), []) for i in range(3))

    if not gas[0]:
        dash["B4"] = "Gas charts to be added: LNG stocks, imports and gas balances (pulls in development)."
    sam.draw_dashboard(dash, "Japan, South Korea and Taiwan energy - gas dashboard", *gas)
    sam.draw_dashboard(dash2, "Japan, South Korea and Taiwan energy - power", *power)

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    for s in sources:
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    for col, w in (("A", 14), ("B", 14), ("C", 38), ("D", 40), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: {len(gas[0])} gas charts, {len(power[0])} power charts; tabs {wb.sheetnames}")
    for label, m in (("gas", gas[2]), ("power", power[2])):
        if m:
            print(f"missing ({label}):", m)


if __name__ == "__main__":
    main()
