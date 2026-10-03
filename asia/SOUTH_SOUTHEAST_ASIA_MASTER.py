"""
South & Southeast Asia master workbook: one file with every dataset this repo pulls for South Asia
(India, Pakistan, Bangladesh, Sri Lanka, Nepal, Bhutan) and Southeast Asia (Thailand, Vietnam,
Philippines, Indonesia, Malaysia, Singapore, Myanmar, Cambodia, Laos, Brunei, Timor-Leste) - NOT
Japan, Taiwan, Korea or China. Same layout as the South & Central America, North America and
Australia/NZ masters (south_america/SOUTH_AMERICA_MASTER.py, whose table/chart/dashboard code this
reuses):

  Dashboard                  - gas: demand by sector, production, LNG imports/exports, gas prices
  Dashboard - Power & Hydro  - regional generation by source (monthly countries summed), then each
                               country: generation by source (raw grid-operator feed where there is one,
                               else Ember, labelled), demand, prices, capacity, hydro reservoirs (water year)
  <CC> <chart> data          - the table each Dashboard chart plots
  <CC> <dataset> raw         - the full data sheet(s) from each source workbook
  Sources                    - where each dataset comes from, units and notes

Raw sources (grid operators, ministries, regulators, statistics offices) are preferred; Ember is a
fallback for countries with no raw feed yet and is labelled as such. A country's raw generation
workbook replaces its Ember chart once it spans a year (sam.MIN_RAW_DAYS). Ember has no monthly
data for Indonesia, Myanmar, Cambodia, Laos, Brunei, Nepal, Bhutan or Timor-Leste, so those come
from Ember's yearly release until a raw feed exists.

Reads (doesn't refetch) the workbooks the scheduled pulls write to "output/Data and Chart Outputs/".
A workbook a pull has not written yet is skipped; one that fails to chart is listed on the Dashboard.

Usage: python3 asia/SOUTH_SOUTHEAST_ASIA_MASTER.py [--out "output/Data and Chart Outputs/south_southeast_asia_master.xlsx"]
"""
import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
import add_charts  # noqa: E402
import capacity_factors  # noqa: E402
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR
FUELS = ["Hydro", "Gas", "Wind", "Solar", "Coal", "Nuclear", "Other"]   # same order/colours as the other dashboards

# (country code, country, workbook, raw sheet / tuple of raw sheets / "*", short dataset name)
DATASETS = [
    ("SG", "Singapore", "singapore_gas.xlsx",
     ("Annual demand by sector", "Annual imports", "Power burn monthly (est)", "Town gas quarterly"), "gas"),
]
# Standard 'Daily' generation workbooks (date + <Fuel>_MWh), one per country
RAW_POWER_DATASETS = [
    ("IN", "India", "india_power_generation_daily.xlsx", "Daily", "power"),
    ("PK", "Pakistan", "pakistan_power_generation_daily.xlsx", "Daily", "power"),
    ("BD", "Bangladesh", "bangladesh_power_generation_daily.xlsx", "Daily", "power"),
    ("LK", "Sri Lanka", "sri_lanka_power_generation_daily.xlsx", "Daily", "power"),
    ("NP", "Nepal", "nepal_power_generation_daily.xlsx", "Daily", "power"),
    ("BT", "Bhutan", "bhutan_power_generation_daily.xlsx", "Daily", "power"),
    ("TH", "Thailand", "thailand_power_generation_daily.xlsx", "Daily", "power"),
    ("VN", "Vietnam", "vietnam_power_generation_daily.xlsx", "Daily", "power"),
    ("PH", "Philippines", "philippines_power_generation_daily.xlsx", "Daily", "power"),
    ("MY", "Malaysia", "malaysia_power_generation_daily.xlsx", "Daily", "power"),
    ("ID", "Indonesia", "indonesia_power_generation_daily.xlsx", "Daily", "power"),
]
# Ember fallback: monthly release (country sheets named as below) and yearly release
EMBER_FILES = [("SSEA", "South & Southeast Asia", "south_southeast_asia_power_by_type.xlsx")]
EMBER_ANNUAL_FILES = [("SSEA", "South & Southeast Asia (annual)", "south_southeast_asia_power_by_type_annual.xlsx")]
EMBER = {f for _, _, f in EMBER_FILES + EMBER_ANNUAL_FILES}
# Other power workbooks: own layouts, charted through add_charts.REGISTRY
OTHER_POWER_DATASETS = [
    ("SG", "Singapore", "singapore_power.xlsx",
     ("Daily demand", "Daily generation by type", "Monthly generation", "Annual consumption", "Annual fuel mix"),
     "power"),
]
CAPACITY_DATASETS = [(code, country, f"{country.lower().replace(' ', '_')}_power_capacity.xlsx", "Monthly",
                      "capacity")
                     for code, country, *_ in RAW_POWER_DATASETS]
HYDRO_DATASETS = []
HYDRO_EXTRA = {}
DASHBOARD_ONLY = {}
MASTER_SPECS = {}

SOURCES = {
    "singapore_gas.xlsx": ("EMA Singapore Energy Statistics (annual); SingStat town gas; power burn ESTIMATED from "
                           "EMC/NEMS metered CCGT generation", "https://www.ema.gov.sg/resources/singapore-energy-statistics"),
    "singapore_power.xlsx": ("EMA half-hourly system demand; EMC/NEMS metered generation by facility type; SingStat; "
                             "EMA SES", "https://www.ema.gov.sg/resources/statistics/half-hourly-system-demand-data"),
    "south_southeast_asia_power_by_type.xlsx": ("Ember monthly electricity data",
                                                "https://ember-energy.org/data/monthly-electricity-data/"),
    "south_southeast_asia_power_by_type_annual.xlsx": ("Ember yearly electricity data",
                                                       "https://ember-energy.org/data/yearly-electricity-data/"),
}
# The grid operator / statistics office Ember compiles each country from (named on Ember-fed charts)
OPERATORS = {"India": "Grid-India / CEA", "Pakistan": "NEPRA / NTDC", "Bangladesh": "PGCB / BPDB",
             "Sri Lanka": "CEB / PUCSL", "Thailand": "EPPO / EGAT", "Vietnam": "EVN / NSMO",
             "Philippines": "DOE / IEMOP", "Malaysia": "Energy Commission (ST)", "Singapore": "EMA",
             "Indonesia": "ESDM / PLN", "Myanmar": "MOEE", "Cambodia": "EAC", "Laos": "EDL / MEM",
             "Brunei": "Department of Electrical Services", "Nepal": "NEA", "Bhutan": "BPC / DGPC",
             "Timor-Leste": "EDTL"}

# Countries summed into the monthly regional total (Ember monthly or raw); the annual-only ones are charted alone
MONTHLY_COUNTRIES = ["India", "Pakistan", "Bangladesh", "Sri Lanka", "Thailand", "Vietnam", "Philippines",
                     "Malaysia", "Singapore"]
LAG_MONTHS = 6   # a country whose data stops more than this before the latest month is left out of the total


def monthly_gwh(path):
    """Standard Daily sheet (MWh) -> complete months, GWh, in the dashboard fuel groups."""
    d = add_charts.by_date(add_charts.read(path, "Daily"), "date")
    d = d[[c for c in d.columns if str(c).endswith("_MWh") and c != "Total_MWh"]].apply(pd.to_numeric,
                                                                                          errors="coerce")
    m = d.resample("MS").sum(min_count=1) / 1000.0
    last = d.dropna(how="all").index.max()
    monthly_rows = len(d) > 2 and d.index.to_series().diff().median().days > 20   # EPPO-style monthly feeds
    if not monthly_rows and last < last + pd.offsets.MonthEnd(0):   # drop the month in progress
        m = m[m.index < last.to_period("M").to_timestamp()]
    m = m[m.index >= "2021-01-01"].rename(columns=lambda c: c.replace("_MWh", "_GWh"))
    m = m.rename(columns={"Oil_GWh": "Other Fossil_GWh", "Other_GWh": "Other Renewables_GWh"})
    return add_charts.power_mix(m)


def regional_generation(data_dir, raw_files, frames_out=None):
    """South & Southeast Asia generation by source, TWh per month: each monthly country from its raw workbook if
    it has one (raw_files: country -> workbook), else Ember. Runs over the months every included country has;
    a country lagging more than LAG_MONTHS is left out (and listed) rather than cutting the total short."""
    ember = os.path.join(data_dir, EMBER_FILES[0][2])
    frames, notes = {}, []
    for country in MONTHLY_COUNTRIES:
        try:
            if country in raw_files:
                fname = raw_files[country]
                m = monthly_gwh(os.path.join(data_dir, fname))
                src = SOURCES.get(fname, (fname,))[0]
            else:
                m = add_charts.power_mix(add_charts.by_date(add_charts.read(ember, country), "Month"))
                src = f"Ember (compiled from {OPERATORS.get(country, 'the grid operator')})"
            m = m.apply(pd.to_numeric, errors="coerce")
            m = m[m.sum(axis=1) > 0]
            frames[country] = m
            notes.append(f"{country}: {src}, {m.index.min():%b/%y}-{m.index.max():%b/%y}")
        except Exception as e:  # noqa: BLE001
            notes.append(f"NOT INCLUDED: {country} ({type(e).__name__}: {e})")
    if frames_out is not None:
        frames_out.update(frames)
    if not frames:
        return pd.DataFrame(), notes
    latest = max(f.index.max() for f in frames.values())
    cutoff = latest - pd.DateOffset(months=LAG_MONTHS)
    for c in [c for c, f in frames.items() if f.index.max() < cutoff]:
        notes.append(f"NOT INCLUDED: {c} (data ends {frames[c].index.max():%b/%y})")
        del frames[c]
    months = sorted(set.intersection(*(set(f.index) for f in frames.values())))
    total = sum(f.reindex(index=months, columns=FUELS).fillna(0) for f in frames.values()) / 1000.0  # GWh -> TWh
    total.index.name = "date"
    return total, notes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "south_southeast_asia_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()
    cfg = sys.modules[__name__]
    exists = lambda d: os.path.exists(os.path.join(args.data_dir, d[2]))  # noqa: E731

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    dash2 = wb.create_sheet("Dashboard - Power & Hydro")
    used = {"Dashboard", "Dashboard - Power & Hydro", "Sources"}
    sources = []

    gas = sam.collect(wb, [d for d in DATASETS if exists(d)], args.data_dir, used, sources, cfg=cfg)

    # raw generation feeds replace Ember once they hold a year of history
    ember_countries = set()
    for _, _, f in EMBER_FILES + EMBER_ANNUAL_FILES:
        try:
            ember_countries |= set(pd.ExcelFile(os.path.join(args.data_dir, f)).sheet_names)
        except Exception:  # noqa: BLE001
            pass
    raw_power, building = [], []
    for d in filter(exists, RAW_POWER_DATASETS):
        days = sam.raw_history_days(os.path.join(args.data_dir, d[2]))
        if days >= sam.MIN_RAW_DAYS or d[1] not in ember_countries:
            raw_power.append(d)
        else:
            building.append(f"{d[1]} raw feed has {days} days so far - Ember shown until it has {sam.MIN_RAW_DAYS}")
    have_raw = {d[1] for d in raw_power}
    print(f"power by type: raw data for {sorted(have_raw) or 'none'}; Ember for the rest")
    parts = [sam.collect(wb, raw_power, args.data_dir, used, sources, cfg=cfg),
             *(sam.collect(wb, [(code, region, f, "*", "power")], args.data_dir, used, sources, skip=have_raw,
                           cfg=cfg)
               for code, region, f in EMBER_FILES + EMBER_ANNUAL_FILES if os.path.exists(os.path.join(args.data_dir, f))),
             sam.collect(wb, [d for d in OTHER_POWER_DATASETS if exists(d)], args.data_dir, used, sources, cfg=cfg),
             sam.collect(wb, [d for d in HYDRO_DATASETS if exists(d)], args.data_dir, used, sources, cfg=cfg)]
    power = tuple(sum((p[i] for p in parts), []) for i in range(3))
    power[2].extend(building)

    gen_frames = {}
    total, notes = regional_generation(args.data_dir, {d[1]: d[2] for d in raw_power}, gen_frames)
    if not total.empty:
        ws = wb.create_sheet(sam.sheet_name("SSEA generation total data", used))
        df, n_bars = xlsx_charts.prepare(total)
        xlsx_charts.write_table(ws, df)
        ws.cell(row=1, column=df.shape[1] + 4, value="Countries summed (only months all of them have):")
        for i, note in enumerate(notes, start=2):
            ws.cell(row=i, column=df.shape[1] + 4, value=note)
        raw_names = sorted(have_raw & set(MONTHLY_COUNTRIES))
        src = ("Sum of the country series on this dashboard (" +
               (f"raw: {', '.join(raw_names)}; " if raw_names else "") + "Ember for the rest)", None)
        power[0].insert(0, (xlsx_charts.build_chart(ws, df, n_bars, "South & Southeast Asia power generation by "
                                                    "source", "TWh per month", "stacked_bar", width=sam.CHART_W,
                                                    height=sam.CHART_H, gridlines=False,
                                                    inner=xlsx_charts.DASHBOARD_INNER), src))
        missing = [n.split(":", 1)[1].split("(")[0].strip() for n in notes if n.startswith("NOT INCLUDED")]
        name = ("South & Southeast Asia power generation by source (" + ", ".join(
            c for c in MONTHLY_COUNTRIES if c in gen_frames and c not in missing) +
            (f"; not included: {', '.join(missing)}" if missing else "") + ")")
        power[1].insert(0, ("South & Southeast Asia", name, df.index.max().strftime("%b/%y"), ws.title, *src))
        print("South & Southeast Asia generation total:", "; ".join(notes))

    cap = sam.collect(wb, [d for d in CAPACITY_DATASETS if exists(d)], args.data_dir, used, sources, cfg=cfg)
    pos = 1 if not total.empty else 0
    power[0][pos:pos] = cap[0]
    power[1][pos:pos] = cap[1]
    power[2].extend(cap[2])
    cf_dash = None
    cap_files = {c: os.path.join(args.data_dir, f) for _, c, f, _, _ in CAPACITY_DATASETS
                 if os.path.exists(os.path.join(args.data_dir, f))}
    if cap_files:
        cf_chart, cf_row, cf_missing, cf_countries = capacity_factors.add_capacity_factor_sheets(
            wb, used, sam.sheet_name, gen_frames, cap_files, sam.CHART_W, sam.CHART_H,
            "Country generation (raw feeds, else Ember) / installed capacity (raw feeds)", "South & Southeast Asia")
        if cf_chart:
            power[0].insert(pos + len(cap[0]), cf_chart)
            power[1].insert(pos + len(cap[1]), cf_row)
        power[2].extend(cf_missing)
        if cf_countries:
            cf_dash = wb.create_sheet("Dashboard - Capacity factors", 2)
            used.add("Dashboard - Capacity factors")
            sam.draw_dashboard(cf_dash, "South & Southeast Asia - capacity factor by generation type "
                               "(generation / capacity x hours)", [c for c, _ in cf_countries],
                               [r for _, r in cf_countries], cf_missing)

    sam.draw_dashboard(dash, "South & Southeast Asia energy (excl. Japan, Taiwan, Korea, China) - gas dashboard", *gas)
    sam.draw_dashboard(dash2, "South & Southeast Asia energy - power generation and hydro", *power)

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    for s in sources:
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    for col, w in (("A", 14), ("B", 14), ("C", 38), ("D", 40), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: {len(gas[0])} gas charts, {len(power[0])} power charts; tabs {wb.sheetnames}")
    for label, m in (("gas", gas[2]), ("power", power[2])):
        if m:
            print(f"missing ({label}):", m)


if __name__ == "__main__":
    main()
