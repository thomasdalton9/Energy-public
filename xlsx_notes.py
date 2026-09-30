"""
Shared helper for writing a highlighted "Units" (or "Notes") tab into
an Excel workbook alongside data tabs. Used by every script in this
repo that writes .xlsx, so units documentation looks and behaves the
same everywhere instead of being re-implemented per script.
"""

import os

import pandas as pd
from openpyxl.styles import Font, PatternFill

_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_BOLD = Font(bold=True)


def highlight_notes_sheet(worksheet, section_titles, column="A", width=110):
    """Bold + light-fill the header cell and any row whose text is an
    exact match in `section_titles`, so section headers stand out from
    the surrounding explanatory text."""
    worksheet[f"{column}1"].font = _BOLD
    worksheet[f"{column}1"].fill = _FILL
    worksheet.column_dimensions[column].width = width
    for row in range(2, worksheet.max_row + 1):
        cell = worksheet[f"{column}{row}"]
        if cell.value in section_titles:
            cell.font = _BOLD
            cell.fill = _FILL


def write_workbook(out_path, sheets, notes_lines, notes_section_titles, notes_sheet_name="Units"):
    """Write one .xlsx: a highlighted notes tab first, then one tab per
    (sheet_name, DataFrame) in `sheets` (a dict, insertion order kept).

    notes_lines: list of strings, one per row (blank strings for
    spacing). notes_section_titles: the subset of those strings that
    should be highlighted as section headers.
    """
    notes_df = pd.DataFrame({"Notes": notes_lines})
    # Written to a temp file in the same directory and swapped into place
    # with os.replace (atomic on the same filesystem), not written
    # directly to out_path - several scripts now checkpoint mid-run
    # (writing the same archive path many times over one run instead of
    # once at the end), so a process kill/timeout landing mid-write would
    # otherwise leave a half-written zip that pd.read_excel can't open on
    # the next run (load_archive's except clause only catches
    # FileNotFoundError/ValueError, not a corrupt-zip error), bricking
    # the pipeline instead of just losing the one in-flight checkpoint.
    # pandas' ExcelWriter picks its engine from the file extension, so the
    # temp name keeps out_path's real extension (e.g. .xlsx) rather than
    # just appending a suffix.
    root, ext = os.path.splitext(out_path)
    tmp_path = f"{root}.tmp{os.getpid()}{ext}"
    try:
        with pd.ExcelWriter(tmp_path, engine="openpyxl") as writer:
            notes_df.to_excel(writer, sheet_name=notes_sheet_name, index=False)
            highlight_notes_sheet(writer.sheets[notes_sheet_name], notes_section_titles)
            for name, df in sheets.items():
                df.to_excel(writer, sheet_name=name)
        os.replace(tmp_path, out_path)
    except PermissionError as exc:
        raise RuntimeError(
            f"Can't write {out_path} - it's probably open in Excel. "
            "Close the workbook and rerun the script."
        ) from exc
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
