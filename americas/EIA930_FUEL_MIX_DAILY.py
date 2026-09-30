"""
Pull EIA-930 (Hourly Electric Grid Monitor) daily net generation by fuel
type for the whole US grid AND for a set of major individual ISOs/RTOs,
and maintain a growing archive - one tab per region in a single
workbook.

Data source: EIA API v2, electricity/rto/daily-fuel-type-data -
https://api.eia.gov/v2/electricity/rto/daily-fuel-type-data/data/ -
"Daily net generation by balancing authority and energy source" per
Form EIA-930. Needs a free EIA API key (unlike MISO/FRED) - read from
the EIA_API_KEY environment variable (a GitHub Actions secret in this
repo, never printed or logged by this script). Found via
EIA930_DISCOVERY.py, including the exact respondent codes below
(confirmed against the API's own respondent facet list, not guessed).

REGIONS: US48 (EIA's own pre-aggregated whole-Lower-48 total) plus PJM,
MISO, CAISO, ERCOT, SPP, NYISO, ISO-NE, Southern Co., and TVA - each is
its own respondent code in EIA-930, reported the same way as US48 (not
summed here from smaller pieces). MISO here is a second, independent,
coarser (daily, not 5-minute) cross-check on MISO_FUEL_MIX_DAILY.py's
own real-time pull - both are kept, not merged.

Unlike MISO's own API (only "today"/"yesterday", no range query), this
dataset supports a real start/end range query with pagination (5000
rows/page) - so, like HENRY_HUB_DAILY.py, one script both backfills the
full history and keeps it current per region: the first run (no archive
yet, or a region's tab missing from it) pages through that region's
entire history back to the dataset's own start (2019-01-01); every run
after that just re-pulls a short rolling window per region (see
ROLLING_WINDOW_DAYS) and upserts, which also naturally catches any
revisions EIA makes to recently-published days.

TIMEZONE FACET: EIA-930 reports each region's total five times over,
once per US timezone (Arizona/Central/Eastern/Mountain/Pacific) - the
"timezone" facet is which timezone's midnight is used to bucket hourly
values into a calendar day, not a regional subset of the data. This
script fixes timezone=Eastern throughout for one consistent daily
series per region - an arbitrary but documented choice, not "the"
official cutoff (and not necessarily each ISO's own local timezone).

Values are net generation in megawatthours (MWh) - a daily total, not
an average MW like MISO's own script (MISO's 5-minute API instead means
a mean MW; there's no equivalent 5-minute read here, only EIA's own
daily aggregate).

Fuel-type categories are read from whatever the API actually reports
each pull (via each row's own "type-name"), not hardcoded, so a new
category EIA adds shows up automatically as a new column - and not
every region reports every category (e.g. ISO-NE has no Coal some
years), which shows up as that column simply being absent or NaN for
that region's tab.
"""

import argparse
import os
import sys
import time
from datetime import date, timedelta

import pandas as pd
import requests

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # repo root, for xlsx_notes

import xlsx_notes

API_KEY = os.environ.get("EIA_API_KEY")
URL = "https://api.eia.gov/v2/electricity/rto/daily-fuel-type-data/data/"
TIMEOUT = (10, 30)

# respondent code -> sheet name. Order here is the order tabs appear in
# the workbook. Codes confirmed via EIA930_DISCOVERY.py's respondent
# facet listing.
REGIONS = {
    "US48": "US_Total",
    "PJM": "PJM",
    "MISO": "MISO",
    "CISO": "CAISO",
    "ERCO": "ERCOT",
    "SWPP": "SPP",
    "NYIS": "NYISO",
    "ISNE": "ISONE",
    "SOCO": "Southern",
    "TVA": "TVA",
}

TIMEZONE = "Eastern"

# The dataset's own startPeriod as of EIA930_DISCOVERY.py's probe.
HISTORY_START = date(2019, 1, 1)

# Every run re-pulls this many trailing days per region (a single
# request each, well under the 5000-row page limit) and upserts them -
# catches EIA revising recently-published days, at negligible cost.
ROLLING_WINDOW_DAYS = 14

PAGE_LENGTH = 5000
FETCH_ATTEMPTS = 5
RETRY_BACKOFF_SECONDS = [5, 15, 30, 60]

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_OUT = os.path.join(REPO_ROOT, "output", "eia930_fuel_mix_daily.xlsx")

# EIA-930 publishes daily, with a lag of a few days for the "final"
# vintage - a few extra days of slack absorb that before this is
# treated as a genuinely stale source. Checked against US48 only (the
# headline series); a single smaller region lagging behind doesn't by
# itself fail the whole run.
STALE_AFTER_DAYS = 10


def fetch_page(respondent, start, end, offset):
    params = {
        "frequency": "daily",
        "data[0]": "value",
        "facets[respondent][]": respondent,
        "facets[timezone][]": TIMEZONE,
        "start": start.isoformat(),
        "end": end.isoformat(),
        "sort[0][column]": "period",
        "sort[0][direction]": "asc",
        "offset": offset,
        "length": PAGE_LENGTH,
        "api_key": API_KEY,
    }
    last_error = None
    for attempt in range(1, FETCH_ATTEMPTS + 1):
        try:
            r = requests.get(URL, params=params, timeout=TIMEOUT)
            r.raise_for_status()
            return r.json()
        except requests.RequestException as e:
            last_error = e
            print(f"    attempt {attempt}/{FETCH_ATTEMPTS} failed: {type(e).__name__}: {e}", file=sys.stderr)
            if attempt < FETCH_ATTEMPTS:
                wait = RETRY_BACKOFF_SECONDS[attempt - 1]
                print(f"    waiting {wait}s before retrying...", file=sys.stderr)
                time.sleep(wait)
    raise last_error


def fetch_range(respondent, start, end):
    """Page through offset/length until every row in [start, end] is
    collected for one respondent."""
    rows = []
    offset = 0
    while True:
        payload = fetch_page(respondent, start, end, offset)
        page_rows = payload["response"]["data"]
        rows.extend(page_rows)
        total = int(payload["response"]["total"])
        offset += len(page_rows)
        print(f"    fetched {offset}/{total} rows...", file=sys.stderr)
        if not page_rows or offset >= total:
            break
    return rows


def sanitize_column(type_name):
    return type_name.strip().replace(" ", "_").replace("-", "_")


def rows_to_wide(rows):
    """Long rows (period, fueltype, type-name, value) -> wide
    DataFrame, index = date, columns = "{type-name}_MWh"."""
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows)
    df["period"] = pd.to_datetime(df["period"]).dt.date
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    df["column"] = df["type-name"].map(sanitize_column) + "_MWh"
    wide = df.pivot_table(index="period", columns="column", values="value", aggfunc="last")
    wide.index.name = "date"
    return wide.sort_index()


CHART_LOOKBACK_MONTHS = 36

# Kept as their own series in the chart (in stacking order); every other
# *_MWh column present (EIA-930 reports up to 16 - battery/pumped/other
# storage variants, Unknown, Petroleum, Geothermal, etc.) is summed into
# a single "Other_MWh" bucket instead. A full-history, full-category
# stacked area chart (2,800+ days x up to 16 series) at the requested
# 5.5x2.8in size renders as unreadable noise - confirmed by rendering it
# and looking at the result - so the embedded chart is a deliberately
# reduced view; the untouched, full-category, full-history data stays in
# the sheet's own columns for anyone who wants it.
MAJOR_FUEL_COLUMNS = ["Nuclear_MWh", "Coal_MWh", "Natural_Gas_MWh", "Hydro_MWh", "Wind_MWh", "Solar_MWh"]


def add_summary_columns(df):
    fuel_cols = [c for c in df.columns if c.endswith("_MWh")]
    df["Total_MWh"] = df[fuel_cols].sum(axis=1, skipna=True)
    # Renewables share computed from known type-names actually present, matched case-
    # insensitively so it works regardless of EIA's exact capitalization, and regardless
    # of which categories a given region happens to report.
    renewable_keywords = ["wind", "solar", "hydro", "geothermal", "biomass"]
    renewable_cols = [c for c in fuel_cols if any(kw in c.lower() for kw in renewable_keywords)]
    renewable_sum = df[renewable_cols].sum(axis=1, skipna=True) if renewable_cols else 0.0
    df["Renewables_Share"] = renewable_sum / df["Total_MWh"].where(df["Total_MWh"] != 0)
    return df


def load_archive(path, sheet_name):
    try:
        df = pd.read_excel(path, sheet_name=sheet_name, index_col=0)
    except (FileNotFoundError, ValueError, KeyError):
        return pd.DataFrame()
    df.index = pd.to_datetime(df.index).date
    df.index.name = "date"
    return df


def upsert(existing, new_df):
    if new_df.empty:
        return existing
    if existing.empty:
        combined = new_df
    else:
        combined = pd.concat([existing, new_df])
        combined = combined[~combined.index.duplicated(keep="last")]
    combined.index.name = "date"
    return combined.sort_index()


def format_date_columns(path, sheet_names):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    for sheet in sheet_names:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for (cell,) in ws.iter_rows(min_row=2, max_col=1):
            cell.number_format = "dd-mmm-yyyy"
        ws.column_dimensions["A"].width = 14
    wb.save(path)


def build_chart_helper(combined_df, months=CHART_LOOKBACK_MONTHS):
    """Reduce a region's full archive to what the embedded chart actually
    plots: the trailing `months` months, major categories kept separate,
    everything else summed into one Other_MWh column."""
    if combined_df.empty:
        return pd.DataFrame()
    max_date = pd.Timestamp(max(combined_df.index))
    cutoff = (max_date - pd.DateOffset(months=months)).date()
    windowed = combined_df[[d >= cutoff for d in combined_df.index]]
    if windowed.empty:
        return pd.DataFrame()

    fuel_cols = [c for c in windowed.columns if c.endswith("_MWh") and c != "Total_MWh"]
    major_present = [c for c in MAJOR_FUEL_COLUMNS if c in fuel_cols]
    other_cols = [c for c in fuel_cols if c not in major_present]

    helper = windowed[major_present].copy() if major_present else pd.DataFrame(index=windowed.index)
    if other_cols:
        helper["Other_MWh"] = windowed[other_cols].sum(axis=1, skipna=True)
    return helper


HELPER_START_COL = 40  # comfortably clear of the archive's own columns (up to ~18 wide)


def write_chart_helper(ws, helper_df, start_col=HELPER_START_COL):
    """Writes a small, clearly-labelled block far to the right of a
    region's real data - date + consolidated categories only, the exact
    slice the embedded chart plots - so the chart's own Reference ranges
    don't depend on the full archive's (currently 16) category columns
    or its ever-growing row count in a way that would put years of daily
    noise into a 2.8-inch-tall chart. Returns the range info
    add_region_chart needs, or None if there's nothing to chart."""
    if helper_df.empty:
        return None

    label_row = 1
    header_row = 2
    ws.cell(row=label_row, column=start_col,
            value=f"Chart data only - trailing {CHART_LOOKBACK_MONTHS} months, minor categories "
                   "combined into Other_MWh. Full history/categories are in the columns to the left.")

    ws.cell(row=header_row, column=start_col, value="date")
    for j, col_name in enumerate(helper_df.columns):
        ws.cell(row=header_row, column=start_col + 1 + j, value=col_name)

    for i, (idx, row) in enumerate(helper_df.iterrows()):
        excel_row = header_row + 1 + i
        date_cell = ws.cell(row=excel_row, column=start_col, value=idx)
        date_cell.number_format = "mmm/yy"
        for j, col_name in enumerate(helper_df.columns):
            val = row[col_name]
            ws.cell(row=excel_row, column=start_col + 1 + j,
                    value=None if pd.isna(val) else float(val))

    return {
        "date_col": start_col,
        "first_data_col": start_col + 1,
        "last_data_col": start_col + len(helper_df.columns),
        "header_row": header_row,
        "n_rows": len(helper_df),
    }


CHART_FONT_SIZE = 1200  # openpyxl font sizes are in hundredths of a point - 1200 = 12pt


def _sized_title(text):
    from openpyxl.chart.text import RichText, Text
    from openpyxl.chart.title import Title
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun

    cp = CharacterProperties(sz=CHART_FONT_SIZE, b=True)
    run = RegularTextRun(t=text, rPr=cp)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[run])
    return Title(tx=Text(rich=RichText(p=[para])))


def _sized_text_props():
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties

    cp = CharacterProperties(sz=CHART_FONT_SIZE)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp))])


# Explicit per-category colors - Excel's automatic palette put several of
# these series (Hydro/Wind/Other especially) at near-zero contrast against
# each other and the plot background, effectively invisible (confirmed by
# actually looking at a rendered chart, not just inspecting the XML). Same
# palette family as MISO_FUEL_MIX_CHART_PREVIEW.py's COLORS.
CHART_COLORS = {
    "Nuclear_MWh": "7B2D8B",
    "Coal_MWh": "4D4D4D",
    "Natural_Gas_MWh": "E8743B",
    "Hydro_MWh": "4A90D9",
    "Wind_MWh": "3F8F4F",
    "Solar_MWh": "F9D71C",
    "Other_MWh": "BFBFBF",
}


def add_region_chart(ws, sheet_name, helper_range, columns):
    """Native (embedded, editable-in-Excel) stacked area chart of a
    region's generation by fuel type - not a static image, and not
    win32com-automated (this pipeline runs on Linux GitHub Actions
    runners, no Excel installed), just openpyxl's chart objects, which
    Excel opens and renders like any chart built by hand. Plots the
    reduced helper block (see build_chart_helper/write_chart_helper),
    not the sheet's full archive columns - a full-history, full-
    category version at this chart's small size renders as unreadable
    noise (confirmed by actually rendering it).

    columns: helper_df.columns in order - used to assign each series its
    fixed color and, in principle, to line series up with names, since a
    region missing a major category (e.g. ISO-NE has no Coal some years)
    still needs its remaining series colored correctly, not shifted."""
    from openpyxl.chart import AreaChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.utils import get_column_letter

    if helper_range is None:
        return

    header_row = helper_range["header_row"]
    last_row = header_row + helper_range["n_rows"]

    chart = AreaChart()
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = _sized_title(f"{sheet_name} daily generation by fuel type")

    # A plain category (text) axis showed NO date labels at all once there
    # were 1,000+ distinct daily categories - Excel's auto tick-skip gave up
    # rather than pick a readable interval (confirmed by actually rendering
    # this chart, not just inspecting its XML). A real DateAxis lets Excel
    # apply its own sensible month-based ticks instead.
    chart.x_axis = DateAxis(axId=10, crossAx=100)
    chart.x_axis.number_format = "mmm/yy"
    chart.x_axis.number_format_source_linked = False
    chart.x_axis.majorTimeUnit = "months"
    chart.x_axis.baseTimeUnit = "days"
    # "Layout 1": title only, no axis titles - but the tick labels
    # themselves (dates, MWh values) must stay on, which needs both
    # `delete = False` and an explicit tick-label position; leaving
    # either unset was enough for Excel to render axes with no labels
    # at all (again, only visible by actually rendering the chart).
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.tickLblPos = "nextTo"
    chart.y_axis.tickLblPos = "nextTo"
    chart.x_axis.txPr = _sized_text_props()
    chart.y_axis.txPr = _sized_text_props()
    # Chart size 5.5 x 2.8 inches - openpyxl uses centimeters (1 in = 2.54 cm).
    chart.width = 5.5 * 2.54
    chart.height = 2.8 * 2.54
    # No border around the chart area.
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))

    data = Reference(ws, min_col=helper_range["first_data_col"], max_col=helper_range["last_data_col"],
                      min_row=header_row, max_row=last_row)
    cats = Reference(ws, min_col=helper_range["date_col"], min_row=header_row + 1, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    for series, col_name in zip(chart.series, columns):
        color = CHART_COLORS.get(col_name)
        if color:
            series.graphicalProperties = GraphicalProperties(
                solidFill=color, ln=LineProperties(noFill=True)
            )

    chart.legend.position = "b"
    if chart.legend is not None:
        chart.legend.txPr = _sized_text_props()

    anchor_col_index = helper_range["last_data_col"] + 3
    ws.add_chart(chart, f"{get_column_letter(anchor_col_index)}2")


def add_charts(path, region_data):
    """region_data: {sheet_name: full combined DataFrame} for each
    region - reduced per-region via build_chart_helper() before being
    written and charted."""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    for sheet_name, combined_df in region_data.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        helper_df = build_chart_helper(combined_df)
        helper_range = write_chart_helper(ws, helper_df)
        add_region_chart(ws, sheet_name, helper_range, list(helper_df.columns))
    wb.save(path)


NOTES_LINES = [
    "UNITS",
    "All *_MWh columns are megawatthours - the total net generation for that fuel category over the "
    "whole day (a daily energy total, not an average MW). Total_MWh is the sum of the category "
    "columns. Renewables_Share is (Wind + Solar + Hydro + Geothermal + Biomass) as a share of "
    "Total_MWh.",
    "",
    "TABS / SCOPE",
    "One tab per region, each EIA-930's own pre-aggregated total for that respondent code - not "
    "summed here from smaller pieces: US_Total (US48, the whole Lower 48), PJM, MISO, CAISO (CISO), "
    "ERCOT (ERCO), SPP (SWPP), NYISO (NYIS), ISONE (ISNE), Southern (SOCO), TVA. The MISO tab here is "
    "a second, independent, coarser (daily, not 5-minute) cross-check on miso_fuel_mix_daily.xlsx's "
    "real-time pull - both are kept, not merged.",
    "",
    "CATEGORIES",
    "Whatever EIA-930 itself reports per region (e.g. Coal, Natural gas, Nuclear, Wind, Solar, Hydro, "
    "Petroleum, Battery storage, Other) - not a fixed list maintained here, so a category EIA adds "
    "shows up as a new column automatically, and a region that doesn't report a given category simply "
    "has no column (or NaN rows) for it.",
    "",
    "TIMESTAMPS",
    f"The 'date' index uses EIA-930's own \"{TIMEZONE}\" timezone facet to bucket hourly readings into "
    "a calendar day - EIA-930 reports each region's total once per US timezone (Arizona/Central/"
    f"Eastern/Mountain/Pacific), each with a different day boundary; {TIMEZONE} was picked as one "
    "consistent choice applied to every tab, not an official EIA cutoff or each ISO's own local zone.",
    "",
    "COVERAGE",
    f"Backed by a real start/end range query (unlike MISO's own API) - each tab is seeded back to "
    f"{HISTORY_START.isoformat()} (the dataset's own start) the first time that tab is built, then "
    f"kept current by re-pulling and upserting the trailing {ROLLING_WINDOW_DAYS} days on every run, "
    "which also catches EIA revising recently-published days.",
    "",
    "SOURCE",
    f"EIA (US Energy Information Administration) API v2, Form EIA-930 (Hourly Electric Grid Monitor): "
    f"{URL} - needs a free EIA API key (api.eia.gov), unlike MISO's public API or FRED's CSV endpoint.",
]
NOTES_SECTION_TITLES = {"UNITS", "TABS / SCOPE", "CATEGORIES", "TIMESTAMPS", "COVERAGE", "SOURCE"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default=DEFAULT_OUT)
    args = parser.parse_args()

    if not API_KEY:
        print("EIA_API_KEY is not set in the environment - aborting.", file=sys.stderr)
        sys.exit(1)

    today = date.today()
    sheets = {}
    us_total_latest = None

    for respondent, sheet_name in REGIONS.items():
        print(f"[{sheet_name}] ({respondent})", file=sys.stderr)
        existing = load_archive(args.out, sheet_name)

        if existing.empty:
            print(f"  no existing tab - backfilling full history {HISTORY_START} to {today} ...",
                  file=sys.stderr)
            start = HISTORY_START
        else:
            start = max(HISTORY_START, today - timedelta(days=ROLLING_WINDOW_DAYS))
            print(f"  existing tab has {len(existing)} days - refreshing rolling window "
                  f"{start} to {today} ...", file=sys.stderr)

        rows = fetch_range(respondent, start, today)
        new_df = rows_to_wide(rows)
        if not new_df.empty:
            new_df = add_summary_columns(new_df)

        before_days = set(existing.index) if not existing.empty else set()
        combined = upsert(existing, new_df)
        new_days = sorted(set(combined.index) - before_days)

        print(f"  added {len(new_days)} new day(s){' - ' + str(new_days[-5:]) if new_days else ''}")
        if not combined.empty:
            print(f"  tab now has {len(combined)} days ({combined.index.min()} to {combined.index.max()})")

        sheets[sheet_name] = combined
        if respondent == "US48" and not combined.empty:
            us_total_latest = max(combined.index)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    xlsx_notes.write_workbook(args.out, sheets, NOTES_LINES, NOTES_SECTION_TITLES)
    format_date_columns(args.out, list(sheets.keys()))
    add_charts(args.out, sheets)

    print(f"Saved to {args.out}")

    age = (date.today() - us_total_latest).days if us_total_latest else None
    if us_total_latest is None or age > STALE_AFTER_DAYS:
        print(f"STALE SOURCE: US_Total's newest saved day is {us_total_latest} ({age} days old) - "
              f"{URL} may have changed or stopped updating.", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
