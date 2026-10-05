"""
Texas gas and power demand REGRESSIONS (monthly), weather-normalised forecast to Dec 2033. A standalone workbook:
it reads texas_gas_monthly.xlsx, eia930_fuel_mix_daily.xlsx, henry_hub_daily.xlsx and (fallback only) macro_drivers.xlsx and
changes none of them or of the Texas forecast scripts; the owner may wire it into the existing forecast later.
Runs in GitHub Actions (.github/workflows/texas_demand_regression.yml, 1st and 15th, after texas_gas); NASA POWER, FRED and the
EIA-930 data are not reachable from the Claude sandbox.

Models (ordinary least squares, numpy; no statsmodels needed)
  1. Texas gas consumption by sector, Bcf/d, monthly from 2015: electric power, residential, commercial, industrial (published
     EIA sectors) on weather (HDD/CDD per day, base 18C, population-weighted over Houston, Dallas-Fort Worth, San Antonio, Austin and
     El Paso; NASA POWER daily T2M, Open-Meteo fallback) OR month dummies, Texas population, a linear trend and (electric power
     only) Henry Hub and the ERCOT wind+solar share. Every combination is a candidate; the chosen one has the lowest error on a
     hold-out of the last 12 months (fit on the months before) among candidates whose signs are plausible (HDD>0 for
     residential/commercial, CDD>0 for electric power, population>=0). All candidates are listed on the 'Candidates' tab.
  2. ERCOT load (EIA-930 ERCOT total net generation as the load proxy, average GW) on HDD, CDD and population. The BASELINE is fit on a
     pre-step-up window (editable, default Jan 2019 - Dec 2022) and extrapolated; actual minus baseline is the 'unexplained load
     growth' (the data-centre signal). A full-sample model with a trend is reported for reference with a hold-out.
  3. Forecast to Dec 2033 at normal weather (average HDD/CDD by calendar month over an editable window of years), population
     growing at an editable rate, Henry Hub and wind+solar share held at editable values, +/-1 prediction standard error band.
  4. LOAD BREAKOUT ('Load breakout' tab + chart): the ERCOT unexplained load growth APPORTIONED into data centres, crypto mining, oil & gas /
     Permian electrification, other industrial (incl. LNG terminal load) and an unattributed remainder by EDITABLE yellow shares per
     case (LOW/BASE/HIGH) and year 2023-26 (live Excel formulas; optional absolute-GW override per category and year). THE SHARES ARE
     JUDGEMENT (seeded 'unverified, from memory'), NOT MEASUREMENT: nothing in the data splits the residual by source. Data-centre
     GW is converted to gas burn with the factors read from texas_production_forecast.xlsx 'Assump - Data centres' (read only) and
     compared with that workbook's LOW/BASE/HIGH data-centre path at Dec 2026.
Edit the yellow cells of the Assumptions and Load breakout tabs (committed workbook is read back on every run) and re-run.

Usage: python3 TEXAS_DEMAND_REGRESSION.py [--out "output/Data and Chart Outputs/texas_demand_regression.xlsx"]
"""
import argparse
import datetime as dt
import io
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from itertools import product

import numpy as np
import pandas as pd
import requests

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import xlsx_notes  # noqa: E402

OUT_DIR = os.path.join(ROOT, "output", "Data and Chart Outputs")
DEFAULT_OUT = os.path.join(OUT_DIR, "texas_demand_regression.xlsx")
GAS_XLSX = os.path.join(OUT_DIR, "texas_gas_monthly.xlsx")
E930_XLSX = os.path.join(OUT_DIR, "eia930_fuel_mix_daily.xlsx")
HH_XLSX = os.path.join(OUT_DIR, "henry_hub_daily.xlsx")
MACRO_XLSX = os.path.join(OUT_DIR, "macro_drivers.xlsx")

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"}
NASA = "https://power.larc.nasa.gov/api/temporal/daily/point"
OPEN_METEO = "https://archive-api.open-meteo.com/v1/archive"
FRED = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=TXPOP"
START = pd.Timestamp("2015-01-01")
END = pd.Timestamp("2033-12-01")
BASE_C = 18.0
# metro populations in millions, approx. 2020 Census (weights only; from memory, unverified)
CITIES = [("Houston", 29.76, -95.37, 7.12), ("Dallas-Fort Worth", 32.78, -96.80, 7.64), ("San Antonio", 29.42, -98.49, 2.56),
          ("Austin", 30.27, -97.74, 2.28), ("El Paso", 31.76, -106.49, 0.87)]
SECTORS = ["Electric power", "Residential", "Commercial", "Industrial"]
HOLD = 12
TEXAS_SHARE_OF_US_2020 = 0.0879   # fallback only: Texas 29.1m / US 331m


def log(*a):
    print(*a, flush=True)


# ------------------------------------------------------------------ data pulls
def get_json(url, params, tries=3, timeout=300):
    last = None
    for i in range(tries):
        try:
            r = requests.get(url, params=params, headers=UA, timeout=timeout)
            r.raise_for_status()
            return r.json()
        except Exception as exc:  # noqa: BLE001
            last = exc
            time.sleep(5 * (i + 1))
    raise last


def fetch_city(job):
    name, lat, lon, start, end = job
    try:
        js = get_json(NASA, {"parameters": "T2M", "community": "RE", "longitude": lon, "latitude": lat,
                             "start": start.strftime("%Y%m%d"), "end": end.strftime("%Y%m%d"), "format": "JSON"})
        t = js["properties"]["parameter"]["T2M"]
        s = pd.Series({pd.Timestamp(k): (np.nan if v is None or v <= -998 else float(v)) for k, v in t.items()}).dropna()
        if s.empty:
            raise RuntimeError("no data")
        return name, s, "NASA POWER"
    except Exception as exc:  # noqa: BLE001
        log(f"  NASA POWER failed for {name}: {exc}; trying Open-Meteo")
    try:
        time.sleep(8)
        js = get_json(OPEN_METEO, {"latitude": lat, "longitude": lon, "start_date": start.date().isoformat(),
                                   "end_date": end.date().isoformat(), "daily": "temperature_2m_mean", "timezone": "UTC"})
        d = js["daily"]
        s = pd.Series(d["temperature_2m_mean"], index=pd.to_datetime(d["time"]), dtype=float).dropna()
        return name, s, "Open-Meteo ERA5"
    except Exception as exc:  # noqa: BLE001
        log(f"  Open-Meteo failed for {name}: {exc}")
        return name, pd.Series(dtype=float), "none"


def update_temps(saved):
    """Incremental city store: fetch only days after the last saved one (minus a 10-day revision window)."""
    end = pd.Timestamp(dt.date.today() - dt.timedelta(days=1))
    jobs = []
    for name, lat, lon, _ in CITIES:
        if name in saved and saved[name].notna().any():
            start = max(START, saved[name].dropna().index.max() - pd.Timedelta(days=10))
        else:
            start = START
        if start <= end:
            jobs.append((name, lat, lon, start, end))
    log(f"weather: {len(jobs)} city requests")
    new, srcs = {}, set()
    with ThreadPoolExecutor(max_workers=5) as ex:
        for name, s, src in ex.map(fetch_city, jobs):
            if not s.empty:
                new[name] = s.round(2)
                srcs.add(src)
    out = pd.DataFrame(new).combine_first(saved) if new else saved.copy()
    out = out.sort_index()
    out.index.name = "date"
    return out[[c[0] for c in CITIES if c[0] in out]], (", ".join(sorted(srcs)) or "saved store only")


def degree_days(temps):
    """Population-weighted daily HDD/CDD (a day counts only when every city has data), summed to months."""
    w = pd.Series({c[0]: c[3] for c in CITIES})
    t = temps[list(w.index)].dropna()
    hdd = ((BASE_C - t).clip(lower=0) * w).sum(axis=1) / w.sum()
    cdd = ((t - BASE_C).clip(lower=0) * w).sum(axis=1) / w.sum()
    m = pd.DataFrame({"HDD_total": hdd.resample("MS").sum(), "CDD_total": cdd.resample("MS").sum(), "Days": hdd.resample("MS").count()})
    m["Days in month"] = m.index.days_in_month
    m["Complete"] = m["Days"] == m["Days in month"]
    m["HDD per day"] = m["HDD_total"] / m["Days"]
    m["CDD per day"] = m["CDD_total"] / m["Days"]
    m.index.name = "Month"
    return m


def get_population(prev, prem=0.9):
    """Annual Texas population (million). FRED TXPOP (Census, July 1 estimates); else the saved store; else IMF/World Bank US x share."""
    try:
        r = requests.get(FRED, headers=UA, timeout=20)
        r.raise_for_status()
        d = pd.read_csv(io.StringIO(r.text))
        d.columns = ["date", "v"]
        d["v"] = pd.to_numeric(d["v"], errors="coerce")
        d = d.dropna()
        s = pd.Series(d["v"].values / 1000.0, index=pd.to_datetime(d["date"]).dt.year)
        s = s[s.index >= 2000]
        if len(s) > 10:
            return s, "FRED TXPOP (US Census Bureau resident population, July 1 estimates)"
    except Exception as exc:  # noqa: BLE001
        log("population: FRED failed:", exc)
    if prev is not None and len(prev) and "FRED" in str(prev["Source"].iloc[-1]) and "unreachable" not in str(prev["Source"].iloc[-1]):
        return prev["Population_m"], "saved copy from an earlier run (FRED not reachable): " + str(prev["Source"].iloc[-1])
    try:
        m = pd.read_excel(MACRO_XLSX, sheet_name="Population_IMF_m").set_index("year")["USA"].dropna()
        g = m.pct_change().fillna(0) + prem / 100
        lvl = {2020: 29.145}   # Texas 2020 Census, million
        for y in range(2021, int(m.index.max()) + 1):
            lvl[y] = lvl[y - 1] * (1 + g[y])
        for y in range(2019, 1999, -1):
            lvl[y] = lvl[y + 1] / (1 + g[y + 1])
        s = pd.Series(lvl).sort_index()
        return s, f"FALLBACK (FRED/Census unreachable): Texas = 2020 Census 29.145m carried with IMF WEO US population growth + {prem} pp a year (editable Texas premium); NOT Census estimates"
    except Exception as exc:  # noqa: BLE001
        raise SystemExit(f"no population source: {exc}")


def pop_monthly(idx, annual, g):
    """Linear interpolation between July-1 annual points; beyond the last point compound growth g per year."""
    pts = pd.Series(annual.values, index=[pd.Timestamp(int(y), 7, 1) for y in annual.index])
    x = np.array([p.toordinal() for p in pts.index], float)
    out = []
    for m in idx:
        mid = m + pd.Timedelta(days=m.days_in_month // 2)
        o = mid.toordinal()
        if o <= x[-1]:
            out.append(float(np.interp(o, x, pts.values)))
        else:
            out.append(float(pts.values[-1] * (1 + g) ** ((o - x[-1]) / 365.25)))
    return pd.Series(out, index=idx)


# ------------------------------------------------------------------ regression helpers
def ols(X, y):
    X, y = np.asarray(X, float), np.asarray(y, float)
    n, k = X.shape
    inv = np.linalg.pinv(X.T @ X)
    b = inv @ X.T @ y
    res = y - X @ b
    dof = max(n - k, 1)
    s2 = float(res @ res) / dof
    se = np.sqrt(np.diag(inv) * s2)
    sst = float(((y - y.mean()) ** 2).sum())
    r2 = 1 - float(res @ res) / sst
    return {"b": b, "se": se, "r2": r2, "adj": 1 - (1 - r2) * (n - 1) / dof, "s": np.sqrt(s2), "n": n, "k": k, "inv": inv, "s2": s2}


def predict(fit, X):
    X = np.asarray(X, float)
    mu = X @ fit["b"]
    se = np.sqrt(fit["s2"] * (1 + np.einsum("ij,jk,ik->i", X, fit["inv"], X)))
    return mu, se


def design(d, weather, pt, extras):
    X = pd.DataFrame({"const": 1.0}, index=d.index)
    if weather == "temp":
        X["HDD per day"] = d["hdd"]
        X["CDD per day"] = d["cdd"]
    else:
        for m in range(2, 13):
            X[f"Month {m:02d} (vs Jan)"] = (d.index.month == m).astype(float)
    if "pop" in pt:
        X["Population (million)"] = d["pop"]
    if "trend" in pt:
        X["Trend (years)"] = d["trend"]
    if "hh" in extras:
        X["Henry Hub ($/MMBtu)"] = d["hh"]
    if "re" in extras:
        X["ERCOT wind+solar share"] = d["re"]
    return X


def plausible(sector, names, b):
    ok, why = True, []
    for n, v in zip(names, b):
        if n == "HDD per day" and sector in ("Residential", "Commercial") and v <= 0:
            ok = False; why.append("HDD coefficient <= 0")
        if n == "CDD per day" and sector == "Electric power" and v <= 0:
            ok = False; why.append("CDD coefficient <= 0")
        if n == "Population (million)" and v < 0:
            ok = False; why.append("population coefficient < 0")
    return ok, "; ".join(why)


def fit_candidate(sector, d, y, weather, pt, extras):
    X = design(d, weather, pt, extras)
    rows = d.index[y.notna() & X.notna().all(axis=1)]
    last = y.dropna().index.max()
    train = rows[rows <= last - pd.DateOffset(months=HOLD)]
    test = rows[rows > last - pd.DateOffset(months=HOLD)]
    if len(train) < len(X.columns) + 12 or len(test) < 6:
        return None
    f_tr = ols(X.loc[train], y.loc[train])
    p, _ = predict(f_tr, X.loc[test])
    err = p - y.loc[test].values
    full = ols(X.loc[rows], y.loc[rows])
    ok, why = plausible(sector, list(X.columns), full["b"])
    return {"weather": weather, "pt": pt, "extras": extras, "X": X, "rows": rows, "fit": full,
            "rmse": float(np.sqrt((err ** 2).mean())), "mape": float(np.mean(np.abs(err) / y.loc[test].values) * 100),
            "bias": float(err.mean()), "ok": ok, "why": why, "test": test, "pred_test": p}


def label(c):
    w = "HDD/CDD" if c["weather"] == "temp" else "month dummies"
    pt = {"pop+trend": "population + trend", "pop": "population", "trend": "trend"}[c["pt"]]
    ex = "".join(f" + {e}" for e in c["extras"]).replace("hh", "Henry Hub").replace("re", "wind+solar share")
    return f"{w}, {pt}{ex}"


MEANING = {
    "const": "intercept (level when every driver is zero; not physical on its own)",
    "HDD per day": "Bcf/d per +1 average daily heating degree-day (18C) in the month; one cold month at 10 HDD/day vs 0 adds 10x this",
    "CDD per day": "Bcf/d per +1 average daily cooling degree-day (18C); a hot month at 10 CDD/day vs 0 adds 10x this",
    "Population (million)": "Bcf/d per +1 million Texans (with a trend in the model, the two share the same growth and are hard to separate)",
    "Trend (years)": "Bcf/d per year beyond the other drivers (efficiency, electrification, new load, anything omitted)",
    "Henry Hub ($/MMBtu)": "Bcf/d per +$1/MMBtu monthly Henry Hub (negative = switching to coal/renewables when gas is dear)",
    "ERCOT wind+solar share": "Bcf/d per +1.0 (=100 percentage points) wind+solar share; divide by 10 for +10 points",
}


def meaning(n):
    return MEANING.get(n, "Bcf/d versus January, all else equal (seasonal shape that weather does not explain)" if n.startswith("Month") else "")


# ------------------------------------------------------------------ load breakout (apportioning the ERCOT unexplained load)
BO, BOV = "Load breakout", "Load breakout values"
CATS = ["Data centres", "Crypto mining", "Oil & gas / Permian electrification", "Other industrial (incl. LNG terminal load)",
        "Unattributed (population/baseline error)"]
BCASES = ["BASE", "LOW", "HIGH"]
BYEARS = [2023, 2024, 2025, 2026]
UNV = "unverified (from memory) - judgement, not measurement"
PROD_XLSX = os.path.join(OUT_DIR, "texas_production_forecast.xlsx")
ERCOT_BURN_XLSX = os.path.join(OUT_DIR, "ercot_gas_burn_daily.xlsx")
# sheet layout (rows 1-based); add_charts.py reads these constants to link its chart sheet to the live formulas
SH_HDR, SH0 = 7, 8                       # share block: 3 cases x (5 categories + sum row + spacer)
OV_HDR, OV0 = 31, 32                     # override block: 4 categories
CF = {"lf": 39, "gas": 40, "hr": 41, "mcf": 42, "pavg": 43, "pen": 44, "stock": 45}
T12_HDR, T12_0 = 48, 49                  # latest-12-month table: 5 categories, total, stack check
DCR = {"u": 59, "slope": 60, "months": 61, "proj": 62, "share": 63, "avg": 64, "inc": 65, "stock": 66, "imp": 67, "scen": 68,
       "diff": 69, "burn12": 70, "burn_inc": 71, "burn_imp": 72, "flag": 73}
MON_HDR, MON0 = 76, 77                   # monthly block from Jan 2022
MON_START = pd.Timestamp("2022-01-01")
CHART_FROM = pd.Timestamp("2023-01-01")
COL_BASE, COL_LOW, COL_HIGH = ["C", "D", "E", "F", "G"], ["H", "I", "J", "K", "L"], ["M", "N", "O", "P", "Q"]
COL_CHK = "R"
COL_TR = ["S", "T", "U", "V", "W", "X", "Y"]     # trailing 12-month means: BASE x5, LOW data centres, HIGH data centres
SPEC_NAMES = ["Data centres (BASE share)", "Crypto mining", "Oil & gas / Permian electrification",
              "Other industrial (incl. LNG terminal load)", "Unattributed (population/baseline error)",
              "Data centres, LOW case", "Data centres, HIGH case"]


def share_row(ci, k):
    return SH0 + ci * 7 + k


def seed_shares():
    base = {0: [.20, .28, .35, .40], 1: [.40, .32, .28, .25], 2: [.20, .20, .20, .20], 3: [.10, .10, .10, .10], 4: [.10, .10, .07, .05]}
    dc = {"BASE": base[0], "LOW": [.12, .20, .25, .30], "HIGH": [.28, .36, .44, .50]}
    out = {}
    for c in BCASES:
        rows = [list(dc[c])]
        for k in (1, 2, 3):
            rows.append([round(base[k][j] * (1 - dc[c][j]) / (1 - base[0][j]), 3) for j in range(4)])
        rows.append([round(1 - sum(rows[k][j] for k in range(4)), 3) for j in range(4)])
        out[c] = rows
    return out


def read_breakout_prior(out):
    """Editable cells of the committed workbook ({'shares','override','stock'} or None); read BEFORE the workbook is rewritten."""
    try:
        from openpyxl import load_workbook
        ws = load_workbook(out, data_only=False)[BO]
        sh = {c: [[ws.cell(share_row(ci, k), 3 + j).value for j in range(4)] for k in range(5)] for ci, c in enumerate(BCASES)}
        ov = [[ws.cell(OV0 + k, 3 + j).value for j in range(4)] for k in range(4)]
        ok = all(isinstance(x, (int, float)) for c in sh.values() for r in c for x in r)
        stock = ws.cell(CF["stock"], 3).value
        return {"shares": sh if ok else None, "override": ov, "stock": stock if isinstance(stock, (int, float)) else None}
    except Exception:  # noqa: BLE001
        return None


def dc_factors():
    """Load factor, gas share, heat rate and MMBtu/Mcf (LOW/BASE/HIGH) read from the data-centre assumption tab of
    texas_production_forecast.xlsx (not duplicated here); the ERCOT heat-rate mean as a fallback."""
    want = {"lf": "Load factor", "gas": "Gas share of marginal supply", "hr": "Effective heat rate", "mcf": "MMBtu per Mcf"}
    default = {"lf": [.85] * 3, "gas": [.4, .5, .6], "hr": [8.5] * 3, "mcf": [1.036] * 3}
    src = "texas_production_forecast.xlsx 'Assump - Data centres' (read only; set by TEXAS_DATACENTRE.py)"
    try:
        a = pd.read_excel(PROD_XLSX, sheet_name="Assump - Data centres", header=None)
        res = {}
        for k, lab in want.items():
            r = a[a[0].astype(str).str.startswith(lab)].iloc[0]
            res[k] = [float(r[1]), float(r[2]), float(r[3])]
        return res, src
    except Exception as exc:  # noqa: BLE001
        log(f"  data-centre factors unreadable ({type(exc).__name__}: {exc}) - defaults used")
        try:
            m = pd.read_excel(ERCOT_BURN_XLSX, sheet_name="Monthly")
            ok = m[(m["ERCOT_days"] >= 28) & ~m["Heat_rate_basis"].astype(str).str.startswith("estimated")].tail(12)
            default["hr"] = [float(ok["Heat_rate_used_MMBtu_per_MWh"].mean())] * 3
        except Exception:  # noqa: BLE001
            pass
        return default, "FALLBACK defaults (texas_production_forecast.xlsx unreadable)"


def scenario_dec26():
    try:
        d = pd.read_excel(PROD_XLSX, sheet_name="Demand to 2033", index_col=0, parse_dates=True)
        return [float(d.loc[pd.Timestamp("2026-12-01"), f"Data centres GW (year-end path), {c}"]) for c in ("LOW", "BASE", "HIGH")]
    except Exception as exc:  # noqa: BLE001
        log(f"  scenario path unreadable ({type(exc).__name__}: {exc})")
        return [np.nan] * 3


def breakout_model(unexp, slope, prior, fac):
    """Python mirror of the Excel formulas: monthly apportionment, trailing 12-month means, latest-12-month table, Dec 2026 comparison."""
    shares = (prior or {}).get("shares") or seed_shares()
    ov = (prior or {}).get("override") or [[None] * 4 for _ in range(4)]
    stock = (prior or {}).get("stock")
    stock = 1.0 if stock is None else float(stock)
    idx = pd.date_range(MON_START, unexp.index.max(), freq="MS")
    u = unexp.reindex(idx)
    cat = {}
    for c in BCASES:
        for k in range(5):
            vals = []
            for m, uv in u.items():
                if pd.isna(uv):
                    vals.append(np.nan)
                elif m.year < 2023:
                    vals.append(0.0)
                else:
                    j = BYEARS.index(min(m.year, 2026))
                    o = ov[k][j] if k < 4 else None
                    vals.append(float(o) if isinstance(o, (int, float)) else uv * shares[c][k][j])
            cat[(c, k)] = pd.Series(vals, index=idx)
    tr = pd.DataFrame({SPEC_NAMES[0]: cat[("BASE", 0)].rolling(12, min_periods=1).mean()})
    for k in range(1, 5):
        tr[SPEC_NAMES[k]] = cat[("BASE", k)].rolling(12, min_periods=1).mean()
    tr[SPEC_NAMES[5]] = cat[("LOW", 0)].rolling(12, min_periods=1).mean()
    tr[SPEC_NAMES[6]] = cat[("HIGH", 0)].rolling(12, min_periods=1).mean()
    tr = tr.where(u.notna())
    last12 = u.dropna().index[-12:]
    t12 = pd.DataFrame({c: [cat[(c, k)].loc[last12].mean() for k in range(5)] for c in BCASES}, index=CATS)
    U12 = float(u.loc[last12].mean())
    per_avg = 1000 * fac["gas"][1] * fac["hr"][1] * 24 / fac["mcf"][1] / 1e6          # Bcf/d per GW of AVERAGE load, BASE factors
    per_en = fac["lf"][1] * per_avg                                                       # Bcf/d per GW energised (= TEXAS_DATACENTRE's per-GW)
    mid = pd.Series(last12).map(pd.Timestamp.toordinal).mean()
    months = (pd.Timestamp("2026-12-15").toordinal() - mid) / 30.4375
    proj = U12 + slope * months / 12
    scen = scenario_dec26()
    dc = {}
    for c in BCASES:
        j = 3
        o = ov[0][j]
        avg = float(o) if isinstance(o, (int, float)) else proj * shares[c][0][j]
        dc[c] = {"avg": avg, "inc": avg / fac["lf"][1], "imp": avg / fac["lf"][1] + stock}
    return {"shares": shares, "override": ov, "stock": stock, "u": u, "tr": tr, "t12": t12, "U12": U12, "per_avg": per_avg, "per_en": per_en,
            "months": months, "proj": proj, "scen": scen, "dc": dc, "last12": last12, "idx": idx, "slope": slope,
            "burn12": float(t12.loc[CATS[0], "BASE"] * per_avg)}


def breakout_flag(m):
    imp = m["dc"]["BASE"]["imp"]
    lo, ba, hi = m["scen"]
    if any(pd.isna(x) for x in m["scen"]):
        return "scenario path not readable - no comparison"
    if imp < lo:
        txt = f"Scenario path looks TOO HIGH: apportioned BASE {imp:.1f} GW is below even the scenario LOW case ({lo:.1f} GW)"
    elif imp > hi:
        txt = f"Scenario path looks TOO LOW: apportioned BASE {imp:.1f} GW is above even the scenario HIGH case ({hi:.1f} GW)"
    elif imp > ba * 1.15:
        txt = f"Scenario BASE path ({ba:.1f} GW) looks somewhat LOW: apportioned BASE {imp:.1f} GW, inside the LOW-HIGH range"
    elif imp < ba * 0.85:
        txt = f"Scenario BASE path ({ba:.1f} GW) looks somewhat HIGH: apportioned BASE {imp:.1f} GW, inside the LOW-HIGH range"
    else:
        txt = f"Scenario BASE path ({ba:.1f} GW) looks about right: apportioned BASE {imp:.1f} GW is within 15%"
    return txt + ". Judgement-based shares: only as good as the shares."


def breakout_values_sheet(m):
    d = pd.DataFrame(index=m["idx"])
    d.index.name = "Month"
    d["Unexplained, monthly (GW)"] = m["u"]
    for c in m["tr"].columns:
        d[c] = m["tr"][c]
    d["Row on 'Load breakout'"] = [MON0 + i for i in range(len(d))]
    return d


def write_breakout(out, m, fac, fsrc):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    yellow = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
    head = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    grey = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")
    bold = Font(bold=True)
    wb = load_workbook(out)
    if BO in wb.sheetnames:
        del wb[BO]
    ws = wb.create_sheet(BO)
    wb.move_sheet(ws, offset=wb.sheetnames.index("Assumptions") + 1 - wb.sheetnames.index(BO))

    def hdr(r, labels, c0=1):
        for j, t in enumerate(labels, start=c0):
            c = ws.cell(r, j, t)
            c.font, c.fill, c.alignment = bold, head, Alignment(wrap_text=True, vertical="top")

    ws["A1"] = "ERCOT unexplained load growth apportioned by source (GW, average load) - yellow cells are editable, formulas recalculate in Excel"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("THE SHARES BELOW ARE JUDGEMENT, NOT MEASUREMENT. Nothing in the repo's data splits the residual by source: 'unexplained' = EIA-930 ERCOT net "
                "generation minus a weather + population baseline fitted 2019-22, so it also holds the baseline's own error. Every seeded share is " + UNV +
                ". ERCOT's large-load interconnection / observed large-load reports could not be read by script (Actions probe, 5 Oct 2026, see Units tab): "
                "paste observed GW per category into the override block (B) and it replaces the share-based value.")
    ws["A3"] = ("Chart and monthly block use the trailing 12-month mean of each category (monthly residuals are weather-noisy). Shares apply by calendar year "
                "(2026 applies to every 2026 month); months before 2023 are the baseline fit window, so nothing is apportioned. Delete this tab to re-seed the shares.")
    for r in (2, 3):
        ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(f"A{r}:H{r}")
        ws.row_dimensions[r].height = 62 if r == 2 else 34
    ws["A5"] = "A. Share of the unexplained GW by category, case and year (EDITABLE; each case-year must sum to 100%)"
    ws["A5"].font = bold
    hdr(SH_HDR, ["Case", "Category"] + BYEARS + ["Status / check"])
    for ci, c in enumerate(BCASES):
        for k in range(5):
            r = share_row(ci, k)
            ws.cell(r, 1, c)
            ws.cell(r, 2, CATS[k])
            for j in range(4):
                x = ws.cell(r, 3 + j, m["shares"][c][k][j])
                x.fill, x.number_format = yellow, "0%"
            ws.cell(r, 7, UNV)
        rs = share_row(ci, 5)
        ws.cell(rs, 2, "Sum of shares").font = bold
        for j in range(4):
            col = L(3 + j)
            x = ws.cell(rs, 3 + j, f"=SUM({col}{share_row(ci, 0)}:{col}{share_row(ci, 4)})")
            x.number_format, x.font = "0%", bold
        ws.cell(rs, 7, f'=IF(AND(ABS(C{rs}-1)<0.0005,ABS(D{rs}-1)<0.0005,ABS(E{rs}-1)<0.0005,ABS(F{rs}-1)<0.0005),"OK: every year sums to 100%","CHECK: shares do not sum to 100%")').font = bold
    ws.cell(OV_HDR - 2, 1, "B. Observed override, GW of average load per category and year (OPTIONAL, EDITABLE; blank = use the share; filled = replaces the share-based value in every case)").font = bold
    ws.cell(OV_HDR - 1, 1, "Paste e.g. ERCOT large-load interconnection / observed large-load or crypto-load GW here. If used, the stack no longer sums to the unexplained load: see 'stack minus unexplained' rows.")
    hdr(OV_HDR, ["", "Category"] + BYEARS + ["Source of the figure"])
    for k in range(4):
        r = OV0 + k
        ws.cell(r, 2, CATS[k])
        for j in range(4):
            ws.cell(r, 3 + j).fill = yellow
            ws.cell(r, 3 + j).number_format = "0.00"
    ws.cell(CF["lf"] - 2, 1, f"C. Data-centre conversion factors, read-only copies from {fsrc}").font = bold
    hdr(CF["lf"] - 1, ["Factor", "", "LOW", "BASE", "HIGH", "Unit", "Note"])
    for key, lab, unit in (("lf", "Load factor (average load / energised capacity)", "ratio"), ("gas", "Gas share of marginal supply", "share"),
                           ("hr", "Effective heat rate (ERCOT, last-12-month mean of ercot_gas_burn_daily.xlsx)", "MMBtu/MWh"), ("mcf", "MMBtu per Mcf", "MMBtu/Mcf")):
        r = CF[key]
        ws.cell(r, 1, lab)
        for j in range(3):
            x = ws.cell(r, 3 + j, fac[key][j])
            x.fill, x.number_format = grey, "0.000"
        ws.cell(r, 6, unit)
    ws.cell(CF["lf"], 7, "Copied at run time; change them on 'Assump - Data centres' in texas_production_forecast.xlsx, not here.")
    ws.cell(CF["pavg"], 1, "Gas burn per GW of AVERAGE data-centre load, BASE factors (formula)").font = bold
    ws.cell(CF["pavg"], 4, f"=1000*D{CF['gas']}*D{CF['hr']}*24/D{CF['mcf']}/1000000").number_format = "0.0000"
    ws.cell(CF["pavg"], 6, "Bcf/d per GW")
    ws.cell(CF["pavg"], 7, "Formula: 1000 MW/GW x gas share x heat rate x 24 h / MMBtu per Mcf / 1e6; apportioned LOW/HIGH lines use BASE factors so only the share differs")
    ws.cell(CF["pen"], 1, "Gas burn per GW of ENERGISED capacity, BASE (formula; = 'Assump - Data centres' value)").font = bold
    ws.cell(CF["pen"], 4, f"=D{CF['lf']}*D{CF['pavg']}").number_format = "0.0000"
    ws.cell(CF["pen"], 6, "Bcf/d per GW")
    ws.cell(CF["stock"], 1, "Data-centre energised GW already in the 2019-22 baseline (EDITABLE)")
    x = ws.cell(CF["stock"], 3, m["stock"])
    x.fill, x.number_format = yellow, "0.0"
    ws.cell(CF["stock"], 6, "GW")
    ws.cell(CF["stock"], 7, "Unexplained load is growth ABOVE the baseline, the scenario path is the whole stock: this adds the pre-2023 stock back. " + UNV)
    # D. latest 12 months
    a0, a1 = m["last12"].min(), m["last12"].max()
    ra, rb = MON0 + (a0.year - 2022) * 12 + a0.month - 1, MON0 + (a1.year - 2022) * 12 + a1.month - 1
    ws.cell(T12_HDR - 2, 1, f"D. Latest 12 months ({a0:%b/%y}-{a1:%b/%y}), GW of average load by category (live: mean of the monthly block, missing months ignored)").font = bold
    hdr(T12_HDR, ["Category", "", "LOW", "BASE", "HIGH"])
    colsets = {"BASE": COL_BASE, "LOW": COL_LOW, "HIGH": COL_HIGH}
    for k in range(5):
        r = T12_0 + k
        ws.cell(r, 1, CATS[k])
        for j, c in enumerate(("LOW", "BASE", "HIGH")):
            cl = colsets[c][k]
            ws.cell(r, 3 + j, f"=AVERAGE({cl}{ra}:{cl}{rb})").number_format = "0.00"
    rt = T12_0 + 5
    ws.cell(rt, 1, "Total = ERCOT unexplained load growth, latest-12-month mean").font = bold
    for j in range(3):
        x = ws.cell(rt, 3 + j, f"=AVERAGE($B{ra}:$B{rb})")
        x.number_format, x.font = "0.00", bold
    ws.cell(rt + 1, 1, "Stack minus unexplained (zero unless an override is used)")
    for j in range(3):
        col = L(3 + j)
        ws.cell(rt + 1, 3 + j, f"=SUM({col}{T12_0}:{col}{T12_0 + 4})-{col}{rt}").number_format = "0.00"
    # E. data centres
    ws.cell(DCR["u"] - 2, 1, "E. Data centres: gas burn and the Dec 2026 energised GW against the scenario path (texas_production_forecast.xlsx, read only)").font = bold
    hdr(DCR["u"] - 1, ["Item", "", "LOW", "BASE", "HIGH", "Unit", "Note"])
    items = [("u", "Unexplained load, latest-12-month mean", "GW"), ("slope", "Unexplained trend since Jan 2023 (this run's OLS)", "GW per year"),
             ("months", "Months from the middle of the latest-12 window to Dec 2026", "months"),
             ("proj", "Unexplained load projected to Dec 2026 (mean + trend x months / 12)", "GW"),
             ("share", "Data-centre share of the unexplained load, 2026", "share"),
             ("avg", "Data-centre average load at Dec 2026 (override if filled)", "GW"),
             ("inc", "Implied energised data-centre GW from the apportioned growth (average load / load factor)", "GW"),
             ("stock", "plus energised GW already in the baseline", "GW"),
             ("imp", "IMPLIED ENERGISED DATA-CENTRE GW, DEC 2026", "GW"),
             ("scen", "Scenario path, Dec 2026 energised GW ('Demand to 2033', year-end path, same case)", "GW"),
             ("diff", "Apportioned minus scenario, same case", "GW"),
             ("burn12", "Data-centre gas burn on the growth apportioned to data centres, latest 12 months (BASE factors)", "Bcf/d"),
             ("burn_inc", "Same, at Dec 2026 (growth only)", "Bcf/d"),
             ("burn_imp", "Gas burn at the implied total energised level, Dec 2026 (stock included)", "Bcf/d")]
    for key, lab, unit in items:
        ws.cell(DCR[key], 1, lab)
        ws.cell(DCR[key], 6, unit)
    ws.cell(DCR["imp"], 1).font = bold
    for j, c in enumerate(("LOW", "BASE", "HIGH")):
        col = L(3 + j)
        ci = BCASES.index(c)
        ws[f"{col}{DCR['u']}"] = f"={col}{rt}"
        ws[f"{col}{DCR['slope']}"] = round(m["slope"], 3)
        ws[f"{col}{DCR['months']}"] = round(m["months"], 1)
        ws[f"{col}{DCR['proj']}"] = f"={col}{DCR['u']}+{col}{DCR['slope']}*{col}{DCR['months']}/12"
        ws[f"{col}{DCR['share']}"] = f"=F{share_row(ci, 0)}"
        ws[f"{col}{DCR['avg']}"] = f"=IF(ISNUMBER($F${OV0}),$F${OV0},{col}{DCR['proj']}*{col}{DCR['share']})"
        ws[f"{col}{DCR['inc']}"] = f"={col}{DCR['avg']}/$D${CF['lf']}"
        ws[f"{col}{DCR['stock']}"] = f"=$C${CF['stock']}"
        ws[f"{col}{DCR['imp']}"] = f"={col}{DCR['inc']}+{col}{DCR['stock']}"
        ws[f"{col}{DCR['scen']}"] = None if pd.isna(m["scen"][j]) else round(float(m["scen"][j]), 3)
        ws[f"{col}{DCR['diff']}"] = f"={col}{DCR['imp']}-{col}{DCR['scen']}"
        ws[f"{col}{DCR['burn12']}"] = f"=$D${T12_0}*$D${CF['pavg']}" if c == "BASE" else None
        ws[f"{col}{DCR['burn_inc']}"] = f"={col}{DCR['avg']}*$D${CF['pavg']}"
        ws[f"{col}{DCR['burn_imp']}"] = f"={col}{DCR['imp']}*$D${CF['pen']}"
        for key, fmt in (("u", "0.00"), ("slope", "0.00"), ("months", "0.0"), ("proj", "0.00"), ("share", "0%"), ("avg", "0.00"), ("inc", "0.00"),
                         ("stock", "0.00"), ("imp", "0.00"), ("scen", "0.00"), ("diff", "+0.00;-0.00"), ("burn12", "0.00"), ("burn_inc", "0.00"), ("burn_imp", "0.00")):
            ws[f"{col}{DCR[key]}"].number_format = fmt
    for col in "CDE":
        ws[f"{col}{DCR['imp']}"].font = bold
    ws.cell(DCR["scen"], 7, "values read from the other workbook when this script ran; LOW/BASE/HIGH there are Texas data-centre scenarios, here they are apportionment cases")
    ws.cell(DCR["flag"], 1, "FLAG: scenario path vs the apportioned figure (BASE)").font = bold
    s = DCR
    ws.cell(DCR["flag"], 3, (
        f'=IF(NOT(ISNUMBER(C{s["scen"]})),"scenario path not readable - no comparison",'
        f'IF(D{s["imp"]}<C{s["scen"]},"Scenario path looks TOO HIGH: apportioned BASE "&TEXT(D{s["imp"]},"0.0")&" GW is below even the scenario LOW case ("&TEXT(C{s["scen"]},"0.0")&" GW)",'
        f'IF(D{s["imp"]}>E{s["scen"]},"Scenario path looks TOO LOW: apportioned BASE "&TEXT(D{s["imp"]},"0.0")&" GW is above even the scenario HIGH case ("&TEXT(E{s["scen"]},"0.0")&" GW)",'
        f'IF(D{s["imp"]}>D{s["scen"]}*1.15,"Scenario BASE path ("&TEXT(D{s["scen"]},"0.0")&" GW) looks somewhat LOW: apportioned BASE "&TEXT(D{s["imp"]},"0.0")&" GW, inside the LOW-HIGH range",'
        f'IF(D{s["imp"]}<D{s["scen"]}*0.85,"Scenario BASE path ("&TEXT(D{s["scen"]},"0.0")&" GW) looks somewhat HIGH: apportioned BASE "&TEXT(D{s["imp"]},"0.0")&" GW, inside the LOW-HIGH range",'
        f'"Scenario BASE path ("&TEXT(D{s["scen"]},"0.0")&" GW) looks about right: apportioned BASE "&TEXT(D{s["imp"]},"0.0")&" GW is within 15%")))))'
        f'&". Judgement-based shares: only as good as the shares."')).font = bold
    # F. monthly block
    ws.cell(MON_HDR - 2, 1, "F. Monthly apportionment, GW of average load (live formulas; columns S-Y = trailing 12-month means used by the chart)").font = bold
    labels = ["Month", "ERCOT unexplained load, monthly (GW)"]
    for c in BCASES:
        labels += [f"{c}: {n}" for n in ("Data centres", "Crypto mining", "Oil & gas", "Other industrial", "Unattributed")]
    labels += ["BASE stack minus unexplained"]
    labels += ["12m mean: " + n for n in SPEC_NAMES]
    hdr(MON_HDR, labels)
    ws.row_dimensions[MON_HDR].height = 62
    for i, mth in enumerate(m["idx"]):
        r = MON0 + i
        ws.cell(r, 1, mth.to_pydatetime()).number_format = "mmm/yy"
        uv = m["u"].iloc[i]
        if pd.notna(uv):
            ws.cell(r, 2, float(uv)).number_format = "0.00"
        mi = f"MATCH(MIN(YEAR($A{r}),2026),$C${SH_HDR}:$F${SH_HDR},0)"
        for ci, c in enumerate(BCASES):
            for k in range(5):
                col = colsets[c][k]
                sh = f"INDEX($C${share_row(ci, k)}:$F${share_row(ci, k)},{mi})"
                if k < 4:
                    ov = f"INDEX($C${OV0 + k}:$F${OV0 + k},{mi})"
                    core = f"IF(ISNUMBER({ov}),{ov},$B{r}*{sh})"
                else:
                    core = f"$B{r}*{sh}"
                ws[f"{col}{r}"] = f'=IF(ISNUMBER($B{r}),IF(YEAR($A{r})>=2023,{core},0),"")'
                ws[f"{col}{r}"].number_format = "0.00"
        ws[f"{COL_CHK}{r}"] = f'=IF(ISNUMBER($B{r}),SUM(C{r}:G{r})-$B{r},"")'
        ws[f"{COL_CHK}{r}"].number_format = "0.00"
        if i >= 11:
            src = COL_BASE + [COL_LOW[0], COL_HIGH[0]]
            for col, sc in zip(COL_TR, src):
                ws[f"{col}{r}"] = f'=IF(ISNUMBER($B{r}),AVERAGE({sc}{r - 11}:{sc}{r}),"")'
                ws[f"{col}{r}"].number_format = "0.00"
    for col, w in zip("ABCDEFGH", (70, 36, 11, 11, 11, 11, 44, 12)):
        ws.column_dimensions[col].width = w
    for cc in range(9, 26):
        ws.column_dimensions[L(cc)].width = 14
    ws.freeze_panes = "A5"
    root, ext = os.path.splitext(out)
    tmp = f"{root}.tmp{os.getpid()}{ext}"
    try:
        wb.save(tmp)
        os.replace(tmp, out)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


# ------------------------------------------------------------------ main
def read_prev(out, sheet, **kw):
    try:
        return pd.read_excel(out, sheet_name=sheet, **kw)
    except Exception:  # noqa: BLE001
        return None


def assumptions_in(out, auto):
    """Yellow cells: keep the user's edit when it differs from last run's auto default, else take the new auto default."""
    prev = read_prev(out, "Assumptions")
    used = {}
    for k, v in auto.items():
        val = v
        if prev is not None and k in set(prev["Assumption"]):
            r = prev[prev["Assumption"] == k].iloc[0]
            try:
                pu, pa = float(r["Value used (edit)"]), float(r["Auto default (this run)"])
                if abs(pu - pa) > 1e-9 * max(1, abs(pa)):
                    val = pu
            except Exception:  # noqa: BLE001
                pass
        used[k] = val
    return used


def run(out):
    t0 = time.time()
    bprior = read_breakout_prior(out)   # editable Load breakout cells, read before the workbook is rewritten
    # ---- history inputs
    cons = pd.read_excel(GAS_XLSX, sheet_name="Consumption by sector")
    cons["Month"] = pd.to_datetime(cons["Month"])
    cons = cons.set_index("Month").sort_index()[SECTORS]
    cons = cons[cons.index >= START].dropna()
    last_gas = cons.index.max()
    log(f"gas sectors {cons.index.min():%b/%y}-{last_gas:%b/%y}")

    prev_t = read_prev(out, "Texas T2M daily", index_col=0)
    saved = pd.DataFrame()
    if prev_t is not None:
        prev_t.index = pd.to_datetime(prev_t.index)
        saved = prev_t.apply(pd.to_numeric, errors="coerce")
    temps, wsrc = update_temps(saved)
    if temps.empty:
        raise SystemExit("no temperature data (NASA POWER and Open-Meteo both failed and no saved store)")
    dd = degree_days(temps)
    ddc = dd[dd["Complete"]]
    log(f"degree days: complete months {ddc.index.min():%b/%y}-{ddc.index.max():%b/%y} ({wsrc})")

    prev_p = read_prev(out, "Texas population")
    prev_p = prev_p.set_index("Year") if prev_p is not None else None
    prem = 0.9
    pa = read_prev(out, "Assumptions")
    if pa is not None:
        r = pa[pa["Assumption"].astype(str).str.startswith("Texas population growth premium")]
        if len(r):
            prem = float(r["Value used (edit)"].iloc[0])
    annual, psrc = get_population(prev_p, prem)
    cagr = float((annual.iloc[-1] / annual.iloc[-6]) ** (1 / 5) - 1)

    hh = pd.read_excel(HH_XLSX, sheet_name="Chart - Monthly")
    hh["Date"] = pd.to_datetime(hh["Date"])
    hh = hh.set_index("Date")["Henry Hub spot"].dropna()
    hh = hh[hh.index.day == 1]

    e = pd.read_excel(E930_XLSX, sheet_name="ERCOT", usecols=["date", "Wind_MWh", "Solar_MWh", "Total_MWh"])
    e["date"] = pd.to_datetime(e["date"])
    e = e.drop_duplicates("date").set_index("date").sort_index()
    e = e[e["Total_MWh"] > 0]
    g = e.resample("MS")
    cnt = g["Total_MWh"].count()
    er = pd.DataFrame({"GWh": g["Total_MWh"].sum() / 1000, "re": (e["Wind_MWh"].fillna(0) + e["Solar_MWh"].fillna(0)).resample("MS").sum() / g["Total_MWh"].sum()})
    er = er[cnt == er.index.days_in_month]
    er["GW"] = er["GWh"] / (er.index.days_in_month * 24)
    log(f"ERCOT load {er.index.min():%b/%y}-{er.index.max():%b/%y}")

    # ---- assumptions
    last_full_year = int(ddc.index.year[ddc.groupby(ddc.index.year)["Days"].transform("count") == 12].max())
    auto = {
        "Texas population growth after the last Census year (% per year)": round(cagr * 100, 3),
        "Texas population growth premium over US, pp a year (used only if Census/FRED unreachable)": 0.9,
        "Henry Hub, forecast months ($/MMBtu)": round(float(hh.iloc[-12:].mean()), 2),
        "ERCOT wind+solar share of generation, forecast months": round(float(er["re"].iloc[-12:].mean()), 3),
        "Normal-weather window, first year": last_full_year - 9,
        "Normal-weather window, last year": last_full_year,
        "ERCOT baseline fit window, last month (YYYYMM)": 202212,
        "Prediction band, number of standard errors": 1.0,
    }
    A = assumptions_in(out, auto)
    gpop = A["Texas population growth after the last Census year (% per year)"] / 100
    hh_f = A["Henry Hub, forecast months ($/MMBtu)"]
    re_f = A["ERCOT wind+solar share of generation, forecast months"]
    y0, y1 = int(A["Normal-weather window, first year"]), int(A["Normal-weather window, last year"])
    base_end = pd.Timestamp(f"{int(A['ERCOT baseline fit window, last month (YYYYMM)']) // 100}-{int(A['ERCOT baseline fit window, last month (YYYYMM)']) % 100:02d}-01")
    nse = A["Prediction band, number of standard errors"]

    # ---- driver frame
    idx = pd.date_range(START, END, freq="MS")
    D = pd.DataFrame(index=idx)
    win = ddc[(ddc.index.year >= y0) & (ddc.index.year <= y1)]
    norm = win.groupby(win.index.month)[["HDD per day", "CDD per day"]].mean()
    D["hdd_act"] = ddc["HDD per day"].reindex(idx)
    D["cdd_act"] = ddc["CDD per day"].reindex(idx)
    last_obs = max(last_gas, er.index.max())
    D["Type"] = np.where(D.index <= last_gas, "History", "Forecast")
    D["hdd"] = [D.loc[m, "hdd_act"] if (m <= last_obs and pd.notna(D.loc[m, "hdd_act"])) else norm.loc[m.month, "HDD per day"] for m in idx]
    D["cdd"] = [D.loc[m, "cdd_act"] if (m <= last_obs and pd.notna(D.loc[m, "cdd_act"])) else norm.loc[m.month, "CDD per day"] for m in idx]
    D["pop"] = pop_monthly(idx, annual, gpop)
    D["trend"] = [(m - START).days / 365.25 for m in idx]
    D["hh"] = hh.reindex(idx)
    D.loc[D.index > hh.index.max(), "hh"] = hh_f
    D["re"] = er["re"].reindex(idx)
    D.loc[D.index > er.index.max(), "re"] = re_f
    # months inside the observed span lacking a driver stay NaN (not filled), except forecast months
    log("normals window", y0, y1, "pop growth", round(gpop * 100, 2), "HH", hh_f, "RE", re_f)

    # ---- model 1: gas sectors
    chosen, cand_rows, fit_rows, hold_rows = {}, [], [], []
    fc_idx = D.index[D.index > last_gas]
    for s in SECTORS:
        y = cons[s].reindex(idx)
        extras_set = [(), ("hh",), ("hh", "re")] if s == "Electric power" else [()]
        cands = []
        for w, pt, ex in product(("temp", "dummy"), ("pop+trend", "pop", "trend"), extras_set):
            c = fit_candidate(s, D, y, w, pt, ex)
            if c:
                cands.append(c)
        ok = [c for c in cands if c["ok"]]
        pick = min(ok or cands, key=lambda c: c["rmse"])
        naive = (y.shift(12)).loc[pick["test"]]
        naive_rmse = float(np.sqrt(((naive - y.loc[pick["test"]]) ** 2).mean()))
        pick["naive"] = naive_rmse
        pick["all_implausible"] = not ok
        chosen[s] = pick
        for c in sorted(cands, key=lambda c: c["rmse"]):
            cand_rows.append({"Sector": s, "Candidate": label(c), "Hold-out RMSE (Bcf/d)": c["rmse"], "Hold-out MAPE %": c["mape"],
                              "Hold-out bias (Bcf/d)": c["bias"], "In-sample R2": c["fit"]["r2"], "Adj R2": c["fit"]["adj"], "n": c["fit"]["n"],
                              "Signs plausible": "yes" if c["ok"] else "NO: " + c["why"], "Chosen": "CHOSEN" if c is pick else ""})
        f = pick["fit"]
        ybar = float(y.dropna().iloc[-12:].mean())
        Xf = design(D.loc[fc_idx], pick["weather"], pick["pt"], pick["extras"])
        mu, se = predict(f, Xf)
        pick["fc_mu"], pick["fc_se"] = mu, se
        pick["fc_flag"] = ""
        a33 = float(pd.Series(mu, index=fc_idx)[fc_idx.year == 2033].mean())
        if abs(a33 / ybar - 1) > 0.5:
            pick["fc_flag"] = f"2033 average {a33:.2f} is {a33 / ybar - 1:+.0%} vs the last 12 months: trend/population extrapolation, treat with caution"
        for n_, b_, se_ in zip(pick["X"].columns, f["b"], f["se"]):
            fit_rows.append({"Model": f"{s} (gas, Bcf/d)", "Specification": label(pick), "Term": n_, "Coefficient": b_, "Std error": se_,
                             "t-stat": b_ / se_ if se_ else np.nan, "What it means": meaning(n_), "n": f["n"], "R2": f["r2"], "Adj R2": f["adj"],
                             "Residual std": f["s"], "Hold-out RMSE (last 12m)": pick["rmse"], "Hold-out MAPE %": pick["mape"],
                             "Naive same-month-last-year RMSE": naive_rmse, "Flags": ("ALL candidates implausible: " if not ok else "") + pick["fc_flag"]})
        for m, a, p in zip(pick["test"], y.loc[pick["test"]].values, pick["pred_test"]):
            hold_rows.append({"Month": m, "Model": s, "Actual": a, "Predicted (fit before the hold-out)": p, "Error": p - a})
        log(f"{s}: {label(pick)} R2={f['r2']:.3f} holdout RMSE={pick['rmse']:.3f} ({pick['mape']:.1f}%) naive={naive_rmse:.3f} {pick['fc_flag']}")

    # fitted / forecast frames
    F = pd.DataFrame(index=idx)
    F.index.name = "Month"
    F["Type"] = D["Type"]
    F["HDD per day (actual; normal in forecast months)"] = D["hdd"]
    F["CDD per day (actual; normal in forecast months)"] = D["cdd"]
    F["Texas population (million)"] = D["pop"]
    F["Henry Hub ($/MMBtu; assumption in forecast months)"] = D["hh"]
    F["ERCOT wind+solar share"] = D["re"]
    var_sum = pd.Series(0.0, index=idx)
    tot_act, tot_fit = cons.sum(axis=1).reindex(idx), pd.Series(np.nan, index=idx)
    parts = []
    for s in SECTORS:
        c = chosen[s]
        Xall = design(D, c["weather"], c["pt"], c["extras"])
        okr = Xall.notna().all(axis=1)
        mu = pd.Series(np.nan, index=idx)
        se = pd.Series(np.nan, index=idx)
        m_, s_ = predict(c["fit"], Xall[okr])
        mu[okr], se[okr] = m_, s_
        F[f"{s} actual"] = cons[s].reindex(idx)
        F[f"{s} fitted (actual weather) / forecast (normal weather)"] = mu
        F[f"{s} lower"] = mu - nse * se
        F[f"{s} upper"] = mu + nse * se
        parts.append((mu, se))
    mu_sum = sum(p[0] for p in parts)
    se_sum = np.sqrt(sum(p[1] ** 2 for p in parts))
    F["Other sectors actual (sum of four)"] = tot_act
    F["Other sectors fitted / forecast (sum of four)"] = mu_sum
    F["Other sectors lower (band in quadrature, sectors independent)"] = mu_sum - nse * se_sum
    F["Other sectors upper (band in quadrature, sectors independent)"] = mu_sum + nse * se_sum
    fv = pd.read_excel(GAS_XLSX, sheet_name="Forecast values")
    fv["Month"] = pd.to_datetime(fv["Month"])
    fv = fv.set_index("Month")
    ex = fv[["Electric power", "Industrial", "Residential", "Commercial"]].sum(axis=1)
    F["Existing damped-trend forecast, four sectors (texas_gas_monthly, to Dec 2028)"] = ex.where(fv["Actual / forecast"].eq("Forecast")).reindex(idx)

    # ---- model 2: ERCOT load
    Er = D.join(er[["GW"]]).copy()
    Er = Er[(Er.index <= er.index.max()) | (Er.index > er.index.max())]
    y = Er["GW"]
    Xa = design(Er, "temp", "pop", ())
    rows_a = Er.index[(Er.index <= base_end) & y.notna() & Xa.notna().all(axis=1)]
    fa = ols(Xa.loc[rows_a], y.loc[rows_a])
    obs = Er.index[y.notna() & Xa.notna().all(axis=1)]
    mu_a, se_a = predict(fa, Xa[Xa.notna().all(axis=1)])
    base = pd.Series(np.nan, index=idx)
    bse = pd.Series(np.nan, index=idx)
    ok_a = Xa.notna().all(axis=1)
    base[ok_a], bse[ok_a] = mu_a, se_a
    unexp = (y - base).where(y.notna())
    unexp_roll = unexp.rolling(12, min_periods=12).mean()
    post = unexp[(unexp.index >= "2023-01-01") & unexp.notna()]
    sl = ols(np.column_stack([np.ones(len(post)), (post.index - pd.Timestamp("2023-01-01")).days / 365.25]), post.values)
    last12 = unexp.dropna().iloc[-12:]
    stepup = float(last12.mean())
    # reference: full-sample model with trend and hold-out
    Xb = design(Er, "temp", "pop+trend", ())
    last_er = y.dropna().index.max()
    rows_b = Er.index[y.notna() & Xb.notna().all(axis=1)]
    trb = rows_b[rows_b <= last_er - pd.DateOffset(months=HOLD)]
    teb = rows_b[rows_b > last_er - pd.DateOffset(months=HOLD)]
    fb_tr = ols(Xb.loc[trb], y.loc[trb])
    pb, _ = predict(fb_tr, Xb.loc[teb])
    rmse_b = float(np.sqrt(((pb - y.loc[teb].values) ** 2).mean()))
    mape_b = float(np.mean(np.abs(pb - y.loc[teb].values) / y.loc[teb].values) * 100)
    fb = ols(Xb.loc[rows_b], y.loc[rows_b])
    pre = rows_a
    err_in = (y.loc[pre] - base.loc[pre])
    corr_pt = float(np.corrcoef(Er["pop"], Er["trend"])[0, 1])
    F["ERCOT load actual (GW, EIA-930 net generation)"] = y
    F["ERCOT baseline (weather + population, fit on window; actual weather, normal weather in forecast months)"] = base
    F["ERCOT baseline lower"] = base - nse * bse
    F["ERCOT baseline upper"] = base + nse * bse
    F["ERCOT unexplained load growth (GW) = actual - baseline"] = unexp
    F["ERCOT unexplained, 12-month average (GW)"] = unexp_roll
    F["ERCOT baseline + unexplained held at last-12-month level (reference)"] = (base + stepup).where(base.index > last_er)
    for n_, b_, se_ in zip(Xa.columns, fa["b"], fa["se"]):
        fit_rows.append({"Model": "ERCOT load baseline (GW)", "Specification": f"HDD/CDD + population, fit {rows_a.min():%b/%y}-{rows_a.max():%b/%y}", "Term": n_,
                         "Coefficient": b_, "Std error": se_, "t-stat": b_ / se_ if se_ else np.nan,
                         "What it means": meaning(n_).replace("Bcf/d", "GW average load"), "n": fa["n"], "R2": fa["r2"], "Adj R2": fa["adj"],
                         "Residual std": fa["s"], "Hold-out RMSE (last 12m)": np.nan, "Hold-out MAPE %": np.nan, "Naive same-month-last-year RMSE": np.nan,
                         "Flags": "baseline is extrapolated beyond the window; the gap to actual is the unexplained growth" + ("; population coefficient negative" if fa["b"][3] < 0 else "")})
    for n_, b_, se_ in zip(Xb.columns, fb["b"], fb["se"]):
        fit_rows.append({"Model": "ERCOT load, full sample with trend (reference)", "Specification": f"HDD/CDD + population + trend, {rows_b.min():%b/%y}-{rows_b.max():%b/%y}", "Term": n_,
                         "Coefficient": b_, "Std error": se_, "t-stat": b_ / se_ if se_ else np.nan,
                         "What it means": meaning(n_).replace("Bcf/d", "GW average load"), "n": fb["n"], "R2": fb["r2"], "Adj R2": fb["adj"],
                         "Residual std": fb["s"], "Hold-out RMSE (last 12m)": rmse_b, "Hold-out MAPE %": mape_b, "Naive same-month-last-year RMSE": np.nan,
                         "Flags": f"population and trend correlation {corr_pt:.3f}: their coefficients are not separately identified; the trend term absorbs the step-up"})
    for m, a, p in zip(teb, y.loc[teb].values, pb):
        hold_rows.append({"Month": m, "Model": "ERCOT load (full-sample model with trend)", "Actual": a, "Predicted (fit before the hold-out)": p, "Error": p - a})
    unexpl = {"slope": float(sl["b"][1]), "slope_se": float(sl["se"][1]), "last12": stepup, "since": post.index.min(), "n": len(post),
              "peak": float(unexp_roll.dropna().iloc[-1]), "first": float(post.iloc[:12].mean())}
    log(f"ERCOT baseline R2={fa['r2']:.3f} resid std={fa['s']:.2f} GW; unexplained last12={stepup:.2f} GW, trend since 2023 {unexpl['slope']:.2f}+/-{unexpl['slope_se']:.2f} GW/yr; full-sample holdout RMSE {rmse_b:.2f} GW ({mape_b:.1f}%)")

    # ---- load breakout (apportioned unexplained load; the shares are judgement)
    fac, fsrc = dc_factors()
    bm = breakout_model(unexp.dropna(), unexpl["slope"], bprior, fac)
    bflag = breakout_flag(bm)
    log("load breakout, latest 12 months (GW, average load):\n" + bm["t12"].round(2).to_string())
    log(f"  unexplained {bm['U12']:.2f} GW; DC gas burn {bm['burn12']:.3f} Bcf/d ({bm['per_avg']:.4f} per average GW); Dec 2026 implied energised DC GW "
        f"LOW/BASE/HIGH {[round(bm['dc'][c]['imp'], 2) for c in BCASES[1:] + BCASES[:1]]} vs scenario {[round(x, 2) for x in bm['scen']]}")
    log("  FLAG: " + bflag)

    # ---- summary numbers
    def ann(col, yr):
        s_ = F[col][F.index.year == yr]
        return float(s_.mean())
    summ = []
    oc, ec = "Other sectors fitted / forecast (sum of four)", "Existing damped-trend forecast, four sectors (texas_gas_monthly, to Dec 2028)"
    for yr in (2026, 2028, 2033):
        dec = pd.Timestamp(f"{yr}-12-01")
        summ.append({"Period": f"Dec {yr}", "Regression other sectors (Bcf/d)": F.loc[dec, oc],
                     "Regression lower": F.loc[dec, "Other sectors lower (band in quadrature, sectors independent)"],
                     "Regression upper": F.loc[dec, "Other sectors upper (band in quadrature, sectors independent)"],
                     "Existing forecast, four sectors (Bcf/d)": F.loc[dec, ec],
                     "ERCOT baseline (GW)": F.loc[dec, "ERCOT baseline (weather + population, fit on window; actual weather, normal weather in forecast months)"]})
        summ.append({"Period": f"Year {yr} average", "Regression other sectors (Bcf/d)": ann(oc, yr),
                     "Existing forecast, four sectors (Bcf/d)": ann(ec, yr) if yr <= 2028 else np.nan,
                     "ERCOT baseline (GW)": ann("ERCOT baseline (weather + population, fit on window; actual weather, normal weather in forecast months)", yr)})
    S = pd.DataFrame(summ).set_index("Period")
    log(S.round(2).to_string())

    # ---- workbook
    fit_df = pd.DataFrame(fit_rows).set_index("Model")
    cand_df = pd.DataFrame(cand_rows).set_index("Sector")
    hold_df = pd.DataFrame(hold_rows).set_index("Month")
    notes_a = {
        "Texas population growth after the last Census year (% per year)": ("% per year", f"Applies after the last annual point ({int(annual.index.max())}). Default = 5-year CAGR of the Census series ({cagr * 100:.2f}%). The IMF outlook for the whole US is far lower; Texas has grown faster than the US."),
        "Texas population growth premium over US, pp a year (used only if Census/FRED unreachable)": ("pp", "Fallback only: Texas grew about 1.4% a year 2015-24 against about 0.5% for the US (from memory, unverified), so the default is 0.9 pp on top of IMF's US growth."),
        "Henry Hub, forecast months ($/MMBtu)": ("$/MMBtu", "Held flat in forecast months; default = mean of the last 12 months of the Henry Hub monthly series. Only the electric-power model uses it, if the chosen candidate includes it."),
        "ERCOT wind+solar share of generation, forecast months": ("share", "Held flat; default = last-12-month mean. Rising wind/solar/battery output would lower gas burn: a flat share is a conservative assumption. Used only if the chosen electric-power candidate includes it."),
        "Normal-weather window, first year": ("year", "Normal weather = average HDD/CDD per day by calendar month over these complete years (default: the last 10)."),
        "Normal-weather window, last year": ("year", "Last complete year in the temperature store."),
        "ERCOT baseline fit window, last month (YYYYMM)": ("YYYYMM", "The ERCOT baseline (weather + population) is fit on data up to this month and extrapolated; default Dec 2022, before the recent step-up in load."),
        "Prediction band, number of standard errors": ("x", "Band = fitted/forecast +/- this x the prediction standard error (residual noise + coefficient uncertainty)."),
    }
    assum = pd.DataFrame([{"Assumption": k, "Value used (edit)": A[k], "Auto default (this run)": auto[k], "Unit": notes_a[k][0], "Note": notes_a[k][1]} for k in auto]).set_index("Assumption")
    popdf = pd.DataFrame({"Population_m": annual, "Source": psrc}); popdf.index.name = "Year"
    dd_out = dd.copy()
    sheets = {"Fit summary": fit_df, "Candidates": cand_df, "Assumptions": assum, "Summary": S, "Forecast": F, "Hold-out": hold_df,
              "Load breakout values": breakout_values_sheet(bm), "Texas degree days": dd_out, "Texas population": popdf, "Texas T2M daily": temps}
    sel = "; ".join(f"{s}: {label(chosen[s])} (R2 {chosen[s]['fit']['r2']:.2f}, hold-out RMSE {chosen[s]['rmse']:.2f} Bcf/d = {chosen[s]['mape']:.1f}%)" for s in SECTORS)
    notes = [
        "Notes", "", "UNITS",
        "Gas sectors: Bcf/d (billion cubic feet per day), monthly average (EIA monthly volumes / days, texas_gas_monthly.xlsx 'Consumption by sector'). ERCOT load: GW, monthly average of EIA-930 daily ERCOT net generation (the repo has no ERCOT demand series; net generation is the proxy). Degree days: HDD/CDD per day, base 18C, population-weighted over Houston, Dallas-Fort Worth, San Antonio, Austin, El Paso (weights are approximate 2020 metro populations). Population: million.",
        "", "METHOD",
        "Ordinary least squares, monthly, from Jan 2015 (gas) / Jan 2019 (ERCOT). Standalone: no existing Texas workbook or script is changed; texas_gas_monthly.xlsx is read for history and, for the comparison chart only, its 'Forecast values' damped-trend forecast.",
        "Gas sectors: for each sector every combination of (HDD+CDD per day | month dummies) x (population+trend | population | trend), and for electric power also +Henry Hub and +Henry Hub + ERCOT wind+solar share, was fitted. The CHOSEN specification has the lowest RMSE on a hold-out of the last 12 months (fit on the months before) among candidates with plausible signs (HDD>0 residential/commercial, CDD>0 electric power, population>=0). One 12-month hold-out is a noisy basis for choosing, and the choice is made on the same data as the fit - read hold-out errors as optimistic. 'Naive' on the Fit summary is the error of simply repeating the same month a year earlier.",
        "Chosen: " + sel,
        "Population and a linear trend are nearly collinear over 2015-2026 (correlation 0.99+), so with both in the model neither coefficient is separately meaningful and the 2033 extrapolation adds both; candidates with population only or trend only are on the Candidates tab, and any chosen forecast whose 2033 average departs more than 50% from the last 12 months is flagged on the Fit summary.",
        f"Population: {psrc}; monthly values are linear interpolation between July-1 annual points, then the editable growth rate after the last point. Weather: {wsrc}; incremental store on 'Texas T2M daily' (this workbook only; macro_drivers.xlsx is untouched); a month enters only when every day has data for all five cities.",
        "ERCOT: the baseline is HDD + CDD + population fitted on the window ending at the Assumptions cell (default Dec 2022) and extrapolated at actual weather; actual minus baseline is the 'unexplained load growth', which includes data centres, crypto mining, electrified industry/oil and gas load, and anything else the baseline omits, plus the population coefficient's own error (4 years of data only). The reference full-sample model with a trend is on the Fit summary with its hold-out.",
        f"ERCOT unexplained load growth: last 12 months {unexpl['last12']:.2f} GW above the baseline; fitted trend since Jan 2023 {unexpl['slope']:+.2f} +/- {unexpl['slope_se']:.2f} GW per year (OLS on {unexpl['n']} months from {unexpl['since']:%b/%y}).",
        "Forecast to Dec 2033: normal weather (editable window), population at the editable growth, Henry Hub and wind+solar share at the editable values; the band is the prediction standard error of each model (the 'Other sectors' band adds the four in quadrature assuming independent errors, which understates it if errors are correlated, as common weather misses are). The forecast is a weather-normalised projection of past relationships; it does NOT add new LNG, data-centre or industrial demand, structural change in power supply or price response beyond what is in the fit. The ERCOT baseline excludes the step-up; the 'baseline + unexplained held flat' column is a reference only.",
        f"LOAD BREAKOUT ('Load breakout' tab, 'Load breakout values' tab, chart 'Load breakout'): the unexplained load growth split into data centres, crypto mining, oil & gas / Permian electrification, other industrial (incl. LNG terminal load) and an unattributed remainder (population/baseline error) by yellow SHARES per case (LOW/BASE/HIGH) and year 2023-26, live Excel formulas, optional observed-GW override. THE SHARES ARE JUDGEMENT, NOT MEASUREMENT (seeded 'unverified, from memory'); the data cannot tell the sources apart. Latest 12 months, BASE (GW of average load): " + "; ".join(f"{k} {v:.2f}" for k, v in bm["t12"]["BASE"].items()) + f"; total {bm['U12']:.2f}. Data-centre gas burn {bm['burn12']:.3f} Bcf/d (average GW x gas share x ERCOT heat rate x 24 / MMBtu per Mcf, factors read from texas_production_forecast.xlsx 'Assump - Data centres', source: {fsrc}). Dec 2026: " + bflag,
        "ERCOT large-load interconnection / observed large-load / crypto figures: one Actions probe (discovery_archive/ERCOT_LARGE_LOAD_PROBE.py, 5 Oct 2026), see the repo CLAUDE.md note for the outcome; paste any observed GW into the override block.",
        "Edit the yellow cells on the Assumptions tab and re-run the workflow (texas_demand_regression.yml, 1st/15th) to recompute; a cell left at its auto default follows the new auto value on the next run, an edited cell is kept.",
        f"Latest history: gas {last_gas:%b/%y}, ERCOT {er.index.max():%b/%y}, weather store to {temps.dropna(how='all').index.max():%d %b %Y}. Run time {time.time() - t0:.0f}s.",
        "", "SOURCES",
        "EIA (gas consumption, via texas_gas_monthly.xlsx); EIA-930 (ERCOT, via eia930_fuel_mix_daily.xlsx); Henry Hub (EIA, via henry_hub_daily.xlsx); NASA POWER / Open-Meteo (temperature); US Census Bureau via FRED TXPOP (population).",
    ]
    os.makedirs(os.path.dirname(out), exist_ok=True)
    xlsx_notes.write_workbook(out, sheets, notes, ["UNITS", "METHOD", "SOURCES"])
    style(out)
    write_breakout(out, bm, fac, fsrc)
    log("wrote", out)
    return {"chosen": chosen, "S": S, "unexpl": unexpl, "fa": fa, "rmse_b": rmse_b, "mape_b": mape_b}


def style(out):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(out)
    yellow = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
    ws = wb["Assumptions"]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).fill = yellow
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["E"].width = 120
    for name, widths in {"Fit summary": {"A": 34, "B": 52, "C": 28, "G": 100, "O": 70}, "Candidates": {"A": 16, "B": 60, "J": 40}}.items():
        w = wb[name]
        for k, v in widths.items():
            w.column_dimensions[k].width = v
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "B2"
    wb.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=DEFAULT_OUT)
    a = ap.parse_args()
    run(a.out)
