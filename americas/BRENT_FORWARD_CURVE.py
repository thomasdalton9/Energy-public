"""
Brent FORWARD CURVE (settlement/close by delivery month) as an incrementally
archived history: output/Data and Chart Outputs/brent_forward_curve.xlsx.

Source (the only free, programmatic source found reachable from GitHub
Actions - see discovery_archive/brent/BRENT_CURVE_SOURCES_PROBE*.py):
Yahoo Finance chart API, one ticker per NYMEX Brent Last Day Financial (BZ)
monthly contract, e.g. BZK27.NYM, daily closes via
https://query1.finance.yahoo.com/v8/finance/chart/<ticker> (curl_cffi with
browser impersonation). BZ settles on the ICE Brent futures final settlement
(its price tracks ICE Brent B: BZ=F equalled the front Brent contract in the
probe). CAVEATS: Yahoo is an unofficial, undocumented API (its terms of use
restrict automated/commercial use; may change or block without notice); it
lists only contracts that have NOT yet expired (expired ones return 404) and
only part of the far end (2028: Jan, Feb, Jun, Dec; 2029: Jun, Dec), so a
past date's curve holds only contracts alive today - the front months of
past dates are missing. This archive fills that in going forward: each run
stores every contract Yahoo still lists, including the front ones, so
complete curves exist from the first run on (runs on the 1st, 15th and 28th
so a contract is last read 2-15 days before it expires). The CME
Settlements JSON endpoint (404 'No endpoint', 403 without TLS impersonation),
ICE report centre (no history endpoint found), Stooq (JS challenge), Nasdaq
Data Link CHRIS (Incapsula 403) and Barchart (WAF) were not usable.

Columns of 'Curve' are contract months YYYY-MM (the BZ contract month: the
Dec 2026 contract is the one that is the front month on 2026-10-06). The
last bar of the current day is dropped (intraday, not a settlement).
Front month on a date = first contract whose ICE last trading day (last
weekday of the second month before delivery; exchange holidays ignored) is
on or after the date. 'By M' numbers contracts from that front month
(M1, M2, ...); a gap means the contract is not in the archive.
"""
import argparse
import os
import sys
import time
from datetime import date, datetime, timedelta, timezone

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import xlsx_notes  # noqa: E402

OUT = "output/Data and Chart Outputs/brent_forward_curve.xlsx"
CODES = "FGHJKMNQUVXZ"
HEAD = {"Accept": "application/json, text/plain, */*"}
SEED_START = date(2018, 1, 1)
LAST_YEAR = 2031


def contract_months(today):
    out = []
    y, m = today.year, today.month - 2
    while m < 1:
        y, m = y - 1, m + 12
    while y <= LAST_YEAR:
        out.append((y, m))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def last_trading_day(y, m):
    """ICE Brent: last weekday of the second month before delivery (holidays ignored)."""
    m2 = m - 2
    y2 = y
    if m2 < 1:
        m2, y2 = m2 + 12, y - 1
    d = date(y2 + (m2 == 12), m2 % 12 + 1, 1) - timedelta(days=1)
    while d.weekday() > 4:
        d -= timedelta(days=1)
    return d


def front_month(d):
    y, m = d.year, d.month
    for _ in range(40):
        if last_trading_day(y, m) >= d:
            return y, m
        y, m = (y, m + 1) if m < 12 else (y + 1, 1)


def fetch(session, ticker, start, end):
    """Daily closes {trade_date: close} for [start, end) in <=1-year windows, or None if the ticker is unknown."""
    rows, d0, known = {}, start, False
    while d0 < end:
        d1 = min(date(d0.year + 1, d0.month, min(d0.day, 28)), end)
        p1 = int(datetime(d0.year, d0.month, d0.day, tzinfo=timezone.utc).timestamp())
        p2 = int(datetime(d1.year, d1.month, d1.day, tzinfo=timezone.utc).timestamp())
        r = None
        for host in ("query1", "query2"):
            for attempt in range(3):
                try:
                    r = session.get(f"https://{host}.finance.yahoo.com/v8/finance/chart/{ticker}?period1={p1}&period2={p2}&interval=1d",
                                    impersonate="chrome", timeout=40, headers=HEAD)
                except Exception as e:  # noqa: BLE001
                    print(f"  {ticker}: {type(e).__name__}", flush=True)
                    r = None
                    time.sleep(2)
                    continue
                if r.status_code in (200, 400, 404):
                    break
                time.sleep(3 * (attempt + 1))
            if r is not None and r.status_code in (200, 400, 404):
                break
        if r is None or r.status_code not in (200, 400, 404):
            raise RuntimeError(f"{ticker}: HTTP {getattr(r, 'status_code', None)}")
        if r.status_code == 400:      # window before the contract's first trade date: no data
            pass
        elif r.status_code == 404:
            if not known:
                return None
        else:
            res = (r.json().get("chart") or {}).get("result")
            if res:
                known = True
                off = res[0]["meta"].get("gmtoffset", 0)
                ts = res[0].get("timestamp") or []
                cl = res[0]["indicators"]["quote"][0].get("close") or []
                for t, c in zip(ts, cl):
                    if c is not None:
                        rows[(datetime.fromtimestamp(t + off, tz=timezone.utc)).date()] = float(c)
        d0 = d1
        time.sleep(0.25)
    return rows


def load_archive(path):
    try:
        df = pd.read_excel(path, sheet_name="Curve", index_col=0)
    except (FileNotFoundError, ValueError):
        return pd.DataFrame()
    df.index = pd.to_datetime(df.index)
    df.columns = [str(c) for c in df.columns]
    return df.sort_index()


NOTES = [
    "UNITS",
    "Prices: US dollars per barrel (USD/bbl), daily close of NYMEX Brent Last Day Financial (BZ) futures by contract month, "
    "via Yahoo Finance (see Source). Not an official settlement file.",
    "",
    "WHAT THIS COVERS",
    "{coverage}",
    "",
    "SHEETS",
    "Curve: rows = trade date, columns = contract month (YYYY-MM). This is the history store: each run adds the new trade dates.",
    "By M: the same prices numbered from the front month on that date (M1 = front contract). Blank = contract not in the archive.",
    "Snapshots: the curve on selected trade dates (latest, 1/3/6/12 months earlier, first trading day of 2026); rows = contract month. "
    "Past columns hold only contracts still listed today.",
    "Spreads: fixed-contract calendar spreads (Dec contract minus the next Dec contract; positive = backwardation) and the M1-M12 spread where both exist.",
    "",
    "SOURCE AND CAVEATS",
    "Yahoo Finance chart API (query1.finance.yahoo.com/v8/finance/chart/BZ<month code><yy>.NYM), unofficial and undocumented; its terms restrict "
    "automated and commercial use and it can change or block without notice. Read from GitHub Actions with TLS impersonation.",
    "Yahoo lists only unexpired contracts, so history of a date shows the contracts alive today (front months of past dates are missing). "
    "Far end listed: all months to Dec 2027, then Jan, Feb, Jun, Dec 2028 and Jun, Dec 2029. The archive keeps everything from the first run on.",
    "The current day's bar is dropped (intraday). Runs: 1st, 15th, 28th; a contract's last days before expiry (up to ~2 weeks) may be missing.",
    "Not reachable: CME Settlements JSON (endpoint 404), ICE report centre history, Stooq, Nasdaq Data Link CHRIS, Barchart. "
    "Details: discovery_archive/brent/BRENT_CURVE_SOURCES_PROBE*.py.",
    "Front-month rule: ICE last trading day = last weekday of the second month before delivery; exchange holidays ignored.",
    "Run log: {runlog}",
]
SECTIONS = ["UNITS", "WHAT THIS COVERS", "SHEETS", "SOURCE AND CAVEATS"]


def build_sheets(curve):
    curve = curve.dropna(how="all").sort_index()
    curve = curve[sorted(curve.columns)]
    curve.index.name = "trade_date"
    # By M
    rows = {}
    for d, r in curve.iterrows():
        fy, fm = front_month(d.date())
        rec = {}
        for n in range(1, 25):
            y, m = divmod(fy * 12 + fm - 1 + n - 1, 12)
            v = r.get(f"{y}-{m + 1:02d}")
            rec[f"M{n}"] = None if v is None or pd.isna(v) else v
        rows[d] = rec
    byM = pd.DataFrame.from_dict(rows, orient="index").sort_index().dropna(how="all", axis=1)
    byM.index.name = "trade_date"
    # Spreads
    sp = pd.DataFrame(index=curve.index)
    decs = [c for c in curve.columns if c.endswith("-12")]
    for a, b in zip(decs, decs[1:]):
        if int(b[:4]) - int(a[:4]) == 1:
            sp[f"Dec {a[:4]} less Dec {b[:4]}"] = curve[a] - curve[b]
    if "M1" in byM and "M12" in byM:
        sp["M1 less M12"] = byM["M1"] - byM["M12"]
    sp = sp.dropna(how="all", axis=1)
    sp.index.name = "trade_date"
    # Snapshots
    last = curve.index.max()
    targets = {"latest": last, "1 month earlier": last - pd.DateOffset(months=1), "3 months earlier": last - pd.DateOffset(months=3),
               "6 months earlier": last - pd.DateOffset(months=6), "12 months earlier": last - pd.DateOffset(months=12),
               "first trading day of 2026": pd.Timestamp("2026-01-01")}
    snap = {}
    for label, t in targets.items():
        if label.startswith("first"):
            c = curve.index[curve.index >= t]
            d = c.min() if len(c) else None
        else:
            c = curve.index[curve.index <= t]
            d = c.max() if len(c) else None
        if d is not None:
            snap[f"{label} ({d.strftime('%d-%b-%Y')})"] = curve.loc[d]
    snaps = pd.DataFrame(snap).dropna(how="all")
    snaps.index = pd.to_datetime([f"{i}-01" for i in snaps.index])
    snaps.index.name = "contract_month"
    return curve, byM, sp, snaps


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=OUT)
    ap.add_argument("--start", help="backfill: earliest trade date to request (YYYY-MM-DD)")
    ap.add_argument("--end", help="latest trade date to request, exclusive (default today)")
    a = ap.parse_args()
    from curl_cffi import requests as cr
    session = cr.Session()
    today = datetime.now(timezone.utc).date()
    end = date.fromisoformat(a.end) if a.end else today          # today's bar is intraday: excluded
    arch = load_archive(a.out)
    if a.start:
        start = date.fromisoformat(a.start)
    elif len(arch):
        start = arch.index.max().date() - timedelta(days=14)
    else:
        start = SEED_START
    print(f"Archive: {len(arch)} dates; requesting {start} to {end}", flush=True)
    got, missing = {}, []
    for y, m in contract_months(today):
        name = f"{y}-{m:02d}"
        rows = fetch(session, f"BZ{CODES[m - 1]}{str(y)[2:]}.NYM", start, end)
        if rows is None:
            missing.append(name)
            continue
        got[name] = pd.Series(rows, dtype=float)
        s = got[name]
        print(f"  {name}: {len(s)} rows" + (f" {s.index.min()}..{s.index.max()}" if len(s) else ""), flush=True)
    if not got:
        sys.exit("no contracts returned - source changed or blocked")
    new = pd.DataFrame(got)
    new.index = pd.to_datetime(new.index)
    comb = new.combine_first(arch) if len(arch) else new
    # fresh values win over archived ones
    if len(arch):
        comb.update(new)
    curve, byM, sp, snaps = build_sheets(comb)
    cnt = curve.notna().sum(axis=1)
    coverage = (f"Brent (NYMEX BZ, tracks ICE Brent) forward curve, {curve.index.min():%d-%b-%Y} to {curve.index.max():%d-%b-%Y}, "
                f"{len(curve)} trade dates, {curve.shape[1]} contract months ({curve.columns[0]} to {curve.columns[-1]}); "
                f"contracts per date: latest {int(cnt.iloc[-1])}, maximum {int(cnt.max())}. "
                "Before the first run the front months of past dates are missing (see Source).")
    notes = [l.format(coverage=coverage, runlog=f"{today}: requested {start}..{end}, contracts found {len(got)}, "
                      f"not listed {len(missing)}") for l in NOTES]
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    xlsx_notes.write_workbook(a.out, {"Curve": curve, "By M": byM, "Spreads": sp, "Snapshots": snaps}, notes, SECTIONS)
    import openpyxl
    wb = openpyxl.load_workbook(a.out)
    for n in ("Curve", "By M", "Spreads"):
        ws = wb[n]
        for (c,) in ws.iter_rows(min_row=2, max_col=1):
            c.number_format = "dd-mmm-yyyy"
        ws.column_dimensions["A"].width = 14
    for (c,) in wb["Snapshots"].iter_rows(min_row=2, max_col=1):
        c.number_format = "mmm/yy"
    wb.save(a.out)
    print(f"Saved {a.out}: {len(curve)} dates {curve.index.min().date()}..{curve.index.max().date()}, {curve.shape[1]} contracts")
    print("contracts per date (last 5):", cnt.tail().to_dict())
    print(snaps.iloc[:14, :3])
    if (today - curve.index.max().date()).days > 7:
        sys.exit("STALE: latest trade date is more than 7 days old")


if __name__ == "__main__":
    main()
