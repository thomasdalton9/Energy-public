"""
APPROXIMATE Haynesville supply cost curve (breakeven $/MMBtu Henry Hub vs cumulative Bcf/d) built from company disclosures and the Bureau of
Economic Geology tiering -> output/Data and Chart Outputs/haynesville_cost_curve.xlsx.   STANDALONE: reads, never writes, gulf_coast_gas_balance.xlsx
('Key question') and henry_hub_daily.xlsx; does not touch GULF_COAST_BALANCE.py or any Texas module.

Question it answers: which breakeven price does the extra Haynesville supply in the Gulf Coast 'Key question' (4.7 Bcf/d by Dec 2030 ...) imply, if
the only constraint were COST (not rigs, takeaway or crews)?

How it is built (every number carries a status in the workbook: SOURCED / SOURCED-CHART-READ / DERIVED / ASSUMPTION):
  TIERS     BEG (Bureau of Economic Geology, UT Austin; Oil & Gas Journal 7 Dec 2015) split the legacy Haynesville into six productivity tiers: EUR per
            4,800-ft well (published), remaining well count per tier and breakeven (both read off the article's charts). Western Haynesville, outside BEG's
            2008-12 well data set, is added from Comstock's disclosed operated inventory (30 Jun 2026) and Expand's ">200 locations" (3Q25 deck).
  ECONOMICS modern costs and opex are company disclosures (Comstock Oct 2026 deck: $/lateral ft drilling + completion; Q2 2026 release: LOE, gathering,
            taxes), productivity is anchored on Expand's 3Q25 deck (12-month cumulative Mcfe per ft, 2026E), tiers 2-6 scale by BEG's EUR ratios, the decline
            shape is a stretched exponential anchored on BEG's text (>80% of EUR in 5 years, >90% in 10). Breakeven = Henry Hub price at which a well's
            PRE-TAX NPV at the discount rate is zero (live formula). The same formula, fed BEG's own 2015 inputs, reproduces BEG's published 10%-IRR breakevens
            within about 10% ('BEG check' tab) - the validation.
  CAPACITY  tier capacity in Bcf/d = the production in year H of a programme that drills the tier's developable inventory evenly over H years
            (live formulas, 'Capacity over time'); H = years from Dec 2025 to the Key-question date, i.e. as fast as cost alone would allow.
  MAPPING   'Mapping' reads the Key-question increments (STEO Haynesville growth + extra supply, base/delayed, with each Louisiana layer) and returns the
            breakeven of the marginal tier from the sorted curve, beside the Henry Hub history.

Editable (yellow, survive reruns): every assumption and every tier on 'Inputs', incl. breakeven / capacity OVERRIDE columns for pasting an own (paid) curve.
Not modelled: replacement of the base decline (maintenance drilling), taxes, hedges, midstream, rig count, takeaway - see the Units tab.

Usage: python3 HAYNESVILLE_COST_CURVE.py [--out "output/Data and Chart Outputs/haynesville_cost_curve.xlsx"]
"""
import argparse
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import xlsx_notes  # noqa: E402

OUT_DIR = os.path.join(ROOT, "output", "Data and Chart Outputs")
DEFAULT_OUT = os.path.join(OUT_DIR, "haynesville_cost_curve.xlsx")
GULF_XLSX = os.path.join(OUT_DIR, "gulf_coast_gas_balance.xlsx")
HH_XLSX = os.path.join(OUT_DIR, "henry_hub_daily.xlsx")

FILL_IN = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
FILL_HEAD = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
FILL_GREY = PatternFill("solid", start_color="EDEDED", end_color="EDEDED")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {"SOURCED": "E2EFDA", "SOURCED-CHART-READ": "E2EFDA", "DERIVED": "DDEBF7", "ASSUMPTION": "FCE4D6", "OWNER": "FFF2CC"}

N_H = 10                       # capacity horizons H = 1..10 years
HS = list(range(1, N_H + 1))
YEARS_20 = list(range(1, 21))
BASE_YEAR = 2025               # the Key question compares with Dec 2025
LABEL_BANNER = "APPROXIMATE - built from company disclosures and BEG tiering; breakevens are pre-tax model results, not a vendor curve"

# ----------------------------------------------------------------------------------------------------------------- sources
BEG_URL = "https://www.beg.utexas.edu/files/content/beg/research/shale/Haynesville%20Shale%20Gas%20Play.pdf"
CRK_PRES = "https://investors.comstockresources.com/static-files/5a596a22-02f6-4b49-a9cc-93ecfe0179c0"
CRK_Q2 = "https://www.sec.gov/Archives/edgar/data/23194/000119312526323767/crk-ex99_1.htm"
CRK_10K = "https://www.sec.gov/Archives/edgar/data/23194/000119312526059001/crk-20251231.htm"
CRK_Q2_25 = "https://www.sec.gov/Archives/edgar/data/23194/000095017025100294/crk-ex99_1.htm"
EXE_PRES = "https://investors.expandenergy.com/static-files/b269b415-dfe3-4eca-b66c-250be02bcbae"
EXE_PR = "https://www.sec.gov/Archives/edgar/data/895126/000089512626000046/exe-ex_991x20260630x8kxpr.htm"
EIA_OGSM = "https://www.eia.gov/outlooks/aeo/assumptions/pdf/OGSM_Assumptions.pdf"
EIA_TIE = "https://www.eia.gov/todayinenergy/detail.php?id=67944"
DAL_Q1 = "https://www.dallasfed.org/research/surveys/des/2026/2601"
DAL_Q3 = "https://www.dallasfed.org/research/surveys/des/2026/2603"
EVAL = "https://info.evaluateenergy.com/haynesville-producers-eye-rising-gas-prices/"
DPR = "https://www.eia.gov/petroleum/drilling/"
OPENED = "OPENED in GitHub Actions 6 Oct 2026"
CHART = "OPENED in GitHub Actions 6 Oct 2026; value READ OFF A CHART IMAGE (rendered page, approx. +/-3%)"

# (id, type, item, value, unit, as-of, publisher, document, url, status, used)
SOURCES = [
    # --- BEG / OGJ
    ("B1", "EUR", "Average EUR per 4,800-ft well, Tier 1 / 2 / 3 / 4 / 5 / 6", "7.9 / 5.9 / 4.5 / 3.3 / 2.4 / 1.4", "Bcf per 4,800-ft well", "Dec 2015 (wells to 2012)", "Bureau of Economic Geology, UT Austin (Browning, Ikonnikova et al.)",
     "Oil & Gas Journal 7 Dec 2015, 'Study forecasts gradual Haynesville production recovery before final decline', Fig. 7 legend", BEG_URL, OPENED, "Inputs (tier productivity ratios), BEG check"),
    ("B2", "inventory", "Remaining potential wells by tier: partly drilled blocks / undrilled blocks / drilled to 2012", "see 'BEG check' (T1 1,700/2,250/450 ... T6 1,170/5,700/180)", "wells (4,800-ft)", "2012 well data", "Bureau of Economic Geology",
     "OGJ 7 Dec 2015, Fig. 6 stacked bars", BEG_URL, CHART, "Inputs (tier inventory)"),
    ("B3", "breakeven", "Breakeven Henry Hub price by tier at 0% / 10% / 20% IRR (2011 $)", "see 'BEG check' (10% IRR: T1 3.5, T2 4.4, T3 5.4, T4 6.8, T5 9.0, T6 15.2)", "$/MMBtu", "2011 $, 2015 study", "Bureau of Economic Geology",
     "OGJ 7 Dec 2015, Fig. 8", BEG_URL, CHART, "BEG check (validation of the formula)"),
    ("B4", "well cost", "Drilling cost 12,000-ft TVD $10.5 million (20% tangible), related capex 13%, expense $82,500/well/yr +13% overhead, gathering/compression/treatment $0.49/Mcf, royalty 25%, severance 4% (2 yrs exempt), basis HH -$0.07, tax 40.2%", "10.5 / 13% / 82,500 / 0.49 / 25% / -0.07", "mixed", "2015 study", "Bureau of Economic Geology",
     "OGJ 7 Dec 2015, Table 1", BEG_URL, OPENED, "Inputs (royalty, basis, related capex), BEG check"),
    ("B5", "inventory", "Development ceilings: 80% of partly drilled blocks, 40% of undrilled blocks; 20-yr well life; 4,800-ft normalised lateral", "80% / 40% / 20 / 4,800", "% / years / ft", "2015 study", "Bureau of Economic Geology",
     "OGJ 7 Dec 2015, Table 2 and text", BEG_URL, OPENED, "Inputs (developable shares, lateral)"),
    ("B6", "EUR", "Share of EUR recovered in first 5 years: >80% (Tier 1: 78%); nearly depleted by year 10 with >90% produced", ">80% (78% T1) / >90%", "% of EUR", "2015 study", "Bureau of Economic Geology",
     "OGJ 7 Dec 2015, 'Well economics'", BEG_URL, OPENED, "Inputs (decline anchors C(5), C(10)) - anchors chosen inside BEG's bounds = ASSUMPTION"),
    ("B7", "breakeven", "'A substantial portion of the reservoir has breakeven gas prices of $4-6/MMbtu'; remaining technically recoverable resource 177 Tcf; 2,527 wells to 2012 EUR 10.3 Tcf", "4-6 / 177 / 10.3", "$/MMBtu / Tcf", "Dec 2015", "Bureau of Economic Geology",
     "OGJ 7 Dec 2015 (also SNL 23 Dec 2015 summary)", BEG_URL, OPENED, "context (2011-15 costs)"),
    # --- Comstock
    ("C1", "inventory", "Legacy Haynesville + Bossier drilling inventory 30 Jun 2026: 1,705 gross / 816 net locations, average lateral 10,153 ft (Haynesville 881, Bossier 824 gross)", "1,705 / 816 / 10,153", "locations / ft", "30 Jun 2026", "Comstock Resources",
     "Investor Presentation October 2026, slide 13", CRK_PRES, CHART, "Cross-check table (inside BEG's legacy inventory, not added)"),
    ("C2", "inventory", "Western Haynesville operated drilling locations 30 Jun 2026: 3,277 gross / 2,528 net (Haynesville 1,206, Bossier 2,071 gross), average lateral 8,875 ft; assumes 70-90% WI", "3,277 / 2,528 / 8,875", "locations / ft", "30 Jun 2026", "Comstock Resources",
     "Investor Presentation October 2026, slide 14", CRK_PRES, CHART, "Inputs (Western Haynesville row)"),
    ("C3", "well cost", "Legacy Haynesville D&C cost, laterals >8,500 ft, Q2 2026: drilling $710 + completion $680 per lateral foot", "1,390", "$ per lateral ft", "Q2 2026", "Comstock Resources",
     "Investor Presentation October 2026, slide 21 (bar labels)", CRK_PRES, CHART, "Inputs (Legacy tiers $/ft)"),
    ("C4", "well cost", "Western Haynesville D&C cost Q2 2026: drilling $1,738 + completion $1,609 per lateral foot (Q1 2026: 1,534 + 1,537; 2025 Q2: 1,817 + 1,305)", "3,347", "$ per lateral ft", "Q2 2026", "Comstock Resources",
     "Investor Presentation October 2026, slide 22 (bar labels)", CRK_PRES, CHART, "Inputs (Western Haynesville $/ft)"),
    ("C5", "opex", "Production cost Q2 2026: gathering & transportation $0.38, lease operating $0.25, production & other taxes $0.06 (+ cash G&A $0.08; total $0.77) per Mcfe; realised $2.55/Mcfe before hedging", "0.38 / 0.25 / 0.06", "$ per Mcfe", "Q2 2026", "Comstock Resources",
     "Q2 2026 results, SEC 8-K Ex. 99.1 (29 Jul 2026)", CRK_Q2, OPENED, "Inputs (variable cost)"),
    ("C6", "IP", "YTD 2026 wells turned to sales: Legacy 22 wells, 12,052 ft average, IP 31 MMcf/d; Western Haynesville 11 wells, 10,331 ft average, IP 31 MMcf/d (per-well tables on slides 17-18)", "31 / 31", "MMcf/d per well", "30 Jun 2026", "Comstock Resources",
     "Investor Presentation October 2026, slides 17-18; Q2 2026 8-K", CRK_PRES, OPENED, "'Cross-checks' (IP per 1,000 ft; Western vs Legacy 1.17x); not used for EUR"),
    ("C7", "inventory", "Dec 2025: 1,848 gross (886 net) Legacy and 3,343 gross (2,561 net) Western Haynesville drilling locations; average lateral 10,077 ft / 8,873 ft; PUD 4.2 Tcf from 332 undeveloped locations (no PUD well below 10% rate of return)", "1,848 / 3,343 / 4.2 Tcf / 332", "locations / Tcf", "31 Dec 2025", "Comstock Resources",
     "Form 10-K FY2025 (filed 19 Feb 2026), Item 1 and reserves note", CRK_10K, OPENED, "'Cross-checks' (PUD Bcf per location)"),
    ("C8", "well cost", "Q2 2025: five Western Haynesville wells, 10,897 ft, IP 36 MMcf/d, drilled and completed at $2,647 per completed lateral foot", "2,647", "$ per lateral ft", "Q2 2025", "Comstock Resources",
     "Q2 2025 results, SEC 8-K Ex. 99.1 (30 Jul 2025)", CRK_Q2_25, OPENED, "reference only (older basis; the Q2 2026 slide is used)"),
    # --- Expand
    ("E1", "breakeven", "Haynesville 'Capital efficiency improvements yielding <$2.75 breakeven' (annual asset-level breakeven, excludes corporate items; hurdle/discount basis not stated)", "<2.75", "$/MMBtu (company basis)", "28 Oct 2025", "Expand Energy",
     "3Q 2025 Earnings presentation, slides 3 and 12", EXE_PRES, OPENED, "'Cross-checks' vs model Tier 1"),
    ("E2", "EUR", "Average Haynesville 12-month cumulative productivity, Mcfe per lateral ft: 2022 806, 2023 812, 2024 783, 2025E ~835, 2026E ~855; 'average EXE well productivity ~40% greater than basin average 2022-25' (peer data Enverus)", "855 (2026E)", "Mcfe/ft, 12 months", "28 Oct 2025", "Expand Energy",
     "3Q 2025 Earnings presentation, slide 12 (chart labels)", EXE_PRES, CHART, "Inputs (productivity anchor)"),
    ("E3", "well cost", "Haynesville well cost per lateral ft: 2022 $1,657, 2023 $1,847, 2024 $1,573, 2025E ~$1,370, 2026E ~$1,340 (Legacy CHK + Legacy SWN actuals)", "~1,340 (2026E)", "$ per lateral ft", "28 Oct 2025", "Expand Energy",
     "3Q 2025 Earnings presentation, slide 12 (chart labels)", EXE_PRES, CHART, "'Cross-checks' vs Comstock Legacy $1,390"),
    ("E4", "inventory", "Haynesville: >2,000 gross locations (31 Dec 2024), ~664,000 net acres, FY25E production ~3,000 MMcfe/d; >5,000 gross locations (all basins) / ~250 TILs a year = 20+ years; Western Haynesville 75,000+ net acres, potential to unlock >200 locations (10k-ft normalised)", ">2,000 / >200", "locations", "28 Oct 2025", "Expand Energy",
     "3Q 2025 Earnings presentation, slides 4 and 13", EXE_PRES, OPENED, "Inputs (Expand Western row); cross-check"),
    ("E5", "breakeven", "'Improved our Haynesville breakevens by approximately 15%' in 2025 (qualitative; no $ figure in the release)", "~15%", "%", "17 Feb 2026", "Expand Energy",
     "4Q 2025 results release, SEC 8-K Ex. 99.1; 2Q 2026 release (28 Jul 2026) repeats the theme", EXE_PR, OPENED, "context"),
    # --- EIA / Fed
    ("A1", "EUR", "AEO2023 Oil & Gas Supply Module: Haynesville-Bossier LA 1,935 mi2, 6 wells/mi2, average EUR 10.532 Bcf/well, TRR 121.9 Tcf; TX 1,561 mi2, 6 wells/mi2, 9.248 Bcf/well, 86.3 Tcf (unproved resources as of 1 Jan 2021)", "10.5 / 9.2 Bcf/well", "Bcf per well", "AEO2023 (Mar 2023)", "U.S. EIA",
     "Assumptions to the AEO2023: Oil and Gas Supply Module, Table 3", EIA_OGSM, OPENED, "'Cross-checks' only (lateral basis not stated)"),
    ("A2", "price context", "EIA forecasts Henry Hub spot $3.44/MMBtu average in 2026; 'at this price drilling in the Haynesville remains economical'; Haynesville production +1.3 Bcf/d (9%) in 2026", "3.44", "$/MMBtu", "Aug 2026 STEO", "U.S. EIA",
     "Today in Energy 'United States on track for record natural gas production in 2026'", EIA_TIE, OPENED, "context"),
    ("A3", "price expectation", "Dallas Fed Energy Survey Q3 2026: respondents expect Henry Hub $3.29 at year-end 2026, $3.82 in two years, $4.28 in five years (Q1 2026: $3.60 / $4.03 / $4.42)", "3.82 / 4.28", "$/MMBtu", "30 Sep 2026", "Federal Reserve Bank of Dallas",
     "Dallas Fed Energy Survey, third quarter 2026 (and first quarter 2026)", DAL_Q3, OPENED, "'Mapping' reference"),
    ("A4", "breakeven", "Dallas Fed annual break-even update is for OIL by basin only ($62-70/bbl; Permian $67/bbl, Q1 2026); no gas or Haynesville breakeven is published in the report text", "n/a", "-", "25 Mar 2026", "Federal Reserve Bank of Dallas",
     "Dallas Fed Energy Survey, first quarter 2026 (and Q1 2025)", DAL_Q1, OPENED, "none - no gas breakeven"),
    ("A5", "context", "EIA Drilling Productivity Report, last edition (May 2024): Haynesville June 2024 gas 15,606 -> 15,339 MMcf/d, new wells +472, legacy decline -739 MMcf/d in one month; new-well gas per rig 13,487 Mcf/d. EIA moved the data into the STEO from June 2024", "-739 MMcf/d per month", "MMcf/d", "May 2024", "U.S. EIA",
     "Drilling Productivity Report, Haynesville Region page 7", DPR, OPENED, "context for the base-decline note; not used"),
    # --- secondary / old
    ("S1", "breakeven (secondary)", "Aethon CEO: 'most producers need $3.50/MMBtu to produce profitably' (Goldman Sachs conference, Jan 2025); article also quotes a '2023 EIA study: at 10% IRR, 13% of Haynesville acreage profitable at $4, 72% at $8' - the EIA study itself was NOT found or opened", "3.50", "$/MMBtu", "15 Apr 2025", "Evaluate Energy (secondary article)",
     "'Haynesville producers eye rising gas prices'", EVAL, "OPENED (secondary quote; underlying documents not opened)", "not used"),
    ("S2", "breakeven (old)", "BTU Analytics (Feb 2016): best Haynesville acreage breakeven $1.50-3.00/MMBtu, plenty of areas above $5; only 20% of sub-$2 locations (177 wells) drilled", "1.5-3.0 / >5", "$/MMBtu", "19 Feb 2016", "NGI Shale Daily (via BEG ext-aff copy)",
     "'Haynesville An Able Price Competitor With Northeast NatGas in Plenty of Areas'", "https://www.beg.utexas.edu/files/content/beg/ext-aff/16-02/Haynesville%20An%20Able%20Price%20Competitor%20With%20Northeast%20NatGas%20in%20'Plenty%20of%20Areas'.pdf", "OPENED (2016 costs, stale)", "not used"),
]

SNIPPETS = [
    ("SNIPPET - unopened", "Expand: 'owns 72% of the lowest breakeven inventory in the Haynesville' and 'breakevens below $3 and <$2.75'", "web-search result summary only (gurufocus / transcripts); the opened Expand deck above gives <$2.75 but not the 72%"),
    ("SNIPPET - unopened", "Comstock Q3 2024 Western Haynesville cost $2,814 per completed lateral foot; horseshoe $1,740 vs $2,270 per lateral foot; D&C 24.8 Bcf EUR for 10,000-ft Gen II wells", "web-search result summaries; the Q3 2024 release (SEC 000095017024119021) was not opened"),
    ("SNIPPET - unopened", "Comstock investor-presentation 'big hole' lateral at $1,306 per lateral foot (Q2 2026 call)", "the Nasdaq call summary was opened (secondary) and states it; not used in the model"),
    ("SNIPPET - unopened", "Haynesville breakeven ~$3.75 / class-1 $3.21, class-2 $3.57 at 30% FCF / 'S&P Global core acreage ~$3'", "search-result snippets from unnamed articles; not opened"),
]

NOT_REACHABLE = [
    ("Comstock IR website www.comstockresources.com/investors/presentations", "HTTP 404 from Actions; ir.comstockresources.com does not resolve. The investor presentations ARE on investors.comstockresources.com/static-files/..., opened with curl_cffi (browser impersonation) - plain requests are refused."),
    ("Expand IR website investors.expandenergy.com", "plain requests read-timed out from Actions; the 3Q 2025 deck opened with curl_cffi. The 4Q25 and 2Q26 decks were not located (only the press releases, which carry no breakeven figure, were opened)."),
    ("IEA Gas Market Report / World Energy Outlook US supply cost charts", "iea.org returned HTTP 403 from Actions (guessed report URL); no IEA document with Haynesville cost data opened."),
    ("Oxford Institute for Energy Studies", "oxfordenergy.org returned HTTP 403; a web search found no OIES Haynesville breakeven publication."),
    ("Baker Institute (Rice)", "site root opened but no Haynesville cost document found."),
    ("Kansas City Fed Energy Survey", "page opened; no Haynesville or gas-basin breakeven in the text. Dallas Fed: oil breakevens by basin only."),
    ("EIA 'Trends in U.S. Oil and Natural Gas Upstream Costs' (2016 edition)", "opened (8 MB PDF); oil-focused, no Haynesville gas breakeven. The AEO OGSM file served is the 2023 edition."),
    ("EIA 2023 study quoted by Evaluate Energy (13% of acreage profitable at $4, 72% at $8)", "underlying study not found; only the secondary quote was opened."),
    ("Aethon, BPX (BP), TG Natural Resources, other private/foreign-owned operators", "no public filings with inventory or breakeven; Aethon appears only through the secondary quote above. BP investor materials were not probed."),
    ("Rystad, Enverus, Kayrros, Wood Mackenzie curves", "paywalled; no public document with data opened. This is the gap the editable override columns are for."),
    ("Goodrich Petroleum", "acquired (2021); no current filings."),
]

# ------------------------------------------------------------------------------------------------------ inputs (defaults)
# key, label, value, unit, status, note
ASSUME = [
    ("disc", "Discount rate for breakeven (pre-tax NPV = 0)", 0.10, "share", "ASSUMPTION", "BEG and the EIA study quoted by Evaluate Energy use a 10% IRR hurdle; Expand's own breakeven basis is not stated."),
    ("roy", "Royalty rate", 0.25, "share of revenue", "SOURCED", "BEG Table 1 (2015). Current lease royalties are not in the opened company documents."),
    ("basis", "Basis, Henry Hub to wellhead (negative = wellhead below HH)", -0.07, "$/MMBtu", "SOURCED", "BEG Table 1 (2015); present-day Gulf Coast basis not sourced. Breakeven is quoted at Henry Hub: HH = needed realised price - basis."),
    ("loe", "Lease operating cost", 0.25, "$/Mcf", "SOURCED", "Comstock Q2 2026 release (per Mcfe, ~96% gas)."),
    ("gt", "Gathering & transportation", 0.38, "$/Mcf", "SOURCED", "Comstock Q2 2026 release (company-owned midstream; third-party rates will differ)."),
    ("ptax", "Production & other taxes", 0.06, "$/Mcf", "SOURCED", "Comstock Q2 2026 release (guidance 2026: 0.10-0.15)."),
    ("oth", "Facilities & other capex as share of D&C", 0.13, "share", "ASSUMPTION", "BEG Table 1 'related capital expenditures 13%' (2015) applied to today's D&C; Comstock's $/ft is drilling + completion only."),
    ("c5", "Cumulative share of EUR recovered after 5 years", 0.80, "share", "ASSUMPTION", "BEG text: 'more than 80%' (Tier 1 78%). Anchor for the decline shape."),
    ("c10", "Cumulative share of EUR recovered after 10 years", 0.92, "share", "ASSUMPTION", "BEG text: 'more than 90%'. Anchor for the decline shape."),
    ("m12", "12-month cumulative production of a modern Tier 1 well", 0.855, "Bcfe per 1,000 lateral ft", "SOURCED", "Expand 3Q25 deck slide 12: ~855 Mcfe/ft (2026E, company estimate; Mcfe ~ Mcf, Expand is ~92% gas)."),
    ("cost_leg", "D&C cost per lateral foot, Legacy tiers", 1390, "$/ft", "SOURCED", "Comstock Oct 2026 deck slide 21, Q2 2026: $710 drilling + $680 completion (laterals >8,500 ft). Expand 2026E ~$1,340 (cross-check)."),
    ("cost_wh", "D&C cost per lateral foot, Western Haynesville", 3347, "$/ft", "SOURCED", "Comstock Oct 2026 deck slide 22, Q2 2026: $1,738 + $1,609."),
    ("wh_mult", "Western Haynesville productivity per ft relative to modern Tier 1", 1.0, "ratio", "ASSUMPTION", "NO sourced Western Haynesville EUR. 1.0 = same EUR per foot as Tier 1. Comstock's YTD-2026 IP per 1,000 ft is 1.17x Legacy ('Cross-checks'), but IP is not EUR."),
    ("dev_p", "Developable share of partly drilled inventory", 0.80, "share", "SOURCED", "BEG Table 2 development ceiling (2015)."),
    ("dev_u", "Developable share of undrilled inventory", 0.40, "share", "SOURCED", "BEG Table 2 development ceiling (2015)."),
    ("inv_share", "Share of BEG's 2012 legacy inventory still undrilled today", 1.0, "share", "ASSUMPTION", "NOT sourced: thousands of Haynesville wells have been drilled since BEG's 2012 count and are not deducted (1.0 = no haircut). Set below 1 to shrink Legacy Tiers 1-6; see 'Sensitivity values' for 50% and 25%."),
    ("lat_beg", "Lateral length of a BEG-inventory well", 4.8, "1,000 ft", "SOURCED", "BEG normalised EUR and inventory to 4,800-ft wells."),
    ("hp", "Programme length for the 'Capacity over time' tab", 8, "years", "ASSUMPTION", "Demonstration only; the curve and mapping use H = years to each Key-question date."),
]
AK = {k: i for i, (k, *_r) in enumerate(ASSUME)}

# tier rows: name, basis, partly, undrilled, share_p, share_u, lateral, rel-productivity (formula key), cost key, source
BEG_EUR = [7.9, 5.9, 4.5, 3.3, 2.4, 1.4]
BEG_DRILLED = [450, 480, 600, 450, 300, 180]
BEG_PARTLY = [1700, 1470, 2550, 2050, 1900, 1170]
BEG_UNDRILLED = [2250, 2000, 3200, 3250, 4500, 5700]
BEG_BE = {0: [2.9, 3.7, 4.6, 6.0, 7.8, 13.4], 10: [3.5, 4.4, 5.4, 6.8, 9.0, 15.2], 20: [4.0, 4.9, 6.0, 7.6, 10.0, 16.8]}
TIERS = [(f"Legacy Tier {i + 1} (BEG)", "BEG tier, legacy Haynesville", BEG_PARTLY[i], BEG_UNDRILLED[i], "dev_p*inv_share", "dev_u*inv_share", "lat_beg", f"beg{i}", "cost_leg",
          "Inventory: BEG Fig. 6 (chart-read, 2012 data, wells drilled since 2012 NOT deducted); productivity: BEG EUR ratio x Expand anchor; cost: Comstock Legacy Q2 2026.") for i in range(6)]
TIERS += [("Western Haynesville - Comstock operated", "Western Haynesville", 3277, 0, 1.0, 0.0, 8.875, "wh_mult", "cost_wh",
           "Inventory: Comstock slide 14 (3,277 gross operated locations, 30 Jun 2026; counted inventory, so shares 1/0); productivity ASSUMED; cost: Comstock slide 22."),
          ("Western Haynesville - Expand prospective", "Western Haynesville", 200, 0, 1.0, 0.0, 10.0, "wh_mult", "cost_wh",
           "Inventory: Expand 3Q25 slide 13 '>200 locations' (10k-ft normalised, prospective acreage, undelineated); productivity and cost as Comstock Western Haynesville (ASSUMED).")]
NT = len(TIERS)

# ------------------------------------------------------------------------------------------------------------ python mirror


def profile(a):
    r = np.log(1 - a["c10"]) / np.log(1 - a["c5"])
    beta = np.log(r) / np.log(2)
    tau = 5 / (-np.log(1 - a["c5"])) ** (1 / beta)
    cum = np.array([1 - np.exp(-((n / tau) ** beta)) for n in YEARS_20])
    f = np.diff(np.concatenate([[0.0], cum]))
    pvf = float(sum(f[n - 1] / (1 + a["disc"]) ** (n - 0.5) for n in YEARS_20))
    return dict(beta=beta, tau=tau, cum=cum, f=f, pvf=pvf, eur_anchor=a["m12"] / f[0])


def model(a, tiers):
    p = profile(a)
    var = a["loe"] + a["gt"] + a["ptax"]
    rows = []
    for t in tiers:
        name, _b, c, d, sp, su, lat, rel, cost = t[:9]
        kft = (c * sp + d * su) * lat
        eur = p["eur_anchor"] * rel
        capex = cost * 1000 * (1 + a["oth"])
        pv = eur * 1e6 * p["pvf"]
        be = (var + capex / pv) / (1 - a["roy"]) - a["basis"]
        rows.append(dict(name=name, kft=kft, eur=eur, capex=capex, pv=pv, be=be, wells=kft / lat, tcf=kft * eur / 1000))
    df = pd.DataFrame(rows)
    df["rank"] = df["be"].rank(method="first").astype(int)
    for H in HS:
        df[f"cap{H}"] = df["kft"] * df["eur"] * p["cum"][H - 1] / H / 365
    s = df.sort_values("rank").reset_index(drop=True)
    for H in HS:
        s[f"cum{H}"] = s[f"cap{H}"].cumsum()
    s["cumwells"] = s["wells"].cumsum()
    return p, df, s


def marginal(s, H, x):
    if x is None or x <= 0:
        return None
    idx = int((s[f"cum{H}"] < x).sum())
    return float(s["be"].iloc[idx]) if idx < len(s) else float("inf")


# --------------------------------------------------------------------------------------------------------------- data in
def read_gulf():
    f = pd.read_excel(GULF_XLSX, "Key question")
    f["Scenario"] = f["Scenario"].ffill()
    f = f[f["Year-end"].astype(str).str.startswith("Dec ")].copy()
    f["Year"] = f["Year-end"].str[-4:].astype(int)
    keep = {"Scenario": "Scenario", "Year": "Year",
            "Incremental Haynesville supply (East Texas + Louisiana, STEO)": "STEO Haynesville growth",
            "Extra supply required to hold the outflow (Haynesville above STEO or other regions)": "Extra supply required",
            "Haynesville growth required to balance (STEO + extra)": "Haynesville growth required (STEO + extra)"}
    out = f[list(keep)].rename(columns=keep)
    for c in f.columns:
        if c.startswith("Extra supply required with layers up to"):
            lab = c.replace("Extra supply required with layers up to ", "").replace(" (cumulative, ILLUSTRATIVE)", "")
            out[f"Extra with layers up to {lab}"] = f[c].values
    asof = dt.datetime.fromtimestamp(os.path.getmtime(GULF_XLSX)).strftime("%d %b %Y")
    return out.reset_index(drop=True), asof


def read_hh():
    d = pd.read_excel(HH_XLSX, "Data")
    d["date"] = pd.to_datetime(d["date"])
    s = d.set_index("date")["Henry_Hub_USD_per_MMBtu"].dropna()
    last = s.index.max()
    return dict(last_date=last, last=float(s.iloc[-1]), m12=float(s[s.index > last - pd.DateOffset(years=1)].mean()),
                y2025=float(s[s.index.year == 2025].mean()), y2024=float(s[s.index.year == 2024].mean()),
                since2021=float(s[s.index >= "2021-01-01"].mean()), min12=float(s[s.index > last - pd.DateOffset(years=1)].min()),
                max12=float(s[s.index > last - pd.DateOffset(years=1)].max()), monthly=s.resample("MS").mean().loc["2021-01-01":])


def read_back(path):
    """Yellow cells of the committed workbook (assumptions by label, tiers by name) so owner edits survive reruns."""
    prev = {"assume": {}, "tiers": {}}
    if not os.path.exists(path):
        return prev
    try:
        wb = load_workbook(path)
        ws = wb["Inputs"]
        for r in range(1, ws.max_row + 1):
            lab = ws.cell(r, 1).value
            if lab in [x[1] for x in ASSUME]:
                v = ws.cell(r, 2).value
                if isinstance(v, (int, float)):
                    prev["assume"][lab] = v
            elif isinstance(lab, str) and lab in [t[0] for t in TIERS]:
                prev["tiers"][lab] = [ws.cell(r, c).value for c in range(3, 13)]
    except Exception as e:  # noqa: BLE001
        print(f"read-back skipped: {type(e).__name__}: {e}")
    return prev


# ---------------------------------------------------------------------------------------------------------------- workbook
def style_head(ws, row, ncol, fill=FILL_HEAD):
    for c in range(1, ncol + 1):
        cell = ws.cell(row, c)
        cell.font = BOLD
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def build(out):
    prev = read_back(out)
    A_vals = {k: prev["assume"].get(lab, v) for k, lab, v, *_ in ASSUME}
    tiers = list(TIERS)
    gulf, gulf_asof = read_gulf()
    hh = read_hh()

    # python mirror with the (possibly owner-edited) inputs
    # tier overrides read back: columns C..L = partly, undrilled, share_p, share_u, lateral, rel, cost, be override, cap override, source
    ov_be, ov_cap = {}, {}
    tiers2 = []
    for t in tiers:
        t = list(t)
        pv = prev["tiers"].get(t[0])
        if pv:
            for i, j in ((0, 2), (1, 3)):          # partly, undrilled counts
                if isinstance(pv[i], (int, float)):
                    t[j] = pv[i]
            if isinstance(pv[4], (int, float)):
                t[6] = pv[4]
            if isinstance(pv[7], (int, float)):
                ov_be[t[0]] = pv[7]
            if isinstance(pv[8], (int, float)):
                ov_cap[t[0]] = pv[8]
        tiers2.append(t)
    A = dict(A_vals)

    def rel_of(t):
        k = t[7]
        return A["wh_mult"] if k == "wh_mult" else BEG_EUR[int(k[3:])] / BEG_EUR[0]
    def num(v):
        if isinstance(v, str):
            r = 1.0
            for part in v.split("*"):
                r *= A[part]
            return r
        return v
    mt = [(t[0], t[1], t[2], t[3], num(t[4]), num(t[5]), num(t[6]), rel_of(t), A[t[8]]) for t in tiers2]
    p, df, s = model(A, mt)
    for n, v in ov_be.items():
        df.loc[df["name"] == n, "be"] = v
    df["rank"] = df["be"].rank(method="first").astype(int)
    for H in HS:
        df[f"cap{H}"] = df["kft"] * df["eur"] * p["cum"][H - 1] / H / 365
        for n, v in ov_cap.items():
            df.loc[df["name"] == n, f"cap{H}"] = v
    s = df.sort_values("rank").reset_index(drop=True)
    for H in HS:
        s[f"cum{H}"] = s[f"cap{H}"].cumsum()
    s["cumwells"] = s["wells"].cumsum()

    # ----- value sheets (script values: used by PNG charts and as the audit copy of the live formulas)
    cv = s[["name", "be", "wells", "cumwells", "kft", "eur", "tcf"] + [f"cap{H}" for H in HS] + [f"cum{H}" for H in HS]].copy()
    cv.columns = ["Tier (sorted by breakeven)", "Breakeven $/MMBtu HH", "Developable locations", "Cumulative locations", "Developable lateral (1,000 ft)",
                  "EUR (Bcf per 1,000 ft)", "Developable EUR (Tcf)"] + [f"Capacity H={H} (Bcf/d)" for H in HS] + [f"Cumulative capacity H={H} (Bcf/d)" for H in HS]
    cv.insert(0, "Rank", range(1, len(cv) + 1))
    cv = cv.set_index("Rank")

    mrows = []
    for _, g in gulf.iterrows():
        if g["Year"] not in (2026, 2028, 2030, 2033):
            continue
        H = int(g["Year"] - BASE_YEAR)
        steo, extra = float(g["STEO Haynesville growth"]), float(g["Extra supply required"])
        lay = [c for c in gulf.columns if c.startswith("Extra with layers")]
        row = {"Scenario": g["Scenario"], "Year-end": f"Dec {int(g['Year'])}", "H (years)": H, "STEO Haynesville growth": steo, "Extra supply required": extra}
        for c in lay:
            row[c] = float(g[c])
        row["Marginal breakeven - extra supply alone"] = marginal(s, H, extra)
        row["Marginal breakeven - STEO growth + extra"] = marginal(s, H, steo + extra)
        for c in lay:
            row["Marginal breakeven - STEO growth + " + c.replace("Extra with layers up to ", "layers up to ")] = marginal(s, H, steo + float(g[c]))
        mrows.append(row)
    mv = pd.DataFrame(mrows)
    mv_num = mv.copy()
    for c in mv_num.columns:
        if c.startswith("Marginal"):
            mv_num[c] = mv_num[c].apply(lambda v: np.nan if v is None or (isinstance(v, float) and np.isinf(v)) else v)

    # ----- sensitivity (script values at the defaults above): marginal breakeven for base-case Dec 2030 / Dec 2033 vs inventory haircut and Western productivity
    sens_rows = []
    for inv in (1.0, 0.5, 0.25):
        for wm in (0.8, 1.0, 1.17):
            A2 = dict(A)
            A2["inv_share"], A2["wh_mult"] = inv, wm
            def num2(v, A2=A2):
                if isinstance(v, str):
                    r = 1.0
                    for part in v.split("*"):
                        r *= A2[part]
                    return r
                return v
            mt2 = [(t[0], t[1], t[2], t[3], num2(t[4]), num2(t[5]), num2(t[6]),
                    wm if t[7] == "wh_mult" else BEG_EUR[int(t[7][3:])] / BEG_EUR[0], A2[t[8]]) for t in tiers2]
            _p, _d, s2 = model(A2, mt2)
            row = {"Share of BEG legacy inventory left": inv, "Western Haynesville productivity vs Tier 1": wm}
            for sc, yr, lab in (("Base", 2030, "Base Dec 2030"), ("Base", 2033, "Base Dec 2033"), ("LNG and takeaway delayed 6 months", 2033, "Delayed Dec 2033")):
                g = gulf[(gulf["Scenario"] == sc) & (gulf["Year"] == yr)].iloc[0]
                H = yr - BASE_YEAR
                row[f"{lab}: STEO + extra"] = marginal(s2, H, float(g["STEO Haynesville growth"] + g["Extra supply required"]))
                lc = [c for c in gulf.columns if c.startswith("Extra with layers")][-1]
                row[f"{lab}: STEO + extra with all layers"] = marginal(s2, H, float(g["STEO Haynesville growth"] + g[lc]))
            sens_rows.append(row)
    sens = pd.DataFrame(sens_rows).replace(float("inf"), np.nan).set_index("Share of BEG legacy inventory left")

    # ----- Units / notes
    sources_df = pd.DataFrame(SOURCES, columns=["ID", "Type", "Item", "Value", "Unit", "As-of", "Publisher", "Document", "URL", "Status", "Used where"]).set_index("ID")
    snip = pd.DataFrame(SNIPPETS, columns=["Status", "Claim seen", "Why not used"])
    nr = pd.DataFrame(NOT_REACHABLE, columns=["Source", "What happened (Actions, 6 Oct 2026)"])
    n_src = sum(1 for r in ASSUME if r[4].startswith("SOURCED"))
    notes = [
        "UNITS",
        "Breakeven: US$ per MMBtu at Henry Hub (HH = needed wellhead price - basis), pre-tax, at the discount rate on 'Inputs'. Capacity: Bcf/d. Inventory: locations and 1,000 lateral ft.",
        "",
        "WHAT THIS IS",
        LABEL_BANNER + ". Question: what breakeven price does the extra Haynesville supply in gulf_coast_gas_balance.xlsx 'Key question' imply if only COST limited it?",
        "The curve is a COST curve. It is NOT constrained by rigs, crews, sand, takeaway pipelines, leasing or midstream: with H = years to the Key-question date it assumes the whole developable inventory of a tier could be drilled evenly over those years.",
        "NOT modelled: replacement of the base-production decline (maintenance drilling; the last EIA DPR, May 2024, showed legacy Haynesville gas falling 739 MMcf/d in one month on 15.6 Bcf/d), taxes, hedges, midstream returns, well-to-well variance inside a tier, inventory used since 2012, interference between tiers. The marginal prices are therefore LOWER bounds for a net increase in output.",
        "",
        "HOW TO READ THE TABS",
        "Chart / Chart - Marginal price: native Excel charts (no borders); their cells are formulas, so editing 'Inputs' redraws them in Excel. Sources: every opened document with value, unit, as-of, publisher, URL and type; then snippet-only claims (never used) and what was not reachable.",
        "Inputs: yellow cells editable and kept on the next run. Block A assumptions (status SOURCED / ASSUMPTION), block B tiers. To use your own (paid) curve overwrite the 'Breakeven override' and 'Capacity override' columns of block B (and rename/add rows by overwriting the tier counts): override cells win over the model.",
        "Profile: decline shape (stretched exponential fitted to BEG's 5- and 10-year recovery anchors), PV factor and the Tier 1 EUR anchor. Curve: live per-tier breakeven, capacity for H = 1..10 years, sorted curve. Capacity over time: Bcf/d build of a programme (live). Mapping: the Key-question increments and the implied marginal breakeven (live) beside Henry Hub. BEG check: BEG's published tiers and the formula run on BEG's own 2015 inputs. Cross-checks: company figures that were NOT used as inputs.",
        "",
        "STATUS COUNTS",
        f"Assumption block: {len(ASSUME)} inputs, {n_src} SOURCED, {len(ASSUME) - n_src} ASSUMPTION. Tiers: 8 rows. Breakeven per tier is a MODEL RESULT; the only published tier breakevens are BEG's (2011 $, chart-read) and Expand's single '<$2.75' (company basis). Western Haynesville EUR is NOT sourced (assumed equal to Tier 1 per foot).",
        "Inventory caveats: BEG's counts are 2012 vintage (thousands of wells since are not deducted); Comstock's counts are operated locations (some 70-90% WI) and sit inside the same basin as BEG's legacy tiers (shown on 'Cross-checks', not added).",
        "",
        "EDITING",
        "Yellow cells on 'Inputs' only. Set the override columns blank to return to the model. Sheets Curve values / Mapping values / Sensitivity values (inventory haircut x Western productivity; NaN = beyond inventory) are the script's own copy of the same arithmetic (used by the PNG charts); after editing in Excel the live tabs and native charts follow, the values tabs do not.",
    ]
    titles = {"UNITS", "WHAT THIS IS", "HOW TO READ THE TABS", "STATUS COUNTS", "EDITING"}
    hh_df = pd.DataFrame({"Henry Hub average (USD/MMBtu)": [hh["last"], hh["m12"], hh["min12"], hh["max12"], hh["y2025"], hh["y2024"], hh["since2021"]]},
                         index=[f"Last daily spot ({hh['last_date']:%d %b %Y})", "Last 12 months average", "Last 12 months min", "Last 12 months max", "2025 average", "2024 average",
                                "Jan 2021 - latest average"])
    sheets = {"Sources": sources_df, "Not reachable": nr.set_index("Source"), "Snippets (unused)": snip.set_index("Status"), "Gulf key question": gulf.set_index("Scenario"),
              "Henry Hub": hh_df, "Curve values": cv, "Mapping values": mv.set_index("Scenario"), "HH monthly values": hh["monthly"].to_frame("Henry Hub monthly average"), "Sensitivity values": sens,
              "Inputs": pd.DataFrame(), "Profile": pd.DataFrame(), "Curve": pd.DataFrame(), "Capacity over time": pd.DataFrame(), "Mapping": pd.DataFrame(),
              "BEG check": pd.DataFrame(), "Cross-checks": pd.DataFrame(), "Chart data": pd.DataFrame()}
    os.makedirs(os.path.dirname(out), exist_ok=True)
    xlsx_notes.write_workbook(out, sheets, notes, titles)

    wb = load_workbook(out)
    for n in ("Inputs", "Profile", "Curve", "Capacity over time", "Mapping", "BEG check", "Cross-checks", "Chart data"):
        del wb[n]
    # sheet order: Units, (charts are added by add_charts.py), Sources ...
    names = ["Inputs", "Profile", "Curve", "Capacity over time", "Mapping", "BEG check", "Cross-checks", "Chart data"]
    ws = {n: wb.create_sheet(n, 2 + i) for i, n in enumerate(names)}
    ref = write_inputs(ws["Inputs"], A_vals, tiers2, ov_be, ov_cap)
    write_profile(ws["Profile"], ref)
    write_curve(ws["Curve"], ref)
    write_capacity(ws["Capacity over time"], ref)
    write_mapping(ws["Mapping"], ref, gulf, hh, mv)
    write_beg_check(ws["BEG check"], ref)
    write_cross(ws["Cross-checks"], ref, A, s, p)
    write_chart_data(ws["Chart data"], ref, hh, mv)
    wb.save(out + ".tmp")
    os.replace(out + ".tmp", out)
    print(f"{out}: Haynesville cost curve built. Tiers sorted by breakeven:")
    print(s[["name", "be", "cum5", "cum8"]].round(2).to_string())
    print(mv_num.round(2).to_string())
    return s, mv, p


def write_inputs(ws, A_vals, tiers, ov_be, ov_cap):
    ws["A1"] = "Inputs - yellow cells are editable and survive reruns. " + LABEL_BANNER
    ws["A1"].font = BOLD
    ws["A2"] = "Block A: assumptions. Block B: tiers (rows can be overwritten with your own curve: use the two override columns)."
    hdr = ["Assumption", "Value", "Unit", "Status", "Source / note"]
    for j, h in enumerate(hdr, 1):
        ws.cell(4, j, h)
    style_head(ws, 4, 5)
    ref = {}
    for i, (k, lab, _v, unit, status, note) in enumerate(ASSUME):
        r = 5 + i
        ws.cell(r, 1, lab)
        c = ws.cell(r, 2, A_vals[k])
        c.fill = FILL_IN
        ws.cell(r, 3, unit)
        st = ws.cell(r, 4, status)
        st.fill = PatternFill("solid", start_color=STATUS_FILL[status], end_color=STATUS_FILL[status])
        ws.cell(r, 5, note).alignment = WRAP
        ref[k] = f"Inputs!$B${r}"
    ra = 5 + len(ASSUME)
    ref["var"] = f"(Inputs!$B${5 + AK['loe']}+Inputs!$B${5 + AK['gt']}+Inputs!$B${5 + AK['ptax']})"
    tb = ra + 2
    ws.cell(tb - 1, 1, "Block B: tiers")
    ws.cell(tb - 1, 1).font = BOLD
    th = ["Tier / block", "Basis", "Partly drilled / counted locations", "Undrilled locations", "Developable share of partly/counted", "Developable share of undrilled",
          "Lateral per location (1,000 ft)", "Relative productivity vs Tier 1 (EUR per ft)", "D&C cost ($ per lateral ft)", "Breakeven override ($/MMBtu HH; blank = model)",
          "Capacity override (Bcf/d, all horizons; blank = model)", "Source / status"]
    for j, h in enumerate(th, 1):
        ws.cell(tb, j, h)
    style_head(ws, tb, len(th))
    ws.row_dimensions[tb].height = 62
    r0 = tb + 1
    for i, t in enumerate(tiers):
        r = r0 + i
        name, basis, c, d, sp, su, lat, rel, cost, src = t
        ws.cell(r, 1, name)
        ws.cell(r, 2, basis)
        ws.cell(r, 3, c).fill = FILL_IN
        ws.cell(r, 4, d).fill = FILL_IN
        ws.cell(r, 5, "=" + "*".join(ref[x] for x in sp.split("*")) if isinstance(sp, str) else sp).fill = FILL_IN
        ws.cell(r, 6, "=" + "*".join(ref[x] for x in su.split("*")) if isinstance(su, str) else su).fill = FILL_IN
        ws.cell(r, 7, f"={ref[lat]}" if isinstance(lat, str) else lat).fill = FILL_IN
        if rel.startswith("beg"):
            k = int(rel[3:])
            ws.cell(r, 8, f"='BEG check'!$C${6 + k}/'BEG check'!$C$6").fill = FILL_IN
        else:
            ws.cell(r, 8, f"={ref[rel]}").fill = FILL_IN
        ws.cell(r, 9, f"={ref[cost]}").fill = FILL_IN
        for cc in (10, 11):
            ws.cell(r, cc).fill = FILL_IN
        if name in ov_be:
            ws.cell(r, 10, ov_be[name])
        if name in ov_cap:
            ws.cell(r, 11, ov_cap[name])
        ws.cell(r, 12, src).alignment = WRAP
    ref["t0"], ref["t1"] = r0, r0 + len(tiers) - 1
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    for col, w in zip("CDEFGHIJKL", (16, 14, 16, 16, 14, 16, 14, 18, 20, 90)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["E"].width = 60 if False else 16
    for r in range(5, 5 + len(ASSUME)):
        ws.row_dimensions[r].height = 30
    return ref


def write_profile(ws, ref):
    ws["A1"] = "Per-well production profile (shared by all tiers): cumulative share of EUR C(n) = 1 - exp(-(n/tau)^beta), fitted to the two BEG recovery anchors on 'Inputs'. Live formulas."
    ws["A1"].font = BOLD
    rows = [("beta = ln( ln(1-C10) / ln(1-C5) ) / ln 2", f"=LN(LN(1-{ref['c10']})/LN(1-{ref['c5']}))/LN(2)"),
            ("tau (years) = 5 / (-ln(1-C5))^(1/beta)", f"=5/(-LN(1-{ref['c5']}))^(1/B3)"),
            ("Year-1 share of EUR f(1)", "=C9"),
            ("PV factor: sum f(n)/(1+d)^(n-0.5)  (mid-year discounting, capex at t=0)", f"=SUMPRODUCT(C9:C28/(1+{ref['disc']})^(A9:A28-0.5))"),
            ("Modern Tier 1 EUR (Bcf per 1,000 lateral ft) = 12-month cumulative / f(1)", f"={ref['m12']}/B5")]
    for i, (lab, f) in enumerate(rows):
        ws.cell(3 + i, 1, lab)
        ws.cell(3 + i, 2, f)
    ws.cell(8, 1, "Year n")
    ws.cell(8, 2, "Cumulative share C(n)")
    ws.cell(8, 3, "Annual share f(n)")
    style_head(ws, 8, 3)
    for n in YEARS_20:
        r = 8 + n
        ws.cell(r, 1, n)
        ws.cell(r, 2, f"=1-EXP(-((A{r}/$B$4)^$B$3))")
        ws.cell(r, 3, f"=B{r}" if n == 1 else f"=B{r}-B{r - 1}")
    ws.column_dimensions["A"].width = 78
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 18
    ref["beta"], ref["tau"], ref["f1"], ref["pvf"], ref["eur_anchor"] = ("Profile!$B$3", "Profile!$B$4", "Profile!$B$5", "Profile!$B$6", "Profile!$B$7")
    ref["cum_rng"] = "Profile!$B$9:$B$28"
    ref["f_rng"] = "Profile!$C$9:$C$28"


def write_curve(ws, ref):
    t0 = ref["t0"]
    ws["A1"] = "Curve - live per-tier breakeven and capacity (formulas). " + LABEL_BANNER
    ws["A1"].font = BOLD
    ws["A2"] = ("Breakeven (HH) = ((LOE + gathering + taxes) + capex per 1,000 ft / PV volume per 1,000 ft) / (1 - royalty) - basis, with capex = $/ft x 1,000 x (1 + other capex) and "
                "PV volume = EUR x 1e6 Mcf x PV factor. Capacity(H) = developable lateral x EUR x C(H) / H / 365 (Bcf/d in year H of a programme drilling the inventory evenly over H years).")
    ws["A2"].alignment = WRAP
    head = ["Tier / block", "Developable lateral (1,000 ft)", "EUR (Bcf per 1,000 ft)", "Capex per 1,000 ft ($)", "PV volume per 1,000 ft (Mcf)", "Breakeven, model ($/MMBtu HH)",
            "Breakeven used", "Source of breakeven", "Rank", "Developable EUR (Tcf)", "Developable locations"]
    for j, h in enumerate(head, 1):
        ws.cell(4, j, h)
    for j, H in enumerate(HS):
        ws.cell(4, 12 + j, H)
    ws.cell(3, 12, "Capacity by horizon H (years) - Bcf/d")
    ws.cell(3, 12).font = BOLD
    style_head(ws, 4, 12 + N_H - 1)
    ws.row_dimensions[4].height = 48
    for i in range(NT):
        r, ir = 5 + i, t0 + i
        ws.cell(r, 1, f"=Inputs!A{ir}")
        ws.cell(r, 2, f"=(Inputs!C{ir}*Inputs!E{ir}+Inputs!D{ir}*Inputs!F{ir})*Inputs!G{ir}")
        ws.cell(r, 3, f"={ref['eur_anchor']}*Inputs!H{ir}")
        ws.cell(r, 4, f"=Inputs!I{ir}*1000*(1+{ref['oth']})")
        ws.cell(r, 5, f"=C{r}*1000000*{ref['pvf']}")
        ws.cell(r, 6, f"=({ref['var']}+D{r}/E{r})/(1-{ref['roy']})-{ref['basis']}")
        ws.cell(r, 7, f'=IF(ISNUMBER(Inputs!J{ir}),Inputs!J{ir},F{r})')
        ws.cell(r, 8, f'=IF(ISNUMBER(Inputs!J{ir}),"OWNER OVERRIDE","MODEL")')
        ws.cell(r, 9, f'=COUNTIF($G$5:$G${4 + NT},"<"&G{r})+COUNTIF($G$5:G{r},G{r})')
        ws.cell(r, 10, f"=B{r}*C{r}/1000")
        ws.cell(r, 11, f"=B{r}/Inputs!G{ir}")
        for j, H in enumerate(HS):
            col = get_column_letter(12 + j)
            ws.cell(r, 12 + j, f"=IF(ISNUMBER(Inputs!$K{ir}),Inputs!$K{ir},$B{r}*$C{r}*INDEX({ref['cum_rng']},{col}$4)/{col}$4/365)")
    s0 = 5 + NT + 3
    ws.cell(s0 - 2, 1, "Sorted curve (cheapest first) - cumulative capacity by horizon H")
    ws.cell(s0 - 2, 1).font = BOLD
    sh = ["Rank", "Row in table above", "Tier / block", "Breakeven used ($/MMBtu HH)", "Developable lateral (1,000 ft)", "Developable EUR (Tcf)", "Cumulative locations"]
    for j, h in enumerate(sh, 1):
        ws.cell(s0 - 1, j, h)
    for j, H in enumerate(HS):
        ws.cell(s0 - 1, 12 + j, H)
    style_head(ws, s0 - 1, 12 + N_H - 1)
    ws.cell(s0 - 3, 12, "Cumulative capacity, Bcf/d, by horizon H (years)")
    ws.cell(s0 - 3, 12).font = BOLD
    for k in range(1, NT + 1):
        r = s0 + k - 1
        ws.cell(r, 1, k)
        ws.cell(r, 2, f"=MATCH(A{r},$I$5:$I${4 + NT},0)")
        ws.cell(r, 3, f"=INDEX($A$5:$A${4 + NT},B{r})")
        ws.cell(r, 4, f"=INDEX($G$5:$G${4 + NT},B{r})")
        ws.cell(r, 5, f"=INDEX($B$5:$B${4 + NT},B{r})")
        ws.cell(r, 6, f"=INDEX($J$5:$J${4 + NT},B{r})")
        ws.cell(r, 7, f"=INDEX($K$5:$K${4 + NT},B{r})" if k == 1 else f"=G{r - 1}+INDEX($K$5:$K${4 + NT},B{r})")
        for j in range(N_H):
            col = get_column_letter(12 + j)
            add = f"INDEX({col}$5:{col}${4 + NT},$B{r})"
            ws.cell(r, 12 + j, f"={add}" if k == 1 else f"={col}{r - 1}+{add}")
    ref["s0"], ref["s1"] = s0, s0 + NT - 1
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["H"].width = 16
    for col in "BDEFGIJK":
        ws.column_dimensions[col].width = 16
    for j in range(N_H):
        ws.column_dimensions[get_column_letter(12 + j)].width = 10
    ws.row_dimensions[2].height = 60
    ref["cum_block"] = f"Curve!$L${ref['s0']}:$U${ref['s1']}"
    ref["h_hdr"] = "Curve!$L$4:$U$4"
    ref["be_sorted"] = f"Curve!$D${ref['s0']}:$D${ref['s1']}"
    ref["name_sorted"] = f"Curve!$C${ref['s0']}:$C${ref['s1']}"


def price_formula(ref, X, H):
    """Marginal tier = first sorted tier whose cumulative capacity (column for horizon H) reaches X. Written as a sum of comparisons (not
    COUNTIF over INDEX(range,0,col)) so every spreadsheet engine evaluates it the same way."""
    m = f"MATCH({H},{ref['h_hdr']},0)"
    n = "1+" + "+".join(f"(INDEX({ref['cum_block']},{k},{m})<{X})" for k in range(1, NT + 1))
    return f'=IF(NOT(ISNUMBER({X})),"-",IF({X}<=0,"-",IF({n}>{NT},"beyond inventory",INDEX({ref["be_sorted"]},{n}))))'


def write_capacity(ws, ref):
    ws["A1"] = ("Capacity over time (live): Bcf/d of a programme that drills each tier's developable inventory evenly over H years (H on 'Inputs'), "
                "year n = wells x EUR x [C(n) - C(n - min(n, H))] / 365. Cost-only upper bound on speed: no rig, crew or takeaway limit.")
    ws["A1"].font = BOLD
    ws["A2"] = "Programme length H (years):"
    ws["B2"] = f"={ref['hp']}"
    ws.cell(4, 1, "Year of programme")
    for i in range(NT):
        ws.cell(4, 2 + i, f"=Curve!A{5 + i}")
    ws.cell(4, 2 + NT, "Total, Bcf/d")
    style_head(ws, 4, 2 + NT)
    ws.row_dimensions[4].height = 48
    for n in YEARS_20:
        r = 4 + n
        ws.cell(r, 1, n)
        for i in range(NT):
            cr = 5 + i
            m = f"MIN($A{r},$B$2)"
            tail = f"IF($A{r}-{m}=0,0,INDEX({ref['cum_rng']},$A{r}-{m}))"
            ws.cell(r, 2 + i, f"=IF(ISNUMBER(Inputs!$K{ref['t0'] + i}),Inputs!$K{ref['t0'] + i},Curve!$B${cr}*Curve!$C${cr}*(INDEX({ref['cum_rng']},$A{r})-{tail})/$B$2/365)")
        ws.cell(r, 2 + NT, f"=SUM(B{r}:{get_column_letter(1 + NT)}{r})")
    ws.column_dimensions["A"].width = 22
    for i in range(NT + 1):
        ws.column_dimensions[get_column_letter(2 + i)].width = 17


def write_mapping(ws, ref, gulf, hh, mv):
    ws["A1"] = ("Mapping - implied marginal breakeven of the Gulf 'Key question' Haynesville increments (live). " + LABEL_BANNER)
    ws["A1"].font = BOLD
    ws["A2"] = (f"Source of the increments: gulf_coast_gas_balance.xlsx 'Key question' (read-only, file dated {read_gulf_asof()}; copy on 'Gulf key question'). Year-end increments are vs Dec 2025, Bcf/d. "
                "Stack position: the curve is filled cheapest first. 'STEO growth' is the Haynesville growth already in EIA STEO / its extension; the EXTRA supply sits on top of it, so the marginal tier for the extra "
                "volume is found at STEO growth + extra (headline); 'extra alone' (curve from zero) is the lower bound. Layers: cumulative extra with each Louisiana LNG layer (ILLUSTRATIVE). "
                "H = years from Dec 2025: capacity on the curve is what a cost-only (not rig / takeaway constrained) programme could add in that time. Price cells show 'beyond inventory' when the stack runs out.")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 92
    lay = [c for c in gulf.columns if c.startswith("Extra with layers")]
    head = ["Scenario", "Year-end", "H (years)", "STEO Haynesville growth", "Extra supply required"] + lay + \
           ["Marginal breakeven - extra supply alone", "Marginal breakeven - STEO growth + extra"] + \
           ["Marginal breakeven - STEO growth + " + c.replace("Extra with layers up to ", "layers up to ") for c in lay] + \
           ["Henry Hub, last 12 months average", "Headline marginal breakeven less Henry Hub"]
    ws.cell(4, 1, "Rows in the same order as 'Gulf key question' (Dec 2026, 2028, 2030, 2033; base then delayed)")
    for j, h in enumerate(head, 1):
        ws.cell(5, j, h)
    style_head(ws, 5, len(head))
    ws.row_dimensions[5].height = 78
    gl = gulf.reset_index(drop=True)
    r = 6
    col_of = {c: 6 + k for k, c in enumerate(lay)}
    pa, pb = 6 + len(lay), 7 + len(lay)
    pl0 = 8 + len(lay)
    hhc, dc = pl0 + len(lay), pl0 + len(lay) + 1
    for gi, g in gl.iterrows():
        if g["Year"] not in (2026, 2028, 2030, 2033):
            continue
        gr = 5 + gi + 1                                       # row on 'Gulf key question' (header row 1)
        ws.cell(r, 1, g["Scenario"])
        ws.cell(r, 2, f"Dec {int(g['Year'])}")
        ws.cell(r, 3, int(g["Year"] - BASE_YEAR))
        ws.cell(r, 4, f"='Gulf key question'!C{gi + 2}")
        ws.cell(r, 5, f"='Gulf key question'!D{gi + 2}")
        for c, cc in col_of.items():
            ws.cell(r, cc, f"='Gulf key question'!{get_column_letter(6 + lay.index(c))}{gi + 2}")
        ws.cell(r, pa, price_formula(ref, f"$E{r}", f"$C{r}"))
        ws.cell(r, pb, price_formula(ref, f"($D{r}+$E{r})", f"$C{r}"))
        for k, c in enumerate(lay):
            ws.cell(r, pl0 + k, price_formula(ref, f"($D{r}+{get_column_letter(col_of[c])}{r})", f"$C{r}"))
        ws.cell(r, hhc, f"='Henry Hub'!$B$3")
        ws.cell(r, dc, f'=IF(ISNUMBER({get_column_letter(pb)}{r}),{get_column_letter(pb)}{r}-{get_column_letter(hhc)}{r},"-")')
        r += 1
    ref["map_rows"] = (6, r - 1)
    ref["map_pb"] = get_column_letter(pb)
    ref["map_hh"] = get_column_letter(hhc)
    r += 1
    ws.cell(r, 1, "Reference prices (sourced): Dallas Fed Energy Survey Q3 2026 (30 Sep 2026) expected Henry Hub $3.29 year-end 2026, $3.82 in two years, $4.28 in five years; EIA STEO (Aug 2026) $3.44 average 2026. "
                  "Haynesville flow is also limited by rigs, pipelines (Gillis, LEAP ...) and midstream - this sheet says only what COST implies.")
    ws.cell(r, 1).alignment = WRAP
    ws.row_dimensions[r].height = 48
    ws.column_dimensions["A"].width = 34
    for j in range(2, len(head) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 17


def read_gulf_asof():
    try:
        return dt.datetime.fromtimestamp(os.path.getmtime(GULF_XLSX)).strftime("%d %b %Y")
    except OSError:
        return "n/a"


def write_beg_check(ws, ref):
    ws["A1"] = "BEG check: Bureau of Economic Geology tiers (OGJ 7 Dec 2015) and this workbook's breakeven formula run on BEG's own 2015 inputs (validation, 2011 $)."
    ws["A1"].font = BOLD
    ws["A2"] = ("Published text: EUR per tier (Fig. 7 legend). READ OFF CHARTS (approx. +/-3%): wells (Fig. 6) and breakevens (Fig. 8). Model: capex = $10.5 million x 1.13 per 4,800-ft well, gathering $0.49/Mcf, "
                "fixed expense $82,500 x 1.13 a year over 20 years, royalty 25%, basis -$0.07, 10% discount, pre-tax, no severance (BEG: 4%, 2 years exempt, and a 40.2% tax rate with a tax shield), BEG's EUR with this "
                "workbook's decline shape. Yellow cells are BEG inputs.")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 62
    hdr = ["Tier", "Fig. 6 drilled to 2012 (wells)", "EUR per 4,800-ft well (Bcf) - published", "Fig. 6 partly drilled (wells)", "Fig. 6 undrilled (wells)",
           "Fig. 8 breakeven 0% IRR", "Fig. 8 breakeven 10% IRR", "Fig. 8 breakeven 20% IRR", "Model breakeven (BEG inputs, 10% pre-tax)", "Model vs Fig. 8 10% IRR"]
    for j, h in enumerate(hdr, 1):
        ws.cell(5, j, h)
    style_head(ws, 5, len(hdr))
    ws.row_dimensions[5].height = 60
    for i in range(6):
        r = 6 + i
        ws.cell(r, 1, f"Tier {i + 1}")
        ws.cell(r, 2, BEG_DRILLED[i]).fill = FILL_IN
        ws.cell(r, 3, BEG_EUR[i]).fill = FILL_IN
        ws.cell(r, 4, BEG_PARTLY[i]).fill = FILL_IN
        ws.cell(r, 5, BEG_UNDRILLED[i]).fill = FILL_IN
        for c, irr in zip((6, 7, 8), (0, 10, 20)):
            ws.cell(r, c, BEG_BE[irr][i]).fill = FILL_IN
        pv = f"(C{r}*1000000*{ref['pvf']})"
        # fixed expense PV: 20-year annuity, mid-year, at the discount rate
        fixed = f"($B$16*1.13*SUMPRODUCT(1/(1+{ref['disc']})^(Profile!$A$9:$A$28-0.5)))"
        ws.cell(r, 9, f"=(($B$15+({ '$B$14*1.13*1000000' }+{fixed})/{pv}))/(1-$B$17)-$B$18")
        ws.cell(r, 10, f"=I{r}/G{r}-1")
        ws.cell(r, 10).number_format = "0%"
    ws.cell(13, 1, "BEG inputs (Table 1)").font = BOLD
    for r, (lab, v) in enumerate([("Drilling cost, 12,000-ft TVD ($ million)", 10.5), ("Gathering, compression, treatment ($/Mcf)", 0.49),
                                  ("Expense per well per year ($)", 82500), ("Royalty", 0.25), ("Basis ($/MMBtu)", -0.07)], 14):
        ws.cell(r, 1, lab)
        ws.cell(r, 2, v).fill = FILL_IN
    ws.cell(20, 1, "Reading: the formula reproduces BEG's 10%-IRR breakevens for Tiers 1-3 within a few percent and runs higher for the deepest tiers; BEG's after-tax model and 2011 costs differ in detail. This validates the structure, not the modern inputs.")
    ws.cell(20, 1).alignment = WRAP
    ws.row_dimensions[20].height = 48
    ws.column_dimensions["A"].width = 46
    for j in range(2, 11):
        ws.column_dimensions[get_column_letter(j)].width = 17
    ws.cell(5, 1).value = "Tier"


def write_cross(ws, ref, A, s, p):
    ws["A1"] = "Cross-checks - company and EIA figures NOT used as inputs, set beside what the model implies (values; see Sources for URLs)."
    ws["A1"].font = BOLD
    t1 = s.loc[s["name"].str.startswith("Legacy Tier 1")].iloc[0]
    wh = s.loc[s["name"].str.startswith("Western Haynesville - Comstock")].iloc[0]
    rows = [
        ("Model breakeven, Legacy Tier 1 (pre-tax, 10%)", "$/MMBtu HH", f"=Curve!F5", "Expand deck: '<$2.75' (annual asset-level breakeven, hurdle not stated) - consistent with the model if Expand's basis is near 10% pre-tax; NOT the same definition."),
        ("Model breakeven, Western Haynesville (Comstock cost, Tier 1 productivity ASSUMED)", "$/MMBtu HH", f"=Curve!F{5 + 6}", "No sourced Western Haynesville EUR. Every +/-20% of EUR per foot moves this by about +/-0.7."),
        ("BEG 10% IRR breakeven, Tier 1 (2011 $, chart-read)", "$/MMBtu", 3.5, "OGJ 2015 Fig. 8; costs were ~$10.5 million per 4,800-ft well."),
        ("Modern Tier 1 EUR used (Expand 12-month cumulative / modelled year-1 share)", "Bcf per 1,000 ft", f"={ref['eur_anchor']}", "DERIVED (depends on the decline shape)."),
        ("BEG Tier 1 EUR per 1,000 ft (published 7.9 Bcf / 4.8)", "Bcf per 1,000 ft", 7.9 / 4.8, "2008-12 wells."),
        ("Comstock SEC proved undeveloped: 4.2 Tcf / 332 locations", "Bcf per location", 4.2e3 / 332, "10-K FY2025; SEC-conservative booking; lateral per location not stated (10,000 ft would be 1.27 Bcf per 1,000 ft)."),
        ("EIA AEO2023 average EUR, Haynesville-Bossier Louisiana / Texas", "Bcf per well", "10.5 / 9.2", "Table 3; lateral length not stated."),
        ("Comstock YTD 2026 IP per 1,000 ft: Legacy 31 MMcf/d at 12,052 ft; Western 31 MMcf/d at 10,331 ft", "MMcf/d per 1,000 ft", "2.57 / 3.00", "Western / Legacy = 1.17x (IP, not EUR; Western decline unknown)."),
        ("Comstock Legacy D&C vs Expand 2026E", "$/ft", "1,390 vs ~1,340", "Within 4%."),
        ("Developable locations, BEG legacy tiers 1-6 (this workbook)", "locations (4,800-ft)", f"=SUM(Curve!K5:K10)", "2012-vintage count; wells since 2012 not deducted."),
        ("Operator-disclosed legacy inventory: Comstock 1,705 gross (30 Jun 2026) + Expand >2,000 gross (31 Dec 2024)", "locations (~10,000 ft)", "3,705+", "Two operators only; these sit inside BEG's legacy area and are not added to the curve."),
        ("Comstock Western Haynesville operated inventory", "locations (8,875 ft)", 3277, "In the curve (Western Haynesville - Comstock row)."),
        ("Dallas Fed Q3 2026 expected Henry Hub in 2 / 5 years", "$/MMBtu", "3.82 / 4.28", "Survey of oil and gas executives, 30 Sep 2026."),
    ]
    for j, h in enumerate(["Check", "Unit", "Value", "Comment"], 1):
        ws.cell(3, j, h)
    style_head(ws, 3, 4)
    for i, (a, b, c, d) in enumerate(rows, 4):
        ws.cell(i, 1, a).alignment = WRAP
        ws.cell(i, 2, b)
        ws.cell(i, 3, c)
        ws.cell(i, 4, d).alignment = WRAP
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 90


def write_chart_data(ws, ref, hh, mv):
    ws["A1"] = "Chart data (live): step points of the cost curve for two horizons (edit H in the yellow cells) and the Henry Hub line; marginal price by Key-question year."
    ws["A1"].font = BOLD
    ws["A3"] = "Horizon A (years)"
    ws["B3"] = 5
    ws["A4"] = "Horizon B (years)"
    ws["B4"] = 8
    ws["B3"].fill = FILL_IN
    ws["B4"].fill = FILL_IN
    ws["A5"] = "Henry Hub, last 12 months average ($/MMBtu)"
    ws["B5"] = "='Henry Hub'!$B$3"
    heads = ["x A: cumulative capacity (Bcf/d)", "y A: breakeven", "x B: cumulative capacity (Bcf/d)", "y B: breakeven", "x HH", "y HH"]
    for j, h in enumerate(heads, 4):
        ws.cell(7, j, h)
    style_head(ws, 7, 9)
    ws["D6"] = '="Supply added by Dec "&(2025+B3)&" if drilled in "&B3&" years (cost only)"'
    ws["F6"] = '="Supply added by Dec "&(2025+B4)&" if drilled in "&B4&" years (cost only)"'
    ws["H6"] = "Henry Hub, last 12 months average"
    s0 = ref["s0"]
    for k in range(1, NT + 1):
        for part in (0, 1):
            r = 7 + 2 * (k - 1) + part + 1
            sr = s0 + k - 1
            for base, hcell in ((4, "$B$3"), (6, "$B$4")):
                colA = f"INDEX({ref['cum_block']},{k},MATCH({hcell},{ref['h_hdr']},0))"
                prevA = "0" if k == 1 else f"INDEX({ref['cum_block']},{k - 1},MATCH({hcell},{ref['h_hdr']},0))"
                ws.cell(r, base, f"={prevA}" if part == 0 else f"={colA}")
                ws.cell(r, base + 1, f"=Curve!$D${sr}")
    ws.cell(8, 8, 0)
    ws.cell(9, 8, f"=MAX(D{7 + 2 * NT},F{7 + 2 * NT})")
    ws.cell(8, 9, "=$B$5")
    ws.cell(9, 9, "=$B$5")
    # marginal price by year
    m0 = 7 + 2 * NT + 3
    ws.cell(m0 - 1, 1, "Marginal breakeven by Key-question date ($/MMBtu HH; STEO growth + extra supply)")
    ws.cell(m0 - 1, 1).font = BOLD
    for j, h in enumerate(["Year-end", "Base case", "LNG and takeaway delayed 6 months", "Henry Hub, last 12 months average"], 1):
        ws.cell(m0, j, h)
    style_head(ws, m0, 4)
    r1, r2 = ref["map_rows"]
    for i, yr in enumerate(("Dec 2026", "Dec 2028", "Dec 2030", "Dec 2033")):
        r = m0 + 1 + i
        ws.cell(r, 1, yr)
        ws.cell(r, 2, f"=IF(ISNUMBER(Mapping!{ref['map_pb']}{r1 + i}),Mapping!{ref['map_pb']}{r1 + i},0)")
        ws.cell(r, 3, f"=IF(ISNUMBER(Mapping!{ref['map_pb']}{r1 + 4 + i}),Mapping!{ref['map_pb']}{r1 + 4 + i},0)")
        ws.cell(r, 4, "=$B$5")
    ref["mchart"] = (m0, m0 + 4)
    ws.column_dimensions["A"].width = 44
    for j in range(2, 10):
        ws.column_dimensions[get_column_letter(j)].width = 20


def chart_layout():
    """Cell positions on 'Chart data' used by add_charts.py (native charts)."""
    m0 = 7 + 2 * NT + 3
    return {"nt": NT, "step_first": 8, "step_last": 7 + 2 * NT, "m0": m0, "banner": LABEL_BANNER}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()
    build(args.out)


if __name__ == "__main__":
    main()
