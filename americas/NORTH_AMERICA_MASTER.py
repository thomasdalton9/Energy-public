"""
North America master workbook: one file with every United States, Canada
and Mexico dataset this repo pulls, and Dashboard front pages carrying
all their charts - the same layout as the South & Central America master
(south_america/SOUTH_AMERICA_MASTER.py, whose table/chart/dashboard code
this reuses).

  Dashboard          - gas: US consumption by sector, production, trade,
                       Lower 48 storage (water-year chart), Henry Hub, LNG
                       feedgas; Canada supply and disposition; Mexico
                       imports from the US
  Dashboard - Power  - North America generation by source (sum of the three
                       countries) and installed capacity (EIA-860M, StatCan,
                       Ember for Mexico), then US Lower 48 and each ISO/RTO region
                       (EIA-930), Canada (StatCan, provincial grid operators, CER; fossil split
                       by province on StatCan's fuel mix, no Ember), Mexico (Ember fallback:
                       no raw CENACE generation feed yet) and Mexico demand
                       (CENACE)
  <CC> <chart> data  - the table each Dashboard chart plots
  <CC> <dataset> raw - the full data sheet(s) from each source workbook
  Sources            - where each dataset comes from, units and notes

Reads (doesn't refetch) the workbooks the scheduled pulls write to
"output/Data and Chart Outputs/". A missing input is listed on the
Dashboard and skipped rather than stopping the rest.

Usage: python3 NORTH_AMERICA_MASTER.py [--out "output/Data and Chart Outputs/Master Outputs/north_america_master.xlsx"]
"""
import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
import add_charts  # noqa: E402
import capacity_factors  # noqa: E402
import fundamentals  # noqa: E402
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR
FUELS = ["Hydro", "Gas", "Wind", "Solar", "Coal", "Nuclear", "Other"]   # same order/colours as the SA dashboards

# (country code, country, workbook, raw sheet / tuple of raw sheets / "*", short dataset name)
DATASETS = [
    ("US", "United States", "us_gas.xlsx", ("Demand by sector", "Supply and trade", "Storage weekly"), "gas"),
    ("US", "United States", "texas_gas_monthly.xlsx", ("Consumption by sector", "Exports", "Mexico by crossing", "Balance", "Forecast values"),
     "Texas gas"),
    ("US", "United States", "texas_production_forecast.xlsx", ("Supply and demand", "Demand to 2033"), "Texas supply and demand"),
    ("US", "United States", "gulf_coast_gas_balance.xlsx", ("Texas", "Louisiana", "Combined"), "Gulf Coast balance"),
    ("US", "United States", "henry_hub_daily.xlsx", "Data", "Henry Hub"),
    ("US", "United States", "lng_feedgas_daily.xlsx", "Best estimate daily", "LNG feedgas"),
    ("CA", "Canada", "canada_gas.xlsx", "Supply and disposition", "gas"),
    ("CA", "Canada", "canada_cer_gas.xlsx", "Gas trade monthly", "gas trade"),
    ("MX", "Mexico", "mexico_gas.xlsx", "Imports from US", "gas"),
    ("MX", "Mexico", "us_mexico_pipeline_capacity.xlsx", ("Export capacity by line", "Import capacity by line"),
     "pipeline capacity"),
]

EIA930_SHEETS = ("US_Total", "ERCOT", "PJM", "MISO", "SPP", "CAISO", "NYISO", "ISONE", "Southern", "TVA")
EMBER_FILES = [("NA", "North America", "north_america_power_by_type.xlsx")]   # Canada and Mexico
EMBER = {f for _, _, f in EMBER_FILES}
RAW_POWER_DATASETS = [
    ("US", "United States", "eia930_fuel_mix_daily.xlsx", EIA930_SHEETS, "power"),
    ("CA", "Canada", "canada_power_generation_daily.xlsx", ("Daily", "Provinces"), "power"),
]
OTHER_POWER_DATASETS = [
    ("CA", "Canada", "canada_provincial_power_daily.xlsx", ("Quebec", "Alberta", "British Columbia", "New Brunswick"),
     "provincial power"),
    ("CA", "Canada", "canada_cer_power.xlsx", ("Electricity trade monthly", "Fossil shares"), "electricity trade"),
    ("MX", "Mexico", "mexico_demanda_nacional_daily.xlsx", "Data", "demand"),
    ("US", "United States", "miso_gas_burn_daily.xlsx", "Monthly", "MISO gas burn"),
    ("US", "United States", "ercot_gas_burn_daily.xlsx", "Monthly", "ERCOT gas burn"),
]
# Installed generating capacity (standard sheet "Monthly"): US monthly, Canada and Mexico annual
CAPACITY_DATASETS = [
    ("US", "United States", "us_power_capacity.xlsx", "Monthly", "capacity"),
    ("CA", "Canada", "canada_power_capacity.xlsx", "Monthly", "capacity"),
    ("MX", "Mexico", "mexico_power_capacity.xlsx", "Monthly", "capacity"),
]
HYDRO_DATASETS = []
HYDRO_EXTRA = {}
# Workbooks whose dashboard shows only some of their charts (by spec name); the rest stay in the workbook
DASHBOARD_ONLY = {"gulf_coast_gas_balance.xlsx": {"Combined Gulf", "Supply growth vs LNG demand"},
                  "texas_production_forecast.xlsx": {"Supply and demand", "Net outflow", "Demand to 2033", "Outflow to 2033"},
                  "us_gas.xlsx": {"Demand", "Production", "Trade", "Storage"},
                  "henry_hub_daily.xlsx": {"Henry Hub"},
                  "miso_gas_burn_daily.xlsx": {"MISO gas burn"},
                  "ercot_gas_burn_daily.xlsx": {"MISO and ERCOT gas burn", "ERCOT gas burn"}}

REGIONS = {"US_Total": "US Lower 48", "ERCOT": "ERCOT (Texas)", "PJM": "PJM (Mid-Atlantic)",
           "MISO": "MISO (Midcontinent)", "SPP": "SPP (Southwest Power Pool)", "CAISO": "CAISO (California)",
           "NYISO": "NYISO (New York)", "ISONE": "ISO-NE (New England)", "Southern": "Southern Company",
           "TVA": "TVA (Tennessee Valley)"}


def complete_months(m, last_day):
    """Drop the month of `last_day` unless it is the month's last day (a daily feed's month in progress)."""
    if last_day < last_day + pd.offsets.MonthEnd(0):
        m = m[m.index < last_day.to_period("M").to_timestamp()]
    return m


def eia930_monthly(path, sheet):
    """EIA-930 daily MWh by fuel -> monthly GWh in the dashboard's fuel groups. Storage (battery, pumped) is left
    out: it shifts energy rather than generating it. EIA's renamed categories (e.g. Solar_Battery ->
    Solar_with_integrated_battery_storage, Nov 2024) never overlap, so they are summed."""
    d = add_charts.by_date(add_charts.read(path, sheet), "date")
    d = d[[c for c in d.columns if str(c).endswith("_MWh")]].apply(pd.to_numeric, errors="coerce")
    z = lambda pat: d[[c for c in d.columns if re.match(pat, c, re.I)]].sum(axis=1, min_count=1)  # noqa: E731
    g = pd.DataFrame({"Hydro": z(r"^Hydro_"), "Gas": z(r"^Natural_Gas_"), "Wind": z(r"^Wind"),
                      "Solar": z(r"^Solar"), "Coal": z(r"^Coal_"), "Nuclear": z(r"^Nuclear_"),
                      "Other": z(r"^(Petroleum|Geothermal|Other|Unknown)_MWh$")})
    m = g.resample("MS").sum(min_count=1) / 1000.0
    m = complete_months(m[m.index >= "2021-01-01"], d.index.max())
    return m.fillna(0)


def eia930(path):
    """EIA-930 workbook draws its own charts (REGISTRY entry is None); the master needs specs: US Lower 48 total in
    TWh, then each region in GWh."""
    out = []
    for sheet in EIA930_SHEETS:
        try:
            m = eia930_monthly(path, sheet)
        except ValueError:   # sheet not in this workbook
            continue
        us = sheet == "US_Total"
        out.append(add_charts.spec("Generation" if us else sheet, m / 1000.0 if us else m,
                                   f"{REGIONS[sheet]} power generation by source (EIA-930)",
                                   "TWh per month" if us else "GWh per month", "stacked_bar"))
    return out


FOSSIL_ROWS = ("Coal", "Gas", "Oil", "Other fossil")


def _fuel_shares(fuel, cer, province, year):
    """(coal, gas, other) share of a province's fossil generation in `year`: StatCan table 25-10-0084 (annual fuel mix)
    where it has the year, else the CER Energy Futures share for that year (its latest historical / projected year),
    else StatCan's last year. None when neither source knows the province."""
    def row(df, key):
        if df is None or df.empty:
            return None
        cols_ = [f"{province}|{g}" for g in key]
        if not all(c in df.columns for c in cols_):
            return None
        return df[cols_]
    sc = row(fuel, FOSSIL_ROWS)
    if sc is not None:
        have = sc[sc.sum(axis=1) > 0]
        if year in have.index.year:
            r = have[have.index.year == year].iloc[0]
            t = r.sum()
            return r["Coal"] / t, r["Gas"] / t, (r["Oil"] + r["Other fossil"]) / t
    ce = None
    if cer is not None and not cer.empty and all(f"{province}|{g}" in cer.columns for g in ("Coal", "Gas", "Oil")):
        ce = cer[[f"{province}|{g}" for g in ("Coal", "Gas", "Oil")]].dropna()
        ce.columns = ["Coal", "Gas", "Oil"]
        if year in ce.index:
            r = ce.loc[year]
            return r["Coal"], r["Gas"], r["Oil"]
    if sc is not None:
        have = sc[sc.sum(axis=1) > 0]
        if not have.empty:
            r = have.iloc[-1]
            t = r.sum()
            return r["Coal"] / t, r["Gas"] / t, (r["Oil"] + r["Other fossil"]) / t
    return None


def canada_monthly(path, cer_path=None):
    """StatCan monthly generation (GWh) in the dashboard's fuel groups. StatCan does not split fossil generation by
    fuel monthly, so each province's monthly fossil total (Provinces fossil sheet) is shared out into coal, gas and
    other fossil on that province's annual fuel mix: StatCan table 25-10-0084 where it has the year, else the CER
    Energy Futures share for the year (canada_cer_power.xlsx). Whatever the provinces do not add up to (suppressed
    provinces) and any province without a fuel mix goes to Other. No Ember data."""
    d = add_charts.by_date(add_charts.read(path, "Daily"), "date").apply(pd.to_numeric, errors="coerce") / 1000.0
    g = pd.DataFrame({f: d.get(f"{f}_MWh") for f in ("Hydro", "Wind", "Solar", "Nuclear")}, index=d.index)
    other = d[add_charts.cols(d, "Bioenergy_MWh", "Other_MWh")].sum(axis=1)
    fossil = d.get("Fossil_MWh", pd.Series(0.0, index=d.index)).fillna(0)
    gas, coal, oth = (pd.Series(0.0, index=d.index) for _ in range(3))
    try:
        pf = add_charts.by_date(add_charts.read(path, "Provinces fossil"), "date").apply(pd.to_numeric, errors="coerce") / 1000.0
    except Exception:  # noqa: BLE001 - workbook from before the provincial fossil sheet
        pf = pd.DataFrame(index=d.index)
    try:
        fuel = add_charts.by_date(add_charts.read(path, "Fuel mix annual"), "date").apply(pd.to_numeric, errors="coerce")
    except Exception:  # noqa: BLE001
        fuel = None
    try:
        cer = add_charts.read(cer_path, "Fossil shares").set_index("Year") if cer_path and os.path.exists(cer_path) else None
    except Exception:  # noqa: BLE001
        cer = None
    covered = pd.Series(0.0, index=d.index)
    for col in pf.columns:
        prov = str(col).replace("_MWh", "")
        series = pf[col].reindex(d.index).fillna(0)
        for year in sorted(set(d.index.year)):
            sh = _fuel_shares(fuel, cer, prov, year)
            if sh is None:
                continue
            m = d.index.year == year
            coal[m] += series[m] * sh[0]
            gas[m] += series[m] * sh[1]
            oth[m] += series[m] * sh[2]
            covered[m] += series[m]
    rest = (fossil - covered).clip(lower=0)
    g["Gas"], g["Coal"] = gas, coal
    g["Other"] = other + oth + rest
    return g[FUELS].fillna(0)


def canada_power_specs(path):
    """Canada generation by source in the dashboard's fuel groups (fossil split by province, see canada_monthly) plus
    StatCan's generation by province."""
    cer = os.path.join(os.path.dirname(path), "canada_cer_power.xlsx")
    m = canada_monthly(path, cer)
    m = m[m.index >= "2021-01-01"]
    specs = [add_charts.spec("Generation", m, "Canada power generation by source (StatCan; fossil split by province)",
                             "GWh per month", "stacked_bar")]
    specs += [s for s in add_charts.canada_power(path) if s["name"] == "Provinces"]
    return specs


def henry_hub(path):
    specs = add_charts.henry_hub(path)
    for s in specs:
        if s["name"] == "Daily":
            s["name"] = "Henry Hub"
    return specs


def lng_feedgas(path):
    """lng_feedgas_daily.xlsx draws its own charts; for the master: monthly average feedgas by plant."""
    d = add_charts.by_date(add_charts.read(path, "Best estimate daily"), "gas_day")
    plants = [c for c in d.columns if c not in ("Total", "Source") and pd.api.types.is_numeric_dtype(d[c])]
    m = add_charts.monthly_mean(d[plants], "2021-01-01")
    m = complete_months(m, d.index.max())
    specs = [add_charts.spec("LNG feedgas", m, "US LNG feedgas by plant (pipeline nominations; EIA before metering)",
                             "Bcf/d, monthly average", "stacked_bar")]
    return specs + lng_vs_nameplate(path)


def lng_vs_nameplate(path):
    """US LNG feedgas against nameplate feedgas capacity: nameplate by plant stacked (each train counts from its FERC
    feed-gas date; trains still to start are the workbook's expected dates and are drawn in lighter bars), total feedgas
    (best estimate) as a line. Same data as the 'Feedgas vs capacity' chart of lng_feedgas_daily.xlsx, as monthly means."""
    c = add_charts.by_date(add_charts.read(path, "Chart data"), "gas_day")
    feed_col = "Total feedgas"
    if feed_col not in c.columns:
        return []
    caps = [x for x in c.columns if x != feed_col and pd.api.types.is_numeric_dtype(c[x])]
    last = c[feed_col].dropna().index.max()
    cap = add_charts.monthly_mean(c[caps], "2021-01-01")
    feed = complete_months(add_charts.monthly_mean(c[[feed_col]], "2021-01-01"), last)
    df = cap.join(feed, how="left")
    sp = add_charts.spec("LNG vs nameplate", df, "US feedgas vs nameplate",
                         "Bcf/d, monthly average", "stacked_bar", line_cols=(feed_col,))
    sp["forecast_from"] = (last + pd.offsets.MonthBegin(1)).normalize() if last.day != 1 else last
    return [sp]


def mexico_demand(path):
    specs = add_charts.mexico(path)
    for s in specs:
        s["df"] = s["df"][s["df"].index >= "2021-01-01"]
    return specs


# Chart specs for workbooks whose add_charts REGISTRY entry is None or that the dashboard shows differently
MASTER_SPECS = {"canada_power_generation_daily.xlsx": canada_power_specs, "eia930_fuel_mix_daily.xlsx": eia930, "henry_hub_daily.xlsx": henry_hub,
                "lng_feedgas_daily.xlsx": lng_feedgas, "mexico_demanda_nacional_daily.xlsx": mexico_demand,
                # combined MISO + ERCOT chart (built from the ERCOT workbook's folder) plus the ERCOT-only chart
                "ercot_gas_burn_daily.xlsx": lambda p: add_charts.miso_ercot_gas_burn(p) + add_charts.ercot_gas_burn(p)}

EIA_GAS = "https://www.eia.gov/naturalgas/data.php"
SOURCES = {
    "us_gas.xlsx": ("EIA (US Energy Information Administration): Natural Gas Monthly (consumption by sector, "
                    "production, trade) and Weekly Natural Gas Storage Report", EIA_GAS),
    "texas_gas_monthly.xlsx": ("EIA Natural Gas Monthly, Texas (state consumption by sector, marketed production, storage; "
                               "exports by port of exit: pipeline to Mexico at Texas crossings and LNG from Corpus Christi, "
                               "Freeport and Golden Pass; Sabine Pass is Louisiana) via EIA API v2; ~2-3 month lag. Forecast to Dec 2028 (lighter bars): "
                               "LNG feedgas = EIA exports x 1.09, existing plants at nameplate x 3-year utilisation, new trains from company-announced start dates "
                               "(NextDecade, Sempra, ExxonMobil, Cheniere; several unverified), other sectors a seasonal-trend base",
                               "https://www.eia.gov/dnav/ng/ng_cons_sum_dcu_STX_m.htm"),
    "texas_production_forecast.xlsx": ("Texas supply and demand to Dec 2028: demand = EIA consumption by sector, Mexico pipeline exports and LNG feedgas "
                                       "(EIA exports x 1.09; forecast from texas_gas_monthly.xlsx); production = EIA dry gas (history) and a forecast from "
                                       "EIA STEO regional marketed production (Permian, Eagle Ford, Haynesville) capped by Permian takeaway capacity "
                                       "(company-announced pipelines, unverified), less NGL extraction loss; outflow = production - demand (own calculation). "
                                       "To Dec 2033 (scenario after Dec 2028, lightest bars): data centres = IEA 'Energy and AI' (Apr 2025) US data-centre electricity (183 TWh 2024, 426 TWh 2030 Base Case; LOW / HIGH = IEA Headwinds / Lift-Off world ratios; years between interpolated) x a Texas share of US growth derived from ERCOT's own data-centre forecast (CDR Dec 2025, 22.2 GW for summer 2030) and the IEA US capacity addition, converted to gas at the calibrated ERCOT heat rate, the IEA US load factor and ERCOT's observed gas share of generation (EIA-930, a proxy for the marginal share); each input is labelled SOURCED / DERIVED / PROXY / ASSUMPTION on the 'Assump - Data centres' tab (cross-checks: LBNL 2024 report, EIA AEO2026, ERCOT large-load queue); LNG train table held at steady utilisation; production extension damped and capped by takeaway (own assumptions, not STEO)",
                                       "https://www.eia.gov/outlooks/steo/data/browser/#/?v=6&f=M&s=0&start=202401&end=202812&id=&linechart=NGMPPM"),
    "gulf_coast_gas_balance.xlsx": ("Gulf Coast (Texas + Louisiana) gas balance: EIA Natural Gas Monthly via API v2 (Louisiana consumption, production, LNG exports of Sabine Pass, "
                                    "Cameron, Calcasieu Pass and Plaquemines; feedgas = exports x 1.09), the Texas workbooks (texas_gas_monthly, texas_production_forecast), EIA STEO "
                                    "Permian / Eagle Ford / Haynesville marketed production (extension after Dec 2027, 2029-30 scenario), Permian takeaway table (unverified) and the "
                                    "EIA U.S. liquefaction capacity file 2026 Q2 for new Louisiana LNG trains (year/half-year sourced, month assumed; pre-FID projects not forecast); "
                                    "outflow and extra-supply figures are own calculations",
                                    "https://www.eia.gov/naturalgas/importsexports/liquefactioncapacity/U.S.liquefactioncapacity_2026_Q2.xlsx"),
    "henry_hub_daily.xlsx": ("EIA, Henry Hub natural gas spot price (RNGWHHD); front-month futures: NYMEX via Yahoo Finance (NG=F)", "https://www.eia.gov/dnav/ng/hist/rngwhhdd.htm"),
    "lng_feedgas_daily.xlsx": ("Interstate pipeline operators' scheduled quantities at each LNG plant (Kinder Morgan, "
                               "Enbridge, Williams, Energy Transfer, Cheniere, ...); EIA monthly LNG exports before "
                               "the daily pull", "https://www.eia.gov/dnav/ng/ng_move_poe2_a_EPG0_ENG_Mmcf_m.htm"),
    "canada_gas.xlsx": ("Statistics Canada, Table 25-10-0055-01 Supply and disposition of natural gas, monthly",
                        "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=2510005501"),
    "mexico_gas.xlsx": ("EIA, US natural gas exports to Mexico (pipeline N9132MX2, LNG N9133MX2) - Mexico's "
                        "imports from the US as measured on the US side", "https://www.eia.gov/dnav/ng/ng_move_expc_s1_m.htm"),
    "us_mexico_pipeline_capacity.xlsx": ("EIA, Natural Gas Pipelines: State to State Capacity (annual capacity by line, "
                                         "US to Mexico and Mexico to US, MMcf/d)", EIA_GAS),
    "eia930_fuel_mix_daily.xlsx": ("EIA-930 Hourly Electric Grid Monitor (balancing authority data, Lower 48)",
                                   "https://www.eia.gov/electricity/gridmonitor/"),
    "miso_gas_burn_daily.xlsx": ("MISO real-time generation fuel mix (gas MWh) x heat rate calibrated on EIA-923 fuel use "
                                 "of MISO balancing-authority gas plants (estimate; EIA-923 months not yet published "
                                 "carry last year's heat rate)", "https://www.eia.gov/electricity/data/eia923/"),
    "ercot_gas_burn_daily.xlsx": ("MISO real-time fuel mix (gas MWh) and ERCOT gas MWh (EIA-930) x heat rates calibrated on EIA-923 "
                                  "fuel use of the MISO and ERCO balancing-authority gas plants (estimates, not metered burn; "
                                  "EIA-923 months not yet published carry last year's heat rate)",
                                  "https://www.eia.gov/electricity/data/eia923/"),
    "canada_power_generation_daily.xlsx": ("Statistics Canada, Table 25-10-0015-01 Electric power generation, "
                                           "monthly generation by type of electricity; fossil generation split by fuel with "
                                           "Table 25-10-0084-01 (annual electricity generated by fuel, by province) and, for "
                                           "years it does not cover yet, CER Canada's Energy Future 2026 provincial shares",
                                           "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=2510001501"),
    "canada_provincial_power_daily.xlsx": ("Provincial grid operators: Hydro-Quebec open data (generation by source), "
                                           "AESO Current Supply and Demand report (Alberta), BC Hydro balancing-authority "
                                           "load, NB Power system information archive (New Brunswick load)",
                                           "https://donnees.hydroquebec.com/explore/dataset/historique-production-electricite-quebec/"),
    "canada_cer_gas.xlsx": ("Canada Energy Regulator, natural gas exports and imports by month (open data)",
                            "https://www.cer-rec.gc.ca/en/data-analysis/energy-commodities/natural-gas/index.html"),
    "canada_cer_power.xlsx": ("Canada Energy Regulator, electricity exports and imports by month (open data); "
                              "Canada's Energy Future 2026 generation by province (fuel shares)",
                              "https://www.cer-rec.gc.ca/en/data-analysis/energy-commodities/electricity/index.html"),
    "mexico_demanda_nacional_daily.xlsx": ("CENACE (Centro Nacional de Control de Energía), Demanda Real del Sistema",
                                           "https://www.cenace.gob.mx/Paginas/SIM/Reportes/EstimacionDemandaReal.aspx"),
    "us_power_capacity.xlsx": ("EIA-860M monthly generator inventory (net summer capacity), via EIA API v2",
                               "https://www.eia.gov/electricity/data/eia860m/"),
    "canada_power_capacity.xlsx": ("Statistics Canada, Tables 25-10-0022-01 (capacity by type) and 25-10-0023-01 "
                                   "(thermal capacity by principal fuel)",
                                   "https://www150.statcan.gc.ca/t1/tbl1/en/tv.action?pid=2510002201"),
    "mexico_power_capacity.xlsx": ("Ember yearly electricity data - FALLBACK, no raw Mexican capacity feed yet",
                                   "https://ember-energy.org/data/yearly-electricity-data/"),
    **{f: ("Ember monthly electricity data", "https://ember-energy.org/data/monthly-electricity-data/") for f in EMBER},
}
# The grid operator / statistics office Ember compiles each country from (named on Ember-fed charts)
OPERATORS = {"Mexico": "CENACE / SENER"}

NA_POWER_COUNTRIES = ["United States", "Canada", "Mexico"]


def north_america_generation(data_dir, have_raw, frames_out=None):
    """North America generation by source, TWh per month: US Lower 48 (EIA-930) + Canada (StatCan, or Ember) +
    Mexico (Ember). Runs to the last month every raw-fed country has in full; an Ember-fed country that lags is left
    out of the months it doesn't have yet, and those gaps are listed rather than estimated."""
    ember = os.path.join(data_dir, EMBER_FILES[0][2])
    frames, notes = {}, []
    for country in NA_POWER_COUNTRIES:
        try:
            if country == "United States":
                fname = "eia930_fuel_mix_daily.xlsx"
                m = eia930_monthly(os.path.join(data_dir, fname), "US_Total")
                src = SOURCES[fname][0] + " (Alaska and Hawaii not included)"
            elif country in have_raw:
                fname = "canada_power_generation_daily.xlsx"
                m = canada_monthly(os.path.join(data_dir, fname), os.path.join(data_dir, "canada_cer_power.xlsx"))
                src = (SOURCES[fname][0] + " (fossil split into gas / coal / other by province on StatCan table 25-10-0084's "
                       "annual fuel mix, CER Energy Futures shares for later years - estimate)")
            else:
                m = add_charts.power_mix(add_charts.by_date(add_charts.read(ember, country), "Month"))
                src = f"Ember (compiled from {OPERATORS.get(country, 'the grid operator')})"
            m = m.apply(pd.to_numeric, errors="coerce")
            frames[country] = m
            notes.append(f"{country}: {src}, {m.index.min():%b/%y}-{m.index.max():%b/%y}")
        except Exception as e:  # noqa: BLE001
            notes.append(f"{country}: not available ({type(e).__name__}: {e})")
    if frames_out is not None:
        frames_out.update(frames)
    if not frames:
        return pd.DataFrame(), notes
    have = {c: set(f.dropna(how="all").index) for c, f in frames.items()}
    raw_have = [h for c, h in have.items() if c == "United States" or c in have_raw] or list(have.values())
    end = min(max(h) for h in raw_have)
    months = sorted(m for m in set.intersection(*raw_have) if m <= end)
    total = sum(f.reindex(index=months, columns=FUELS).fillna(0) for f in frames.values()) / 1000.0   # GWh -> TWh
    total.index.name = "date"
    for c, h in have.items():
        gap = [m for m in months if m not in h]
        if gap:
            notes.append(f"NOT INCLUDED: {c} in {gap[0]:%b/%y}-{gap[-1]:%b/%y} ({len(gap)} months, no data yet)")
    return total, notes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "Master Outputs", "north_america_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()
    cfg = sys.modules[__name__]

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    dash2 = wb.create_sheet("Dashboard - Power")
    used = {"Dashboard", "Dashboard - Power", "Sources"}
    sources = []

    gas = sam.collect(wb, DATASETS, args.data_dir, used, sources, cfg=cfg)
    ember_countries = set()
    for _, _, f in EMBER_FILES:
        try:
            ember_countries |= set(pd.ExcelFile(os.path.join(args.data_dir, f)).sheet_names)
        except Exception:  # noqa: BLE001
            pass
    raw_power, building = [], []
    for d in RAW_POWER_DATASETS:
        path = os.path.join(args.data_dir, d[2])
        if not os.path.exists(path):
            continue
        days = sam.raw_history_days(path) if d[2] != "eia930_fuel_mix_daily.xlsx" else sam.MIN_RAW_DAYS
        if days >= sam.MIN_RAW_DAYS or d[1] not in ember_countries or d[1] == "Canada":   # Canada: no Ember fallback
            raw_power.append(d)
        else:
            building.append(f"{d[1]} raw feed has {days} days so far - Ember shown until it has {sam.MIN_RAW_DAYS}")
    have_raw = {d[1] for d in raw_power}
    print(f"power by type: raw data for {sorted(have_raw) or 'none'}; Ember for the rest")
    parts = [sam.collect(wb, raw_power, args.data_dir, used, sources, cfg=cfg),
             *(sam.collect(wb, [(code, region, f, "*", "power")], args.data_dir, used, sources, skip=have_raw, cfg=cfg)
               for code, region, f in EMBER_FILES),
             sam.collect(wb, OTHER_POWER_DATASETS, args.data_dir, used, sources, cfg=cfg)]
    power = tuple(sum((p[i] for p in parts), []) for i in range(3))
    power[2].extend(building)

    gen_frames = {}
    total, notes = north_america_generation(args.data_dir, have_raw, gen_frames)
    cap_total, cap_notes = sam.south_america_capacity(args.data_dir, cfg=cfg)
    CAP_LINE = "Total installed capacity (TWh at full output)"
    if not total.empty and not cap_total.empty:
        # Overlay line on the generation chart: installed GW x hours in the month, so it shares the TWh axis
        idx = pd.to_datetime(total.index)
        gw = cap_total.sum(axis=1)
        gw.index = pd.to_datetime(gw.index)
        gw = gw.reindex(gw.index.union(idx)).sort_index().ffill().reindex(idx)
        total = total.copy()
        total[CAP_LINE] = (gw * idx.days_in_month * 24 / 1000.0).values
        notes = notes + [f"{CAP_LINE}: installed capacity (storage excluded) x hours in the month; last capacity "
                         f"month ({cap_total.index.max():%b/%y}) held for later months"]
    if not total.empty:
        ws = wb.create_sheet(sam.sheet_name("NA generation total data", used))
        total, gen_units = xlsx_charts.monthly_energy_to_gw(total, "TWh per month", "power generation")
        cap_line = CAP_LINE.replace("(TWh at full output)", "(GW)")
        df, n_bars = xlsx_charts.prepare(total, (cap_line,))
        xlsx_charts.write_table(ws, df)
        ws.cell(row=1, column=df.shape[1] + 4, value="Countries summed (only months all of them have):")
        for i, note in enumerate(notes, start=2):
            ws.cell(row=i, column=df.shape[1] + 4, value=note)
        src = ("Sum of the country series on this dashboard (EIA-930, StatCan + CER; Ember for Mexico)", None)
        power[0].insert(0, (xlsx_charts.build_chart(ws, df, n_bars, "North America power generation by source",
                                                    gen_units, "stacked_bar", width=sam.CHART_W,
                                                    height=sam.CHART_H, gridlines=False,
                                                    inner=xlsx_charts.DASHBOARD_INNER), src))
        missing = [n.split(":", 1)[1].split(" in ")[0].strip() for n in notes if n.startswith("NOT INCLUDED")]
        name = "North America power generation by source (US Lower 48, Canada, Mexico" + (
            f"; {', '.join(missing)} missing in latest months)" if missing else ")")
        power[1].insert(0, ("North America", name, df.index.max().strftime("%b/%y"), ws.title, *src))
        print("North America generation total:", "; ".join(notes))

    # Capacity: combined North America chart (GW) right after the generation total, then each country
    cap = sam.collect(wb, [d for d in CAPACITY_DATASETS if os.path.exists(os.path.join(args.data_dir, d[2]))],
                      args.data_dir, used, sources, cfg=cfg)
    pos = 1 if not total.empty else 0
    if not cap_total.empty:
        ws = wb.create_sheet(sam.sheet_name("NA capacity total data", used))
        df, n_bars = xlsx_charts.prepare(cap_total)
        xlsx_charts.write_table(ws, df)
        ws.cell(row=1, column=df.shape[1] + 4, value="Countries summed (annual figures carried forward; storage "
                                                      "excluded):")
        for i, note in enumerate(cap_notes, start=2):
            ws.cell(row=i, column=df.shape[1] + 4, value=note)
        src = ("Sum of the country capacity workbooks (EIA-860M, StatCan; Ember for Mexico)", None)
        power[0].insert(pos, (xlsx_charts.build_chart(ws, df, n_bars, "North America installed generating capacity",
                                                      "GW installed", "stacked_bar", width=sam.CHART_W,
                                                      height=sam.CHART_H, gridlines=False,
                                                      inner=xlsx_charts.DASHBOARD_INNER), src))
        missing = [n.split(":", 1)[1].split("(")[0].strip() for n in cap_notes if n.startswith("NOT INCLUDED")]
        power[1].insert(pos, ("North America", "North America installed generating capacity" +
                              (f" (missing: {', '.join(missing)})" if missing else ""),
                              df.index.max().strftime("%b/%y"), ws.title, *src))
        pos += 1
        print("North America capacity total:", "; ".join(cap_notes))
    power[0][pos:pos] = cap[0]
    power[1][pos:pos] = cap[1]
    power[2].extend(cap[2])
    # One regional capacity-factor chart (generation by type / capacity x hours), after the capacity charts
    cf_chart, cf_row, cf_missing, cf_countries = capacity_factors.add_capacity_factor_sheets(
        wb, used, sam.sheet_name, gen_frames,
        {c: os.path.join(args.data_dir, f) for _, c, f, _, _ in CAPACITY_DATASETS}, sam.CHART_W, sam.CHART_H,
        "Country generation (EIA-930 Lower 48, StatCan, Ember for Mexico) / capacity (EIA-860M, StatCan, Ember)",
        "North America")
    if cf_chart:
        power[0].insert(pos + len(cap[0]), cf_chart)
        power[1].insert(pos + len(cap[1]), cf_row)
    power[2].extend(cf_missing)

    sam.draw_dashboard(dash, "North America energy (US, Canada, Mexico) - gas dashboard", *gas)
    sam.draw_dashboard(dash2, "North America energy (US, Canada, Mexico) - power generation", *power)
    cf_dash = wb.create_sheet("Dashboard - Capacity factors", 2)
    used.add("Dashboard - Capacity factors")
    cf_items = cf_countries   # regional chart first, then each country
    sam.draw_dashboard(cf_dash, "North America - capacity factor by generation type (generation / capacity x hours)", [c for c, _ in cf_items], [r for _, r in cf_items], cf_missing)

    # long-term fundamentals: annual history (Ember yearly, EI Statistical Review, World Bank, IMF, degree days)
    fundamentals.add_long_term_dashboard(wb, used, sam.sheet_name, "North America", args.data_dir, sam.CHART_W, sam.CHART_H,
                                         rows_per_chart=sam.ROWS_PER_CHART, degree_days=("CDD_18", "HDD_18"),
                                         index=sum(s.startswith("Dashboard") for s in wb.sheetnames))

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    for s in sources:
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    for col, w in (("A", 14), ("B", 14), ("C", 38), ("D", 40), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: {len(gas[0])} gas charts, {len(power[0])} power charts; tabs {wb.sheetnames}")
    for label, m in (("gas", gas[2]), ("power", power[2])):
        if m:
            print(f"missing ({label}):", m)


if __name__ == "__main__":
    main()
