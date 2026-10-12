"""
East Mediterranean master workbook: Israel, Egypt, Cyprus (plus Jordan and Lebanon on the long-term page) - every
dataset this repo pulls for the region from a raw source, with Dashboard front pages carrying all their charts. Same
layout as the other masters (south_america/SOUTH_AMERICA_MASTER.py, whose table/chart/dashboard code this reuses).

  Dashboard - Gas        Israel gas sales by field (Leviathan, Tamar, Karish; company reports, quarterly), Leviathan
                         sales by destination, exports to Egypt and Jordan; annual balance from the Natural Gas
                         Authority reviews (supply by source, domestic use by sector, exports by country, power by
                         producer); Egypt monthly balance from JODI (fallback) with implied net LNG, CAPMAS/MOEE annual
                         Jordan gas supply vs burn (Leviathan sales to Jordan, NewMed, quarterly, beside MEMR's annual
                         gas consumed for power; Egyptian gas, Aqaba LNG and Risha volumes are unpublished)
  Dashboard - Power      Egypt daily peak / minimum grid load (Ministry of Electricity, from Oct 2026), Cyprus
                         generation mix (TSOC), Jordan daily peak load and monthly maximum (NEPCO, from Oct 2019),
                         Jordan annual statistics (NEPCO / MEMR annual reports: energy by fuel, fuel burned, gas for
                         power, imports / exports, peak and minimum load)
  Dashboard - Long-term  annual fundamentals (fundamentals.py region 'East Med': Ember yearly power, Energy Institute
                         gas, World Bank / IMF, degree days)
  <CC> <chart> data      the table each Dashboard chart plots;  <CC> <dataset> raw: the source workbook's sheets
  Sources                where each dataset comes from, units and notes

Egypt's grid load is also on the Africa master and Cyprus on the Europe master; the workbooks are only read here.
Reads (doesn't refetch) the workbooks the scheduled pulls write to "output/Data and Chart Outputs/". A missing input
is listed on the Dashboard and skipped rather than stopping the rest.

TO ADD A DATASET: append (country code, country, workbook, raw sheet / tuple / "*", short name) to GAS_DATASETS or
POWER_DATASETS, add a SOURCES entry and register the workbook's charts in add_charts.py REGISTRY.

Usage: python3 middle_east/EAST_MED_MASTER.py [--out "output/Data and Chart Outputs/Master Outputs/east_med_master.xlsx"]
"""
import argparse
import os
import sys
import types

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
import fundamentals  # noqa: E402
import pandas as pd  # noqa: E402
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR

GAS_DATASETS = [
    ("IL", "Israel", "israel_gas_fields_quarterly.xlsx", ("Quarterly", "Leviathan", "Tamar", "Karish YTD"), "gas by field"),
    ("IL", "Israel", "israel_gas_annual_nga.xlsx", ("Annual", "Checks"), "annual gas balance"),
    ("EG", "Egypt", "egypt_gas_monthly.xlsx", ("Monthly JODI (fallback)", "JODI (fallback) all rows", "Annual official",
                                               "Checks", "Checks FY"), "gas balance (JODI fallback)"),
    ("JO", "Jordan", "jordan_power_annual.xlsx", ("MEMR sentences",), "gas supply vs burn"),
]
POWER_DATASETS = [
    ("EG", "Egypt", "egypt_grid_load_daily.xlsx", "Daily", "grid load"),
    ("CY", "Cyprus", "cyprus_generation_mix_daily.xlsx", "*", "power"),
    ("JO", "Jordan", "jordan_peak_load_daily.xlsx", ("Daily", "Monthly"), "peak load"),
    ("JO", "Jordan", "jordan_power_annual.xlsx", ("Annual", "Checks", "Daily vs annual peak"), "annual power"),
]
SOURCES = {
    "israel_gas_fields_quarterly.xlsx": (
        "Field owners' reports: NewMed Energy (Leviathan), Tamar Petroleum (Tamar), Energean plc via investegate.co.uk "
        "(Karish) - gas sold, 100% of each field, quarterly", "https://newmedenergy.com/investor-relations/financial-reports/"),
    "israel_gas_annual_nga.xlsx": (
        "Israel Natural Gas Authority (Ministry of Energy and Infrastructure): annual 'Natural Gas Sector Review' "
        "(Hebrew PDF, gov.il BlobFolder) - supply by source, domestic use by sector, exports by country, BCM per year, "
        "read from the reviews' tables and chart data labels", "https://www.gov.il/BlobFolder/reports/ng-2025/he/ng-2025.pdf"),
    "egypt_gas_monthly.xlsx": (
        "JODI Gas World database (FALLBACK: Egypt's own JODI submission, no official monthly Egyptian gas statistics are "
        "reachable; LNG IMPLIED from the balance, not JODI's reported LNG) + official annual layer: CAPMAS (energy "
        "balance, petroleum-sector statistics) and Ministry of Electricity (MOEE, gas used in power stations, fiscal years)",
        "http://www.jodidb.org/TableViewer/tableView.aspx?ReportId=38673"),
    "egypt_grid_load_daily.xlsx": (
        "Ministry of Electricity and Renewable Energy (MOEE), Egypt: home-page 'Grid Loads' panel (no history at source; "
        "series from Oct 2026)", "http://www.moee.gov.eg/english_new/home.aspx"),
    "cyprus_generation_mix_daily.xlsx": ("Transmission System Operator Cyprus (TSOC)", "https://tsoc.org.cy/"),
    "jordan_peak_load_daily.xlsx": (
        "National Electric Power Company (NEPCO), Jordan: Max Load Archive (daily maximum load as posted; unposted days "
        "are gaps)", "https://www.nepco.com.jo/en/MaxLoadArchive.aspx"),
    "jordan_power_annual.xlsx": (
        "National Electric Power Company (NEPCO) annual reports (English 2003-2024; 2025 from the Arabic report, used "
        "only where its tables agree with the English reports) and Ministry of Energy and Mineral Resources (MEMR) "
        "annual reports (gas consumed for power, from the text)", "https://www.nepco.com.jo/en/AnnualReports.aspx"),
}
# the Jordan gas chart on the gas Dashboard has its own source line (the workbook's source above is for the power charts)
JORDAN_GAS_SOURCE = (
    "NewMed Energy (Leviathan sales to Jordan, quarterly, israel_gas_fields_quarterly.xlsx) and MEMR annual reports "
    "(natural gas consumed for power, annual average shown in each quarter of its year). Egyptian pipeline gas, Aqaba "
    "LNG and Risha: no annual volumes by source are published (MEMR states Risha only for 2016-17, 2019 and 2022); "
    "no residual is computed", "https://www.memr.gov.jo/En/List/Annual_Reports")


def jordan_gas_supply(path):
    """Jordan gas supply beside burn: Leviathan sales to Jordan (NewMed, quarterly bcm -> Bcf/d = bcm x 35.3147 / days
    in the quarter) as bars, MEMR's annual gas consumed for power (Bcf/d) as a line over the quarters of its year.
    Supply from Egypt, Aqaba LNG and Risha has no published volumes, so nothing is netted or inferred."""
    q = pd.read_excel(os.path.join(os.path.dirname(path), "israel_gas_fields_quarterly.xlsx"), sheet_name="Quarterly")
    q["date"] = pd.to_datetime(q["date"])
    q = q.set_index("date").sort_index()
    lev = pd.to_numeric(q["Leviathan_Jordan_bcm"], errors="coerce").dropna()
    days = (lev.index + pd.offsets.QuarterBegin(startingMonth=1) - lev.index).days
    lev_bcfd = lev * 35.3147 / days
    a = pd.read_excel(path, sheet_name="Annual")
    memr = pd.to_numeric(a.set_index("Year")["MEMR_Gas_power_Bcf_d"], errors="coerce").dropna()
    idx = pd.date_range(lev.index.min(), lev.index.max(), freq="QS")   # from Leviathan's first sales to Jordan
    df = pd.DataFrame(index=idx)
    df["Leviathan sales to Jordan (NewMed)"] = lev_bcfd.reindex(idx)
    df["Gas consumed for power (MEMR, annual average)"] = [memr.get(d.year) for d in idx]
    df = df.apply(pd.to_numeric, errors="coerce").dropna(how="all")
    return [{"name": "Jordan gas supply vs power burn", "df": df, "kind": "stacked_bar", "date_format": "%Y-%m",
             "title": "Jordan gas: Leviathan sales to Jordan vs gas burned for power (other sources unpublished)",
             "units": "Bcf/d", "line_cols": ("Gas consumed for power (MEMR, annual average)",)}]

# settings sam.collect reads from the master module it is given
MASTER_SPECS = {}
HYDRO_DATASETS = []
HYDRO_EXTRA = {}
DASHBOARD_ONLY = {}
EMBER = set()
OPERATORS = {}
CAVEATS = [
    "Israel: quarterly SALES from the field owners' reports (100% of field); the national total is only for quarters "
    "all three fields report; no monthly Israeli gas source is reachable (gov.il is Cloudflare-blocked). Annual balance "
    "from the Natural Gas Authority reviews (2004-latest; segments a review prints no label for are drawn as "
    "'Not labelled in the review')",
    "Egypt gas: MONTHLY from JODI (fallback - Egypt's own JODI submission; no official monthly source is reachable); "
    "net LNG is IMPLIED = consumption (gross inland deliveries) + pipeline exports - production - pipeline imports "
    "(positive = net imports), not JODI's reported LNG rows; it carries JODI's statistical difference (large in 2011-13). "
    "Official annual layer: CAPMAS (to 2022; units as the API states them) and MOEE gas for power (fiscal years to 2022/23)",
    "Jordan gas: Leviathan sales to Jordan (NewMed, quarterly) beside MEMR's gas consumed for power (annual average, "
    "drawn over each quarter of its year; 2016 = LNG only as MEMR words it; 2018 and 2025 not published). Egyptian "
    "pipeline gas, Aqaba LNG and Risha have no published annual volumes by source, so the difference is NOT supply "
    "from them and no residual is computed",
    "Egypt power: grid load only (daily peak and minimum); generation by type is annual (MOEE) and not yet built",
    "Jordan power: NEPCO daily peak load only, as posted (often 14-27 of 31 days a month in 2021-24; unposted days are "
    "gaps, implausible figures are blanked and flagged on the raw tab); the monthly maximum is charted only for months "
    "with at least 20 days posted",
    "Jordan annual (NEPCO annual reports): energy PURCHASED by NEPCO by fuel (sent-out, 2015 on) and the older national "
    "production basis (2007-2017) are separate charts, never mixed; 2025 is from the Arabic report (no English one yet), "
    "used only where its tables agree with the English reports; fuel burned in thousand toe as published; peak load "
    "before 2013 from older tables (interconnected system, basis varies). The daily archive is dated one day late from "
    "2023 and showed the daytime peak before then (see the workbook's 'Daily vs annual peak' tab)",
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "Master Outputs", "east_med_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()
    cfg = sys.modules[__name__]

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard - Gas"
    dash2 = wb.create_sheet("Dashboard - Power")
    used = {"Dashboard - Gas", "Dashboard - Power", "Dashboard - Long-term", "Sources"}
    sources = []

    # gas Dashboard: the Jordan annual workbook gives only the supply-vs-burn chart, with its own source line
    gas_cfg = types.SimpleNamespace(**{k: getattr(cfg, k) for k in ("HYDRO_DATASETS", "HYDRO_EXTRA", "DASHBOARD_ONLY",
                                                                      "EMBER", "OPERATORS")})
    gas_cfg.MASTER_SPECS = {"jordan_power_annual.xlsx": jordan_gas_supply}
    gas_cfg.SOURCES = {**SOURCES, "jordan_power_annual.xlsx": JORDAN_GAS_SOURCE}
    gas = sam.collect(wb, GAS_DATASETS, args.data_dir, used, sources, cfg=gas_cfg)
    power = sam.collect(wb, POWER_DATASETS, args.data_dir, used, sources, cfg=cfg)
    gas[2].extend(CAVEATS[:3])
    power[2].extend(CAVEATS[3:])
    sam.draw_dashboard(dash, "East Mediterranean - gas (Israel by field, exports to Egypt and Jordan; Egypt balance, JODI fallback; Jordan supply vs burn)", *gas)
    sam.draw_dashboard(dash2, "East Mediterranean - power (Egypt grid load, Cyprus generation, Jordan peak load and "
                              "annual statistics)", *power)

    notes = fundamentals.add_long_term_dashboard(wb, used, sam.sheet_name, "East Med", args.data_dir, sam.CHART_W,
                                                 sam.CHART_H, rows_per_chart=sam.ROWS_PER_CHART,
                                                 degree_days=("CDD_18", "HDD_18"),
                                                 index=sum(s.startswith("Dashboard") for s in wb.sheetnames))

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    for s in sources:
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    for col, w in (("A", 14), ("B", 24), ("C", 38), ("D", 40), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: {len(gas[0])} gas charts, {len(power[0])} power charts; tabs {wb.sheetnames}")
    for label, m in (("gas", gas[2]), ("power", power[2]), ("long-term", notes or [])):
        if m:
            print(f"notes ({label}):", m)


if __name__ == "__main__":
    main()
