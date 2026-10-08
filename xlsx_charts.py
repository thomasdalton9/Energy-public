"""
Shared helper: add a chart sheet with a native Excel chart to a workbook
that has already been written (e.g. by xlsx_notes.write_workbook).

Every xlsx this repo pulls should carry a chart of its data (see
CLAUDE.md). The chart sheet holds its own copy of the plotted table, so
it works whatever layout the data sheets use (wide, long, with an
index column...): pass a wide DataFrame indexed by date, one column per
series.

    import xlsx_charts
    xlsx_charts.add_chart_sheet(path, wide_df, "Bolivia gas demand by sector",
                                "million m3/day", kind="stacked_bar")

kind: "line", "stacked_area", "stacked_bar".
For storage/level series use water_year_chart.add_water_year_chart instead.
"""
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

# Validated categorical order (dataviz palette): fixed order, never cycled.
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]
OTHER_GREY = "9A9A9A"


def _fold_to_palette(df):
    """More series than palette slots -> keep the largest, fold the rest into 'Other'."""
    if df.shape[1] <= len(PALETTE):
        return df
    order = df.abs().mean().sort_values(ascending=False).index
    keep = list(order[: len(PALETTE) - 1])
    out = df[keep].copy()
    out["Other"] = df.drop(columns=keep).sum(axis=1, min_count=1)
    return out


def rotated_labels(axis, degrees=-45):
    """Category-axis tick labels at a fixed angle (Excel and LibreOffice otherwise pick their own), so every chart's
    date axis looks the same."""
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
    axis.txPr = RichText(bodyPr=RichTextProperties(rot=int(degrees * 60000), vert="horz"),
                         p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties()),
                                      endParaRPr=CharacterProperties())])


def tidy_layout(chart, gridlines=True, inner=None):
    """Excel-safe layout: title, legend and axis titles never overlay the plot (openpyxl leaves <c:overlay>
    unset, which newer Excel draws on top of the plot), the category axis sits at the bottom, optional
    removal of horizontal gridlines, and an optional manual inner plot area (x, y, w, h as fractions of
    the chart) so tick labels and axis titles get fixed room around the plot."""
    from openpyxl.chart.layout import Layout, ManualLayout
    if chart.title is not None:
        chart.title.overlay = False
    if chart.legend is not None:
        chart.legend.overlay = False
    chart.x_axis.axPos = "b"
    for ax in (chart.x_axis, chart.y_axis):
        if getattr(ax, "title", None) is not None:
            ax.title.overlay = False
    if not gridlines:
        chart.y_axis.majorGridlines = None
        chart.x_axis.majorGridlines = None
    if inner:
        x, y, w, h = inner
        chart.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                                                  x=x, y=y, w=w, h=h))
    return chart


# Plot-area box used on the master dashboards: room for a one-line title, y-axis labels and title on the
# left, rotated mmm/yy labels and a two-line legend underneath.
DASHBOARD_INNER = (0.12, 0.15, 0.83, 0.52)


GW_TITLE_KEYS = ("generation", "production by technology", "power balance", "net electricity imports",
                 "grid batteries", "hydro: domestic use")


def monthly_energy_to_gw(df, units, title):
    """Power charts: an energy total per month or per year (GWh or TWh) -> average power in GW (energy / hours in the
    period), so periods of different length are comparable. Gas and other charts (title without a power keyword) are
    returned unchanged. Returns (df, units)."""
    import calendar
    import re as _re
    m = _re.match(r"^([GT])Wh per (month|year)$", str(units))
    if not m or not any(k in str(title).lower() for k in GW_TITLE_KEYS) or df is None or df.empty:
        return df, units
    if pd.api.types.is_integer_dtype(df.index):   # annual frames indexed by calendar year
        years = pd.Index(df.index)
        idx = pd.to_datetime(years.astype(str), format="%Y")
    else:
        idx = pd.to_datetime(df.index)
    if m.group(2) == "month":
        hours = idx.days_in_month * 24.0
    else:
        hours = pd.Index([(366 if calendar.isleap(y) else 365) * 24.0 for y in idx.year])
    out = df.apply(pd.to_numeric, errors="coerce").div(pd.Series(hours, index=df.index), axis=0) * (
        1000.0 if m.group(1) == "T" else 1.0)
    out = out.rename(columns=lambda c: str(c).replace("(TWh at full output)", "(GW)"))   # capacity line = GW itself
    return out, "GW (monthly average)" if m.group(2) == "month" else "GW (annual average)"


def prepare(df, line_cols=()):
    """Clean a wide frame for charting; returns (frame, number of bar/area series before overlay lines)."""
    df = df.copy()
    df.index = pd.to_datetime(df.index)
    df = df.sort_index()
    df = df.apply(pd.to_numeric, errors="coerce")
    df = df.loc[:, df.notna().any()]
    if df.empty:
        return df, 0
    lines_df = df[[c for c in line_cols if c in df.columns]]
    body = _fold_to_palette(df.drop(columns=list(lines_df.columns)))
    return pd.concat([body, lines_df], axis=1), body.shape[1]


def write_table(ws, df, date_format="%Y-%m", start_row=1, start_col=1):
    """Write Date + series columns; dates are real Excel dates shown as mmm/yy (annual: yyyy)."""
    excel_fmt = "yyyy" if date_format == "%Y" else "mmm/yy"
    ws.cell(row=start_row, column=start_col, value="Date")
    for j, c in enumerate(df.columns):
        ws.cell(row=start_row, column=start_col + 1 + j, value=str(c))
    for i, (ts, row) in enumerate(df.iterrows(), start=1):
        cell = ws.cell(row=start_row + i, column=start_col, value=ts.to_pydatetime())
        cell.number_format = excel_fmt
        for j, v in enumerate(row.values):
            ws.cell(row=start_row + i, column=start_col + 1 + j, value=None if pd.isna(v) else float(v))
    ws.column_dimensions[ws.cell(row=1, column=start_col).column_letter].width = 11


def _lighten(hex_colour, share=0.55):
    """Blend a hex colour with white (forecast bars)."""
    rgb = [int(hex_colour[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{int(c + (255 - c) * share):02X}" for c in rgb)


def build_chart(ws, df, n_bars, title, y_title, kind="line", date_format="%Y-%m", start_row=1, start_col=1,
                width=28, height=13, gridlines=True, inner=None, forecast_from=None, scenario_from=None, line_styles=None):
    """Native chart over a table written by write_table on sheet ws (the chart can be placed on any sheet)."""
    n = start_row + len(df)
    excel_fmt = "yyyy" if date_format == "%Y" else "mmm/yy"
    if kind == "stacked_bar":
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.gapWidth = 30
    elif kind == "stacked_area":
        chart = AreaChart()
        chart.grouping = "stacked"
    else:
        chart = LineChart()
    spacing = df.index.to_series().diff().median().days if len(df) > 1 else 30
    if spacing <= 2:
        # Daily: a category axis can only skip evenly, so a "label every 30 days" drifts off month starts. Write a
        # text label column beside the table holding mmm/yy on the 1st of each (every k-th) month, blank otherwise,
        # and show every label.
        months = pd.period_range(df.index.min(), df.index.max(), freq="M")
        k = max(1, -(-len(months) // 18))   # at most ~18 labels
        label_col = start_col + df.shape[1] + 1
        ws.cell(row=start_row, column=label_col, value="Axis label")
        seen = set()
        for i, d in enumerate(df.index, start=start_row + 1):
            m = d.year * 12 + d.month - 1
            first = m not in seen and m % k == 0   # first day present in the month (a missing 1st still gets a label)
            seen.add(m)
            ws.cell(row=i, column=label_col, value=d.strftime("%b/%y") if first else None)
        cats = Reference(ws, min_col=label_col, min_row=start_row + 1, max_row=n)
    else:
        cats = Reference(ws, min_col=start_col, min_row=start_row + 1, max_row=n)
    chart.add_data(Reference(ws, min_col=start_col + 1, max_col=start_col + n_bars, min_row=start_row, max_row=n),
                   titles_from_data=True)
    chart.set_categories(cats)
    for i, s in enumerate(chart.series):
        colour = OTHER_GREY if str(df.columns[i]) == "Other" else PALETTE[i % len(PALETTE)]
        if kind == "line":
            s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=colour, w=22225))
            s.smooth = False
            s.marker.symbol = "none"
        else:
            s.graphicalProperties = GraphicalProperties(solidFill=colour, ln=LineProperties(noFill=True))
    if forecast_from is not None and kind == "stacked_bar":
        # forecast months in a lighter shade of each series' colour
        from openpyxl.chart.marker import DataPoint
        idxs = [i for i, d in enumerate(df.index) if d >= pd.Timestamp(forecast_from)]
        for si, s in enumerate(chart.series):
            colour = OTHER_GREY if str(df.columns[si]) == "Other" else PALETTE[si % len(PALETTE)]
            for i in idxs:
                # scenario_from: months beyond the forecast horizon (e.g. Texas 2029-33) get a still lighter shade
                far = scenario_from is not None and df.index[i] >= pd.Timestamp(scenario_from)
                s.dPt.append(DataPoint(idx=i, invertIfNegative=False,
                                       spPr=GraphicalProperties(solidFill=_lighten(colour, 0.8 if far else 0.55),
                                                                ln=LineProperties(noFill=True))))
    chart.title = f"{title} (to {df.index.max().strftime('%Y' if date_format == '%Y' else '%b/%y')})"
    if forecast_from is not None:
        chart.title = f"{title} (to {df.index.max().strftime('%b/%y')}; forecast from {pd.Timestamp(forecast_from).strftime('%b/%y')}, lighter bars)"
        if scenario_from is not None:
            chart.title = (f"{title} (to {df.index.max().strftime('%b/%y')}; forecast from {pd.Timestamp(forecast_from).strftime('%b/%y')}, "
                           f"scenario from {pd.Timestamp(scenario_from).strftime('%b/%y')}" + (", lightest bars)" if kind == "stacked_bar" else ")"))
    chart.x_axis.number_format = excel_fmt
    chart.y_axis.title = y_title
    if str(y_title).startswith("Bcf/d") and "MISO" in title:   # MISO gas burn: Bcf/d to 2 decimals
        chart.y_axis.number_format = "0.00"
        chart.y_axis.numFmt.sourceLinked = False
    chart.x_axis.tickLblPos = "low"
    rotated_labels(chart.x_axis)
    step = 1 if spacing <= 2 else max(1, len(df) // 12)   # daily: every (mostly blank) label is shown
    chart.x_axis.tickLblSkip = step
    chart.x_axis.tickMarkSkip = step
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend.position = "b"
    if df.shape[1] == 1:   # a single series needs no legend box: the title names it
        chart.legend = None
    # no borders: chart frame and plot area
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    if df.shape[1] > n_bars and kind != "line":
        over = LineChart()
        over.add_data(Reference(ws, min_col=start_col + n_bars + 1, max_col=start_col + df.shape[1],
                                min_row=start_row, max_row=n), titles_from_data=True)
        over.set_categories(cats)
        for j, s in enumerate(over.series):
            # overlay lines: solid, dashed, dotted (then repeat) so they stay distinguishable
            colour, dash = ((line_styles or {}).get(str(df.columns[n_bars + j])) or ("252525", ("solid", "dash", "sysDot")[j % 3]))
            s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=colour, w=28575, prstDash=dash))
            s.smooth = False
            s.marker.symbol = "none"
        over.y_axis.delete = True   # shares the bar chart's axis
        if not gridlines:
            over.y_axis.majorGridlines = None
        chart += over
    chart.width, chart.height = width, height
    return tidy_layout(chart, gridlines, inner)


def save_atomic(wb, path):
    root, ext = os.path.splitext(path)
    tmp = f"{root}.tmp{os.getpid()}{ext}"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def add_chart_sheet(path, df, title, y_title, kind="line", sheet_name="Chart", date_format="%Y-%m",
                    width=28, height=13, line_cols=(), wb=None, forecast_from=None, scenario_from=None, line_styles=None):
    """line_cols: columns of df drawn as lines over the bars/areas, on the same axis (same units only).
    wb: an open workbook to add the sheet to (the caller saves it). openpyxl drops chart formatting (axis
    scaling, titles, number formats) of charts it reads back in, so several charts must be added to one
    workbook in a single load/save - add_charts.py does that."""
    df, n_bars = prepare(df, line_cols)
    if df.empty:
        return
    own = wb is None
    if own:
        wb = load_workbook(path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 1 if len(wb.sheetnames) > 1 else None)
    write_table(ws, df, date_format)
    chart = build_chart(ws, df, n_bars, title, y_title, kind, date_format, width=width, height=height,
                        forecast_from=forecast_from, scenario_from=scenario_from, line_styles=line_styles)
    ws.add_chart(chart, f"{chr(ord('A') + min(df.shape[1] + 2, 20))}2")
    if own:
        save_atomic(wb, path)
