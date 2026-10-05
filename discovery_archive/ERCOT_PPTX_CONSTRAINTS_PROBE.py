"""One-off probe (5 Oct 2026): native chart data inside the ERCOT March TAC 'Updated' pptx (queue by type), the Dec 2025 Report on Existing and Potential Electric System Constraints and Needs,
the Jan 2026 Annual Report on ERCOT Demand Response, the Summer 2025 Operational and Market Review. Log only."""
import io, re, requests, pymupdf
H = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/124 Safari/537.36"}
E = "https://www.ercot.com/files/docs/"
KEY = re.compile(r"crypto|data.?cent|large.?load|hydrogen|oil and gas|industrial|CLR|controllable|flexible", re.I)
NUM = re.compile(r"\d[\d,.]*\s?(GW|MW|gigawatt|megawatt)", re.I)
def get(u):
    try:
        r = requests.get(u, headers=H, timeout=180); print("GET", u, r.status_code, len(r.content), r.headers.get("content-type", "")[:50]); return r if r.status_code == 200 else None
    except Exception as e:
        print("FAIL", u, type(e).__name__); return None
# pptx
for u in [E + "2026/03/27/March-TAC-Report-Updated_03262026.pptx", E + "2026/03/27/March-TAC-Report-Updated_03262026.pdf"]:
    r = get(u)
    if r is None: continue
    if u.endswith(".pptx"):
        import zipfile
        z = zipfile.ZipFile(io.BytesIO(r.content)); names = z.namelist()
        print("charts:", [n for n in names if "charts/chart" in n and n.endswith(".xml")][:40], "embeddings:", [n for n in names if "embeddings" in n][:20])
        from pptx import Presentation
        prs = Presentation(io.BytesIO(r.content))
        for i, sl in enumerate(prs.slides, 1):
            for sh in sl.shapes:
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    ch = sh.chart
                    print(f"--- slide {i} chart {ch.chart_type} title={(ch.chart_title.text_frame.text if ch.has_title else '')!r}")
                    try:
                        cats = list(ch.plots[0].categories)
                        print("   cats:", cats[:40])
                        for pl in ch.plots:
                            for se in pl.series: print("   series", se.name, [round(v, 2) if v is not None else None for v in se.values][:40])
                    except Exception as e: print("   err", e)
                elif sh.has_text_frame and re.search(r"type", sh.text_frame.text, re.I):
                    print(f"--- slide {i} text:", re.sub(r"\s+", " ", sh.text_frame.text)[:200])
    else:
        doc = pymupdf.open(stream=r.content, filetype="pdf")
        for i, pg in enumerate(doc): 
            t = re.sub(r"\s+", " ", pg.get_text())
            if re.search(r"by type", t, re.I): print(f"p{i+1}: {t[:600]}")
# pdfs
for u in [E + "2025/12/23/2025-Report-on-Existing-and-Potential-Electric-System-Constraints-and-Needs.pdf",
          "https://www.ercot.com/misdownload/servlets/mirDownload?doclookupId=1188179018",
          E + "2025/09/15/12-Summer-2025-Operational-and-Market-Review.pdf"]:
    r = get(u)
    if r is None: continue
    try: doc = pymupdf.open(stream=r.content, filetype="pdf")
    except Exception as e: print("not pdf", e, r.content[:200]); continue
    print("pages", len(doc))
    for i, pg in enumerate(doc):
        t = re.sub(r"\s+", " ", pg.get_text())
        for s in re.split(r"(?<=[.;])\s", t):
            if KEY.search(s) and NUM.search(s): print(f"  p{i+1}: {s.strip()[:450]}")
print("DONE")
