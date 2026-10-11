"""Ghana probe 3: Energy Commission WEM weekly page + Energy Statistics downloads (type, sheets, electricity tables)."""
import re, subprocess, requests, io
H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"}
EC = "https://www.energycom.gov.gh"
def txt(h):
    b = re.sub(r"<script.*?</script>|<style.*?</style>", " ", h, flags=re.S)
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", b))
r = requests.get(EC + "/index.php/planning/weekly-wholesale-electricity-market-wem-statistics", headers=H, timeout=30)
print("WEM weekly", r.status_code, len(r.text), flush=True)
t = txt(r.text); i = t.find("Weekly Wholesale"); j = t.find("Weekly Wholesale", i + 50)
print(t[j:j + 2500] if j > 0 else t[2500:5000], flush=True)
for l in sorted(set(re.findall(r'href=["\']([^"\']*(?:download|\.xls|\.pdf|wem|market-watch)[^"\']*)["\']', r.text)))[:120]: print("  ", l, flush=True)
r2 = requests.get(EC + "/index.php/planning/energy-statistics", headers=H, timeout=30)
t2 = txt(r2.text); k = t2.find("2026 Energy Statistics"); print("ES page:", t2[max(0,k-600):k+900], flush=True)
import os
os.makedirs("/tmp/dl", exist_ok=True)
for name, dl in [("2026", "855:2026-energy-statistics"), ("2025", "774:2025-energy-statistics"), ("2025key", "784:2025-key-energy-statistics")]:
    path = "/index.php/planning/" + ("key-energy-statistics" if "key" in name else "energy-statistics") + "?download=" + dl
    print("=== DL", name, path, flush=True)
    try:
        r = requests.get(EC + path, headers=H, timeout=(10, 60), allow_redirects=True)
        print(r.status_code, len(r.content), r.headers.get("content-type"), r.headers.get("content-disposition"), r.url, flush=True)
        fn = f"/tmp/dl/{name}.bin"; open(fn, "wb").write(r.content)
        print("magic", r.content[:8], flush=True)
        if r.content[:4] == b"%PDF":
            subprocess.run(["pip", "install", "-q", "pdfplumber"], check=False)
            import pdfplumber
            with pdfplumber.open(fn) as pdf:
                print("pages", len(pdf.pages), flush=True)
                for n, p in enumerate(pdf.pages):
                    tx = p.extract_text() or ""
                    if re.search(r"(?i)(electricity generation|grid electricity|generation by|akosombo)", tx) and n < 400:
                        print(f"--- page {n+1}", tx[:1800].replace("\n", " | "), flush=True)
        elif r.content[:2] == b"PK":
            import openpyxl
            wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
            for ws in wb.worksheets:
                print("SHEET", ws.title, ws.max_row, ws.max_column, flush=True)
    except Exception as e:
        print("ERR", type(e).__name__, str(e)[:200], flush=True)
