"""Ghana probe 4: Energy Commission weekly WEM statistics categories + Market Watch + Bui table text."""
import re, requests, os
H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"}
EC = "https://www.energycom.gov.gh"
def txt(h):
    b = re.sub(r"<script.*?</script>|<style.*?</style>", " ", h, flags=re.S)
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", b))
def page(path):
    r = requests.get(EC + path, headers=H, timeout=30)
    t = txt(r.text); k = t.find("Warning"); 
    print("=====", path, r.status_code, len(r.text), flush=True)
    i = t.rfind("Right To Information Upcoming Event")
    print(t[i:i + 2500] if i > 0 else t[-2500:], flush=True)
    dl = sorted(set(re.findall(r'href=["\']([^"\']*download=[^"\']*)["\']', r.text)))
    print("downloads:", len(dl), flush=True)
    for d in dl[:12]: print("  ", d, flush=True)
    return dl
for cat in ["76-2026", "72-2025", "48-2018"]:
    dl = page(f"/index.php/planning/weekly-wholesale-electricity-market-wem-statistics/category/{cat}")
    if cat == "76-2026" and dl:
        os.makedirs("/tmp/dl", exist_ok=True)
        for d in dl[:2]:
            try:
                r = requests.get(EC + d if d.startswith("/") else d, headers=H, timeout=(10, 60))
                print("DL", d, r.status_code, len(r.content), r.headers.get("content-type"), r.headers.get("content-disposition"), r.content[:6], flush=True)
                fn = "/tmp/dl/wem.bin"; open(fn, "wb").write(r.content)
                if r.content[:2] == b"PK":
                    import openpyxl
                    wb = openpyxl.load_workbook(fn, data_only=True)
                    for ws in wb.worksheets:
                        print("SHEET", ws.title, ws.max_row, ws.max_column, flush=True)
                        for row in ws.iter_rows(min_row=1, max_row=14, values_only=True):
                            print("   ", [c for c in row if c is not None][:14], flush=True)
                elif r.content[:4] == b"%PDF":
                    os.system("pip install -q pdfplumber")
                    import pdfplumber
                    with pdfplumber.open(fn) as pdf:
                        print("pages", len(pdf.pages), flush=True)
                        for p in pdf.pages[:3]: print((p.extract_text() or "")[:1800].replace("\n", " | "), flush=True)
            except Exception as e:
                print("ERR", d, type(e).__name__, str(e)[:150], flush=True)
page("/index.php/planning/ghana-wholesale-electricity-market-watch")
# Bui table
os.system("pip install -q pdfplumber")
import pdfplumber
r = requests.get(EC + "/index.php/planning/energy-statistics?download=855:2026-energy-statistics", headers=H, timeout=(10, 90))
open("/tmp/dl/es26.pdf", "wb").write(r.content)
with pdfplumber.open("/tmp/dl/es26.pdf") as pdf:
    for n, p in enumerate(pdf.pages):
        tx = p.extract_text() or ""
        if re.search(r"Table 3\.(15|12|5|4|3|8|9): ", tx) and n > 8:
            print(f"--- page {n+1}", tx[:2600].replace("\n", " | "), flush=True)
