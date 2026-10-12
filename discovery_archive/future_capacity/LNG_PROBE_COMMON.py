"""Shared fetch helper for the global LNG liquefaction probes (runs in GitHub Actions only)."""
import io
import os
import re
import sys

import requests

H = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36",
     "Accept": "text/html,application/pdf,application/xhtml+xml,*/*;q=0.8", "Accept-Language": "en"}
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "results", "future_capacity")
os.makedirs(OUT, exist_ok=True)


def get(url, timeout=60):
    try:
        r = requests.get(url, headers=H, timeout=timeout)
        return r
    except Exception as e:  # noqa: BLE001
        try:
            from curl_cffi import requests as cr
            return cr.get(url, impersonate="chrome", timeout=timeout)
        except Exception as e2:  # noqa: BLE001
            print("   FAIL", url, type(e).__name__, str(e)[:80], "|", type(e2).__name__, str(e2)[:80])
            return None


def slug(u):
    return re.sub(r"[^A-Za-z0-9]+", "_", u.split("://")[-1])[:90]


def to_text(r, url):
    ct = (r.headers.get("content-type") or "").lower()
    b = r.content
    if b[:4] == b"%PDF" or "pdf" in ct:
        try:
            import pdfplumber
            parts = []
            with pdfplumber.open(io.BytesIO(b)) as pdf:
                for i, p in enumerate(pdf.pages[:150]):
                    parts.append(f"\n=== page {i+1}\n" + (p.extract_text() or ""))
            return "".join(parts), "pdf"
        except Exception as e:  # noqa: BLE001
            return f"pdf read failed {e}", "pdf"
    if b[:2] == b"PK":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"\n=== sheet {ws.title} {ws.max_row}x{ws.max_column}\n")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        parts.append(" | ".join("" if v is None else str(v) for v in row) + "\n")
            return "".join(parts), "xlsx"
        except Exception as e:  # noqa: BLE001
            return f"xlsx read failed {e}", "zip"
    t = r.text
    t = re.sub(r"(?s)<(script|style).*?</\1>", " ", t)
    t = re.sub(r"<[^>]+>", " ", t)
    return re.sub(r"[ \t\r\f]+", " ", re.sub(r"\n\s*\n+", "\n", t)), "html"


def probe(url, kw=None, maxkeep=400000, links=True):
    r = get(url)
    if r is None:
        return None
    print(f"{r.status_code} {len(r.content):>9} {url}")
    if r.status_code != 200:
        return r
    txt, kind = to_text(r, url)
    with open(os.path.join(OUT, slug(url) + ".txt"), "w", encoding="utf-8") as f:
        f.write(f"URL {url}\nKIND {kind}\n" + txt[:maxkeep])
    if kind == "html" and links:
        for h, t in sorted(set(re.findall(r'href="([^"#]+)"[^>]*>\s*([^<]{3,110})', r.text))):
            if kw and re.search(kw, h + " " + t, re.I):
                print("     link", h[:170], "|", re.sub(r"\s+", " ", t).strip()[:90])
    return r
