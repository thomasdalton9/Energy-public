"""Probe 2: open EPDK's July 2026 and Jan 2025 gas sector report Excel annex and PDF; show sheet names, tables with
sector consumption. Saves the files to discovery_archive/results/turkey/."""
import io, os, re, sys
from curl_cffi import requests as cr
import openpyxl

OUT = "discovery_archive/results/turkey"
S = cr.Session()
DOCS = {
    "epdk_2026_07_ek": "/Detay/DownloadDocument?id=UefnmoQDbG4=",
    "epdk_2026_07_pdf": "/Detay/DownloadDocument?id=VpZSyBUYkTY=",
    "epdk_2026_05_ek": "/Detay/DownloadDocument?id=cqqA3zj0Y0M=",
    "epdk_2025_01": "/Detay/DownloadDocument?id=Bm5q8kh1uGI=",
}
for name, path in DOCS.items():
    r = S.get("https://www.epdk.gov.tr" + path, impersonate="chrome", timeout=60)
    ct = r.headers.get("content-type", ""); cd = r.headers.get("content-disposition", "")
    print(f"\n=== {name}: {r.status_code} {len(r.content)}B {ct[:50]} {cd[:80]}")
    if r.status_code != 200:
        continue
    b = r.content
    kind = "pdf" if b[:4] == b"%PDF" else "xlsx" if b[:2] == b"PK" else "xls" if b[:4] == b"\xd0\xcf\x11\xe0" else "bin"
    open(f"{OUT}/{name}.{kind}", "wb").write(b)
    print("  kind", kind)
    if kind == "xlsx":
        wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True)
        for ws in wb.worksheets:
            print(f"  sheet '{ws.title}' {ws.max_row}x{ws.max_column}")
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 14), values_only=True):
                vals = [str(v)[:22] for v in row if v is not None]
                if vals:
                    print("     ", " | ".join(vals[:10]))
    elif kind == "pdf":
        try:
            import pypdf
            rd = pypdf.PdfReader(io.BytesIO(b))
            print("  pages", len(rd.pages))
            for pn in range(min(len(rd.pages), 12)):
                tx = rd.pages[pn].extract_text() or ""
                if re.search(r"tüketim|Tüketim|sektör", tx):
                    print(f"  --- page {pn+1}:", re.sub(r"\s+", " ", tx)[:900])
        except Exception as e:  # noqa: BLE001
            print("  pdf parse", e)
