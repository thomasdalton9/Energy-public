"""EU biomethane probe D (prints only): ET 4.2 Month sheet biomethane column, Eurostat TI_BNG_E fixed parse, Energimyndigheten biogas tables."""
import io
import json
import re
import sys

import requests

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"
H = {"User-Agent": UA}


def get(url, **kw):
    try:
        return requests.get(url, headers=H, timeout=kw.pop("timeout", 30), **kw)
    except Exception as e:  # noqa: BLE001
        print(f"  ERR {url[:100]}: {type(e).__name__}")


def et():
    print("== ET 4.2")
    r = get("https://www.gov.uk/government/statistics/gas-section-4-energy-trends")
    u = re.search(r'href="(https://assets\.publishing\.service\.gov\.uk/media/[^"]+ET_4\.2[^"]+\.xlsx)"', r.text).group(1)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(get(u, timeout=90).content), data_only=True)
    for sh in ("Month (GWh)", "Annual (GWh)"):
        ws = wb[sh]
        hdr = None
        for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
            if row and any(c and "Biomethane" in str(c) for c in row):
                hdr = row
                break
        ci = [i for i, c in enumerate(hdr) if c and "Biomethane" in str(c)][0]
        print(sh, "biomethane col", ci, "hdr", [str(c)[:20] for c in hdr[:12]])
        rows = [(r[0], r[ci]) for r in ws.iter_rows(min_row=9, values_only=True) if r[0] is not None]
        print(" n", len(rows), "first", rows[:3], "last", rows[-14:])
        print(" 2019-22 sample", [x for x in rows if re.search(r"20(19|22|23|24)", str(x[0]))][:40])
        print(" row types", type(rows[0][0]), type(rows[-1][0]))
    ws = wb["Notes"]
    for row in ws.iter_rows(values_only=True):
        s = " ".join(str(c) for c in row if c)
        if "iomethane" in s or "estimate" in s.lower():
            print("  NOTE", s[:400])
    ws = wb["Commentary"]
    for row in ws.iter_rows(values_only=True):
        s = " ".join(str(c) for c in row if c)
        if "iomethane" in s:
            print("  COMM", s[:600])


def eurostat():
    print("== Eurostat")
    B = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/"
    geos = ["AT", "BE", "BG", "HR", "CY", "CZ", "DK", "EE", "FI", "FR", "DE", "EL", "HU", "IE", "IT", "LV", "LT", "LU", "MT", "NL", "PL", "PT", "RO", "SK", "SI", "ES", "SE", "UK", "NO", "CH", "EU27_2020"]
    for siec, nb in (("R5300", "TI_BNG_E"), ("G3000", "TO_BNG"), ("R5300", "IPRD"), ("R5300", "PPRD")):
        r = get(B + "nrg_bal_c", params={"format": "JSON", "lang": "EN", "geo": geos, "siec": siec, "nrg_bal": nb, "unit": "GWH", "freq": "A", "sinceTimePeriod": "2012"}, timeout=90)
        if r is None or not r.ok:
            print(siec, nb, r and r.status_code, r and r.text[:150])
            continue
        j = r.json()
        dims = j["id"]
        idx = {k: list(j["dimension"][k]["category"]["index"]) for k in dims}
        sizes = j["size"]
        out = {}
        for pos, v in j["value"].items():
            pos = int(pos)
            co = {}
            for s, k in reversed(list(zip(sizes, dims))):
                co[k] = idx[k][pos % s]
                pos //= s
            out.setdefault(co["geo"], {})[co["time"]] = round(v)
        print(siec, nb, "geos", len(out), "years", idx["time"])
        for g in sorted(out):
            print("   ", g, out[g])


def se():
    print("== Energimyndigheten PX")
    base = "https://pxexternal.energimyndigheten.se/api/v1/en/Energimyndighetens_statistikdatabas/Officiell_energistatistik/Produktion_av_biogas_och_rotrester"
    r = get(base)
    print(base[-60:], r and r.status_code, r and r.text[:1500])
    if r is not None and r.ok:
        for x in r.json():
            print("  ", x)
            if x["type"] == "t":
                m = get(base + "/" + x["id"])
                if m is not None and m.ok:
                    j = m.json()
                    for v in j["variables"]:
                        print("     var", v["code"], v["text"], len(v["values"]), v["valueTexts"][:12])
            elif x["type"] == "l":
                rr = get(base + "/" + x["id"])
                print("    sub", rr and rr.text[:800])


if __name__ == "__main__":
    for f in (et, eurostat, se):
        try:
            f()
        except Exception as e:  # noqa: BLE001
            print("FAILED", f.__name__, type(e).__name__, e)
        sys.stdout.flush()
