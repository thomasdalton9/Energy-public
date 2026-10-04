"""EU biomethane probe C (prints only): DESNZ ET 4.2 layout, Ofgem GGSS data page, Eurostat TI_BNG_E biogases + biomethane indicators,
SCB/Energimyndigheten PX tables, GIE/EBA map PDF, AGGM SummeBioOesterreich history, Italy pages."""
import io
import json
import re
import sys
import itertools

import requests

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124.0 Safari/537.36"
H = {"User-Agent": UA}


def get(url, **kw):
    try:
        return requests.get(url, headers=kw.pop("headers", H), timeout=kw.pop("timeout", 45), **kw)
    except Exception as e:  # noqa: BLE001
        print(f"  ERR {url[:100]}: {type(e).__name__} {str(e)[:80]}")


def links(url, pat, maxn=30):
    r = get(url)
    if r is None:
        return
    print(f"GET {url[:110]} -> {r.status_code} len={len(r.text)}")
    seen = set()
    for m in re.finditer(r'href="([^"]+)"[^>]*>([^<]{0,100})', r.text):
        h, tx = m.group(1), m.group(2).strip()
        if re.search(pat, h + " " + tx, re.I) and h not in seen:
            seen.add(h)
            print("   LINK", h[:170], "|", tx[:70])
            if len(seen) >= maxn:
                break


def desnz():
    print("== DESNZ ET 4.2")
    r = get("https://www.gov.uk/government/statistics/gas-section-4-energy-trends")
    u = re.search(r'href="(https://assets\.publishing\.service\.gov\.uk/media/[^"]+ET_4\.2[^"]+\.xlsx)"', r.text).group(1)
    print(u)
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(get(u, timeout=90).content), data_only=True)
    print(wb.sheetnames)
    ws = wb["Month (GWh)"]
    print(ws.max_row, ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 8 or i > ws.max_row - 6:
            print(i + 1, [str(c)[:14] for c in row[:14]])
    ws = wb["Annual (GWh)"]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 6 or i > ws.max_row - 12:
            print(i + 1, [str(c)[:12] for c in row[:12]])
    for n in ("Notes",):
        for row in wb[n].iter_rows(values_only=True):
            s = " ".join(str(c) for c in row if c)
            if "iomethane" in s:
                print("  NOTE", s[:500])


def ofgem():
    print("== Ofgem")
    links("https://www.ofgem.gov.uk/environmental-and-social-schemes/green-gas-support-scheme-and-green-gas-levy/green-gas-support-scheme-and-green-gas-levy-guidance-resources-data", r"xls|csv|data|report|statist|public|biomethane", 40)
    for u in ["https://www.ofgem.gov.uk/environmental-and-social-schemes/renewable-heat-incentive-rhi/non-domestic-renewable-heat-incentive-rhi",
              "https://www.ofgem.gov.uk/environmental-and-social-schemes/non-domestic-rhi/non-domestic-rhi-public-reports-and-statistics"]:
        links(u, r"xls|csv|data|report|statist|biomethane", 25)


def eurostat():
    print("== Eurostat")
    B = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/"

    def dump(ds, params, show=lambda *a: True, maxn=400):
        r = get(B + ds, params={"format": "JSON", "lang": "EN", **params}, timeout=90)
        if r is None or not r.ok:
            print(" ", ds, params, r.status_code if r is not None else None, r.text[:200] if r is not None else "")
            return
        j = r.json()
        dims = j["id"]
        idx = {k: list(j["dimension"][k]["category"]["index"]) for k in dims}
        sizes = j["size"]
        n = 0
        rows = {}
        for pos, v in j["value"].items():
            pos = int(pos)
            co = []
            for s, k in reversed(list(zip(sizes, dims))):
                co.append(idx[k][pos % s])
                pos //= s
            co = dict(zip(reversed(dims), reversed(co)))
            rows[tuple(co[k] for k in dims if k not in ("freq", "unit"))] = v
        print(" ", ds, params, "n=", len(rows), "dims", dims)
        return rows, j

    geos = ["EU27_2020", "AT", "BE", "BG", "HR", "CY", "CZ", "DK", "EE", "FI", "FR", "DE", "EL", "HU", "IE", "IT", "LV", "LT", "LU", "MT", "NL", "PL", "PT", "RO", "SK", "SI", "ES", "SE", "UK", "NO", "CH"]
    for nb in ("TI_BNG_E", "TO_BNG"):
        for siec in ("R5300", "G3000", "R5310"):
            out = dump("nrg_bal_c", {"geo": geos, "siec": siec, "nrg_bal": nb, "unit": "GWH", "freq": "A", "sinceTimePeriod": "2016"})
            if out:
                rows, j = out
                byg = {}
                for (nbv, siv, g, t), v in rows.items():
                    byg.setdefault(g, {})[t] = round(v)
                for g in sorted(byg):
                    print("    ", nb, siec, g, byg[g])
    for ds in ("nrg_ind_ebmg", "nrg_ind_ibmg"):
        out = dump(ds, {"geo": geos[:6], "sinceTimePeriod": "2018"})
        if out:
            print("    sample", list(out[0].items())[:12])
    # nrg_cb_gas / nrg_cb_gasm flows mentioning bio
    for ds in ("nrg_cb_gas", "nrg_cb_gasm"):
        r = get(B + ds, params={"format": "JSON", "lang": "EN", "geo": "DE", "sinceTimePeriod": "2024"} if ds == "nrg_cb_gas" else {"format": "JSON", "lang": "EN", "geo": "DE", "sinceTimePeriod": "2025-01"}, timeout=90)
        if r is not None and r.ok:
            j = r.json()
            for k in j["dimension"]:
                labs = j["dimension"][k]["category"]["label"]
                print("  ", ds, k, len(labs), [(a, b) for a, b in labs.items() if re.search("bio|blend|inject", b, re.I)][:10])
        else:
            print("  ", ds, r.status_code if r is not None else None)


def sweden():
    print("== Sweden PX")
    for base in ("https://api.scb.se/OV0104/v1/doris/en/ssd/EN/EN0107/EN0107X", "https://api.scb.se/OV0104/v1/doris/en/ssd/EN/EN0105/EN0105A",
                 "https://api.scb.se/OV0104/v1/doris/en/ssd/EN/EN0120"):
        r = get(base)
        print(base, r.status_code if r is not None else None)
        if r is not None and r.ok:
            for x in r.json():
                print("   ", x["id"], x["type"], x["text"][:110])
    for base in ("https://pxexternal.energimyndigheten.se/api/v1/en/Energimyndighetens_statistikdatabas",):
        def walk(path, depth):
            r = get(path)
            if r is None or not r.ok:
                return
            for x in r.json():
                t = x.get("text", "")
                print("   " * depth, x["id"], x["type"], t[:100])
                if x["type"] == "l" and depth < 2:
                    walk(path + "/" + x["id"], depth + 1)
        walk(base, 1)


def gie():
    print("== GIE / EBA map")
    for u in ("https://www.gie.eu/wp-content/uploads/filr/14334/GIE_Press-Release_30062026.pdf", "https://www.gie.eu/wp-content/uploads/filr/14330/GIE_EBA_BIO_2026_A0_FULL_115.pdf"):
        r = get(u, timeout=120)
        print(u, r.status_code if r is not None else None, len(r.content) if r is not None else 0)
        if r is not None and r.ok:
            from pypdf import PdfReader
            rd = PdfReader(io.BytesIO(r.content))
            print(" pages", len(rd.pages))
            for pg in rd.pages[:3]:
                tx = re.sub(r"\s+", " ", pg.extract_text() or "")
                print("  TEXT", tx[:3500])


def aggm():
    print("== AGGM SummeBioOesterreich")
    url = "https://platform.aggm.at/vis-service/api/ts/values"
    for frm, to in (("2010-01-01", "2011-01-01"), ("2014-01-01", "2015-01-01"), ("2017-01-01", "2018-01-01"), ("2020-01-01", "2021-01-01"), ("2023-01-01", "2024-01-01"), ("2025-01-01", "2026-01-01")):
        body = {"rangeType": "individual", "from": f"{frm}T06:00:00", "to": f"{to}T06:00:00", "granularity": "month", "timeseries": ["SummeBioOesterreich", "EntryBiogasOst_MGM-Allokationen"]}
        r = requests.post(url, json=body, headers={"User-Agent": UA, "Content-Type": "application/json", "Accept": "application/json"}, timeout=90)
        try:
            for cd in r.json()["timeSeriesData"]["chartData"]:
                ys = [(p["x"], p["y"]) for p in cd["dataSet"] if p.get("y") is not None]
                print(frm, cd["header"]["name"], len(ys), "sum GWh", round(sum(y for _, y in ys) / 1e6, 1), [round(y / 1e6, 1) for _, y in ys][:12])
        except Exception as e:  # noqa: BLE001
            print(frm, r.status_code, e, r.text[:200])


def italy():
    print("== Italy")
    for u in ["https://www.snam.it/en/our-businesses/market-solutions/biomethane.html", "https://www.snam.it/it/energy-transition/biometano/",
              "https://www.snam.it/it/trasporto/dati-operativi/", "https://www.snam.it/en/transportation/operational-data/",
              "https://dgsaie.mase.gov.it/gas_naturale.php", "https://dgsaie.mise.gov.it/gas_naturale.php", "https://www.mase.gov.it/energia/statistiche-energetiche",
              "https://www.ilbiometano.it/", "https://www.gse.it/dati-e-scenari/open-data"]:
        links(u, r"biomet|bilancio|xls|csv|mensile|statist", 12)
        r = get(u)
        if r is not None and r.ok:
            tx = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", r.text))
            for m in list(re.finditer(r"biomet[a-z]+", tx, re.I))[:3]:
                print("    TXT", tx[max(0, m.start() - 120):m.end() + 200])


if __name__ == "__main__":
    for f in (desnz, aggm, eurostat, sweden, gie, ofgem, italy):
        try:
            f()
        except Exception as e:  # noqa: BLE001
            print("FAILED", f.__name__, type(e).__name__, e)
        sys.stdout.flush()
