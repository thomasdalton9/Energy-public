"""
Probe 8 (SE Europe capacity). Output: results/se_probe8.txt
EPDK monthly reports before 2026 are Word (.docx): dump the table rows that hold the installed-capacity-by-source table for several years
(python-docx), to write the parser.
"""
import os, re, sys, io, html
import requests
from urllib.parse import urljoin
H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36", "Accept-Language": "tr,en;q=0.8"}
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "results"); os.makedirs(OUT, exist_ok=True)
LOG = open(os.path.join(OUT, "se_probe8.txt"), "w")
def p(*a):
    s = " ".join(str(x) for x in a); print(s, flush=True); LOG.write(s + "\n"); LOG.flush()
import docx
U = "https://www.epdk.gov.tr/Detay/Icerik/3-0-23-3/elektrikaylik-sektor-raporlar"
t = requests.get(U, headers=H, timeout=60).text
named = []
for m in re.finditer(r'href="([^"]*DownloadDocument\?id=[^"]+)"', t):
    pre = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", t[max(0, m.start() - 700):m.start()])))
    mm = re.findall(r"(\d{4}) Yılı Elektrik Piyasası (\w+) Ayı Sektör Raporu", pre)
    named.append((mm[-1] if mm else None, m.group(1)))
p("named", sum(1 for n in named if n[0]), "of", len(named))
want = [("2025", "Ocak"), ("2023", "Haziran"), ("2021", "Haziran"), ("2019", "Haziran"), ("2018", "Ocak"), ("2016", "Haziran")]
for w in want:
    for nm, u in named:
        if nm == w:
            r = requests.get(urljoin("https://www.epdk.gov.tr/", html.unescape(u)), headers=H, timeout=120)
            p("-----", w, r.status_code, len(r.content))
            try:
                d = docx.Document(io.BytesIO(r.content))
            except Exception as e:
                p("docx err", e); break
            p("tables", len(d.tables))
            shown = 0
            for ti, tb in enumerate(d.tables):
                txt = " ".join(c.text for row in tb.rows for c in row.cells).upper()
                if ("HİDROLİK" in txt or "HIDROLIK" in txt or "HİDROLİK" in txt) and "TOPLAM" in txt and ("KURULU" in txt):
                    p(f" table {ti}: {len(tb.rows)} rows x {len(tb.columns)}")
                    for row in tb.rows[:30]:
                        cells = []
                        for c in row.cells:
                            x = re.sub(r"\s+", " ", c.text).strip()
                            if not cells or cells[-1] != x: cells.append(x)
                        p("    ", " | ".join(cells)[:260])
                    shown += 1
                    if shown >= 2: break
            if not shown:
                for pi, para in enumerate(d.paragraphs):
                    if re.search(r"(?i)kurulu\s*g[üu]ç", para.text): p("  para", pi, para.text[:200])
            break
