"""
Round 8 (Oct-2026), Honduras. NORTH_CENTRAL_AMERICA_POWER_PROBE7.py found
the ODS 'Informe Diario' list keeps only the last 250 days (26-Jan-2026 on),
each a Power BI PDF whose page 2 has 'PRODUCCION BRUTA POR TIPO DE
GENERACION (MWh)' as a donut chart (labels and values come out of the text
layer in no fixed order). Here:
  - the words of that page with coordinates (to tie values to labels),
  - the other report lists: informes diarios OM (p4_id 6), operacion del
    mercado (10), anual (41) - format of their newest file,
  - the real-time 'otr' page's JET chart data (generation by technology,
    last 24 h) via the APEX plugin AJAX call.

Usage: python3 NORTH_CENTRAL_AMERICA_POWER_PROBE8.py ods
"""

print("STARTING", flush=True)

import html
import io
import json
import re
import ssl
import sys

import requests

H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                   "Chrome/124.0 Safari/537.36", "Accept-Language": "es-ES,es;q=0.9"}
S = requests.Session()
S.headers.update(H)
BASE = "https://appcnd.enee.hn:3200/odsprd/"


def aia_bundle(host, port=443):
    import certifi
    from cryptography import x509
    from cryptography.hazmat.primitives.serialization import Encoding, pkcs7
    from cryptography.x509.oid import AuthorityInformationAccessOID, ExtensionOID
    pem = ssl.get_server_certificate((host, port))  # read the leaf only to find its AIA URL
    cert = x509.load_pem_x509_certificate(pem.encode())
    extra = []
    for _ in range(3):
        try:
            aia = cert.extensions.get_extension_for_oid(ExtensionOID.AUTHORITY_INFORMATION_ACCESS).value
        except x509.ExtensionNotFound:
            break
        urls = [d.access_location.value for d in aia if d.access_method == AuthorityInformationAccessOID.CA_ISSUERS]
        if not urls:
            break
        der = requests.get(urls[0], timeout=30).content
        certs = pkcs7.load_der_pkcs7_certificates(der) if urls[0].endswith(".p7c") else [x509.load_der_x509_certificate(der)]
        extra += [c.public_bytes(Encoding.PEM).decode() for c in certs]
        cert = certs[0]
        if cert.issuer == cert.subject:
            break
    path = f"/tmp/{host}_bundle.pem"
    open(path, "w").write(open(certifi.where()).read() + "\n" + "\n".join(extra))
    return path


def page_ctx(url, item=None, value=None, page=None):
    t = S.get(url, timeout=60).text
    g = lambda p: (re.search(p, t).group(1) if re.search(p, t) else "")  # noqa: E731
    return {"t": t, "sess": g(r'name="p_instance" value="(\d+)"'),
            "step": g(r'name="p_flow_step_id" value="(\d+)"'), "flow": g(r'name="p_flow_id" value="(\d+)"'),
            "ajax": g(r'"ajaxIdentifier":"([^"]+)"').encode().decode("unicode_escape"),
            "ws": g(r'_worksheet_id" value="(\d+)"'), "rep": g(r'_report_id" value="(\d+)"'),
            "salt": g(r'value="(\d+)" id="pSalt"'),
            "prot": g(r'id="pPageItemsProtected" value="([^"]+)"').replace("&#x2F;", "/"),
            "ck": g(r'data-for="' + (item or "X").upper() + r'" value="([^"]+)"'), "item": item, "value": value}


def ir_rows(c):
    data = {"p_flow_id": c["flow"], "p_flow_step_id": c["step"], "p_instance": c["sess"], "p_debug": "",
            "p_request": "PLUGIN=" + c["ajax"], "p_widget_name": "worksheet", "p_widget_mod": "ACTION",
            "p_widget_action": "LAZY_LOAD", "p_widget_num_return": "250", "x01": c["ws"], "x02": c["rep"],
            "p_json": json.dumps({"pageItems": {"itemsToSubmit": [{"n": c["item"].upper(), "v": str(c["value"]), "ck": c["ck"]}],
                                                "protected": c["prot"], "rowVersion": "", "formRegionChecksums": []},
                                  "salt": c["salt"]})}
    htm = S.post(BASE + "wwv_flow.ajax", data=data, timeout=60).text
    out = []
    for m in re.finditer(r"<tr[^>]*>(.*?)</tr>", htm, re.S):
        cells = [html.unescape(re.sub(r"<[^>]+>|\s+", " ", x)).strip() for x in re.findall(r"<td[^>]*>(.*?)</td>", m.group(1), re.S)]
        links = [html.unescape(h) for h in re.findall(r'href="([^"]+)"', m.group(1)) if "get_blob" in h]
        if cells:
            out.append((cells, links))
    return out


def show_file(content, label):
    print(f"  FILE {label}: {len(content):,} B magic={content[:8]!r}", flush=True)
    if content[:4] == b"%PDF":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            print(f"   pages={len(pdf.pages)}", flush=True)
            for i, p in enumerate(pdf.pages[:3]):
                print(f"   --- page {i + 1}: {(p.extract_text() or '')[:1500]}", flush=True)
    elif content[:2] == b"PK":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for ws in wb.worksheets[:8]:
            print(f"   == sheet {ws.title!r} {ws.max_row}x{ws.max_column}", flush=True)
            for i, row in enumerate(ws.iter_rows(max_row=25, max_col=20, values_only=True)):
                vals = [str(v)[:16] for v in row if v not in (None, "")]
                if vals:
                    print(f"     r{i}: " + " | ".join(vals[:16]), flush=True)


def ods():
    S.verify = aia_bundle("cnd.enee.hn")
    # 1) daily report, page 2 words with coordinates
    c = page_ctx(f"{BASE}f?p=110:6:::::p6_id:2", "p6_id", 2)
    rows = ir_rows(c)
    for cells, links in rows[:1] + rows[-1:]:
        b = S.get(BASE + links[0], timeout=120).content
        import pdfplumber
        with pdfplumber.open(io.BytesIO(b)) as pdf:
            p = pdf.pages[1]
            print(f"  == {cells[0]} page 2 size={p.width}x{p.height}", flush=True)
            for w in p.extract_words(keep_blank_chars=True, x_tolerance=2):
                if re.search(r"\d|Hidro|E[oó]lica|Solar|Geot|Biomasa|Carb|Interc|T[eé]rmica|BESS", w["text"]):
                    print(f"    ({w['x0']:.0f},{w['top']:.0f})-({w['x1']:.0f},{w['bottom']:.0f}) {w['text']}", flush=True)
            # any embedded text annotations / chart data?
            print("   chars total", len(p.chars), "rects", len(p.rects), "curves", len(p.curves), flush=True)
    # 2) other lists
    for pid in [6, 10, 41]:
        c = page_ctx(f"{BASE}f?p=110:4:::::p4_id:{pid}", "p4_id", pid)
        rows = ir_rows(c)
        print(f"  LIST p4_id={pid}: {len(rows)} rows; first={rows[0][0] if rows else None}; last={rows[-1][0] if rows else None}", flush=True)
        if rows and rows[0][1]:
            show_file(S.get(BASE + rows[0][1][0], timeout=120).content, rows[0][0][0])
    # 3) real-time page charts
    c = page_ctx(f"{BASE}ods_prd/r/operador-del-sistema-ods/otr")
    for m in re.finditer(r'apex\.widget\.jetChart\.init\("([^"]+)".*?"(UkVHSU9O[^"]+)"\)', c["t"], re.S):
        region, ajax = m.group(1), m.group(2).encode().decode("unicode_escape")
        data = {"p_flow_id": c["flow"], "p_flow_step_id": c["step"], "p_instance": c["sess"], "p_debug": "",
                "p_request": "PLUGIN=" + ajax,
                "p_json": json.dumps({"pageItems": {"itemsToSubmit": [], "protected": c["prot"], "rowVersion": "",
                                                    "formRegionChecksums": []}, "salt": c["salt"]})}
        r = S.post(BASE + "wwv_flow.ajax", data=data, timeout=60)
        print(f"  JET {region}: [{r.status_code}] {len(r.text)} chars: {r.text[:2500]}", flush=True)


if __name__ == "__main__":
    {"ods": ods}[sys.argv[1]]()
    print("DONE", flush=True)
