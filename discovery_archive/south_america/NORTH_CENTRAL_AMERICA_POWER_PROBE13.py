"""
Round 13 (Oct-2026), El Salvador via SIGET. PROBE12 found SIGET's
'Estadisticas del Mercado Electrico' page linking:
  1. Informes del Mercado Electrico        https://www.siget.gob.sv/mercado-electrico/
  3. Boletines de Estadisticas Electricas  .../informe-de-mercado-y-estadisticas-electricas/estadisticas-electricas/
  4. Visualizador dinamico (BI)            .../estadisticas-electricas-bi/
(electricity duties moved from SIGET to DGEHM on 17-Jul-2026).
List each page's downloads and open the newest market report / bulletin.

Usage: python3 NORTH_CENTRAL_AMERICA_POWER_PROBE13.py siget
"""

print("STARTING", flush=True)

import io
import re
import sys
from urllib.parse import urljoin

import requests

H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                   "Chrome/124.0 Safari/537.36", "Accept-Language": "es-ES,es;q=0.9"}
S = requests.Session()
S.headers.update(H)
PAGES = ["https://www.siget.gob.sv/mercado-electrico/",
         "https://www.siget.gob.sv/gerencias/electricidad/informe-de-mercado-y-estadisticas-electricas/estadisticas-electricas/",
         "https://www.siget.gob.sv/gerencias/electricidad/informe-de-mercado-y-estadisticas-electricas/estadisticas-electricas-bi/"]


def get(url):
    try:
        r = S.get(url, timeout=90)
        print(f"[{r.status_code}] {r.url[:180]} {r.headers.get('content-type')} {len(r.content):,} B "
              f"{r.headers.get('content-disposition', '')[:90]}", flush=True)
        return r
    except requests.RequestException as e:
        print(f"ERR {url}: {e}", flush=True)
        return None


def show(content, label):
    if content[:4] == b"%PDF":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            print(f"  PDF {label}: {len(pdf.pages)} pages", flush=True)
            hits = 0
            for i, p in enumerate(pdf.pages):
                t = p.extract_text() or ""
                if re.search(r"inyecci|generaci[oó]n|recurso|t[eé]rmic|gas natural|geot", t, re.I) and re.search(r"\d{2,}", t):
                    print(f"  --- page {i + 1}\n{t[:2000]}", flush=True)
                    hits += 1
                    if hits >= 4:
                        break
    elif content[:2] == b"PK":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for ws in wb.worksheets[:15]:
            print(f"  == sheet {ws.title!r} {ws.max_row}x{ws.max_column}", flush=True)
            for row in ws.iter_rows(max_row=12, max_col=16, values_only=True):
                vals = [str(v)[:14] for v in row if v not in (None, "")]
                if vals:
                    print("     ", " | ".join(vals), flush=True)
    else:
        print(f"  {label}: magic {content[:8]!r}", flush=True)


def siget():
    for page in PAGES:
        r = get(page)
        if r is None:
            continue
        t = r.text
        main = t[t.find("Inicio &gt;") if "Inicio &gt;" in t else 0:]
        for m in re.finditer(r"<iframe[^>]*src=[\"']([^\"']+)", main, re.I):
            print("   IFRAME", m.group(1)[:250], flush=True)
        links = []
        for m in re.finditer(r"<a\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", main, re.I | re.S):
            h = urljoin(r.url, m.group(1).strip())
            tx = re.sub(r"<[^>]+>|\s+", " ", m.group(2)).strip()
            if re.search(r"wpdmdl=|/download/|wp-content/uploads|\.pdf|\.xls|powerbi|mercado|boletin|estadistic", h, re.I) \
                    and "facebook" not in h:
                links.append((h, tx))
        seen = set()
        for h, tx in links:
            if h in seen:
                continue
            seen.add(h)
            print(f"   LINK {h[:220]} | {tx[:90]}", flush=True)
        files = [(h, tx) for h, tx in links if re.search(r"wpdmdl=|/download/|\.pdf|\.xls", h, re.I)]
        for h, tx in files[:2]:
            d = get(h)
            if d is None or not d.ok:
                continue
            if "html" in (d.headers.get("content-type") or ""):
                for m in re.finditer(r"href=[\"']([^\"']*(?:wpdmdl=|\.pdf|\.xlsx?)[^\"']*)", d.text):
                    f = urljoin(d.url, m.group(1).replace("&amp;", "&"))
                    print("     FILE", f[:220], flush=True)
                    dd = get(f)
                    if dd is not None and dd.ok and "html" not in (dd.headers.get("content-type") or ""):
                        show(dd.content, tx)
                    break
            else:
                show(d.content, tx)


if __name__ == "__main__":
    {"siget": siget}[sys.argv[1]]()
    print("DONE", flush=True)
