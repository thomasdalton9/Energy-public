"""
Round 15 (Oct-2026), El Salvador.
  siget  SIGET 'Boletin de Estadisticas Electricas' 2021-2024 (WP Download
         Manager pages /download/boletin-...): file type and the pages with
         monthly generation / injections by resource.
  dgehm  DGEHM statistics portal https://estadisticas.dgehm.gob.sv/ (found in
         PROBE12: annual energy balance shows gas natural imports 18,749 TJ in
         2023): look for monthly hydrocarbon imports (LNG) and electricity
         statistics.

Usage: python3 NORTH_CENTRAL_AMERICA_POWER_PROBE15.py siget|dgehm
"""

print("STARTING", flush=True)

import io
import re
import ssl
import sys
from urllib.parse import urljoin, urlparse

import requests

H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                   "Chrome/124.0 Safari/537.36", "Accept-Language": "es-ES,es;q=0.9"}
S = requests.Session()
S.headers.update(H)
MONTHS = r"(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre|ene|feb|mar|abr|may|jun|jul|ago|sep|oct|nov|dic)"


def get(url, **kw):
    try:
        r = S.get(url, timeout=120, **kw)
        print(f"[{r.status_code}] {r.url[:180]} {r.headers.get('content-type')} {len(r.content):,} B "
              f"{r.headers.get('content-disposition', '')[:90]}", flush=True)
        return r
    except requests.RequestException as e:
        print(f"ERR {url}: {e}", flush=True)
        return None


def show(content, label, want=r"inyecci|generaci[oó]n|gas natural|GNL|importaci"):
    if content[:4] == b"%PDF":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            print(f"  PDF {label}: {len(pdf.pages)} pages", flush=True)
            hits = 0
            for i, p in enumerate(pdf.pages):
                t = p.extract_text() or ""
                if re.search(want, t, re.I) and len(re.findall(MONTHS, t, re.I)) >= 6:
                    print(f"  --- page {i + 1}\n{t[:2200]}", flush=True)
                    hits += 1
                    if hits >= 4:
                        break
    elif content[:2] == b"PK":
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            import zipfile
            print("  zip:", zipfile.ZipFile(io.BytesIO(content)).namelist()[:20], e, flush=True)
            return
        print(f"  XLSX {label}: sheets {[ws.title for ws in wb.worksheets]}", flush=True)
        for ws in wb.worksheets[:40]:
            rows = list(ws.iter_rows(max_row=10, max_col=16, values_only=True))
            blob = " ".join(str(v) for r in rows for v in r if v is not None)
            if re.search(want, blob, re.I):
                print(f"  == sheet {ws.title!r} {ws.max_row}x{ws.max_column}", flush=True)
                for row in rows:
                    vals = [str(v)[:14] for v in row if v not in (None, "")]
                    if vals:
                        print("     ", " | ".join(vals), flush=True)
    else:
        print(f"  {label}: magic {content[:8]!r}", flush=True)


def siget():
    for slug in ["boletin-estadisticas-electricas-ano-2024", "boletin-estadisticas-electricas-ano-2023",
                 "boletin-de-estadisticas-electricas-ano-2021"]:
        r = get(f"https://www.siget.gob.sv/download/{slug}/")
        if r is None:
            continue
        links = {urljoin(r.url, m.group(1).replace("&amp;", "&"))
                 for m in re.finditer(r"(?:href|data-downloadurl)=[\"']([^\"']*wpdmdl=\d+[^\"']*)", r.text)}
        print("  wpdm links:", list(links)[:4], flush=True)
        for f in list(links)[:1]:
            d = get(f)
            if d is not None and d.ok:
                show(d.content, slug)


def aia_bundle(host, port=443):
    import certifi
    from cryptography import x509
    from cryptography.hazmat.primitives.serialization import Encoding, pkcs7
    from cryptography.x509.oid import AuthorityInformationAccessOID, ExtensionOID
    pem = ssl.get_server_certificate((host, port))  # read the leaf only to find its AIA URL
    cert = x509.load_pem_x509_certificate(pem.encode())
    extra = []
    for _ in range(3):
        try:
            aia = cert.extensions.get_extension_for_oid(ExtensionOID.AUTHORITY_INFORMATION_ACCESS).value
        except x509.ExtensionNotFound:
            break
        urls = [d.access_location.value for d in aia if d.access_method == AuthorityInformationAccessOID.CA_ISSUERS]
        if not urls:
            break
        der = requests.get(urls[0], timeout=30).content
        certs = pkcs7.load_der_pkcs7_certificates(der) if urls[0].endswith(".p7c") else [x509.load_der_x509_certificate(der)]
        extra += [c.public_bytes(Encoding.PEM).decode() for c in certs]
        cert = certs[0]
        if cert.issuer == cert.subject:
            break
    path = f"/tmp/{host}_bundle.pem"
    open(path, "w").write(open(certifi.where()).read() + "\n" + "\n".join(extra))
    return path


def dgehm():
    try:
        r = S.get("https://estadisticas.dgehm.gob.sv/", timeout=30)
        print("plain verify ok", r.status_code, flush=True)
    except requests.exceptions.SSLError:
        S.verify = aia_bundle("estadisticas.dgehm.gob.sv")
    queue, seen, data = ["https://estadisticas.dgehm.gob.sv/"], set(), {}
    follow = r"hidrocarb|import|gas|electric|generac|estad|mensual|serie|boletin|energ|balance|petrol|combust"
    while queue and len(seen) < 60:
        u = queue.pop(0)
        if u in seen:
            continue
        seen.add(u)
        r = get(u)
        if r is None or "html" not in (r.headers.get("content-type") or ""):
            continue
        title = re.search(r"<title[^>]*>(.*?)</title>", r.text, re.S)
        print("   title:", re.sub(r"\s+", " ", title.group(1)).strip()[:90] if title else "", flush=True)
        for m in re.finditer(r"<iframe[^>]*src=[\"']([^\"']+)", r.text, re.I):
            print("   IFRAME", m.group(1)[:200], flush=True)
        for m in re.finditer(r"<a\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", r.text, re.I | re.S):
            h = urljoin(r.url, m.group(1).strip())
            tx = re.sub(r"<[^>]+>|\s+", " ", m.group(2)).strip()
            if "injuve" in h.lower() or "convenio" in h.lower():
                continue
            if re.search(r"\.(pdf|xlsx?|csv|zip)(\?|$)|wpdmdl=|/download/|powerbi", h, re.I):
                data.setdefault(h, tx[:90])
            elif urlparse(h).netloc.endswith("dgehm.gob.sv") and re.search(follow, h + " " + tx, re.I) and h not in seen:
                queue.append(h)
    print(f"-- {len(data)} data links", flush=True)
    for h, tx in data.items():
        print(f"   DATA {h[:220]} [{tx}]", flush=True)
    picks = [h for h, tx in data.items() if re.search(r"gas natural|GNL|import.*hidrocarb|hidrocarb.*import|mensual", h + tx, re.I)]
    for h in picks[:3]:
        d = get(h)
        if d is not None and d.ok:
            show(d.content, h)


if __name__ == "__main__":
    {"siget": siget, "dgehm": dgehm}[sys.argv[1]]()
    print("DONE", flush=True)
