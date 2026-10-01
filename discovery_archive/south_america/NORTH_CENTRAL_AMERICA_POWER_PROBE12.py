"""
Round 12 (Oct-2026) after NORTH_CENTRAL_AMERICA_POWER_PROBE11.py:
  siget  'informe-de-mercado-y-estadisticas-electricas' page exists - list its
         downloads (WP Download Manager '/download/...' links) and open the newest
  dgehm  www.dgehm.gob.sv fails verification (missing intermediate): fetch it
         from the leaf's AIA URL (verification stays on) and crawl for
         electricity / hydrocarbon (LNG import) statistics
  puc    PUC Belize (BEL annual reports are 40 MB PDFs with annual totals only)

Usage: python3 NORTH_CENTRAL_AMERICA_POWER_PROBE12.py siget|dgehm|puc
"""

print("STARTING", flush=True)

import io
import re
import ssl
import sys
from urllib.parse import urljoin, urlparse

import requests

H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                   "Chrome/124.0 Safari/537.36", "Accept-Language": "es-ES,es;q=0.9,en;q=0.8"}
S = requests.Session()
S.headers.update(H)


def get(url, **kw):
    try:
        r = S.get(url, timeout=60, **kw)
        print(f"[{r.status_code}] {r.url[:180]} {r.headers.get('content-type')} {len(r.content):,} B "
              f"{r.headers.get('content-disposition', '')[:80]}", flush=True)
        return r
    except requests.RequestException as e:
        print(f"ERR {url}: {type(e).__name__}: {str(e)[:200]}", flush=True)
        return None


def links(t, base):
    for m in re.finditer(r"<a\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>", t, re.I | re.S):
        yield urljoin(base, m.group(1).strip()), re.sub(r"<[^>]+>|\s+", " ", m.group(2)).strip()


def show(content, label):
    if content[:4] == b"%PDF":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            print(f"  PDF {label}: {len(pdf.pages)} pages", flush=True)
            hits = 0
            for i, p in enumerate(pdf.pages):
                t = p.extract_text() or ""
                if re.search(r"gas natural|GNL|t[eé]rmic|generaci[oó]n por|recurso|Energ[ií]a del Pac", t, re.I) and re.search(r"\d", t):
                    print(f"  --- page {i + 1}\n{t[:2200]}", flush=True)
                    hits += 1
                    if hits >= 4:
                        break
    elif content[:2] == b"PK":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for ws in wb.worksheets[:10]:
            print(f"  == sheet {ws.title!r} {ws.max_row}x{ws.max_column}", flush=True)
            for i, row in enumerate(ws.iter_rows(max_row=15, max_col=16, values_only=True)):
                vals = [str(v)[:14] for v in row if v not in (None, "")]
                if vals:
                    print("     ", " | ".join(vals), flush=True)


def siget():
    for u in ["https://www.siget.gob.sv/gerencias/electricidad/informe-de-mercado-y-estadisticas-electricas/",
              "https://www.siget.gob.sv/gerencias/electricidad/"]:
        r = get(u)
        if r is None:
            continue
        t = r.text
        main = t[t.find("<main") if "<main" in t else 0:]
        text = re.sub(r"\s+", " ", re.sub(r"<script.*?</script>|<style.*?</style>|<[^>]+>", " ", main, flags=re.S))
        print("  TEXT", text[:2500], flush=True)
        found = []
        for h, tx in links(main, r.url):
            if re.search(r"download|wp-content/uploads|\.pdf|\.xls|estad|mercado|boletin|informe|wpdmc", h, re.I):
                print("   LINK", h[:200], "|", tx[:90], flush=True)
                found.append((h, tx))
        for m in re.finditer(r"<iframe[^>]*src=[\"']([^\"']+)", main, re.I):
            print("   IFRAME", m.group(1)[:200], flush=True)
        for h, tx in found[:3]:
            d = get(h)
            if d is not None and d.ok:
                if "html" in (d.headers.get("content-type") or ""):
                    for h2, tx2 in links(d.text, d.url):
                        if re.search(r"wpdmdl=|\.pdf|\.xls|download", h2, re.I):
                            print("     SUB", h2[:200], "|", tx2[:80], flush=True)
                else:
                    show(d.content, tx)


def aia_bundle(host, port=443):
    import certifi
    from cryptography import x509
    from cryptography.hazmat.primitives.serialization import Encoding, pkcs7
    from cryptography.x509.oid import AuthorityInformationAccessOID, ExtensionOID
    pem = ssl.get_server_certificate((host, port))  # read the leaf only to find its AIA URL
    cert = x509.load_pem_x509_certificate(pem.encode())
    print("  leaf:", cert.subject.rfc4514_string(), "issuer:", cert.issuer.rfc4514_string(), flush=True)
    extra = []
    for _ in range(3):
        try:
            aia = cert.extensions.get_extension_for_oid(ExtensionOID.AUTHORITY_INFORMATION_ACCESS).value
        except x509.ExtensionNotFound:
            break
        urls = [d.access_location.value for d in aia if d.access_method == AuthorityInformationAccessOID.CA_ISSUERS]
        if not urls:
            break
        der = requests.get(urls[0], timeout=30).content
        certs = pkcs7.load_der_pkcs7_certificates(der) if urls[0].endswith(".p7c") else [x509.load_der_x509_certificate(der)]
        extra += [c.public_bytes(Encoding.PEM).decode() for c in certs]
        cert = certs[0]
        if cert.issuer == cert.subject:
            break
    path = f"/tmp/{host}_bundle.pem"
    open(path, "w").write(open(certifi.where()).read() + "\n" + "\n".join(extra))
    return path


def crawl(starts, follow, budget, hosts):
    queue, seen, data = list(starts), set(), {}
    while queue and len(seen) < budget:
        u = queue.pop(0)
        if u in seen:
            continue
        seen.add(u)
        r = get(u)
        if r is None or "html" not in (r.headers.get("content-type") or ""):
            continue
        for h, tx in links(r.text, r.url):
            if re.search(r"facebook|twitter|instagram|youtube|mailto:|javascript:", h):
                continue
            if re.search(r"\.(pdf|xlsx?|csv|zip)(\?|$)|wpdmdl=|/download/", h, re.I):
                data.setdefault(h, tx[:90])
            elif any(urlparse(h).netloc.endswith(x) for x in hosts) and re.search(follow, h + " " + tx, re.I) and h not in seen:
                queue.append(h)
    print(f"-- {len(data)} data links", flush=True)
    for h, tx in data.items():
        print(f"   DATA {h[:220]} [{tx}]", flush=True)
    return data


def dgehm():
    S.verify = aia_bundle("www.dgehm.gob.sv")
    data = crawl(["https://www.dgehm.gob.sv/"], r"estad|electric|hidrocarb|gas|import|public|document|informe|energ|"
                 r"boletin|mercado|precio|transparen", 60, ["dgehm.gob.sv"])
    for h, tx in list(data.items()):
        if re.search(r"gas natural|GNL|import|estad", h + tx, re.I):
            d = get(h)
            if d is not None and d.ok:
                show(d.content, tx)
                break


def puc():
    crawl(["https://www.puc.bz/", "https://puc.bz/"], r"electric|statistic|report|annual|tariff|review|data|generat|decision",
          40, ["puc.bz"])


if __name__ == "__main__":
    {"siget": siget, "dgehm": dgehm, "puc": puc}[sys.argv[1]]()
    print("DONE", flush=True)
