"""
Round 6 (Oct-2026) after NORTH_CENTRAL_AMERICA_POWER_PROBE5.py found:
  AMM  GraficaPW graficaCombustible (hourly MW by fuel) starts 2025-04-01;
       the daily Posdespacho 'Carga Horaria' sheet (hourly MW per unit,
       same layout since 2021) has no fuel per unit. Look inside the
       monthly Posdespacho (IPM<yyyymm>.zip) for a daily-by-fuel table.
  ODS  'Listado website' pages are APEX 23.2 Interactive Reports with
       lazyLoading=true: try the IR CSV download and the plugin AJAX call.

Usage: python3 NORTH_CENTRAL_AMERICA_POWER_PROBE6.py amm|ods
"""

print("STARTING", flush=True)

import io
import json
import re
import ssl
import sys
import zipfile

import requests

H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                   "Chrome/124.0 Safari/537.36", "Accept-Language": "es-ES,es;q=0.9"}
S = requests.Session()
S.headers.update(H)


def get(url, **kw):
    try:
        r = S.get(url, timeout=120, **kw)
        print(f"[{r.status_code}] {r.url[:200]} {r.headers.get('content-type')} {len(r.content):,} B", flush=True)
        return r
    except requests.RequestException as e:
        print(f"ERR {url}: {type(e).__name__}: {str(e)[:300]}", flush=True)
        return None


def amm():
    import openpyxl
    for ym in ["202608", "202101"]:
        r = get(f"https://www.amm.org.gt/pdfs2/post_despacho/POSDESPACHO_MENSUAL/{ym[:4]}/IPM{ym}.zip")
        if r is None or r.content[:2] != b"PK":
            continue
        z = zipfile.ZipFile(io.BytesIO(r.content))
        name = z.namelist()[0]
        wb = openpyxl.load_workbook(io.BytesIO(z.read(name)), read_only=True, data_only=True)
        print(f"  {name}: sheets", [(ws.title, ws.max_row, ws.max_column) for ws in wb.worksheets], flush=True)
        if ym != "202608":
            continue
        for ws in wb.worksheets:
            text = []
            for i, row in enumerate(ws.iter_rows(max_row=14, max_col=40, values_only=True)):
                vals = [str(v)[:18] for v in row if v not in (None, "")]
                if vals:
                    text.append(f"r{i}: " + " | ".join(vals[:16]))
            blob = " ".join(text).lower()
            if re.search(r"recurso|combustible|tecnolog|tipo|carbón|bunker|diaria|diario", blob):
                print(f"  == sheet {ws.title!r}", flush=True)
                for t in text:
                    print("    ", t[:300], flush=True)


def aia_bundle(host, port=443):
    import certifi
    from cryptography import x509
    from cryptography.hazmat.primitives.serialization import Encoding, pkcs7
    from cryptography.x509.oid import AuthorityInformationAccessOID, ExtensionOID
    pem = ssl.get_server_certificate((host, port))  # read the leaf only to find its AIA URL
    cert = x509.load_pem_x509_certificate(pem.encode())
    extra = []
    for _ in range(3):
        try:
            aia = cert.extensions.get_extension_for_oid(ExtensionOID.AUTHORITY_INFORMATION_ACCESS).value
        except x509.ExtensionNotFound:
            break
        urls = [d.access_location.value for d in aia if d.access_method == AuthorityInformationAccessOID.CA_ISSUERS]
        if not urls:
            break
        der = requests.get(urls[0], timeout=30).content
        certs = pkcs7.load_der_pkcs7_certificates(der) if urls[0].endswith(".p7c") else [x509.load_der_x509_certificate(der)]
        extra += [c.public_bytes(Encoding.PEM).decode() for c in certs]
        cert = certs[0]
        if cert.issuer == cert.subject:
            break
    path = f"/tmp/{host}_bundle.pem"
    open(path, "w").write(open(certifi.where()).read() + "\n" + "\n".join(extra))
    return path


BASE = "https://appcnd.enee.hn:3200/odsprd/"


def ir_rows(page, item, value):
    """Open an APEX IR page, then try to get its rows (CSV download, then plugin AJAX)."""
    r = get(f"{BASE}f?p=110:{page}:::::{item}:{value}")
    t = r.text
    sess = re.search(r'id="pInstance"\s*value="?(\d+)', t) or re.search(r'name="p_instance" value="(\d+)"', t)
    sess = sess.group(1) if sess else ""
    ajax = re.search(r'"ajaxIdentifier":"([^"]+)"', t)
    ajax = ajax.group(1).encode().decode("unicode_escape") if ajax else ""
    ws = re.search(r'_worksheet_id" value="(\d+)"', t)
    rep = re.search(r'_report_id" value="(\d+)"', t)
    salt = re.search(r'id="pSalt"\s*value="?([^"]+)"', t) or re.search(r'value="(\d+)" id="pSalt"', t)
    prot = re.search(r'id="pPageItemsProtected" value="([^"]+)"', t)
    ck = re.search(r'data-for="' + item.upper() + r'" value="([^"]+)"', t)
    print(f"  page {page} {item}={value}: session={sess} ws={ws and ws.group(1)} report={rep and rep.group(1)}", flush=True)
    for req in ["CSV", "IR[R%s]_CSV" % "", "IR_CSV"]:
        c = get(f"{BASE}f?p=110:{page}:{sess}:{req}:::")
        if c is not None and c.ok:
            print("    ", c.text[:1500].replace("\n", " || "), flush=True)
    for action in ["LAZY_LOAD", "PAGE", "SORT", "CONTROL"]:
        data = {"p_flow_id": "110", "p_flow_step_id": str(page), "p_instance": sess, "p_debug": "",
                "p_request": "PLUGIN=" + ajax, "p_widget_name": "worksheet", "p_widget_mod": "ACTION",
                "p_widget_action": action, "p_widget_num_return": "250",
                "x01": ws.group(1) if ws else "", "x02": rep.group(1) if rep else "",
                "p_json": json.dumps({"pageItems": {"itemsToSubmit": [{"n": item.upper(), "v": str(value), "ck": ck.group(1) if ck else ""}],
                                                    "protected": prot.group(1).replace("&#x2F;", "/") if prot else "",
                                                    "rowVersion": "", "formRegionChecksums": []},
                                      "salt": salt.group(1) if salt else ""})}
        try:
            a = S.post(BASE + "wwv_flow.ajax", data=data, timeout=60)
            txt = re.sub(r"\s+", " ", a.text)
            print(f"    AJAX {action}: [{a.status_code}] {len(a.text)} chars: {txt[:200]}", flush=True)
            if a.ok and ("<table" in a.text or "a-IRR-table" in a.text):
                for m in re.finditer(r"<tr[^>]*>(.*?)</tr>", a.text, re.S):
                    cells = [re.sub(r"<[^>]+>|\s+", " ", c).strip() for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", m.group(1), re.S)]
                    links = re.findall(r'href="([^"]+)"', m.group(1))
                    if any(cells):
                        print("      ROW", cells[:8], links[:2], flush=True)
                break
        except requests.RequestException as e:
            print("    AJAX err", e, flush=True)


def ods():
    S.verify = aia_bundle("cnd.enee.hn")
    ir_rows(6, "p6_id", 2)       # informe diario
    ir_rows(4, "p4_id", 504)     # informe de operacion mensual
    ir_rows(4, "p4_id", 21)      # trimestral


if __name__ == "__main__":
    {"amm": amm, "ods": ods}[sys.argv[1]]()
    print("DONE", flush=True)
