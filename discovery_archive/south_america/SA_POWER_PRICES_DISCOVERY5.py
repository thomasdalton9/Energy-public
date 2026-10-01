"""
Discovery probe, round 5, for south_america/SA_POWER_PRICES_DAILY.py (compact output).

Round 2 (SA_POWER_PRICES_DISCOVERY2.py) found:
  - Brazil: every CCEE host (www, dadosabertos, pda-download) answers 403 "Acesso bloqueado ... nao atender as
    politicas de seguranca da CCEE" to GitHub runners, whatever the user agent. ONS open data (S3, no key) has
    'cmo-semi-horario' (CMO_SEMIHORARIO_YYYY.csv, from 2020) and 'cmo-semanal'.
  - Peru: COES menu links /Portal/mercadomayorista/costosmarginales/index (and /revisados).
  - Chile: cne.cl media has monthly Precio_Medio_de_Mercado.xlsx (2024-2026) and old BarCMg zips (node-price
    studies); no marginal-cost statistics files.
  - Uruguay: ADME menu 'Precio Spot sancionado' (/mme_admin/sancionado.php), /detalleejecucionhoraria/,
    /datosabiertos.html.
  - Bolivia: CNDC API /cm-diario (arg dias, default 30), /precmp (arg fecha), /dashboard/precios (modo, anio, mes),
    /historico/monomicos (desde, hasta).
Round 3 (SA_POWER_PRICES_DISCOVERY3.py) found:
  - Peru: COES page /Portal/mercadomayorista/costosmarginales/index, script costomarginal.js: controller
    siteRoot + 'mercadomayorista/costosmarginales/' with actions lista, mapa, Exportar?fecha=, ExportarMasivo?fechaInicio=.
  - Uruguay: spot price page https://adme.com.uy/mmee/spot/sancionado.php?remota=1; /detalleejecucionhoraria/ lists
    monthly YYYYMM_despachoEjecutado.xlsx files from 2015.
  - Bolivia: /cm-diario?dias=N -> {"fechas": ["dd/mm"...], "valores": [...]} (only ~6 months back, no year);
    /dashboard/precios?modo=mes&anio=&mes= -> monthly precio_energia / potencia / monomico (US$/MWh) back to 1996;
    /historico/monomicos/detalle?anio= monthly energy/power/transport components.
Round 4 (SA_POWER_PRICES_DISCOVERY4.py) found:
  - Peru: GET .../costosmarginales/ExportarMasivo?fechaInicio=dd/mm/yyyy&fechaFin=dd/mm/yyyy returns
    CostosMarginalesNodales.xlsx, sheet COSTOMARGINAL: FECHA HORA (half-hourly), NODO EMD, NOMBRE BARRA, ENERGIA,
    CONGESTION, TOTAL (3 days = 0.9 MB, 28.7k rows).
  - Chile: Precio_Medio_de_Mercado.xlsx sheets 'PMM SEN', 'PMM SIC', 'PMM SING' (monthly, $/kWh, 4-month windows).
    reportediario.cne.cl does not resolve; energiaabierta.cl fails TLS verification.
  - Uruguay: sancionado.php?remota=1 links spotSancionadoSeleccion.php?anio=YYYY&remota=1 (hourly spot sancionado);
    YYYYMM_despachoEjecutado.xlsx sheet 'CMO' = daily rows of 24 hourly CMO values (USD/MWh) + 'Total' (their sum).
  - Bolivia: cm-diario?dias=400 returns 180 days (28/03-29/09), values ~15-23.
"""
import io
import re
import sys
import traceback

import requests

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0 Safari/537.36"}
ONLY = set(sys.argv[1:])


def flat(s, n):
    return re.sub(r"\s+", " ", str(s))[:n]


def get(url, n=300, quiet=False, **kw):
    kw.setdefault("timeout", 90)
    kw.setdefault("headers", UA)
    try:
        r = requests.get(url, **kw)
        if not quiet:
            body = r.text if "html" not in str(r.headers.get("content-type")) else "[html]"
            print(f"  {r.status_code} {str(r.headers.get('content-type'))[:25]} {len(r.content)}B {r.url[:150]}"
                  f"\n     {flat(body, n)}")
        return r
    except Exception as e:  # noqa: BLE001
        print(f"  !! {url[:150]}: {type(e).__name__}: {str(e)[:200]}")
        return None


def hrefs(html, pat, n=25):
    out = sorted(set(re.findall(r"""(?:href|src|action)\s*=\s*["']([^"']+)["']""", html, re.I)))
    out = [u for u in out if re.search(pat, u, re.I)]
    for u in out[:n]:
        print("     link", u[:200])
    return out


def run(name, fn):
    if ONLY and name not in ONLY:
        return
    print("\n==== " + name, flush=True)
    try:
        fn()
    except Exception:  # noqa: BLE001
        traceback.print_exc(limit=2)


def brazil():
    r = get("https://ons-aws-prod-opendata.s3.amazonaws.com/dataset/cmo_tm/CMO_SEMIHORARIO_2026.csv", quiet=True)
    lines = r.content.decode("utf-8", "replace").splitlines()
    print(f"  2026 csv {r.status_code} {len(r.content)}B rows={len(lines)}")
    for line in lines[:4] + ["..."] + lines[-4:]:
        print("    ", line[:160])
    r = get("https://ons-aws-prod-opendata.s3.amazonaws.com/dataset/cmo_tm/CMO_SEMIHORARIO_2021.csv", quiet=True)
    lines = r.content.decode("utf-8", "replace").splitlines()
    print(f"  2021 csv {r.status_code} {len(r.content)}B rows={len(lines)}")
    for line in lines[:3]:
        print("    ", line[:160])


def xlsx_peek(content, rows=8, width=200):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for ws in wb.worksheets[:6]:
        print(f"   sheet {ws.title!r} max_row={ws.max_row} max_col={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= rows:
                break
            vals = [v for v in row if v is not None]
            if vals:
                print("      ", flat(vals, width))


def peek(r):
    ct = str(r.headers.get("content-type"))
    print(f"  {r.status_code} {ct[:60]} {len(r.content)}B cd={r.headers.get('content-disposition')} {r.url[-90:]}")
    if r.content[:2] == b"PK":
        try:
            xlsx_peek(r.content)
        except Exception as e:  # noqa: BLE001
            import zipfile
            print("     zip members:", zipfile.ZipFile(io.BytesIO(r.content)).namelist()[:10], e)
    else:
        print("     ", flat(r.text, 400))


def peru():
    base = "https://www.coes.org.pe/Portal/mercadomayorista/costosmarginales/"
    import time
    import openpyxl
    for a, b in (("01/01/2021", "02/01/2021"), ("01/08/2026", "31/08/2026")):
        t = time.time()
        r = requests.get(base + f"ExportarMasivo?fechaInicio={a}&fechaFin={b}", headers=UA, timeout=600)
        print(f"  {a}-{b}: {r.status_code} {len(r.content)}B {time.time() - t:.0f}s")
        if r.content[:2] != b"PK":
            print("     ", flat(r.text, 300))
            continue
        ws = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True).worksheets[0]
        names, first, last, n = {}, None, None, 0
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row or row[0] is None:
                continue
            n += 1
            first = first or row[0]
            last = row[0]
            if re.search(r"SANTA ROSA|SROSA|CHAVARR|SAN JUAN", str(row[2]) + str(row[1]), re.I):
                names.setdefault((row[1], row[2]), []).append(row[5])
        print(f"     rows={n} first={first} last={last}")
        for k, v in sorted(names.items()):
            vals = [x for x in v if isinstance(x, (int, float))]
            print(f"     {k}: n={len(v)} mean={sum(vals) / max(len(vals), 1):.2f} sample={v[:4]}")


def chile():
    import openpyxl
    r = get("https://www.cne.cl/wp-content/uploads/2026/09/Precio_Medio_de_Mercado.xlsx", quiet=True)
    ws = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)["PMM SEN"]
    rows = [[v for v in row] for row in ws.iter_rows(values_only=True)]
    print("  header:", rows[1])
    for row in rows[-4:]:
        print("   ", flat([v for v in row if v is not None], 300))


def uruguay():
    base = "https://adme.com.uy/mmee/spot/"
    r = get(base + "spotSancionadoSeleccion.php?anio=2026&remota=1", quiet=True)
    print("  sel 2026:", r.status_code, len(r.content))
    t = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", re.sub(r"(?is)<(script|style).*?</\1>", " ", r.text)))
    print("     text:", t[:700])
    hrefs(r.text, r".", n=25)
    for m in re.findall(r"(?is)<form.*?</form>", r.text)[:2]:
        print("     form", flat(m, 700))


run("brazil", brazil)
run("peru", peru)
run("chile", chile)
run("uruguay", uruguay)
print("\nDONE")
