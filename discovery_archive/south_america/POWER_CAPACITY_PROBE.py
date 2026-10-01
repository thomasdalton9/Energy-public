"""
Probe raw official sources of monthly installed generation capacity by
technology for Brazil, Argentina, Chile and Colombia, plus Ember's yearly
capacity figures for validation. Prints what each source returns (fields,
value sets, date ranges) so the production pulls can be written against
real names. Run in GitHub Actions (external sites are blocked from the
Claude sandbox):

    python3 POWER_CAPACITY_PROBE.py brazil|argentina|chile|colombia|ember

Found (Oct 2026): ANEEL SIGA has DatEntradaOperacao for every operating
plant but no retirement date (retired plants are dropped); MMGD is a separate
parquet. CAMMESA's nemo API has no public capacity publication (INFORME_MENSUAL,
POTENCIA_INSTALADA, ... 'no es publica'). CNE uploads Capacidad_Instalada_
Generacion.xlsx monthly (current snapshot only). XM CapEfecNeta per plant per
day works back to 2021 (/daily, Entity Recurso, kW). Ember yearly capacity in GW
by fuel to 2025. -> BRAZIL_ANEEL_CAPACITY_MONTHLY.py, CHILE_CNE_CAPACITY.py,
COLOMBIA_XM_CAPACITY.py.
"""
import datetime as dt
import io
import json
import re
import sys
import zipfile

import requests

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"}
T = (15, 120)


def hdr(s):
    print(f"\n{'=' * 78}\n{s}\n{'=' * 78}", flush=True)


def get(url, **kw):
    try:
        r = requests.get(url, headers=kw.pop("headers", UA), timeout=kw.pop("timeout", T), **kw)
        print(f"GET {r.url[:200]} -> {r.status_code} {r.headers.get('content-type')} {len(r.content):,} B", flush=True)
        return r
    except requests.RequestException as e:
        print(f"GET {url} FAILED {type(e).__name__}: {e}", flush=True)
        return None


# ---------------------------------------------------------------- Brazil
def brazil():
    import pandas as pd
    hdr("ANEEL SIGA CSV")
    url = ("https://dadosabertos.aneel.gov.br/dataset/6d90b77c-c5f5-4d81-bdec-7bc619494bb9/"
           "resource/11ec447d-698d-4ab8-977f-b424d5deee6a/download/siga-empreendimentos-geracao.csv")
    r = get(url, timeout=(15, 300))
    df = pd.read_csv(io.BytesIO(r.content), sep=";", encoding="utf-8", low_memory=False, dtype=str)
    print("columns:", list(df.columns))
    print(df.head(3).T.to_string())
    for c in ["DscFaseUsina", "SigTipoGeracao", "DscOrigemCombustivel", "DscFonteCombustivel",
              "NomFonteCombustivel", "DscTipoOutorga", "DscPropriRegimePariticipacao", "DscTipoConexao"]:
        if c in df.columns:
            print(f"\n{c}:\n{df[c].value_counts(dropna=False).head(40).to_string()}")
    date_cols = [c for c in df.columns if c.lower().startswith("dat") or "data" in c.lower()]
    print("\ndate-like columns:", date_cols)
    for c in date_cols:
        s = df[c].dropna()
        print(f"  {c}: non-null {len(s):,}; sample {s.head(5).tolist()}; min {s.min()} max {s.max()}")
    op = df[df["DscFaseUsina"].str.contains("Opera", na=False)]
    kw = pd.to_numeric(op["MdaPotenciaFiscalizadaKw"].str.replace(",", ".", regex=False), errors="coerce")
    print(f"\noperating rows {len(op):,}; fiscalized total {kw.sum() / 1e6:,.2f} GW; "
          f"missing DatEntradaOperacao {op['DatEntradaOperacao'].isna().sum() if 'DatEntradaOperacao' in op else 'n/a'}")
    if "DatEntradaOperacao" in op:
        d = pd.to_datetime(op["DatEntradaOperacao"], errors="coerce")
        print("operating capacity by start year (GW):")
        print((kw.groupby(d.dt.year).sum() / 1e6).tail(12).round(2).to_string())
        print("missing-date capacity GW:", round(kw[d.isna()].sum() / 1e6, 2))
    print(op.groupby("SigTipoGeracao").apply(
        lambda g: pd.to_numeric(g["MdaPotenciaFiscalizadaKw"].str.replace(",", "."), errors="coerce").sum() / 1e6
    ).round(2).to_string())
    non = df[~df["DscFaseUsina"].str.contains("Opera", na=False)]
    print("\nnon-operating phases sample:")
    print(non[[c for c in ["NomEmpreendimento", "DscFaseUsina", "DatEntradaOperacao", "MdaPotenciaFiscalizadaKw"]
               if c in non]].head(10).to_string())

    hdr("ANEEL CKAN search: geracao distribuida")
    base = "https://dadosabertos.aneel.gov.br/api/3/action/"
    for q in ["geracao distribuida", "empreendimentos geracao distribuida", "SIGA"]:
        r = get(base + "package_search", params={"q": q, "rows": 8})
        if r is None or not r.ok:
            continue
        for pkg in r.json().get("result", {}).get("results", []):
            print(f"  PKG {pkg.get('name')}: {pkg.get('title')}")
            for res in pkg.get("resources", []):
                print(f"     res {res.get('name')} fmt={res.get('format')} size={res.get('size')} "
                      f"modified={res.get('last_modified') or res.get('metadata_modified')} url={res.get('url')}")
    # stream the head of the GD CSV(s) found
    r = get(base + "package_search", params={"q": "geracao distribuida", "rows": 8})
    if r is not None and r.ok:
        for pkg in r.json().get("result", {}).get("results", []):
            for res in pkg.get("resources", []):
                u = res.get("url") or ""
                if u.lower().endswith(".csv") and "distribu" in (pkg.get("name", "") + u).lower():
                    try:
                        with requests.get(u, headers=UA, timeout=(15, 60), stream=True) as s:
                            print(f"\nHEAD of {u}: status {s.status_code} length {s.headers.get('content-length')}")
                            raw = b""
                            for chunk in s.iter_content(65536):
                                raw += chunk
                                if len(raw) > 200000:
                                    break
                            for line in raw.decode("utf-8", errors="replace").splitlines()[:4]:
                                print("   ", line[:600])
                    except requests.RequestException as e:
                        print("   stream failed", e)


# ---------------------------------------------------------------- Argentina
LOOKUP = "https://api.cammesa.com/pub-svc/public/findDocumentosByNemoRango"
ATTACH = "https://api.cammesa.com/pub-svc/public/findAttachmentByNemoId"
TF = "%Y-%m-%dT%H:%M:%S.000Z"


def cammesa_docs(nemo, a, b):
    r = get(LOOKUP, params={"fechadesde": a.strftime(TF), "fechahasta": b.strftime(TF), "nemo": nemo},
            headers={"User-Agent": "gas-demand-scripts/1.0"}, timeout=(15, 60))
    if r is None or not r.ok:
        return []
    try:
        docs = r.json()
    except ValueError:
        return []
    if not isinstance(docs, list):
        print("   ", str(docs)[:200])
        return []
    for d in docs[:6]:
        print(f"   doc {d.get('id')} {d.get('titulo')!r} fecha={d.get('fecha')} "
              f"adj={[(x.get('id'), x.get('nombre') or x.get('name')) for x in d.get('adjuntos', [])][:8]}")
    print(f"   ({len(docs)} docs)")
    return docs


def argentina():
    hdr("cammesaweb pages: look for nemo names")
    pages = ["https://cammesaweb.cammesa.com/informe-sintesis-mensual/",
             "https://cammesaweb.cammesa.com/informe-mensual/",
             "https://cammesaweb.cammesa.com/informes-mensuales/",
             "https://cammesaweb.cammesa.com/informe-anual/",
             "https://cammesaweb.cammesa.com/potencia-instalada/",
             "https://cammesaweb.cammesa.com/download/potencia-instalada/",
             "https://cammesaweb.cammesa.com/datos-relevantes-del-mem/",
             "https://cammesaweb.cammesa.com/"]
    nemos = set()
    for p in pages:
        r = get(p)
        if r is None or not r.ok:
            continue
        t = r.text
        found = set(re.findall(r"nemo[=\"':\s]+([A-Z][A-Z0-9_]{3,})", t))
        nemos |= found
        print("   nemos:", sorted(found))
        links = sorted(set(re.findall(r"https?://[^\"'\s<>]+(?:xlsx?|zip|pdf|csv)", t)))
        print("   file links:", links[:25])
        hits = sorted(set(re.findall(r"href=\"(https?://cammesaweb[^\"]*(?:informe|potencia|mensual|sintesis)[^\"]*)\"",
                                      t, re.I)))
        print("   page links:", hits[:40])
        for js in re.findall(r"src=\"([^\"]+\.js[^\"]*)\"", t)[:30]:
            if "cammesa" in js and ("custom" in js or "app" in js or "main" in js or "docs" in js):
                j = get(js)
                if j is not None and j.ok:
                    jn = set(re.findall(r"[\"']([A-Z][A-Z0-9_]{5,})[\"']", j.text))
                    print("   js consts:", sorted(jn)[:60])
    hdr("CAMMESA nemo candidates (last 120 days)")
    b = dt.datetime.utcnow()
    a = b - dt.timedelta(days=120)
    cands = sorted(nemos | {"INFORME_MENSUAL", "INFORME_SINTESIS_MENSUAL", "INFORME_SINTESIS", "SINTESIS_MENSUAL",
                            "INFORME_MENSUAL_PRINCIPALES_VARIABLES", "INF_MENSUAL", "INFORME_MENSUAL_MEM",
                            "POTENCIA_INSTALADA", "BASE_INFORME_MENSUAL", "DATOS_RELEVANTES", "INFORME_ANUAL",
                            "PRINCIPALES_VARIABLES", "INFORME_SINTETICO_MENSUAL", "MENSUAL", "ANUARIO",
                            "INFORME_ANUAL_MEM", "BASE_DATOS_INFORME_MENSUAL", "INFORME_MENSUAL_DATOS"})
    found_docs = {}
    for n in cands:
        print(f"\n-- {n}")
        docs = cammesa_docs(n, a, b)
        if docs:
            found_docs[n] = docs
    for n, docs in found_docs.items():
        d = docs[0]
        for att in d.get("adjuntos", [])[:6]:
            r = get(ATTACH, params={"attachmentId": att["id"], "docId": d["id"], "nemo": d.get("nemo") or n},
                    headers={"User-Agent": "gas-demand-scripts/1.0"}, timeout=(15, 180))
            if r is None or not r.ok:
                continue
            c = r.content
            print(f"   attachment {att['id']}: {c[:8]!r}")
            if c[:2] == b"PK":
                try:
                    zf = zipfile.ZipFile(io.BytesIO(c))
                    names = zf.namelist()
                    print("   zip/xlsx members:", names[:30])
                    if any(x.startswith("xl/") for x in names):
                        import openpyxl
                        wb = openpyxl.load_workbook(io.BytesIO(c), read_only=True, data_only=True)
                        print("   sheets:", wb.sheetnames)
                        for ws in wb.worksheets:
                            if re.search(r"pot|inst|ofert|capac", ws.title, re.I):
                                for i, row in enumerate(ws.iter_rows(values_only=True)):
                                    if i > 40:
                                        break
                                    vals = [v for v in row if v is not None]
                                    if vals:
                                        print("     ", vals[:14])
                except Exception as e:  # noqa: BLE001
                    print("   unzip failed", e)

    hdr("Secretaria de Energia / datos.gob.ar CKAN")
    for base in ["http://datos.energia.gob.ar/api/3/action/package_search",
                 "https://datos.gob.ar/api/3/action/package_search"]:
        for q in ["potencia instalada", "capacidad instalada", "generadores potencia"]:
            r = get(base, params={"q": q, "rows": 8})
            if r is None or not r.ok:
                continue
            try:
                res = r.json()["result"]["results"]
            except Exception:  # noqa: BLE001
                continue
            for pkg in res:
                print(f"  PKG {pkg.get('name')}: {pkg.get('title')}")
                for x in pkg.get("resources", [])[:8]:
                    print(f"     res {x.get('name')} fmt={x.get('format')} url={x.get('url')}")


# ---------------------------------------------------------------- Chile
def chile():
    hdr("cne.cl media search")
    seen = {}
    for q in ["Capacidad", "capacidad_instalada", "Capacidad_Instalada", "Capacidad-instalada", "instalada",
              "Potencia", "capacidad generacion"]:
        r = get("https://www.cne.cl/wp-json/wp/v2/media",
                params={"search": q, "per_page": 100, "_fields": "date,modified,source_url,title"})
        if r is None or not r.ok:
            continue
        for i in r.json():
            seen[i["source_url"]] = (i.get("date"), i.get("modified"))
    for u, (d, m) in sorted(seen.items(), key=lambda x: x[1][0] or ""):
        print(f"  {d} {m} {u}")
    hdr("cne.cl stats page links")
    r = get("https://www.cne.cl/estadisticas/electricidad/")
    if r is not None and r.ok:
        for u in sorted(set(re.findall(r"https?://www\.cne\.cl/wp-content/uploads/[^\"'\s<>]+", r.text))):
            print("  ", u)
    xl = [u for u in seen if re.search(r"capacidad", u, re.I) and u.lower().endswith((".xlsx", ".xls"))]
    if not xl:
        return
    newest = sorted(xl, key=lambda u: seen[u][0] or "")[-1]
    hdr(f"inspect {newest}")
    r = get(newest, timeout=(15, 300))
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    print("sheets:", wb.sheetnames)
    for ws in wb.worksheets[:8]:
        print(f"\n--- sheet {ws.title} dims={ws.max_row}x{ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 25:
                break
            vals = [v for v in row if v is not None]
            if vals:
                print("   ", vals[:16])


# ---------------------------------------------------------------- Colombia
def colombia():
    hdr("XM ListadoMetricas (capacity metrics)")
    r = requests.post("https://servapibi.xm.com.co/lists", json={"MetricId": "ListadoMetricas"}, timeout=T)
    print("status", r.status_code)
    try:
        for item in r.json().get("Items", []):
            for e in item.get("ListEntities", []):
                v = e.get("Values", {})
                if re.search(r"cap|potencia|efect", json.dumps(v), re.I):
                    print("  ", v)
    except Exception as e:  # noqa: BLE001
        print("  parse failed", e, r.text[:300])
    end = dt.date.today() - dt.timedelta(days=5)
    for ep in ["daily", "monthly", "hourly"]:
        hdr(f"XM /{ep} CapEfecNeta Recurso")
        body = {"MetricId": "CapEfecNeta", "Entity": "Recurso",
                "StartDate": (end - dt.timedelta(days=3 if ep != "monthly" else 70)).isoformat(),
                "EndDate": end.isoformat()}
        try:
            r = requests.post(f"https://servapibi.xm.com.co/{ep}", json=body, timeout=T)
            print("status", r.status_code, len(r.content))
            j = r.json()
            items = j.get("Items", [])
            print("items", len(items))
            for it in items[:2]:
                keys = list(it.keys())
                print("  keys", keys)
                for k, v in it.items():
                    if isinstance(v, list):
                        print(f"  {k}: {len(v)} entities; first 3: {v[:3]}")
                        tot = 0.0
                        for e in v:
                            vv = e.get("Values", {}).get("value") or e.get("Values", {}).get("Value")
                            try:
                                tot += float(vv)
                            except (TypeError, ValueError):
                                pass
                        print(f"  sum of values: {tot:,.1f}")
                    else:
                        print(f"  {k}: {v}")
        except Exception as e:  # noqa: BLE001
            print("  failed", type(e).__name__, e)
    hdr("XM /daily CapEfecNeta Recurso, 2021-01-01..03 (history depth)")
    r = requests.post("https://servapibi.xm.com.co/daily",
                      json={"MetricId": "CapEfecNeta", "Entity": "Recurso", "StartDate": "2021-01-01",
                            "EndDate": "2021-01-03"}, timeout=T)
    try:
        items = r.json().get("Items", [])
        print("items", len(items), [len(it.get("DailyEntities", [])) for it in items])
        if items:
            print("  first", items[0].get("DailyEntities", [])[:2])
    except Exception as e:  # noqa: BLE001
        print("  failed", e, r.text[:300])
    hdr("XM ListadoRecursos fields")
    r = requests.post("https://servapibi.xm.com.co/lists", json={"MetricId": "ListadoRecursos", "Entity": "Sistema"},
                      timeout=T)
    ents = [e.get("Values", {}) for it in r.json().get("Items", []) for e in it.get("ListEntities", [])]
    print(len(ents), "resources; first:", ents[:2])
    from collections import Counter
    for k in ["EnerSource", "Type", "Disp", "State", "RecType", "CompanyCode"]:
        print(k, Counter(str(e.get(k)) for e in ents).most_common(20))


# ---------------------------------------------------------------- Ember
def ember():
    import pandas as pd
    hdr("Ember yearly capacity")
    url = "https://storage.googleapis.com/emb-prod-bkt-publicdata/public-downloads/yearly_full_release_long_format.csv"
    r = get(url, timeout=(15, 300))
    d = pd.read_csv(io.BytesIO(r.content), low_memory=False)
    print(list(d.columns))
    c = d[(d["Category"] == "Capacity") & d["Area"].isin(["Brazil", "Argentina", "Chile", "Colombia"])]
    print(c["Unit"].unique(), c["Subcategory"].unique(), c["Variable"].unique())
    for area, g in c.groupby("Area"):
        p = g[g["Subcategory"].isin(["Fuel", "Total"])].pivot_table(index="Year", columns="Variable", values="Value")
        print(f"\n{area}\n{p.tail(6).round(2).to_string()}")


if __name__ == "__main__":
    {"brazil": brazil, "argentina": argentina, "chile": chile, "colombia": colombia, "ember": ember}[sys.argv[1]]()
