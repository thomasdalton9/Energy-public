"""
Discovery probe, round 4, for south_america/SA_POWER_PRICES_DAILY.py (compact output).

Round 2 (SA_POWER_PRICES_DISCOVERY2.py) found:
  - Brazil: every CCEE host (www, dadosabertos, pda-download) answers 403 "Acesso bloqueado ... nao atender as
    politicas de seguranca da CCEE" to GitHub runners, whatever the user agent. ONS open data (S3, no key) has
    'cmo-semi-horario' (CMO_SEMIHORARIO_YYYY.csv, from 2020) and 'cmo-semanal'.
  - Peru: COES menu links /Portal/mercadomayorista/costosmarginales/index (and /revisados).
  - Chile: cne.cl media has monthly Precio_Medio_de_Mercado.xlsx (2024-2026) and old BarCMg zips (node-price
    studies); no marginal-cost statistics files.
  - Uruguay: ADME menu 'Precio Spot sancionado' (/mme_admin/sancionado.php), /detalleejecucionhoraria/,
    /datosabiertos.html.
  - Bolivia: CNDC API /cm-diario (arg dias, default 30), /precmp (arg fecha), /dashboard/precios (modo, anio, mes),
    /historico/monomicos (desde, hasta).
Round 3 (SA_POWER_PRICES_DISCOVERY3.py) found:
  - Peru: COES page /Portal/mercadomayorista/costosmarginales/index, script costomarginal.js: controller
    siteRoot + 'mercadomayorista/costosmarginales/' with actions lista, mapa, Exportar?fecha=, ExportarMasivo?fechaInicio=.
  - Uruguay: spot price page https://adme.com.uy/mmee/spot/sancionado.php?remota=1; /detalleejecucionhoraria/ lists
    monthly YYYYMM_despachoEjecutado.xlsx files from 2015.
  - Bolivia: /cm-diario?dias=N -> {"fechas": ["dd/mm"...], "valores": [...]} (only ~6 months back, no year);
    /dashboard/precios?modo=mes&anio=&mes= -> monthly precio_energia / potencia / monomico (US$/MWh) back to 1996;
    /historico/monomicos/detalle?anio= monthly energy/power/transport components.
"""
import io
import re
import sys
import traceback

import requests

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0 Safari/537.36"}
ONLY = set(sys.argv[1:])


def flat(s, n):
    return re.sub(r"\s+", " ", str(s))[:n]


def get(url, n=300, quiet=False, **kw):
    kw.setdefault("timeout", 90)
    kw.setdefault("headers", UA)
    try:
        r = requests.get(url, **kw)
        if not quiet:
            body = r.text if "html" not in str(r.headers.get("content-type")) else "[html]"
            print(f"  {r.status_code} {str(r.headers.get('content-type'))[:25]} {len(r.content)}B {r.url[:150]}"
                  f"\n     {flat(body, n)}")
        return r
    except Exception as e:  # noqa: BLE001
        print(f"  !! {url[:150]}: {type(e).__name__}: {str(e)[:200]}")
        return None


def hrefs(html, pat, n=25):
    out = sorted(set(re.findall(r"""(?:href|src|action)\s*=\s*["']([^"']+)["']""", html, re.I)))
    out = [u for u in out if re.search(pat, u, re.I)]
    for u in out[:n]:
        print("     link", u[:200])
    return out


def run(name, fn):
    if ONLY and name not in ONLY:
        return
    print("\n==== " + name, flush=True)
    try:
        fn()
    except Exception:  # noqa: BLE001
        traceback.print_exc(limit=2)


def brazil():
    r = get("https://ons-aws-prod-opendata.s3.amazonaws.com/dataset/cmo_tm/CMO_SEMIHORARIO_2026.csv", quiet=True)
    lines = r.content.decode("utf-8", "replace").splitlines()
    print(f"  2026 csv {r.status_code} {len(r.content)}B rows={len(lines)}")
    for line in lines[:4] + ["..."] + lines[-4:]:
        print("    ", line[:160])
    r = get("https://ons-aws-prod-opendata.s3.amazonaws.com/dataset/cmo_tm/CMO_SEMIHORARIO_2021.csv", quiet=True)
    lines = r.content.decode("utf-8", "replace").splitlines()
    print(f"  2021 csv {r.status_code} {len(r.content)}B rows={len(lines)}")
    for line in lines[:3]:
        print("    ", line[:160])


def xlsx_peek(content, rows=8, width=200):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for ws in wb.worksheets[:6]:
        print(f"   sheet {ws.title!r} max_row={ws.max_row} max_col={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= rows:
                break
            vals = [v for v in row if v is not None]
            if vals:
                print("      ", flat(vals, width))


def peek(r):
    ct = str(r.headers.get("content-type"))
    print(f"  {r.status_code} {ct[:60]} {len(r.content)}B cd={r.headers.get('content-disposition')} {r.url[-90:]}")
    if r.content[:2] == b"PK":
        try:
            xlsx_peek(r.content)
        except Exception as e:  # noqa: BLE001
            import zipfile
            print("     zip members:", zipfile.ZipFile(io.BytesIO(r.content)).namelist()[:10], e)
    else:
        print("     ", flat(r.text, 400))


def peru():
    base = "https://www.coes.org.pe/Portal/mercadomayorista/costosmarginales/"
    s = requests.Session()
    s.headers.update({**UA, "X-Requested-With": "XMLHttpRequest"})
    s.get(base + "index", timeout=60)
    j = s.get("https://www.coes.org.pe/Portal/Areas/MercadoMayorista/Content/Scripts/costomarginal.js?v=4.7", timeout=60).text
    for key in ("ExportarMasivo", "Exportar?fecha", "'lista'", "'mapa'", "GetDateNormalFormat"):
        i = j.find(key)
        print(f"  js[{key}]:", flat(j[max(0, i - 500): i + 300], 800) if i >= 0 else "none")
    for method, path, data in (("POST", "lista", {"fecha": "01/09/2026"}), ("POST", "lista", {"fecha": "2026-09-01"}),
                               ("POST", "mapa", {"fecha": "01/09/2026", "correlativo": 1, "defecto": 1}),
                               ("GET", "Exportar?fecha=01/09/2026", None),
                               ("GET", "ExportarMasivo?fechaInicio=01/09/2026&fechaFin=03/09/2026", None)):
        try:
            r = s.request(method, base + path, data=data, timeout=120)
            print(f"  {method} {path} {data}")
            peek(r)
        except Exception as e:  # noqa: BLE001
            print("  !!", path, e)


def chile():
    r = get("https://www.cne.cl/wp-content/uploads/2026/09/Precio_Medio_de_Mercado.xlsx", quiet=True)
    print(f"  PMM xlsx {r.status_code} {len(r.content)}B")
    xlsx_peek(r.content, rows=14)
    for u in ("http://reportediario.cne.cl", "https://www.energiaabierta.cl/"):
        r = get(u, n=150)
        if r is not None and r.ok:
            hrefs(r.text, r"\.js|api|marginal|cmg|costo", n=12)


def uruguay():
    r = get("https://adme.com.uy/mmee/spot/sancionado.php?remota=1", quiet=True)
    print("  sancionado remota:", r.status_code, len(r.content))
    t = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", re.sub(r"(?is)<(script|style).*?</\1>", " ", r.text)))
    print("     text:", t[:1500])
    hrefs(r.text, r"\.(xls|xlsx|ods|csv|pdf|zip)|php|cgi", n=20)
    for m in re.findall(r"(?is)<form.*?</form>", r.text)[:2]:
        print("     form", flat(m, 500))
    r = get("https://adme.com.uy/detalleejecucionhoraria/", quiet=True)
    files = sorted(set(re.findall(r"(\d{6}_despachoEjecutado\.xlsx)", r.text)))
    print("  despacho files:", len(files), files[-4:])
    if files:
        x = get("https://adme.com.uy/detalleejecucionhoraria/" + files[-1], quiet=True)
        peek(x)


def bolivia():
    api = "https://www.cndc.bo/wp-json/cndc/v1/"
    r = get(api + "cm-diario?dias=400", n=200)
    j = r.json()
    print("  n:", len(j["fechas"]), j["fechas"][:3], j["fechas"][-3:], j["valores"][:3])


run("brazil", brazil)
run("peru", peru)
run("chile", chile)
run("uruguay", uruguay)
run("bolivia", bolivia)
print("\nDONE")
