"""
Round 7 (Oct-2026), Honduras. NORTH_CENTRAL_AMERICA_POWER_PROBE6.py found
that the ODS/CND 'Informe Diario' list (APEX app 110 page 6, P6_ID=2) comes
out as CSV with request 'CSV' (f?p=110:6:<session>:CSV) and that the
plugin AJAX call 'LAZY_LOAD' returns the report HTML with one
apex_util.get_blob link per day (~1.8 MB each). Here: how far back the list
goes, and what one daily file contains.

Usage: python3 NORTH_CENTRAL_AMERICA_POWER_PROBE7.py ods
"""

print("STARTING", flush=True)

import html
import io
import json
import re
import ssl
import sys
import zipfile

import requests

H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                   "Chrome/124.0 Safari/537.36", "Accept-Language": "es-ES,es;q=0.9"}
S = requests.Session()
S.headers.update(H)
BASE = "https://appcnd.enee.hn:3200/odsprd/"


def aia_bundle(host, port=443):
    import certifi
    from cryptography import x509
    from cryptography.hazmat.primitives.serialization import Encoding, pkcs7
    from cryptography.x509.oid import AuthorityInformationAccessOID, ExtensionOID
    pem = ssl.get_server_certificate((host, port))  # read the leaf only to find its AIA URL
    cert = x509.load_pem_x509_certificate(pem.encode())
    extra = []
    for _ in range(3):
        try:
            aia = cert.extensions.get_extension_for_oid(ExtensionOID.AUTHORITY_INFORMATION_ACCESS).value
        except x509.ExtensionNotFound:
            break
        urls = [d.access_location.value for d in aia if d.access_method == AuthorityInformationAccessOID.CA_ISSUERS]
        if not urls:
            break
        der = requests.get(urls[0], timeout=30).content
        certs = pkcs7.load_der_pkcs7_certificates(der) if urls[0].endswith(".p7c") else [x509.load_der_x509_certificate(der)]
        extra += [c.public_bytes(Encoding.PEM).decode() for c in certs]
        cert = certs[0]
        if cert.issuer == cert.subject:
            break
    path = f"/tmp/{host}_bundle.pem"
    open(path, "w").write(open(certifi.where()).read() + "\n" + "\n".join(extra))
    return path


def page_ctx(page, item, value):
    t = S.get(f"{BASE}f?p=110:{page}:::::{item}:{value}", timeout=60).text
    g = lambda p: (re.search(p, t).group(1) if re.search(p, t) else "")  # noqa: E731
    return {"sess": g(r'name="p_instance" value="(\d+)"'),
            "ajax": g(r'"ajaxIdentifier":"([^"]+)"').encode().decode("unicode_escape"),
            "ws": g(r'_worksheet_id" value="(\d+)"'), "rep": g(r'_report_id" value="(\d+)"'),
            "salt": g(r'value="(\d+)" id="pSalt"'),
            "prot": g(r'id="pPageItemsProtected" value="([^"]+)"').replace("&#x2F;", "/"),
            "ck": g(r'data-for="' + item.upper() + r'" value="([^"]+)"'), "page": page, "item": item, "value": value}


def ir_ajax(c, action, extra=None):
    data = {"p_flow_id": "110", "p_flow_step_id": str(c["page"]), "p_instance": c["sess"], "p_debug": "",
            "p_request": "PLUGIN=" + c["ajax"], "p_widget_name": "worksheet", "p_widget_mod": "ACTION",
            "p_widget_action": action, "p_widget_num_return": "250", "x01": c["ws"], "x02": c["rep"],
            "p_json": json.dumps({"pageItems": {"itemsToSubmit": [{"n": c["item"].upper(), "v": str(c["value"]), "ck": c["ck"]}],
                                                "protected": c["prot"], "rowVersion": "", "formRegionChecksums": []},
                                  "salt": c["salt"]})}
    data.update(extra or {})
    return S.post(BASE + "wwv_flow.ajax", data=data, timeout=60).text


def rows(htm):
    out = []
    for m in re.finditer(r"<tr[^>]*>(.*?)</tr>", htm, re.S):
        cells = [html.unescape(re.sub(r"<[^>]+>|\s+", " ", c)).strip() for c in re.findall(r"<td[^>]*>(.*?)</td>", m.group(1), re.S)]
        links = [html.unescape(h) for h in re.findall(r'href="([^"]+)"', m.group(1)) if "get_blob" in h]
        if cells:
            out.append((cells, links))
    return out


def ods():
    S.verify = aia_bundle("cnd.enee.hn")
    c = page_ctx(6, "p6_id", 2)
    csv = S.get(f"{BASE}f?p=110:6:{c['sess']}:CSV:::", timeout=60).content.decode("latin-1")
    lines = csv.strip().splitlines()
    print(f"  CSV rows: {len(lines) - 1}; first: {lines[1][:80]!r}; last: {lines[-1][:80]!r}", flush=True)
    # pagination: try the next page of the IR
    first = rows(ir_ajax(c, "LAZY_LOAD"))
    print("  page1 rows:", len(first), first[0][0] if first else None, flush=True)
    for mod in ["pgR_min_row=26max_rows=25rows_fetched=25", "pgR_min_row=226max_rows=25rows_fetched=25"]:
        r = rows(ir_ajax(c, "PAGE", {"p_widget_action_mod": mod}))
        print(f"  PAGE {mod}: {len(r)} rows; first={r[0][0] if r else None}; last={r[-1][0] if r else None}", flush=True)
    # search for an old date
    for q in ["2021", "01/2023", "2024"]:
        r = rows(ir_ajax(c, "QUICK_FILTER", {"f01": "", "f02": q, "p_widget_action_mod": ""}))
        print(f"  filter {q}: {len(r)} rows; e.g. {[x[0][0] for x in r[:3]]}", flush=True)
    # download the newest file
    if first and first[0][1]:
        link = BASE + first[0][1][0]
        b = S.get(link, timeout=120)
        print(f"  BLOB {first[0][0][0]}: [{b.status_code}] {b.headers.get('content-type')} "
              f"{b.headers.get('content-disposition')} {len(b.content):,} B magic={b.content[:8]!r}", flush=True)
        content = b.content
        if content[:2] == b"PK":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                for ws in wb.worksheets:
                    print(f"  == sheet {ws.title!r} {ws.max_row}x{ws.max_column}", flush=True)
                    for i, row in enumerate(ws.iter_rows(max_row=60, max_col=30, values_only=True)):
                        vals = [str(v)[:16] for v in row if v not in (None, "")]
                        if vals:
                            print(f"     r{i}: " + " | ".join(vals[:18]), flush=True)
            except Exception as e:  # noqa: BLE001
                z = zipfile.ZipFile(io.BytesIO(content))
                print("  zip entries:", z.namelist()[:30], e, flush=True)
        elif content[:4] == b"%PDF":
            import pdfplumber
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                print(f"  PDF pages={len(pdf.pages)}", flush=True)
                for i, p in enumerate(pdf.pages[:5]):
                    print(f"  --- page {i + 1}\n{(p.extract_text() or '')[:3000]}", flush=True)


if __name__ == "__main__":
    {"ods": ods}[sys.argv[1]]()
    print("DONE", flush=True)
