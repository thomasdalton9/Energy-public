"""One-off probe (5 Oct 2026): by-type ERCOT large-load numbers hidden in images / charts of Operational Overview, TAC decks; prior CDR editions;
LTLF files. Output is the log only. Uses pymupdf (text spans, embedded images, page render) + tesseract OCR."""
import io, re, sys, os
import requests, fitz
from PIL import Image
import pytesseract
H = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/124 Safari/537.36", "Accept": "*/*"}
E = "https://www.ercot.com/files/docs/"
KEY = re.compile(r"large.?load|crypto|data.?cent|hydrogen|oil|industrial|LLIS|energiz|queue|co-?locat", re.I)

def get(u, t=120):
    try:
        return requests.get(u, headers=H, timeout=t)
    except Exception as e:
        print("  FAILED", u, type(e).__name__, str(e)[:100]); return None

def ocr(img):
    try:
        return pytesseract.image_to_string(img, config="--psm 6")
    except Exception as e:
        return f"OCR-ERR {e}"

def pdf_dig(u, maxpages=120):
    r = get(u)
    if r is None or r.status_code != 200:
        print("== PDF", u, None if r is None else r.status_code); return
    doc = fitz.open(stream=r.content, filetype="pdf")
    print(f"\n==== PDF {u} pages={len(doc)}")
    for i, pg in enumerate(doc):
        if i >= maxpages: break
        tx = pg.get_text()
        if not re.search(r"large.?load|LLIS|crypto|data.?cent", tx, re.I): continue
        imgs = pg.get_images(full=True)
        print(f"--- page {i+1}: {len(tx)} chars, {len(imgs)} images, {len(pg.get_drawings())} drawings")
        print("   TEXT:", re.sub(r"\s+", " ", tx)[:1800])
        # render the whole page and OCR it
        pix = pg.get_pixmap(dpi=200)
        im = Image.open(io.BytesIO(pix.tobytes("png")))
        o = ocr(im)
        print("   PAGE-OCR:", re.sub(r"[ \t]+", " ", o).replace("\n", " | ")[:2500])
        for k, x in enumerate(imgs[:6]):
            try:
                p = fitz.Pixmap(doc, x[0])
                if p.n > 4: p = fitz.Pixmap(fitz.csRGB, p)
                if p.width < 200: continue
                im2 = Image.open(io.BytesIO(p.tobytes("png")))
                print(f"   IMG{k} {p.width}x{p.height} OCR:", re.sub(r"[ \t]+", " ", ocr(im2)).replace("\n", " | ")[:1500])
            except Exception as e:
                print("   img err", e)

for f in ["2026/09/16/ERCOT-Monthly-Operational-Overview-August-2026.pdf",
          "2025/12/16/ERCOT-Monthly-Operational-Overview-November-2025.pdf",
          "2026/03/12/March-TAC-Report.pdf"]:
    pdf_dig(E + f, 60)

# listings
def links(u, pat):
    r = get(u, 60)
    print("\n== LIST", u, None if r is None else r.status_code)
    if r is None or r.status_code != 200: return []
    out = sorted(set(re.findall(r'href=["\']([^"\']+)["\']', r.text)))
    from urllib.parse import urljoin
    out = [urljoin(u, l) for l in out if re.search(pat, l, re.I)]
    for l in out: print("  ", l)
    return out
cdr = links("https://www.ercot.com/gridinfo/resource", r"CapacityDemandandReserves|CDR|Capacity.?Demand")
ltl = links("https://www.ercot.com/gridinfo/load/forecast", r"forecast|ltlf|adjust|llwg")
links("https://www.ercot.com/services/rq/large-load-integration", r"\.(pdf|xlsx|docx)|large|llis|status")
links("https://www.ercot.com/gridinfo/load/llwg", r"\.(pdf|xlsx|docx)|large|llis")
links("https://www.ercot.com/committees/tac", r"large|llis")
from openpyxl import load_workbook
for u in [x for x in cdr if x.lower().endswith(".xlsx")][:14]:
    r = get(u)
    if r is None or r.status_code != 200: print("  CDR fail", u); continue
    try:
        wb = load_workbook(io.BytesIO(r.content), data_only=True)
    except Exception as e:
        print("  CDR parse fail", u, e); continue
    print("\n==== CDR", u, wb.sheetnames[:12])
    for ws in wb.worksheets:
        if not re.search(r"LoadResourceScen|Load", ws.title, re.I): continue
        print("  SHEET", ws.title)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            v = [(str(round(x, 1)) if isinstance(x, float) else str(x).replace("\n", " ")[:120]) for x in row if x is not None]
            if v and (KEY.search(" ".join(v)) or i < 4): print(f"   {i+1:3d}|", " ; ".join(v)[:400])
print("DONE")
