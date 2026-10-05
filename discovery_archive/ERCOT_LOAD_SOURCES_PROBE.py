"""One-off multi-source probe (6 Oct 2026): sourced ERCOT large-load figures (GW by category) for the Load breakout tab of
americas/TEXAS_DEMAND_REGRESSION.py. Crawls ERCOT/PUCT/EIA/CBECI/LBNL/Comptroller pages, lists data/report links, downloads the
promising ones (pdf/xlsx/csv) and prints sentences / rows with MW/GW near large-load keywords. Output is the log only."""
import io, re, sys, json
import requests

H = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/124 Safari/537.36", "Accept": "*/*"}
KEY = re.compile(r"large.?load|crypto|bitcoin|data.?cent|LLIS|controllable load|\bCLR\b|\bLFL\b|flexible load|hydrogen|batch zero|interconnection (status|queue)|co-?located|behind.the.meter", re.I)
NUM = re.compile(r"\b\d[\d,.]*\s?(GW|MW|gigawatt|megawatt|TWh|EH/s)", re.I)
LINKKEY = re.compile(r"large.?load|llis|crypto|bitcoin|data.?cent|load.?forecast|ltlf|batch|flexible|controllable|clr|cdr|sb.?6|58317|interconnection", re.I)
seen_files = set()


def get(u, t=45):
    try:
        r = requests.get(u, headers=H, timeout=t, allow_redirects=True)
        return r
    except Exception as e:  # noqa
        print(f"   FAILED {u}: {type(e).__name__}: {str(e)[:120]}")
        return None


def abs_url(base, l):
    from urllib.parse import urljoin
    return urljoin(base, l)


def show_text(t, label, n=14):
    t = re.sub(r"\s+", " ", t)
    hits = [s.strip() for s in re.split(r"(?<=[.;])\s", t) if KEY.search(s) and NUM.search(s)]
    for s in hits[:n]:
        print(f"   * [{label}] {s[:380]}")
    return len(hits)


def read_file(u, r):
    ct = r.headers.get("content-type", "")
    b = r.content
    print(f"   FILE {u[:150]} {len(b)} bytes {ct}")
    try:
        if u.lower().split("?")[0].endswith(".pdf") or b[:4] == b"%PDF":
            from pypdf import PdfReader
            rd = PdfReader(io.BytesIO(b))
            print(f"     pdf pages {len(rd.pages)}")
            for i, p in enumerate(rd.pages[:60]):
                tx = p.extract_text() or ""
                if KEY.search(tx) and NUM.search(tx):
                    print(f"     --- page {i+1}")
                    for ln in tx.split("\n"):
                        if KEY.search(ln) or NUM.search(ln):
                            print("       |", ln.strip()[:200])
        elif b[:2] == b"PK":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(b), read_only=True, data_only=True)
            for ws in wb.worksheets[:12]:
                print(f"     sheet {ws.title} dims {ws.dimensions}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 25: break
                    v = [str(x)[:25] for x in row if x is not None]
                    if v: print("       |", " ; ".join(v)[:220])
        else:
            tx = b.decode("utf8", "ignore")
            print("     head:", tx[:600].replace("\n", " | "))
    except Exception as e:  # noqa
        print(f"     parse failed {type(e).__name__}: {e}")


def crawl(u, depth_files=6, follow_pages=0):
    print(f"\n== {u}")
    r = get(u)
    if r is None: return
    print(f"   status {r.status_code} {len(r.content)} bytes {r.headers.get('content-type')}")
    if r.status_code != 200: return
    ct = r.headers.get("content-type", "")
    if "html" not in ct and "json" not in ct and "text" not in ct:
        read_file(u, r); return
    show_text(re.sub(r"<[^>]+>", " ", r.text), "page")
    links = sorted(set(re.findall(r'href=["\']([^"\']+)["\']', r.text)))
    cand = [abs_url(u, l) for l in links if (LINKKEY.search(l) or re.search(r"\.(xlsx?|csv|pdf|zip)(\?|$)", l, re.I))]
    docs = [c for c in cand if re.search(r"\.(xlsx?|csv|pdf)(\?|$)", c, re.I)]
    pages = [c for c in cand if c not in docs]
    for c in (docs + pages)[:40]:
        print("   link:", c[:200])
    n = 0
    for c in docs:
        if c in seen_files or n >= depth_files: continue
        seen_files.add(c); n += 1
        rr = get(c, 90)
        if rr is not None and rr.status_code == 200: read_file(c, rr)
        elif rr is not None: print("   file status", rr.status_code, c[:120])
    for c in pages[:follow_pages]:
        crawl(c, 3, 0)


ERC = "https://www.ercot.com"
for u in [
    ERC + "/services/rq/large-load-integration", ERC + "/services/rq/integration", ERC + "/gridinfo/load/forecast", ERC + "/gridinfo/load",
    ERC + "/gridinfo/resource", ERC + "/gridinfo/generation", ERC + "/news/mediakit/factsheets", ERC + "/about/fact", ERC + "/mktinfo/loads",
    ERC + "/mktrules/issues/NPRR1234", ERC + "/services/comm/mkt_notices", ERC + "/committees/tac", ERC + "/committees/rpg",
    ERC + "/services/rq/batch-zero", ERC + "/services/rq/large-load-integration/batch-zero",
    ERC + "/mktinfo/services/ancillary", ERC + "/gridmktinfo/dashboards", ERC + "/gridinfo/load/demand-response",
]:
    crawl(u, 4, 2 if "large-load" in u else 0)

# ERCOT MIS document lists (public JSON list endpoints) for a few report type ids
for rid in (15801, 16001, 12315, 13057, 14836, 17005, 14954, 11485):
    u = f"{ERC}/misapp/servlets/IceDocListJsonWS?reportTypeId={rid}&_={rid}"
    r = get(u)
    if r is not None:
        print(f"\n== MIS list {rid}: {r.status_code} {r.text[:400].replace(chr(10), ' ')}")
for q in ("Large Load Interconnection Status", "Long-Term Load Forecast", "Controllable Load Resource", "Large Flexible Load"):
    r = get(f"{ERC}/search?keywords={q.replace(' ', '+')}")
    if r is not None and r.status_code == 200:
        print(f"\n== ercot search {q}: {len(r.text)}")
        for l in sorted(set(re.findall(r'href=["\']([^"\']+)["\']', r.text)))[:80]:
            if LINKKEY.search(l) or re.search(r"\.(xlsx?|pdf)$", l, re.I): print("   link:", l[:200])

# CBECI / Cambridge
for u in ["https://ccaf.io/cbnsi/cbeci/mining_map", "https://ccaf.io/cbnsi/cbeci", "https://ccaf.io/cbnsi/cbeci/mining_map/methodology",
          "https://ccaf.io/cbnsi/api/mining_map", "https://ccaf.io/cbnsi/cbeci/api/mining_map",
          "https://raw.githubusercontent.com/cambridge-centre-alternative-finance/cbeci/main/README.md",
          "https://api.github.com/search/repositories?q=cbeci+mining+map",
          "https://api.github.com/search/code?q=%22Bitcoin+mining+map%22+Texas",
          "https://www.cambridge.org/ccaf-bitcoin-mining-map", "https://www.jbs.cam.ac.uk/insight/2025/bitcoin-mining-map"]:
    crawl(u, 3, 0)

# EIA
for u in ["https://www.eia.gov/electricity/data/eia860m/", "https://www.eia.gov/todayinenergy/detail.php?id=61364",
          "https://www.eia.gov/todayinenergy/", "https://www.eia.gov/todayinenergy/detail.php?id=61603",
          "https://www.eia.gov/analysis/studies/ ", "https://www.eia.gov/electricity/monthly/epm_table_grapher.php?t=epmt_5_6_a"]:
    crawl(u, 2, 0)
r = get("https://api.eia.gov/v2/electricity/operating-generator-capacity/?api_key=DEMO_KEY")
print("EIA v2 operating-generator-capacity:", None if r is None else (r.status_code, r.text[:300]))

# PUCT / SB6 / Texas
for u in ["https://interchange.puc.texas.gov/search/filings/?ControlNumber=58317&ItemMatch=Equal&UtilityType=A&ItemNumber=1",
          "https://www.puc.texas.gov/industry/electric/sb6.aspx", "https://www.puc.texas.gov/", "https://interchange.puc.texas.gov/Search/Documents?controlNumber=58317",
          "https://comptroller.texas.gov/economy/economic-data/data-centers/", "https://comptroller.texas.gov/economy/fiscal-notes/",
          "https://www.dallasfed.org/research/economics/2025/", "https://eta-publications.lbl.gov/sites/default/files/2024-12/lbnl-2024-united-states-data-center-energy-usage-report.pdf",
          "https://www.energy.gov/sites/default/files/2024-12/doe-data-center-report.pdf"]:
    crawl(u, 3, 0)

# GridStatus / third-party mirrors of ERCOT large-load numbers (data only if reachable)
for u in ["https://www.gridstatus.io/insights/ercot-large-load", "https://www.gridstatus.io/", "https://www.potomaceconomics.com/reports/",
          "https://www.ercot.com/files/docs/2025/12/", "https://www.ercot.com/files/docs/2026/"]:
    crawl(u, 2, 0)
print("\nDONE")
