"""Cambodia probe 10: final test of asia/CAMBODIA_EAC.py - run it twice to a temp workbook (the second run must
download nothing), print the sheets, then draw the proposed add_charts registry entry into the temp copy."""
import os
import subprocess
import sys
import tempfile

import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)
tmp = os.path.join(tempfile.mkdtemp(), "cambodia_power_generation.xlsx")
for run in (1, 2):
    print(f"######## run {run}", flush=True)
    subprocess.run([sys.executable, os.path.join(ROOT, "asia", "CAMBODIA_EAC.py"), "--out", tmp], check=False)
pd.set_option("display.width", 250)
for sh in ("Daily", "Monthly", "Basis"):
    d = pd.read_excel(tmp, sheet_name=sh)
    print(f"######## {sh} {d.shape}")
    print(d.tail(8).to_string(max_colwidth=200))

import add_charts as A  # noqa: E402


def cambodia_power(p):
    """Proposed REGISTRY entry: annual generation by type, supply = domestic + imports by country, capacity."""
    out = A.power_annual("Cambodia power generation by type (EAC, annual)")(p)
    d = A.by_date(A.read(p, "Daily"), "date")
    d = d[d.index >= "2010-01-01"]
    s = pd.DataFrame({"Domestic generation": d.get("Total_MWh"),
                      "Imports from Vietnam": d.get("Imports_Vietnam_MWh"),
                      "Imports from Thailand": d.get("Imports_Thailand_MWh"),
                      "Imports from Laos": d.get("Imports_Laos_MWh")}) / 1000
    out.append(A.spec("Supply", s, "Cambodia electricity supply: domestic generation and imports by country (EAC)",
                      "GWh per year", "stacked_bar", "%Y"))
    return out + A.power_capacity("Cambodia installed capacity by type (EAC, end of year)")(p)


A.REGISTRY["cambodia_power_generation.xlsx"] = cambodia_power
print("charts:", A.add_charts(tmp))
import openpyxl  # noqa: E402
print("sheets:", openpyxl.load_workbook(tmp).sheetnames)
