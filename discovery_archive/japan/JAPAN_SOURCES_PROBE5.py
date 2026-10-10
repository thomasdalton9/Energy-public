"""Japan sources probe 5 (manual): Chubu / Kansai CSV link scripts, METI table 1 totals and table 4 LNG rows, MOF
Japanese press-release PDF list and LNG row across years."""
import io
import re
import unicodedata
import requests
import pymupdf
import openpyxl

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"}


def get(u):
    try:
        return requests.get(u, headers=UA, timeout=(8, 40))
    except Exception as e:  # noqa: BLE001
        print("   FAIL", u, type(e).__name__, str(e)[:80])


for u in ("https://powergrid.chuden.co.jp/denkiyoho/resource/js/create-download-link.js",
          "https://www.kansai-td.co.jp/denkiyoho/download/js/jisseki.js"):
    r = get(u)
    print("JS", u, None if r is None else (r.status_code, len(r.content)))
    if r is not None and r.status_code == 200:
        print(r.content.decode("utf-8", "replace")[:3500])
r = get("https://powergrid.chuden.co.jp/denkiyoho/eriajukyu_data/")
if r is not None:
    t = r.content.decode("utf-8", "replace")
    for m in list(re.finditer(r"p-link__list", t))[:4]:
        print("HTML~", t[m.start() - 200: m.start() + 1400].replace("\n", " ").replace("\t", ""))
    print("occurrences of 'eria' in html:", [t[max(0, m.start() - 60): m.end() + 80].replace("\n", " ") for m in re.finditer(r"eria_", t)][:6])
# METI table 1 totals
r = get("https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/2026/1-1-2026n.xlsx")
wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
ws = wb.worksheets[1]
names = {}
for row in ws.iter_rows(min_row=6, values_only=True):
    if row[2]:
        names[row[2]] = row[:3] + (row[50] if len(row) > 50 else None,)
tot = {k: v for k, v in names.items() if "計" in str(k) or "全" in str(k)}
print("n names", len(names), "total-like", tot)
print("first names", list(names)[:5], "last names", list(names)[-8:])
ssum = sum((v[3] or 0) for v in names.values() if isinstance(v[3], (int, float)))
print("sum of col50 kW over rows:", ssum)
# METI table 4 LNG rows
for fy in (2021, 2022, 2023, 2024, 2025):
    for suffix in ("", "n"):
        u = f"https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/{fy}/4-{fy}{suffix}.xlsx"
        r = get(u)
        if r is None:
            continue
        print(u, r.status_code, len(r.content))
        if r.status_code == 200:
            wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
            print("  sheets", [w.title for w in wb.worksheets][:14])
            ws = wb.worksheets[1]
            for row in ws.iter_rows(values_only=True):
                if row and any("LNG" in unicodedata.normalize("NFKC", str(c)) for c in row[:4] if c):
                    print("  ", row[:9])
            break
# MOF Japanese list
r = get("https://www.customs.go.jp/toukei/shinbun/happyou.htm")
print("MOF happyou.htm", None if r is None else (r.status_code, len(r.content)))
if r is not None and r.status_code == 200:
    t = r.content.decode(r.apparent_encoding or "utf-8", "replace")
    ls = list(dict.fromkeys(re.findall(r'href="([^"]*trade-st/[^"]*\.pdf)"', t)))
    print(len(ls), ls[:60])
for f in ("2026/2026085.pdf", "2025/2025127.pdf", "2024/2024127.pdf", "2024/2024017.pdf"):
    r = get("https://www.customs.go.jp/toukei/shinbun/trade-st/" + f)
    print("MOF", f, None if r is None else (r.status_code, len(r.content)))
    if r is not None and r.status_code == 200:
        d = pymupdf.open(stream=r.content, filetype="pdf")
        for i, pg in enumerate(d):
            tx = pg.get_text()
            j = tx.find("液化天然ガス")
            if j >= 0:
                print("  page", i, repr(tx[max(0, j - 120): j + 150]))
                print("  header:", repr(d[0].get_text()[:250]))
                break
