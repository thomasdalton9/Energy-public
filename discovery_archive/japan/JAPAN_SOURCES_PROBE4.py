"""Japan sources probe 4 (manual): Chubu and Kansai area supply-demand CSV locations, Tohoku Aug file, MOF English
press-release LNG row, METI electricity-survey workbook layouts (fuel stocks, capacity)."""
import io
import re
import requests
import pymupdf
import openpyxl

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"}


def get(u, headers=None):
    h = dict(UA)
    h.update(headers or {})
    try:
        return requests.get(u, headers=h, timeout=(8, 40))
    except Exception as e:  # noqa: BLE001
        print("   FAIL", type(e).__name__, str(e)[:80])


def csvhead(u, n=3):
    r = get(u)
    if r is None:
        return
    print(f"{u}\n   HTTP {r.status_code} {len(r.content)}b {r.headers.get('content-type')}")
    if r.status_code == 200 and r.content:
        for enc in ("cp932", "utf-8"):
            try:
                for line in r.content.decode(enc).splitlines()[:n]:
                    print("   |", line[:260])
                return
            except Exception:  # noqa: BLE001
                pass


def page(u, pat, maxn=30, ctx=False):
    r = get(u)
    if r is None:
        return
    print(f"PAGE {u}\n   HTTP {r.status_code} {len(r.content)}b")
    if r.status_code != 200:
        return
    t = r.content.decode(r.apparent_encoding or "utf-8", "replace")
    ls = [l for l in dict.fromkeys(re.findall(r'(?:href|src)=["\']?([^"\'#\s>]+)', t)) if re.search(pat, l, re.I)]
    for l in ls[:maxn]:
        print("    ", l)
    if ctx:
        for m in list(re.finditer(r"csv", t, re.I))[:8]:
            print("    ~", t[max(0, m.start() - 100):m.end() + 100].replace("\n", " "))
    return t


# Tohoku
csvhead("https://setsuden.nw.tohoku-epco.co.jp/common/demand/eria_jukyu_202608_02.csv")
# Chubu
t = page("https://powergrid.chuden.co.jp/denkiyoho/eriajukyu_data/", r"csv|jukyu|download|zip|xls", ctx=True)
for u in ("https://powergrid.chuden.co.jp/denki_yoho_content_data/eria_jukyu_202608_04.csv",
          "https://powergrid.chuden.co.jp/denki_yoho_content_data/eriajukyu_data/eria_jukyu_202608_04.csv",
          "https://powergrid.chuden.co.jp/denkiyoho/eriajukyu_data/eria_jukyu_202608_04.csv",
          "https://powergrid.chuden.co.jp/denki_yoho_content_data/areajuyo_current.csv"):
    csvhead(u)
# Kansai
t = page("https://www.kansai-td.co.jp/denkiyoho/download/", r"csv|jukyu|js|zip", ctx=True)
for u in ("https://www.kansai-td.co.jp/denkiyoho/csv/eria_jukyu_202608_06.csv",
          "https://www.kansai-td.co.jp/denkiyoho/download/csv/eria_jukyu_202608_06.csv",
          "https://www.kansai-td.co.jp/denkiyoho/download/eria_jukyu_202608_06.csv",
          "https://www.kansai-td.co.jp/yamasou/eria_jukyu_202608_06.csv",
          "https://www.kansai-td.co.jp/denkiyoho/area-performance/eria_jukyu_202608_06.csv"):
    csvhead(u)
# MOF English PDF
for f in ("2026/2026085e.pdf", "2025/2025127e.pdf", "2024/2024127e.pdf"):
    r = get("https://www.customs.go.jp/toukei/shinbun/trade-st_e/" + f)
    print("MOF", f, None if r is None else (r.status_code, len(r.content)))
    if r is not None and r.status_code == 200:
        d = pymupdf.open(stream=r.content, filetype="pdf")
        print("  pages", len(d))
        for i, pg in enumerate(d):
            tx = pg.get_text()
            j = tx.find("iquefied natural gas")
            if j >= 0:
                print("  page", i, repr(tx[max(0, j - 700): j + 400]))
                break
        print("  first page:", repr(d[0].get_text()[:300]))
# METI xlsx
for u in ("https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/2025/4-2025.xlsx",
          "https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/2025/4-2025n.xlsx",
          "https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/2024/4-2024.xlsx",
          "https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/2021/4-2021.xlsx",
          "https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/2016/4-2016.xlsx"):
    r = get(u)
    print(u, None if r is None else (r.status_code, len(r.content)))
    if r is not None and r.status_code == 200:
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
        print("  sheets", [ws.title for ws in wb.worksheets][:40])
        ws = wb.worksheets[1] if len(wb.worksheets) > 1 else wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            if row and row[2] in ("ＬＮＧ",):
                print("  ", row[:9])
        break
r = get("https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/2026/1-1-2026n.xlsx")
wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
ws = wb.worksheets[1]
rows = list(ws.iter_rows(min_row=1, max_row=6, values_only=True))
for c in range(10, 54):
    print(c, [rows[i][c] if c < len(rows[i]) else None for i in range(5)])
print("last rows:")
allr = list(ws.iter_rows(min_row=1700, values_only=True))
for row in allr[-6:]:
    print(row[:14])
