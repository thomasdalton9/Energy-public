"""Japan sources probe 3 (manual): compact re-check of the TSO area-supply-demand CSV locations (Tokyo, Tohoku, Chubu,
Hokuriku, Kansai, Kyushu, Okinawa), MOF customs download page and CSV heads, METI electricity survey xlsx heads."""
import io
import re
import requests

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"}


def get(u, headers=None):
    h = dict(UA)
    h.update(headers or {})
    try:
        return requests.get(u, headers=h, timeout=(8, 30))
    except Exception as e:  # noqa: BLE001
        print("   FAIL", type(e).__name__, str(e)[:80])


def csvhead(u, n=3, headers=None):
    r = get(u, headers)
    if r is None:
        return
    print(f"{u}\n   HTTP {r.status_code} {len(r.content)}b {r.headers.get('content-type')}")
    if r.status_code == 200 and r.content:
        for enc in ("cp932", "utf-8"):
            try:
                t = r.content.decode(enc)
                for line in t.splitlines()[:n]:
                    print("   |", line[:300])
                return
            except Exception:  # noqa: BLE001
                pass


def pagelinks(u, pat, maxn=25, headers=None):
    r = get(u, headers)
    if r is None:
        return
    print(f"PAGE {u}\n   HTTP {r.status_code} {len(r.content)}b")
    if r.status_code != 200:
        return
    t = r.content.decode(r.apparent_encoding or "utf-8", "replace")
    ls = [l for l in dict.fromkeys(re.findall(r'(?:href|src)=["\']?([^"\'#\s>]+)', t)) if re.search(pat, l, re.I)]
    for l in ls[:maxn]:
        print("    ", l)


# Tokyo with other approaches
for u in ["https://www.tepco.co.jp/forecast/html/images/eria_jukyu_202609_03.csv",
          "https://www.tepco.co.jp/forecast/html/images/juyo-2026.csv"]:
    csvhead(u, 3, {"Referer": "https://www.tepco.co.jp/forecast/html/area_jukyu-j.html", "Accept-Language": "ja"})
csvhead("https://www.tepco.co.jp/forecast/html/images/eria_jukyu_202609_03.csv", 3,
        {"User-Agent": "python-requests/2.31", "Accept": "*/*"})
# Tohoku
csvhead("https://setsuden.nw.tohoku-epco.co.jp/common/demand/eria_jukyu_202609_02.csv")
pagelinks("https://setsuden.nw.tohoku-epco.co.jp/download.html", r"csv|zip|xls")
pagelinks("https://www.tohoku-epco.co.jp/nw/", r"setsuden|jukyu|juyo")
# Chubu
csvhead("https://powergrid.chuden.co.jp/denki_yoho_content_data/download_csv/eria_jukyu_202609_04.csv")
pagelinks("https://powergrid.chuden.co.jp/denkiyoho/", r"csv|jukyu|download")
pagelinks("https://powergrid.chuden.co.jp/denki_yoho_content_data/", r"csv|jukyu")
# Hokuriku
csvhead("https://www.rikuden.co.jp/nw/denki-yoho/csv/eria_jukyu_202609_05.csv")
pagelinks("https://www.rikuden.co.jp/nw/denki-yoho/results_jukyu.html", r"csv|jukyu")
pagelinks("https://www.rikuden.co.jp/nw/denki-yoho/index.html", r"csv|jukyu|result")
# Kansai
csvhead("https://www.kansai-td.co.jp/yamasou/eria_jukyu_202609_06.csv")
pagelinks("https://www.kansai-td.co.jp/denkiyoho/download/", r"csv|jukyu|zip|xls")
pagelinks("https://www.kansai-td.co.jp/denkiyoho/area-performance/index.html", r"csv|jukyu|download")
# Kyushu
pagelinks("https://www.kyuden.co.jp/td_area_jukyu/jukyu.html", r"csv|jukyu|zip|xls")
csvhead("https://www.kyuden.co.jp/td_area_jukyu/csv/eria_jukyu_202609_09.csv")
csvhead("https://www.kyuden.co.jp/td_area_jukyu/csv/eria_jukyu_202609_09.csv", 3, {"Referer": "https://www.kyuden.co.jp/td_area_jukyu/jukyu.html"})
# Okinawa
for h in ("https://www.okinawa-epco.co.jp/", "https://www.okiden.co.jp/", "https://www.okiden.co.jp/shared/pdf/"):
    pagelinks(h, r"jukyu|csv|supply|yoho|denki", 10)
csvhead("https://www.okiden.co.jp/denki/csv/eria_jukyu_202609_10.csv")
# MOF customs
pagelinks("https://www.customs.go.jp/toukei/info/tsdl_e.htm", r"\.csv|\.zip|\.xls|\.htm", 80)
csvhead("https://www.customs.go.jp/toukei/suii/html/data/d41ma.csv", 6)
csvhead("https://www.customs.go.jp/toukei/suii/html/data/d42ma001.csv", 6)
pagelinks("https://www.customs.go.jp/toukei/shinbun/happyou_e.htm", r"\.csv|\.zip|\.xls|\.pdf|\.htm", 40)
pagelinks("https://www.customs.go.jp/toukei/info/tsdl.htm", r"\.csv|\.zip|\.xls|\.htm", 80)
# METI electricity survey xlsx
for f in ("3-1-2026n", "3-2-2026n", "1-1-2026n", "4-2026n"):
    u = f"https://www.enecho.meti.go.jp/statistics/electric_power/ep002/xls/2026/{f}.xlsx"
    r = get(u)
    if r is not None:
        print(u, r.status_code, len(r.content))
        if r.status_code == 200:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
            for ws in wb.worksheets[:2]:
                print("  sheet", ws.title, ws.max_row, ws.max_column)
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 14:
                        break
                    print("   ", [c for c in row[:12]])
