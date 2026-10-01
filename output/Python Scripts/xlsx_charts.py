"""
Shared helper: add a chart sheet with a native Excel chart to a workbook
that has already been written (e.g. by xlsx_notes.write_workbook).

Every xlsx this repo pulls should carry a chart of its data (see
CLAUDE.md). The chart sheet holds its own copy of the plotted table, so
it works whatever layout the data sheets use (wide, long, with an
index column...): pass a wide DataFrame indexed by date, one column per
series.

    import xlsx_charts
    xlsx_charts.add_chart_sheet(path, wide_df, "Bolivia gas demand by sector",
                                "million m3/day", kind="stacked_bar")

kind: "line", "stacked_area", "stacked_bar".
For storage/level series use water_year_chart.add_water_year_chart instead.
"""
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

# Validated categorical order (dataviz palette): fixed order, never cycled.
PALETTE = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]
OTHER_GREY = "9A9A9A"


def _fold_to_palette(df):
    """More series than palette slots -> keep the largest, fold the rest into 'Other'."""
    if df.shape[1] <= len(PALETTE):
        return df
    order = df.abs().mean().sort_values(ascending=False).index
    keep = list(order[: len(PALETTE) - 1])
    out = df[keep].copy()
    out["Other"] = df.drop(columns=keep).sum(axis=1, min_count=1)
    return out


def prepare(df, line_cols=()):
    """Clean a wide frame for charting; returns (frame, number of bar/area series before overlay lines)."""
    df = df.copy()
    df.index = pd.to_datetime(df.index)
    df = df.sort_index()
    df = df.apply(pd.to_numeric, errors="coerce")
    df = df.loc[:, df.notna().any()]
    if df.empty:
        return df, 0
    lines_df = df[[c for c in line_cols if c in df.columns]]
    body = _fold_to_palette(df.drop(columns=list(lines_df.columns)))
    return pd.concat([body, lines_df], axis=1), body.shape[1]


def write_table(ws, df, date_format="%Y-%m", start_row=1, start_col=1):
    """Write Date + series columns; dates are real Excel dates shown as mmm/yy (annual: yyyy)."""
    excel_fmt = "yyyy" if date_format == "%Y" else "mmm/yy"
    ws.cell(row=start_row, column=start_col, value="Date")
    for j, c in enumerate(df.columns):
        ws.cell(row=start_row, column=start_col + 1 + j, value=str(c))
    for i, (ts, row) in enumerate(df.iterrows(), start=1):
        cell = ws.cell(row=start_row + i, column=start_col, value=ts.to_pydatetime())
        cell.number_format = excel_fmt
        for j, v in enumerate(row.values):
            ws.cell(row=start_row + i, column=start_col + 1 + j, value=None if pd.isna(v) else float(v))
    ws.column_dimensions[ws.cell(row=1, column=start_col).column_letter].width = 11


def build_chart(ws, df, n_bars, title, y_title, kind="line", date_format="%Y-%m", start_row=1, start_col=1,
                width=28, height=13):
    """Native chart over a table written by write_table on sheet ws (the chart can be placed on any sheet)."""
    n = start_row + len(df)
    excel_fmt = "yyyy" if date_format == "%Y" else "mmm/yy"
    if kind == "stacked_bar":
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.gapWidth = 30
    elif kind == "stacked_area":
        chart = AreaChart()
        chart.grouping = "stacked"
    else:
        chart = LineChart()
    cats = Reference(ws, min_col=start_col, min_row=start_row + 1, max_row=n)
    chart.add_data(Reference(ws, min_col=start_col + 1, max_col=start_col + n_bars, min_row=start_row, max_row=n),
                   titles_from_data=True)
    chart.set_categories(cats)
    for i, s in enumerate(chart.series):
        colour = OTHER_GREY if str(df.columns[i]) == "Other" else PALETTE[i % len(PALETTE)]
        if kind == "line":
            s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=colour, w=22225))
            s.smooth = False
            s.marker.symbol = "none"
        else:
            s.graphicalProperties = GraphicalProperties(solidFill=colour, ln=LineProperties(noFill=True))
    chart.title = f"{title} (to {df.index.max().strftime('%Y' if date_format == '%Y' else '%b/%y')})"
    chart.x_axis.number_format = excel_fmt
    chart.y_axis.title = y_title
    chart.x_axis.tickLblPos = "low"
    spacing = df.index.to_series().diff().median().days if len(df) > 1 else 30
    if spacing <= 2:   # daily: label once every whole number of months so mmm/yy labels don't repeat
        step = 30 * max(1, -(-len(df) // (30 * 12)))
    else:
        step = max(1, len(df) // 12)
    chart.x_axis.tickLblSkip = step
    chart.x_axis.tickMarkSkip = step
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend.position = "b"
    # no borders: chart frame and plot area
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    if df.shape[1] > n_bars and kind != "line":
        over = LineChart()
        over.add_data(Reference(ws, min_col=start_col + n_bars + 1, max_col=start_col + df.shape[1],
                                min_row=start_row, max_row=n), titles_from_data=True)
        over.set_categories(cats)
        for s in over.series:
            s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="252525", w=28575))
            s.smooth = False
            s.marker.symbol = "none"
        over.y_axis.delete = True   # shares the bar chart's axis
        chart += over
    chart.width, chart.height = width, height
    return chart


def save_atomic(wb, path):
    root, ext = os.path.splitext(path)
    tmp = f"{root}.tmp{os.getpid()}{ext}"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def add_chart_sheet(path, df, title, y_title, kind="line", sheet_name="Chart", date_format="%Y-%m",
                    width=28, height=13, line_cols=()):
    """line_cols: columns of df drawn as lines over the bars/areas, on the same axis (same units only)."""
    df, n_bars = prepare(df, line_cols)
    if df.empty:
        return
    wb = load_workbook(path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 1 if len(wb.sheetnames) > 1 else None)
    write_table(ws, df, date_format)
    chart = build_chart(ws, df, n_bars, title, y_title, kind, date_format, width=width, height=height)
    ws.add_chart(chart, f"{chr(ord('A') + min(df.shape[1] + 2, 20))}2")
    save_atomic(wb, path)
