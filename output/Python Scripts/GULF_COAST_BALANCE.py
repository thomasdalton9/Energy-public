"""
Gulf Coast gas balance, Texas + Louisiana, monthly Bcf/d, history from 2021 and forecast to Dec 2030 -> gulf_coast_gas_balance.xlsx.
Question it answers: is there enough Permian gas for the new Gulf LNG, or does the Gulf need more Haynesville (or other-region) supply?
Runs in GitHub Actions (.github/workflows/gulf_coast_balance.yml; EIA_API_KEY secret). Reads, never writes, the Texas workbooks.

  TEXAS       texas_production_forecast.xlsx ('Supply and demand' to 2023, 'Demand to 2033' from 2024: consumption by sector, Mexico pipeline,
              BASE data-centre gas burn, dry production base / takeaway-delayed, STEO basin series, takeaway table) and the LNG feedgas model of
              TEXAS_GAS_FORECAST.py run to Dec 2030 (base and 6-month delayed). All Texas inputs and their status labels are the Texas workbooks'.
  LOUISIANA   EIA API v2 (natural-gas/cons/sum, prod/sum, move/poe2; area SLA): consumption by sector, marketed / dry production, LNG exports
              of Sabine Pass, Cameron, Calcasieu Pass and Plaquemines (all four are Louisiana ports; feedgas = exports x 1.09 as for Texas).
              Forecast: existing plants = peak nameplate (EIA liquefaction capacity file) x 1.09 x 3-year seasonal utilisation; new trains from the
              editable 'Assump - LA LNG' tab (date status SOURCED only where an opened document gives it, else UNVERIFIED); base and 6-month delay;
              consumption = seasonal trend; production = STEO Haynesville x a Louisiana share + calibrated remainder.
  HAYNESVILLE STEO Haynesville (NGMPHA) is East Texas + Louisiana. Split: Texas 30% (the Texas workbook's assumption, kept so the two states
              add up to STEO) and Louisiana 70%; EIA state data give the implied Louisiana share (LA marketed / STEO Haynesville) as a check.
  COMBINED    production, demand incl. LNG and implied net outflow of the two states together, and the key-question table: growth in Gulf LNG
              feedgas (and other Gulf demand) from Dec 2025 against growth in Permian (takeaway-capped), Eagle Ford and Haynesville supply, and the
              extra supply needed to hold the outflow, base and delayed.
Beyond Dec 2027 the basin supply is an extension of EIA STEO (2028 by the STEO 2027/2026 growth, 2029-30 damped x0.5 a year, as the Texas
workbook), not an EIA forecast; 2029-30 is lighter on the charts.

Usage: python3 GULF_COAST_BALANCE.py [--out "output/Data and Chart Outputs/gulf_coast_gas_balance.xlsx"] [--full] [--fetch-only]
"""
import argparse
import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import xlsx_notes  # noqa: E402
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import TEXAS_GAS as tg  # noqa: E402  (EIA helpers, read-only use)
import TEXAS_GAS_FORECAST as gf  # noqa: E402  (LNG model, read-only use)

OUT_DIR = os.path.join(ROOT, "output", "Data and Chart Outputs")
DEFAULT_OUT = os.path.join(OUT_DIR, "gulf_coast_gas_balance.xlsx")
TEXAS_XLSX = os.path.join(OUT_DIR, "texas_gas_monthly.xlsx")
PROD_XLSX = os.path.join(OUT_DIR, "texas_production_forecast.xlsx")
LNG_XLSX = os.path.join(OUT_DIR, "lng_feedgas_daily.xlsx")
START, END = pd.Timestamp("2021-01-01"), pd.Timestamp("2030-12-01")
EXT_FROM = pd.Timestamp("2029-01-01")          # basin supply beyond the Texas workbook's 2028 horizon: damped extension
RAW = "Raw MMcf"
SECT = ["Electric power", "Industrial", "Residential", "Commercial", "Vehicle fuel"]
RATIO = 1.09                                    # feedgas / LNG exports, as TEXAS_GAS_FORECAST.py
LA_PORTS = {"YSPL": "Sabine Pass", "YCAM": "Cameron", "YCCPL": "Calcasieu Pass", "YPLAQ": "Plaquemines"}
SRC_EIA = "EIA U.S. liquefaction capacity file 2026 Q2 (opened 5 Oct 2026)"
AL = "Assump - LA LNG"
FILL_IN = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
FILL_HEAD = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
BOLD = Font(bold=True)

# EIA peak nameplate capacity of the operating Louisiana plants, Bcf/d of LNG (EIA liquefaction capacity file 2026 Q2: per-train peak x trains)
EXISTING = {"Sabine Pass": (6 * 0.7588767123287671, "6 trains x 0.759 Bcf/d peak (EIA 2026 Q2); EIA Today in Energy 15 Sep 2026: 3.6 nominal / 4.6 peak"),
            "Cameron": (3 * 0.66, "3 trains x 0.66 Bcf/d peak (EIA 2026 Q2)"),
            "Calcasieu Pass": (2 * 0.7903397260273972, "2 train groups x 0.790 Bcf/d peak (EIA 2026 Q2)"),
            "Plaquemines": (1.89 + 1.96, "Phase 1 1.89 + Phase 2 1.96 Bcf/d peak (EIA 2026 Q2; Venture Global: peak > 28 MTPA once complete)")}

# New Louisiana projects: (plant, train, EIA peak nameplate Bcf/d LNG, first-LNG month or None, ramp months, ramp-start util, steady util,
# date status, note). Feedgas nameplate = peak x 1.09 (how the Texas table's feedgas figures relate to EIA peak: Golden Pass 0.795 -> 0.865).
LA_TRAINS = [
    ("CP2 (Venture Global)", "Phase 1, Trains 1-26", 2.86, "2027-10", 12, 0.15, 0.95, "SOURCED",
     "EIA 2026 Q2 file: under construction, in service 2H2027 (month = mid-half, ASSUMED); ramp 12 months ASSUMED (Plaquemines took ~1 year from first LNG to full output; EIA 1 Sep 2026: exporting at full capacity)"),
    ("CP2 (Venture Global)", "Phase 2, Trains 27-36", 1.10, "2029-04", 8, 0.15, 0.95, "SOURCED",
     "EIA 2026 Q2 file: under construction, 1H2029 (month = mid-half, ASSUMED); ramp ASSUMED"),
    ("Woodside Louisiana LNG", "Phase 1, Trains 1-3", 2.326520547945205, "2029-07", 18, 0.10, 0.95, "SOURCED",
     "EIA 2026 Q2 file: under construction, 2029 (month = mid-year, ASSUMED); three trains staggered, ramp 18 months ASSUMED (a search snippet of Woodside's Q2 2026 report says first LNG 2029, 28% complete - not opened)"),
    ("Commonwealth LNG", "Trains 1-6", 1.1845263157894736, "2030-07", 12, 0.10, 0.95, "SOURCED",
     "EIA 2026 Q2 file: under construction, 2030 (month = mid-year, ASSUMED); six trains, ramp 12 months ASSUMED"),
    ("Delfin FLNG", "Vessel 1", 0.6, "2030-07", 4, 0.30, 0.95, "SOURCED",
     "EIA 2026 Q2 file: under construction, 2030 (month = mid-year, ASSUMED); floating vessel; DOE extended its start deadline to 1 Jun 2031"),
    # pre-FID or without a sourced date: no first-LNG month, so nothing is forecast until the owner types one
    ("Cameron LNG", "Train 4", 0.96, None, 8, 0.15, 0.95, "UNVERIFIED",
     "EIA 2026 Q2 'Approved' tab: new single larger train, EPC awarded to Bechtel, FID TBA; DOE start deadline Mar 2033. No date -> not in the forecast"),
    ("Woodside Louisiana LNG", "Phase 2, Trains 4-5", 1.4542250958, None, 12, 0.10, 0.95, "UNVERIFIED",
     "EIA 2026 Q2 'Approved' tab: limited notice to proceed, FID TBA. No date -> not in the forecast"),
    ("Lake Charles LNG (Energy Transfer)", "Trains 1-3", 2.1734342465753427, None, 12, 0.10, 0.95, "UNVERIFIED",
     "EIA 2026 Q2 'Approved' tab: Energy Transfer announced on 18 Dec 2025 it suspended development (open to third parties). No date -> not in the forecast"),
    ("Delfin FLNG", "Vessels 2-3", 1.0533333333, None, 4, 0.30, 0.95, "UNVERIFIED",
     "EIA 2026 Q2 'Approved' tab: EPC awarded to Samsung Heavy / Black & Veatch, FID TBA. No date -> not in the forecast"),
    ("Sabine Pass Stage 5 (Cheniere)", "Expansion", None, None, 12, 0.10, 0.95, "UNVERIFIED",
     "FERC page (opened 5 Oct 2026): FERC staff issued the final EIS on 25 Sep 2026. Not in EIA's 2026 Q2 file; capacity, FID and dates not found -> not in the forecast"),
    ("Plaquemines expansion (Venture Global)", "Expansion", None, None, 12, 0.10, 0.95, "UNVERIFIED",
     "Venture Global site (opened 5 Oct 2026): 'additional plans for expansion'. Not in EIA's file; no capacity or dates -> not in the forecast"),
]


# --------------------------------------------------------------------------- Louisiana EIA pull
def fetch(store, full):
    new = {}

    def put(col, period, v):
        if v is not None:
            new.setdefault(col, {})[pd.Timestamp(period + "-01")] = v

    for prefix, route, procs in (("cons|", "cons/sum/data/", tg.CONS), ("prod|", "prod/sum/data/", tg.PROD),
                                 ("stor|", "stor/sum/data/", tg.STOR)):
        st = tg.start_for(store, prefix, full)
        data = tg.rows(route, {"duoarea": ["SLA"], "process": list(procs)}, st)
        print(f"  LA {prefix} from {st}: {len(data)} rows", flush=True)
        for r in data:
            put(prefix + procs[r["process"]], r["period"], tg.volume(r))
    st = tg.start_for(store, "ENG|", full)
    sids = [f"NGM_EPG0_ENG_{k}-Z00_MMCF" for k in LA_PORTS]
    data = tg.rows("move/poe2/data/", {"series": sids}, st)
    print(f"  LA ENG| from {st}: {len(data)} rows", flush=True)
    for r in data:
        m = re.match(r"NGM_EPG0_ENG_(Y\w+?)-Z00_MMCF", str(r.get("series", "")))
        if m and m[1] in LA_PORTS:
            put("ENG|" + LA_PORTS[m[1]], r["period"], tg.volume(r))
    add = pd.DataFrame({c: pd.Series(v) for c, v in new.items()}).sort_index()
    if store.empty:
        merged = add
    else:
        merged = store.reindex(store.index.union(add.index)).reindex(columns=store.columns.union(add.columns))
        merged.update(add)
    merged.index.name = "Month"
    return merged.sort_index().dropna(how="all")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--full", action="store_true")
    ap.add_argument("--fetch-only", action="store_true")
    ap.add_argument("--no-fetch", action="store_true", help="use the saved Louisiana store (local testing)")
    a = ap.parse_args()
    store = tg.load_store(a.out)
    raw = store if a.no_fetch else fetch(store, a.full)
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    if a.fetch_only:
        xlsx_notes.write_workbook(a.out, {RAW: raw}, ["Notes", "Fetch-only run of americas/GULF_COAST_BALANCE.py: Louisiana EIA store only."], set())
        print(raw.tail(6).T.to_string())
        return
    run(a.out, raw)


# --------------------------------------------------------------------------- helpers
def _clean(df):
    return df.loc[:, [c for c in df.columns if not str(c).startswith("Unnamed")]]


def read_sheet(path, sheet, month_col="Month"):
    d = _clean(pd.read_excel(path, sheet_name=sheet))
    d[month_col] = pd.to_datetime(d[month_col])
    return d.set_index(month_col).sort_index()


def smooth121(vals):
    return [(vals[i - 1] + 2 * vals[i] + vals[(i + 1) % 12]) / 4 for i in range(12)]


def month_profile(util):
    """Calendar-month mean utilisation (12 values), smoothed 1-2-1 like the Texas model."""
    return smooth121([float(util[util.index.month == m].mean()) for m in range(1, 13)])


def prior_table(out, sheet, keys):
    """Editable table of the committed workbook as {key tuple: row Series}; empty when absent."""
    try:
        d = pd.read_excel(out, sheet_name=sheet)
    except Exception:  # noqa: BLE001
        return {}
    return {tuple(r[k] for k in keys): r for _, r in d.iterrows()}


# --------------------------------------------------------------------------- Texas (read only)
def texas_inputs():
    sd = read_sheet(PROD_XLSX, "Supply and demand")
    d33 = read_sheet(PROD_XLSX, "Demand to 2033")
    fcast = read_sheet(PROD_XLSX, "Forecast")
    assum = _clean(pd.read_excel(PROD_XLSX, sheet_name="Assumptions"))
    take = _clean(pd.read_excel(PROD_XLSX, sheet_name="Takeaway")).dropna(subset=["Pipeline"])
    take = take[["Pipeline", "Capacity Bcf/d", "In-service month", "Status", "Source / verification"]]
    par = {r["Assumption"]: float(r["Value"]) for _, r in assum.iterrows()}
    g = lambda k: [v for kk, v in par.items() if kk.startswith(k)][0]   # noqa: E731
    shares = dict(permian=g("Permian Texas share"), hay=g("Haynesville Texas share"), ef=g("Eagle Ford Texas share"),
                  local=g("Local Permian demand"), delay=int(g("Takeaway delay")))
    wb = load_workbook(PROD_XLSX, read_only=False)
    a = wb["Assump - Data centres"]
    damp = next(float(a.cell(r, 2).value) for r in range(1, 80) if str(a.cell(r, 1).value or "").startswith("Production extension after 2028"))
    dci = str(d33["Data-centre inputs"].dropna().iloc[0]) if "Data-centre inputs" in d33 and d33["Data-centre inputs"].notna().any() else ""
    tx = read_sheet(TEXAS_XLSX, "Balance")
    ratio = 1 - float((tx["Dry production"] / tx["Marketed production"]).dropna().iloc[-12:].mean())
    other = float(fcast["Other Texas / calibration (held flat)"].dropna().iloc[-1])
    return dict(sd=sd, d33=d33, fcast=fcast, take=take, shares=shares, damp=damp, dc_inputs=dci, loss_tx=ratio, other_tx=other)


def texas_lng(delay_check):
    """LNG feedgas, base and delayed, to Dec 2030 from the Texas model (TEXAS_GAS_FORECAST.py, run read-only to END)."""
    cons = read_sheet(TEXAS_XLSX, "Consumption by sector")
    exp = read_sheet(TEXAS_XLSX, "Exports")
    bal = read_sheet(TEXAS_XLSX, "Balance")
    old = gf.LAST
    gf.LAST = END
    try:
        f, ctx = gf.build(TEXAS_XLSX, cons, exp, bal)
    finally:
        gf.LAST = old
    v = gf.values_sheet(f)
    return v["LNG feedgas total - base"].astype(float), v["LNG feedgas total - delayed"].astype(float), ctx["fs"], v


def texas_block(T, idx):
    sd, d33 = T["sd"], T["d33"]
    lb, ld, fs_tx, _ = texas_lng(None)
    tx = pd.DataFrame(index=idx)
    old = idx < d33.index.min()
    for c in SECT + ["Pipeline exports to Mexico"]:
        tx[c] = d33[c].reindex(idx).where(~old, sd[c].reindex(idx))
    tx["Data centres (BASE added gas burn)"] = d33["Data centres added gas burn, BASE"].reindex(idx).fillna(0.0)
    tx["LNG feedgas (base)"] = lb.reindex(idx)
    tx["LNG feedgas (delayed)"] = ld.reindex(idx)
    tx["Dry production, base"] = d33["Dry production, base"].reindex(idx).where(~old, sd["Dry production, base"].reindex(idx))
    tx["Dry production, takeaway delayed"] = d33["Dry production, takeaway delayed"].reindex(idx).where(
        ~old, sd["Dry production, takeaway delayed"].reindex(idx))
    tx["Type"] = np.where(idx < fs_tx, "Actual", np.where(idx < EXT_FROM, "Forecast", "Scenario"))
    return tx, fs_tx


def texas_check(tx, d33):
    t = d33["Total demand incl. data centres, BASE"].reindex(tx.index)
    mine = tx[SECT + ["Pipeline exports to Mexico", "Data centres (BASE added gas burn)", "LNG feedgas (base)"]].sum(axis=1)
    e = (mine - t).dropna()
    return f"Texas demand incl. LNG and data centres vs 'Demand to 2033' BASE total 2024-30: max abs gap {e.abs().max():.3f} Bcf/d"


# --------------------------------------------------------------------------- basin supply (Texas workbook inputs, damped extension)
def basin_supply(T, idx, la_share, loss_la):
    f, sh = T["fcast"], T["shares"]
    cols = {"Permian": "STEO Permian marketed (Bcf/d)", "Eagle Ford": "STEO Eagle Ford marketed (Bcf/d)",
            "Haynesville": "STEO Haynesville marketed (Bcf/d)"}
    s = pd.DataFrame({k: f[c].reindex(idx) for k, c in cols.items()})
    for k in list(s.columns):
        g = s.loc["2027-01-01":"2027-12-01", k].mean() / s.loc["2026-01-01":"2026-12-01", k].mean() - 1
        for m in idx[idx >= EXT_FROM]:
            s.loc[m, k] = s.loc[m - pd.DateOffset(months=1), k] * (1 + g * T["damp"] ** (m.year - 2028)) ** (1 / 12)
    take = T["take"]
    add = take[~take["Pipeline"].str.contains("Existing", case=False) & take["Status"].astype(str).str.lower().ne("in service")]
    base_ex = float(take[take["Pipeline"].str.contains("Existing", case=False)]["Capacity Bcf/d"].sum())
    caps = {}
    for nm, d in (("base", 0), ("delayed", sh["delay"])):
        c = pd.Series(base_ex + sh["local"], index=idx, dtype=float)
        for _, r in add.iterrows():
            c[c.index >= pd.Timestamp(str(r["In-service month"])[:7] + "-01") + pd.DateOffset(months=d)] += float(r["Capacity Bcf/d"])
        caps[nm] = c
    out = pd.DataFrame(index=idx)
    for k in s:
        out[f"STEO {k} marketed"] = s[k]
    for sc in ("base", "delayed"):
        out[f"Permian takeaway + local demand, {sc}"] = caps[sc]
        out[f"Permian marketed capped, {sc}"] = pd.concat([s["Permian"], caps[sc]], axis=1).min(axis=1)
    k_tx, k_la = 1 - T["loss_tx"], 1 - loss_la
    for sc in ("base", "delayed"):
        out[f"Permian dry (Texas share), {sc}"] = sh["permian"] * out[f"Permian marketed capped, {sc}"] * k_tx
    out["Eagle Ford dry"] = sh["ef"] * s["Eagle Ford"] * k_tx
    out["Haynesville dry, East Texas"] = sh["hay"] * s["Haynesville"] * k_tx
    out["Haynesville dry, Louisiana"] = la_share * s["Haynesville"] * k_la
    out["Haynesville dry, total"] = out["Haynesville dry, East Texas"] + out["Haynesville dry, Louisiana"]
    return out, s


def basin_check(T, bs, idx):
    sh = T["shares"]
    chk = []
    for sc, col in (("base", "Dry production, base"), ("delayed", "Dry production, takeaway delayed")):
        mine = (bs[f"Permian dry (Texas share), {sc}"] + bs["Eagle Ford dry"] + bs["Haynesville dry, East Texas"]
                + T["other_tx"] * (1 - T["loss_tx"]))
        ref = T["d33"][col].reindex(idx)
        fc = (idx >= pd.Timestamp("2026-09-01"))
        e = (mine - ref)[fc].dropna()
        chk.append(f"{sc}: max abs gap {e.abs().max():.3f} Bcf/d")
    return "Texas dry production rebuilt from the basin series vs 'Demand to 2033' (Sep 2026-Dec 2030) " + "; ".join(chk)


# --------------------------------------------------------------------------- Louisiana
def la_history(raw):
    d = tg.bcfd(raw)
    c = lambda p: d[[x for x in d.columns if x.startswith(p)]].rename(columns=lambda x: x.split("|", 1)[1])   # noqa: E731
    cons, eng, prod = c("cons|"), c("ENG|"), c("prod|")
    return cons, eng, prod, c("stor|")


def la_trains(out):
    """Train table: script defaults, then the committed workbook's editable columns win (rows the owner added are kept)."""
    prior = prior_table(out, AL, ["Plant", "Train"])
    rows = {}
    for p, t, peak, first, ramp, rs, st, status, note in LA_TRAINS:
        rows[(p, t)] = dict(peak=peak, first=pd.Timestamp(first + "-01") if first else pd.NaT, ramp=ramp, rs=rs, st=st, status=status, note=note)
    for k, r in prior.items():
        if k in rows:
            o = rows[k]
            for key, col in (("peak", "EIA peak nameplate, Bcf/d LNG"), ("ramp", "Ramp months"), ("rs", "Ramp-start utilisation"), ("st", "Steady utilisation")):
                if col in r and pd.notna(r[col]):
                    o[key] = float(r[col])
            if "First-LNG month" in r:
                v = pd.to_datetime(r["First-LNG month"], errors="coerce")
                if pd.notna(v) != pd.notna(o["first"]) or (pd.notna(v) and v != o["first"]):
                    o["first"] = v          # owner edited the date; status stays the script's unless the owner changes it
                    if "Date status" in r and pd.notna(r["Date status"]):
                        o["status"] = str(r["Date status"])
        elif pd.notna(r.get("Plant")):
            rows[k] = dict(peak=float(r["EIA peak nameplate, Bcf/d LNG"]) if pd.notna(r["EIA peak nameplate, Bcf/d LNG"]) else np.nan,
                           first=pd.to_datetime(r["First-LNG month"], errors="coerce"), ramp=float(r["Ramp months"]), rs=float(r["Ramp-start utilisation"]),
                           st=float(r["Steady utilisation"]), status=str(r["Date status"]), note=str(r["Source / note"]))
    return rows


def la_lng(eng, trains, idx, fs, delay):
    """LNG feedgas by plant to END: history EIA exports x 1.09; forecast = existing plants (peak x 1.09 x utilisation profile) + new trains."""
    hist = eng.reindex(idx) * RATIO
    win = hist.index[(hist.index < fs)][-36:]
    prof, util, cap = {}, {}, {}
    for pl, (peak, _) in EXISTING.items():
        cap[pl] = peak * RATIO
    nb_src = []
    for pl in ("Sabine Pass", "Cameron", "Calcasieu Pass"):
        w = win[win >= pd.Timestamp("2023-11-01")] if pl == "Calcasieu Pass" else win
        u = hist.loc[w, pl] / cap[pl]
        util[pl] = u
        prof[pl] = month_profile(u)
        nb_src.append((prof[pl], cap[pl]))
    nb = [sum(p[i] * c for p, c in nb_src) / sum(c for _, c in nb_src) for i in range(12)]
    mnb = float(np.mean(nb))
    nb = [v / mnb for v in nb]
    pw = hist.index[(hist.index >= pd.Timestamp("2026-01-01")) & (hist.index < fs)]
    plaq_u = float((hist.loc[pw, "Plaquemines"] / cap["Plaquemines"]).mean())
    util["Plaquemines"] = hist.loc[pw, "Plaquemines"] / cap["Plaquemines"]
    prof["Plaquemines"] = [plaq_u * v for v in nb]
    out = {}
    for sc, dly in (("base", 0), ("delayed", delay)):
        f = pd.DataFrame(0.0, index=idx, columns=list(EXISTING) + ["New Louisiana plants"])
        for m in idx:
            if m < fs:
                for pl in EXISTING:
                    f.at[m, pl] = hist.at[m, pl] if pd.notna(hist.at[m, pl]) else 0.0
                continue
            for pl in EXISTING:
                f.at[m, pl] = cap[pl] * prof[pl][m.month - 1]
            tot = 0.0
            for k, t in trains.items():
                st = t["first"]
                if pd.isna(st) or pd.isna(t["peak"]):
                    continue
                if dly and st >= fs:
                    st = st + pd.DateOffset(months=dly)
                if m < st:
                    continue
                n = (m.year - st.year) * 12 + m.month - st.month
                tot += t["peak"] * RATIO * (t["rs"] + (t["st"] - t["rs"]) * min(1.0, n / max(t["ramp"], 1))) * nb[m.month - 1]
            f.at[m, "New Louisiana plants"] = tot
        out[sc] = f
    return out, cap, prof, util, nb


def la_block(raw, T, out_path, delay, idx):
    cons, eng, prod, stor = la_history(raw)
    trains = la_trains(out_path)
    must = [cons["Electric power"], cons["Industrial"], prod["Marketed production"]] + [eng[p] for p in EXISTING]
    last_full = min(s.dropna().index.max() for s in must)
    fs = last_full + pd.offsets.MonthBegin(1)
    lng, cap, prof, util, nb = la_lng(eng, trains, idx, fs, delay)
    la = pd.DataFrame(index=idx)
    est = []
    for c in SECT:
        s = cons[c].reindex(idx)
        last = s.dropna().index.max()
        fut = idx[idx > last]
        fcs, g = gf.seasonal_trend(s, fut)
        s = s.interpolate(limit_area="inside")          # EIA leaves single months blank (withheld): bridge them
        la[c] = s.where(idx <= last).fillna(fcs)
        if last < fs - pd.offsets.MonthBegin(1):
            est.append(f"{c} after {last:%b/%y}")
    la["Type"] = np.where(idx < fs, "Actual", np.where(idx < EXT_FROM, "Forecast", "Scenario"))
    la["LNG feedgas (base)"] = lng["base"].sum(axis=1)
    la["LNG feedgas (delayed)"] = lng["delayed"].sum(axis=1)
    # production: marketed history (EIA), dry = EIA to its last month, then marketed x (1 - loss share)
    p = prod.reindex(idx)
    ratio = (1 - (p["Dry production"] / p["Marketed production"]).dropna().iloc[-12:].mean())
    la["Marketed production (history)"] = p["Marketed production"].where(idx < fs)
    la["Dry production (history)"] = p["Dry production"].where(idx < fs).fillna(p["Marketed production"].where(idx < fs) * (1 - ratio))
    stor_ = stor.reindex(idx)["Net withdrawals"]
    la["Net storage withdrawal (memo, not in balance)"] = stor_.where(idx < fs)
    return la, lng, trains, fs, ratio, cons, eng, prod, prof, util, cap, est


def la_supply(la, steo_hay, idx, fs, la_share, ratio, s_other_months=6):
    """Louisiana marketed = share x STEO Haynesville + remainder calibrated on the last months of history (held flat)."""
    comp = la_share * steo_hay
    ov = idx[(idx < fs)][-s_other_months:]
    other = float((la["Marketed production (history)"].reindex(ov) - comp.reindex(ov)).mean())
    mk = la["Marketed production (history)"].fillna(comp + other)
    return mk, other


# --------------------------------------------------------------------------- the whole model
def run(out, raw):
    idx = pd.date_range(START, END, freq="MS")
    T = texas_inputs()
    delay = int(T["shares"]["delay"])
    prior_a = prior_table(out, "Assumptions", ["Item"])

    def edit(item, default):
        r = prior_a.get((item,))
        return float(r["Value"]) if r is not None and pd.notna(r.get("Value")) else default

    la, lng, trains, fs_la, loss_la, cons, eng, prod, prof, util, cap, est = la_block(raw, T, out, delay, idx)
    tx, fs_tx = texas_block(T, idx)
    fs = max(fs_tx, fs_la)
    sh_tx = T["shares"]["hay"]
    la_share = edit("Louisiana share of STEO Haynesville", round(1 - sh_tx, 4))
    out_chg = edit("Outflow change allowed vs Dec 2025 (Bcf/d; 0 = hold)", 0.0)
    floor = edit("Outflow floor for the 'outflow declines' case (Bcf/d)", 0.0)
    bs, steo = basin_supply(T, idx, la_share, loss_la)
    la["Marketed production"], other_la = la_supply(la, steo["Haynesville"], idx, fs_la, la_share, loss_la)
    la["Dry production"] = la["Dry production (history)"].fillna(la["Marketed production"] * (1 - loss_la))
    la["Type"] = np.where(idx < fs_la, "Actual", np.where(idx < EXT_FROM, "Forecast", "Scenario"))
    # implied Louisiana share of STEO Haynesville from EIA state marketed production
    imp = (prod["Marketed production"].reindex(idx) / steo["Haynesville"]).where(idx >= pd.Timestamp("2024-01-01")).dropna()
    imp = imp[imp.index < fs_la]
    imp12 = float(imp.iloc[-12:].mean())
    # ---------------- frames
    oth = lambda d: d[["Residential", "Commercial", "Vehicle fuel"]].sum(axis=1)   # noqa: E731
    la["Residential, commercial, vehicle fuel"] = oth(la)
    tx["Residential, commercial, vehicle fuel"] = oth(tx)
    la["Demand excl. LNG"] = la[SECT].sum(axis=1)
    la["Demand incl. LNG, base"] = la["Demand excl. LNG"] + la["LNG feedgas (base)"]
    la["Demand incl. LNG, LNG delayed"] = la["Demand excl. LNG"] + la["LNG feedgas (delayed)"]
    la["Implied net outflow, base"] = la["Dry production"] - la["Demand incl. LNG, base"]
    la["Implied net outflow, delayed"] = la["Dry production"] - la["Demand incl. LNG, LNG delayed"]
    tx["Demand excl. LNG"] = tx[SECT].sum(axis=1) + tx["Pipeline exports to Mexico"] + tx["Data centres (BASE added gas burn)"]
    tx["Demand incl. LNG, base"] = tx["Demand excl. LNG"] + tx["LNG feedgas (base)"]
    tx["Demand incl. LNG, LNG delayed"] = tx["Demand excl. LNG"] + tx["LNG feedgas (delayed)"]
    tx["Implied net outflow, base"] = tx["Dry production, base"] - tx["Demand incl. LNG, base"]
    tx["Implied net outflow, delayed"] = tx["Dry production, takeaway delayed"] - tx["Demand incl. LNG, LNG delayed"]
    cb = pd.DataFrame(index=idx)
    cb["Type"] = np.where(idx < fs, "Actual", np.where(idx < EXT_FROM, "Forecast", "Scenario"))
    cb["Electric power"] = tx["Electric power"] + la["Electric power"]
    cb["Industrial"] = tx["Industrial"] + la["Industrial"]
    cb["Residential, commercial, vehicle fuel"] = oth(tx) + oth(la)
    cb["Pipeline exports to Mexico (Texas)"] = tx["Pipeline exports to Mexico"]
    cb["Data centres (Texas, BASE added burn)"] = tx["Data centres (BASE added gas burn)"]
    cb["LNG feedgas, Texas (base)"] = tx["LNG feedgas (base)"]
    cb["LNG feedgas, Louisiana (base)"] = la["LNG feedgas (base)"]
    cb["LNG feedgas, Texas (delayed)"] = tx["LNG feedgas (delayed)"]
    cb["LNG feedgas, Louisiana (delayed)"] = la["LNG feedgas (delayed)"]
    cb["LNG feedgas, total (base)"] = cb["LNG feedgas, Texas (base)"] + cb["LNG feedgas, Louisiana (base)"]
    cb["LNG feedgas, total (delayed)"] = cb["LNG feedgas, Texas (delayed)"] + cb["LNG feedgas, Louisiana (delayed)"]
    cb["Demand excl. LNG"] = cb[["Electric power", "Industrial", "Residential, commercial, vehicle fuel", "Pipeline exports to Mexico (Texas)",
                                 "Data centres (Texas, BASE added burn)"]].sum(axis=1)
    cb["Demand incl. LNG, base"] = cb["Demand excl. LNG"] + cb["LNG feedgas, total (base)"]
    cb["Demand incl. LNG, LNG delayed"] = cb["Demand excl. LNG"] + cb["LNG feedgas, total (delayed)"]
    cb["Dry production, Texas (base)"] = tx["Dry production, base"]
    cb["Dry production, Louisiana"] = la["Dry production"]
    cb["Dry production, base"] = tx["Dry production, base"] + la["Dry production"]
    cb["Dry production, takeaway delayed"] = tx["Dry production, takeaway delayed"] + la["Dry production"]
    cb["Implied net outflow, base"] = cb["Dry production, base"] - cb["Demand incl. LNG, base"]
    cb["Implied net outflow, delayed"] = cb["Dry production, takeaway delayed"] - cb["Demand incl. LNG, LNG delayed"]
    # ---------------- key question
    kq, kchart, kmeta = key_question(cb, bs, tx, la, T, out_chg, floor, la_share, loss_la, idx, steo)
    # ---------------- sheets
    write(out, raw, idx, T, tx, la, cb, bs, steo, kq, kchart, kmeta, trains, prof, util, cap, est, imp, imp12, la_share, other_la,
          loss_la, fs, fs_tx, fs_la, out_chg, floor, edit)


def key_question(cb, bs, tx, la, T, out_chg, floor, la_share, loss_la, idx, steo):
    d0 = pd.Timestamp("2025-12-01")
    rows = []
    o0 = float(cb.loc[d0, "Implied net outflow, base"])
    recent = cb[cb["Type"].eq("Actual")]["Implied net outflow, base"].iloc[-12:].mean()
    for sc, lt, lc, perm in (("Base", "base", "base", "base"), ("LNG and takeaway delayed 6 months", "delayed", "delayed", "delayed")):
        for y in (2026, 2027, 2028, 2029, 2030):
            m = pd.Timestamp(f"{y}-12-01")
            dl = lambda s: float(s.loc[m] - s.loc[d0])   # noqa: E731
            d_lng_tx = dl(cb[f"LNG feedgas, Texas ({lt})"])
            d_lng_la = dl(cb[f"LNG feedgas, Louisiana ({lt})"])
            d_oth = dl(cb["Demand excl. LNG"])
            d_dem = d_lng_tx + d_lng_la + d_oth
            d_perm = dl(bs[f"Permian dry (Texas share), {perm}"])
            d_ef = dl(bs["Eagle Ford dry"])
            d_hay = dl(bs["Haynesville dry, total"])
            d_sup = d_perm + d_ef + d_hay
            gap = d_sup - d_dem                      # + surplus / - shortfall at an unchanged outflow
            extra = max(0.0, out_chg - gap)      # outflow must end at Dec-2025 level + out_chg (0 = hold; negative = may fall)
            out_y = o0 + gap
            cap_ = bs[f"Permian marketed capped, {perm}"]
            nm = (1 - T["shares"]["permian"]) * dl(cap_) * (1 - T["loss_tx"])
            rows.append({"Scenario": sc, "Year-end": f"Dec {y}", "Incremental LNG feedgas, Texas": d_lng_tx,
                         "Incremental LNG feedgas, Louisiana": d_lng_la, "Incremental LNG feedgas, Gulf total": d_lng_tx + d_lng_la,
                         "Incremental other Gulf demand (sectors, Mexico, data centres)": d_oth, "Incremental total Gulf demand": d_dem,
                         "Incremental Permian supply (Texas share, takeaway-capped)": d_perm, "Incremental Eagle Ford supply": d_ef,
                         "Incremental Haynesville supply (East Texas + Louisiana, STEO)": d_hay, "Incremental basin supply, total": d_sup,
                         "Supply growth less demand growth (+ surplus / - shortfall)": gap,
                         "Extra supply required to hold the outflow (Haynesville above STEO or other regions)": extra,
                         "Haynesville growth required to balance (STEO + extra)": d_hay + extra,
                         "Implied net outflow at Dec if no extra supply (outflow declines)": out_y,
                         "Fall in outflow vs Dec 2025 (other regions must replace it)": max(0.0, -gap),
                         "Extra supply needed to keep outflow above the floor": max(0.0, floor - out_y),
                         "Memo: Permian New Mexico growth (not counted)": nm})
    kq = pd.DataFrame(rows)
    # chart table: Dec of each year, base case, supply stack = demand growth when there is a shortfall
    b = kq[kq["Scenario"].eq("Base")].copy()
    kc = pd.DataFrame({
        "Permian (Texas share, takeaway-capped)": b["Incremental Permian supply (Texas share, takeaway-capped)"].values,
        "Eagle Ford": b["Incremental Eagle Ford supply"].values,
        "Haynesville (East Texas + Louisiana, STEO)": b["Incremental Haynesville supply (East Texas + Louisiana, STEO)"].values,
        "Extra supply required (Haynesville above STEO or other regions)": b["Extra supply required to hold the outflow (Haynesville above STEO or other regions)"].values,
        "Incremental Gulf LNG feedgas": b["Incremental LNG feedgas, Gulf total"].values,
        "Incremental total Gulf demand (LNG + other)": b["Incremental total Gulf demand"].values},
        index=[2026, 2027, 2028, 2029, 2030])
    kc.index.name = "Year"
    return kq, kc, dict(o0=o0, recent=float(recent))


# --------------------------------------------------------------------------- workbook
def assumption_rows(T, la_share, imp12, loss_la, loss_tx, other_la, out_chg, floor, delay, trains, n_src, n_tr, cap, fs_la):
    sh = T["shares"]
    R = []   # (item, value, unit, status, editable, source / note)

    def add(item, value, unit, status, editable, note):
        R.append((item, value, unit, status, "yes" if editable else "", note))
    add("Delayed scenario: months of delay (LNG trains not yet started; Permian takeaway pipelines)", delay, "months", "ASSUMPTION", False,
        "Same 6-month case as the Texas workbooks (value read from texas_production_forecast.xlsx 'Assumptions'; edit it there).")
    add("Feedgas / LNG export ratio", RATIO, "x", "ASSUMPTION", False,
        "Liquefaction fuel and shrink; the ratio TEXAS_GAS_FORECAST.py and lng_feedgas_daily.xlsx use (kept so Texas and Louisiana are like for like).")
    add("Louisiana share of STEO Haynesville", la_share, "share", "ASSUMPTION", True,
        f"= 1 - Texas share ({sh['hay']:.0%}, the Texas workbook's assumption) so the two states add up to STEO Haynesville. Check from EIA state data: Louisiana marketed production / STEO Haynesville "
        f"= {imp12:.0%} (mean of the last 12 actual months; an upper bound for Louisiana, which also holds non-Haynesville onshore gas).")
    add("Implied Louisiana share of STEO Haynesville (EIA state marketed production / STEO)", round(imp12, 4), "share", "DERIVED", False,
        "EIA Natural Gas Monthly Louisiana marketed production (SOURCED) over STEO Haynesville (SOURCED). Monthly series on 'Haynesville split'.")
    add("Louisiana marketed production remainder (not Haynesville-driven), held flat", round(other_la, 4), "Bcf/d", "DERIVED", False,
        "Mean gap between Louisiana marketed production and share x STEO Haynesville over the last 6 actual months.")
    add("Louisiana extraction loss (1 - dry / marketed)", round(loss_la, 4), "share", "DERIVED", False,
        "EIA dry / marketed production, last 12 published months (EIA dry production stops after Dec 2024); later dry = marketed x (1 - this).")
    add("Texas extraction loss", round(loss_tx, 4), "share", "DERIVED", False, "As TEXAS_PRODUCTION_FORECAST.py (EIA, last 12 published months).")
    add("Permian Texas share of STEO Permian", sh["permian"], "share", "ASSUMPTION", False,
        "Texas workbook. STEO Permian = Texas + southeast New Mexico; the New Mexico part is not counted in the Gulf supply (memo column on 'Key question').")
    add("Eagle Ford Texas share", sh["ef"], "share", "ASSUMPTION", False, "Texas workbook.")
    add("Local Permian demand added to takeaway in the cap", sh["local"], "Bcf/d", "UNVERIFIED", False, "Texas workbook (unverified).")
    add("Production extension after Dec 2028: growth multiplier per year", T["damp"], "x", "ASSUMPTION", False,
        "Texas workbook ('Assump - Data centres'): STEO 2027/2026 growth damped x0.5 a year from 2029. Extension, not STEO.")
    add("Outflow change allowed vs Dec 2025 (Bcf/d; 0 = hold)", out_chg, "Bcf/d", "ASSUMPTION", True,
        "Key question, 'outflow held' case: the combined outflow to other states must end each year at its Dec 2025 level plus this number (negative lets it fall).")
    add("Outflow floor for the 'outflow declines' case (Bcf/d)", floor, "Bcf/d", "ASSUMPTION", True,
        "Key question, 'outflow declines' case: extra supply is needed only if the implied outflow would fall below this.")
    for pl, (peak, note) in EXISTING.items():
        add(f"{pl}: EIA peak nameplate (LNG output)", round(peak, 3), "Bcf/d", "SOURCED", False,
            f"{note}. Feedgas nameplate = x {RATIO}; forecast = nameplate x 3-year seasonal utilisation of EIA exports x {RATIO} (Plaquemines: mean of Jan-Jul 2026, once at full output).")
    add("Louisiana new-train start dates: sourced / forecast trains", f"{n_src} of {n_tr}", "trains", "SOURCED" if n_src == n_tr else "UNVERIFIED", False,
        f"{SRC_EIA}: year or half-year per project; the month inside it is ASSUMED. Pre-FID projects have no date and are not forecast. Detail on '{AL}'.")
    add("Louisiana consumption forecast", "seasonal trend", "", "DERIVED", False,
        "Each sector: last 3 years' monthly profile x last-12-month level x damped trend (TEXAS_GAS_FORECAST.seasonal_trend). Not weather-adjusted.")
    add("Texas consumption, Mexico pipeline, LNG model, data centres, dry production (base / delayed)", "see Texas workbooks", "", "ASSUMPTION", False,
        f"Read from texas_gas_monthly.xlsx / texas_production_forecast.xlsx unchanged; their own status labels apply (data centres: {T['dc_inputs'] or 'see Assump - Data centres'}; the Permian takeaway table is UNVERIFIED except Matterhorn).")
    add("STEO regional marketed production to Dec 2027 (Permian, Eagle Ford, Haynesville)", "EIA STEO", "Bcf/d", "SOURCED", False,
        "EIA Short-Term Energy Outlook via API v2 (NGMPPM / NGMPEF / NGMPHA), as saved in texas_production_forecast.xlsx 'STEO raw'.")
    add("Basin supply Jan 2028 - Dec 2028", "STEO 2027 x STEO growth", "Bcf/d", "DERIVED", False, "Same month of 2027 x STEO 2027/2026 annual growth (Texas workbook).")
    add("Basin supply 2029-2030", "damped extension", "Bcf/d", "ASSUMPTION", False, "Month-on-month growth damped x0.5 a year; Permian capped by the takeaway table.")
    return R


def style_sheet(ws, widths, header_row=1):
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for c in ws[header_row]:
        c.font = BOLD
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(header_row + 1, 1)


def write(out, raw, idx, T, tx, la, cb, bs, steo, kq, kchart, kmeta, trains, prof, util, cap, est, imp, imp12, la_share, other_la, loss_la,
          fs, fs_tx, fs_la, out_chg, floor, edit):
    delay = T["shares"]["delay"]
    dated = {k: t for k, t in trains.items() if pd.notna(t["first"]) and pd.notna(t["peak"])}
    n_src = sum(1 for t in dated.values() if t["status"].upper().startswith("SOURCED"))
    n_tr = len(dated)
    n_pre = len(trains) - n_tr
    rows = assumption_rows(T, la_share, imp12, loss_la, 1 - T["loss_tx"] if False else T["loss_tx"], other_la, out_chg, floor, delay, trains, n_src, n_tr, cap, fs_la)
    adf = pd.DataFrame(rows, columns=["Item", "Value", "Unit", "Status", "Editable", "Source / note"])
    counts = adf["Status"].value_counts()
    cnt_txt = ", ".join(f"{int(counts.get(k, 0))} {k}" for k in ("SOURCED", "DERIVED", "ASSUMPTION", "UNVERIFIED"))
    tdf = pd.DataFrame([{"Plant": k[0], "Train": k[1], "EIA peak nameplate, Bcf/d LNG": t["peak"],
                         "Feedgas nameplate, Bcf/d (peak x 1.09)": t["peak"] * RATIO if pd.notna(t["peak"]) else np.nan,
                         "First-LNG month": t["first"], "Ramp months": t["ramp"], "Ramp-start utilisation": t["rs"],
                         "Steady utilisation": t["st"], "Date status": t["status"],
                         "In forecast": "yes" if k in dated else "no (no sourced date / pre-FID)", "Source / note": t["note"]}
                        for k, t in trains.items()])
    ex = pd.DataFrame({"Peak nameplate, Bcf/d LNG (EIA)": {p: v[0] for p, v in EXISTING.items()},
                       "Feedgas nameplate, Bcf/d": {p: v[0] * RATIO for p, v in EXISTING.items()},
                       "Source": {p: v[1] for p, v in EXISTING.items()}})
    for m in range(12):
        ex[f"Utilisation {pd.Timestamp(2001, m + 1, 1):%b}"] = {p: prof[p][m] for p in EXISTING}
    ex.index.name = "Plant"
    tot_note = f"x of y sourced: {n_src} of {n_tr} forecast Louisiana trains ({n_pre} pre-FID/undated projects not forecast)"
    # Haynesville split sheet
    hs = pd.DataFrame({"STEO Haynesville marketed": steo["Haynesville"].reindex(imp.index),
                       "Louisiana marketed production (EIA)": la["Marketed production (history)"].reindex(imp.index),
                       "Implied Louisiana share (upper bound)": imp,
                       "Implied East Texas share (lower bound)": 1 - imp,
                       "Texas share in the Texas workbook": T["shares"]["hay"]})
    hs.index.name = "Month"
    sens = []
    kq_base = kq[(kq["Scenario"] == "Base") & (kq["Year-end"] == "Dec 2030")].iloc[0]
    for s_ in (0.60, 0.65, round(1 - T["shares"]["hay"], 2), 0.75, 0.80):
        dh = (s_ - la_share) * (steo.loc["2030-12-01", "Haynesville"] - steo.loc["2025-12-01", "Haynesville"]) * (1 - loss_la)
        sens.append({"Louisiana share of STEO Haynesville": s_, "Change in Dec 2030 Haynesville supply growth vs the default, Bcf/d": dh,
                     "Extra supply required Dec 2030 (base, outflow held), Bcf/d": max(0.0, out_chg - (kq_base["Supply growth less demand growth (+ surplus / - shortfall)"] + dh))})
    sens = pd.DataFrame(sens).set_index("Louisiana share of STEO Haynesville")
    # frames for the data tabs
    tx_out = tx.copy()
    tx_out.insert(0, "Type", tx_out.pop("Type"))
    la_out = la.copy()
    la_out.insert(0, "Type", la_out.pop("Type"))
    for pl in list(EXISTING) + ["New Louisiana plants"]:
        pass
    lng_cols = {}
    notes = [
        "Notes", "", "UNITS",
        "All flows are Bcf/d (billion cubic feet per day), monthly averages; EIA monthly MMcf / days / 1000. Dry gas for production, consumption and LNG feedgas for demand.",
        "", "WHAT IT ANSWERS",
        "Whether Permian (takeaway-capped), Eagle Ford and Haynesville supply growth covers the growth in Gulf Coast gas demand (Texas + Louisiana: sectors, Mexico pipeline exports, data centres and above all LNG feedgas) from Dec 2025, or whether Haynesville growth above STEO (or supply from other regions) is needed. See 'Key question' and chart 'Supply growth vs LNG demand'.",
        "", "TEXAS",
        "Read unchanged (read only) from texas_gas_monthly.xlsx and texas_production_forecast.xlsx: consumption by sector, Mexico pipeline exports, BASE data-centre gas burn, dry production base / takeaway-delayed, LNG feedgas (TEXAS_GAS_FORECAST.py run to Dec 2030). " + texas_check(tx, T["d33"]) + ".",
        "", "LOUISIANA",
        "EIA Natural Gas Monthly via API v2, area SLA: consumption by sector (cons/sum), marketed and dry production (prod/sum), LNG exports of the four Louisiana ports Sabine Pass, Cameron, Calcasieu Pass and Plaquemines (move/poe2; Plaquemines is in Plaquemines Parish, Calcasieu Pass in Cameron Parish). Feedgas = exports x 1.09 (as Texas). Pipeline exports to Mexico: none from Louisiana.",
        f"Residential, commercial and vehicle fuel are small and EIA leaves single months blank (bridged by interpolation); months after a sector's last published month are the seasonal-trend estimate ({'; '.join(est) if est else 'none'}). Dry production: EIA to Dec 2024, then marketed x (1 - {loss_la:.1%}).",
        f"First forecast month: Texas {fs_tx:%b/%y}, Louisiana {fs_la:%b/%y}; combined {fs:%b/%y}. Louisiana consumption forecast = seasonal trend (not weather-adjusted).",
        f"LNG forecast: existing plants = EIA peak nameplate x {RATIO} x the monthly utilisation of the last 36 EIA months (smoothed 1-2-1); Plaquemines = mean utilisation Jan-Jul 2026 (full output) x the seasonal index; new trains from '{AL}' (ramp from first LNG to steady utilisation); delayed = trains not yet started slip {delay} months. {tot_note}.",
        "Date status: SOURCED means an opened document gives the year or half-year (EIA liquefaction capacity file 2026 Q2); the month within it is assumed (mid-period). Anything else is UNVERIFIED and has no first-LNG month, so it is not in the forecast until a date is typed in the yellow cell.",
        "", "SUPPLY",
        f"STEO regional marketed production (Permian, Eagle Ford, Haynesville) as in the Texas workbook, to Dec 2027; 2028 and 2029-30 are extensions (see Assumptions), not EIA forecasts. Haynesville is split Texas {T['shares']['hay']:.0%} / Louisiana {la_share:.0%}; EIA state data imply a Louisiana share of {imp12:.0%} (upper bound). Total Haynesville is STEO's either way. " + basin_check(T, bs, idx) + ".",
        "Permian: Texas share (85%) of STEO Permian, capped by the Texas workbook's takeaway table (all but Matterhorn UNVERIFIED), base and delayed. Dry = marketed x (1 - extraction loss). New Mexico's Permian gas is not counted (memo column on 'Key question').",
        "", "KEY QUESTION",
        "For each year-end: incremental LNG feedgas (Texas, Louisiana) and other Gulf demand since Dec 2025 against incremental supply by basin. Supply growth less demand growth > 0 is a surplus (the outflow to other states rises); < 0 is a shortfall. 'Outflow held': extra supply needed so the combined outflow ends at its Dec 2025 level (+ the editable allowance). 'Outflow declines': the outflow is allowed to fall; the fall is the gas other regions must replace, and extra supply is needed only if it would go below the floor.",
        f"Combined implied outflow = dry production - demand incl. LNG; Dec 2025 base level {kmeta['o0']:.1f} Bcf/d, last 12 actual months mean {kmeta['recent']:.1f}. It includes lease/plant and pipeline fuel EIA withholds and ignores storage; Louisiana storage net withdrawal is shown on the Louisiana tab as a memo.",
        "", "STATUS COUNTS",
        f"Assumptions tab ({len(adf)} items): {cnt_txt}.",
        "", "EDITING",
        "Yellow cells (Assumptions column Value where Editable = yes, and the train table) are read back from the committed workbook on every run; edit them and dispatch gulf_coast_balance.yml. Forecast values are computed by the script (no live Excel formulas).",
        "Beyond Dec 2027 supply is an extension of EIA STEO and beyond Dec 2028 it is a lighter-shaded scenario. Not an EIA or company forecast; a transparent scenario tool.",
    ]
    sheets = {"Assumptions": adf.set_index("Item"), AL: tdf.set_index("Plant"), "LA existing plants": ex,
              "Texas": tx_out, "Louisiana": la_out, "Combined": cb, "Key question": kq.set_index(["Scenario", "Year-end"]),
              "Key chart data": kchart, "Basin supply": pd.concat([steo.add_prefix("STEO (extended) "), bs], axis=1),
              "Haynesville split": hs, "Haynesville sensitivity": sens, RAW: raw}
    for n in ("Texas", "Louisiana", "Combined", "Basin supply"):
        sheets[n].index.name = "Month"
    kchart = kchart.rename(columns={"Permian (Texas share, takeaway-capped)": "Permian (Texas, capped)",
                                    "Haynesville (East Texas + Louisiana, STEO)": "Haynesville (STEO)",
                                    "Extra supply required (Haynesville above STEO or other regions)": "Extra supply required",
                                    "Incremental Gulf LNG feedgas": "Incremental Gulf LNG", "Incremental total Gulf demand (LNG + other)": "Incremental Gulf demand"})
    sheets["Key chart data"] = kchart
    xlsx_notes.write_workbook(out, sheets, notes, {"UNITS", "WHAT IT ANSWERS", "TEXAS", "LOUISIANA", "SUPPLY", "KEY QUESTION", "STATUS COUNTS", "EDITING"})
    wb = load_workbook(out)
    a = wb["Assumptions"]
    style_sheet(a, [70, 22, 9, 13, 9, 140])
    for r in range(2, a.max_row + 1):
        if a.cell(r, 5).value == "yes":
            a.cell(r, 2).fill = FILL_IN
    a["H1"], a["H2"] = "Status counts", cnt_txt
    a["H1"].font = BOLD
    a.column_dimensions["H"].width = 60
    t = wb[AL]
    style_sheet(t, [34, 24, 14, 14, 13, 9, 12, 11, 14, 30, 150])
    for r in range(2, t.max_row + 1):
        for c in (3, 5, 6, 7, 8, 9):
            t.cell(r, c).fill = FILL_IN
        t.cell(r, 5).number_format = "mmm/yy"
    for n, w in (("Texas", 14), ("Louisiana", 14), ("Combined", 14), ("Basin supply", 14), ("Key question", 18), ("Key chart data", 20)):
        ws = wb[n]
        ws.freeze_panes = "C2"
        for c in ws[1]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[1].height = 75
        for j in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        if n in ("Texas", "Louisiana", "Combined", "Basin supply", "Key chart data"):
            for r in range(2, ws.max_row + 1):
                ws.cell(r, 1).number_format = "0" if n == "Key chart data" else "mmm/yy"
    wb["Key question"].column_dimensions["A"].width = 32
    root, ext = os.path.splitext(out)
    tmp = f"{root}.tmp{os.getpid()}{ext}"
    try:
        wb.save(tmp)
        os.replace(tmp, out)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    # printouts for the log
    pd.set_option("display.width", 250, "display.max_columns", 30)
    print(tot_note)
    print(kq.set_index(["Scenario", "Year-end"]).round(2).T.to_string())
    print(cb.loc[[m for m in idx if m.month == 12 and m.year >= 2025], ["LNG feedgas, Texas (base)", "LNG feedgas, Louisiana (base)", "Demand incl. LNG, base",
                                                                         "Dry production, base", "Implied net outflow, base", "Implied net outflow, delayed"]].round(2).to_string())
    print(la.loc[la.index[la.index < fs_la][-12:], ["Electric power", "Industrial", "Dry production", "LNG feedgas (base)"]].round(2).to_string())
    print(f"Saved {out}")


if __name__ == "__main__":
    main()
