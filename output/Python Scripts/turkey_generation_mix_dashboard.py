"""
Pull Turkey's generation-by-fuel-type from EPIAS's PUBLIC, unauthenticated
dashboard endpoint and maintain a growing daily archive.

Data source: https://seffaflik.epias.com.tr/electricity-service/v1/dashboard/realtime-generation
Found via IRELAND_TURKEY_POWER_PLAYWRIGHT_DISCOVERY.py's Playwright
network capture of EPIAS's own homepage widget - confirmed via
TURKEY_EPIAS_DASHBOARD_INSPECT.py to return real hourly generation by
fuel type for the current day, as a plain GET, no credentials needed.

This is a DIFFERENT, unauthenticated endpoint from the one
turkey_generation_mix.py uses (generation/data/realtime-generation,
which requires a free EPIAS account's username/password via CAS login -
still blocked here on an empty password in api_keys.py). This dashboard
endpoint only ever returns "today so far" (like MISO's /FuelMix/Today),
not an arbitrary historical range, so - like MISO_FUEL_MIX_DAILY.py -
this upserts by date into a local archive on every run, building up
history from whenever this script first started running, rather than
backfilling the past.

Categories mapped the same way as turkey_generation_mix.py's
PRODUCTION_MAPPING, for consistency between the two scripts if the
authenticated one ever gets unblocked.
"""

import argparse
import os
import sys
from datetime import datetime
from zoneinfo import ZoneInfo

import pandas as pd
import requests

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # repo root, for xlsx_notes

import xlsx_notes

URL = "https://seffaflik.epias.com.tr/electricity-service/v1/dashboard/realtime-generation"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Accept": "application/json",
    "Referer": "https://seffaflik.epias.com.tr/",
}
TIMEOUT = (10, 30)

PRODUCTION_MAPPING = {
    "biomass": ["biomass"],
    "solar": ["sun"],
    "geothermal": ["geothermal"],
    "oil": ["fueloil", "naphta"],
    "gas": ["naturalGas", "lng"],
    "wind": ["wind"],
    "coal": ["blackCoal", "asphaltiteCoal", "lignite", "importCoal"],
    "hydro": ["river", "dammedHydro"],
    "unknown": ["wasteheat"],
}
FIELD_TO_CATEGORY = {field: category for category, fields in PRODUCTION_MAPPING.items() for field in fields}
CATEGORIES = list(PRODUCTION_MAPPING.keys())
IGNORED_KEYS = {"total", "date", "hour", "importExport"}

TR_ZONE = ZoneInfo("Europe/Istanbul")

EXPECTED_INTERVALS_PER_DAY = 24
MIN_INTERVALS_PER_DAY = 18

RENEWABLE_CATEGORIES = ["wind", "solar", "hydro", "geothermal", "biomass"]

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_OUT = os.path.join(REPO_ROOT, "output", "turkey_generation_mix_dashboard_daily.xlsx")


def fetch_today():
    r = requests.get(URL, headers=HEADERS, timeout=TIMEOUT)
    r.raise_for_status()
    return r.json().get("items", [])


def parse_items(items):
    rows = []
    for item in items:
        row = {category: 0.0 for category in CATEGORIES}
        for key, value in item.items():
            if key in IGNORED_KEYS or value is None:
                continue
            category = FIELD_TO_CATEGORY.get(key)
            if category is None:
                continue
            row[category] += max(float(value), 0.0)
        row["datetime"] = pd.to_datetime(item["date"])
        rows.append(row)
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows).set_index("datetime").sort_index()


def to_daily_row(wide, expected_day):
    if wide.empty:
        return None, 0
    on_day = wide[wide.index.date == expected_day]
    n_intervals = len(on_day)
    if n_intervals < MIN_INTERVALS_PER_DAY:
        return None, n_intervals
    means = on_day.mean()
    row = {f"{cat}_MW": means[cat] for cat in means.index}
    total = means.sum()
    row["Total_MW"] = total
    renewable = sum(means.get(cat, 0.0) for cat in RENEWABLE_CATEGORIES)
    row["Renewables_Share"] = renewable / total if total else None
    row["Intervals_Reported"] = n_intervals
    return pd.Series(row, name=expected_day), n_intervals


def load_archive(path):
    try:
        df = pd.read_excel(path, sheet_name="Data", index_col=0)
    except (FileNotFoundError, ValueError):
        return pd.DataFrame()
    df.index = pd.to_datetime(df.index).date
    df.index.name = "date"
    return df


def upsert(existing, new_row):
    if new_row is None:
        return existing
    new_df = new_row.to_frame().T
    new_df.index.name = "date"
    if existing.empty:
        combined = new_df
    else:
        combined = pd.concat([existing, new_df])
        combined = combined[~combined.index.duplicated(keep="last")]
    combined.index.name = "date"
    return combined.sort_index()


def format_date_column(path, sheet="Data"):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    if sheet not in wb.sheetnames:
        return
    ws = wb[sheet]
    for (cell,) in ws.iter_rows(min_row=2, max_col=1):
        cell.number_format = "dd-mmm-yyyy"
    ws.column_dimensions["A"].width = 14
    wb.save(path)


NOTES_LINES = [
    "UNITS",
    "All *_MW columns are MW, the daily mean of EPIAS's hourly generation readings for that fuel "
    "category - 'average MW for that day', not total daily energy. Total_MW is the sum of the "
    "category means.",
    "",
    "CATEGORIES",
    "biomass, solar, geothermal, wind: as published (solar = 'sun').",
    "oil: fuel oil and naphtha summed together.",
    "gas: natural gas and LNG summed together.",
    "coal: black coal, asphaltite coal, lignite, and imported coal summed together.",
    "hydro: run-of-river and dammed hydro summed together.",
    "unknown: waste heat.",
    "Renewables_Share: (wind + solar + hydro + geothermal + biomass) daily mean, as a share of "
    "Total_MW.",
    "Intervals_Reported: how many of the expected 24 hourly intervals that day actually came back; "
    f"a day is only saved if at least {MIN_INTERVALS_PER_DAY} came back.",
    "",
    "TIMESTAMPS",
    "The 'date' index is Turkey local time (Europe/Istanbul, UTC+3 year-round, no DST since 2016).",
    "",
    "COVERAGE",
    "This is a PUBLIC, unauthenticated dashboard endpoint found on EPIAS's own homepage widget - "
    "different from turkey_generation_mix.py's generation/data/realtime-generation endpoint, which "
    "needs a free EPIAS account (still blocked on an unset password as of this script's writing). "
    "It only ever returns 'today so far', not an arbitrary historical range (confirmed by inspection) "
    "- like MISO's /FuelMix/Today, this upserts by date into the Data tab on every run, building up "
    "history from whenever this script first started running, not backfilling the past.",
    "",
    "SOURCE",
    f"EPIAS (Turkey's energy exchange) Transparency Platform's public dashboard widget API: {URL}",
]
NOTES_SECTION_TITLES = {"UNITS", "CATEGORIES", "TIMESTAMPS", "COVERAGE", "SOURCE"}

STALE_AFTER_DAYS = 4


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default=DEFAULT_OUT)
    args = parser.parse_args()

    target_day = datetime.now(TR_ZONE).date()
    print(f"Fetching {URL} (expecting {target_day}) ...", file=sys.stderr)
    items = fetch_today()
    wide = parse_items(items)
    new_row, n_intervals = to_daily_row(wide, target_day)

    existing = load_archive(args.out)
    before_days = set(existing.index) if not existing.empty else set()
    combined = upsert(existing, new_row)
    is_new = bool(new_row is not None and new_row.name not in before_days)

    if new_row is None:
        print(f"Not enough intervals to save a day ({n_intervals}/{EXPECTED_INTERVALS_PER_DAY}) - "
              "keeping the archive as-is.", file=sys.stderr)
    else:
        print(f"Day {new_row.name}: {n_intervals}/{EXPECTED_INTERVALS_PER_DAY} intervals "
              f"({'new' if is_new else 'refreshed'})")

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    xlsx_notes.write_workbook(args.out, {"Data": combined}, NOTES_LINES, NOTES_SECTION_TITLES)
    format_date_column(args.out)
    print(f"Saved to {args.out}")


if __name__ == "__main__":
    main()
