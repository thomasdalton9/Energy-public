"""
Pull the daily Brent and WTI crude oil spot prices and maintain a growing
archive (output/Data and Chart Outputs/brent_wti_daily.xlsx): sheet 'Data'
(date index + Brent_USD_per_bbl + WTI_USD_per_bbl + Brent_less_WTI) and
sheet 'Monthly' (monthly means).

Source: EIA only, tried in order (first that answers wins):
  1. EIA API v2, route petroleum/pri/spt, series RBRTE (Europe Brent Spot
     Price FOB) and RWTC (WTI, Cushing) - needs the free EIA key, read from
     the EIA_API_KEY environment variable (repo secret EIA_API_KEY).
     Incremental: once the archive exists only the last weeks are re-read.
  2. EIA's keyless history workbooks
     https://www.eia.gov/dnav/pet/hist_xls/RBRTEd.xls and RWTCd.xls.
Modelled on HENRY_HUB_DAILY.py.

History: Brent from 1987-05-20, WTI from 1986-01-02. Prices publish once
per business day with a lag of a few days - no weekend/holiday rows, so
gaps in the date index are expected. USD per barrel; missing observations
are dropped, never filled. Negative values are real and kept (WTI settled
at -36.98 on 2020-04-20). Brent_less_WTI exists only where both are present.
Charts: add_charts.py (registry entry brent_wti_daily.xlsx) after the pull.
"""

import argparse
import io
import os
import sys
import time
from datetime import date, timedelta

import pandas as pd
import requests

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # repo root, for xlsx_notes

import xlsx_notes

EIA_API = "https://api.eia.gov/v2/petroleum/pri/spt/data/"
SERIES = {  # column -> (EIA series id, keyless history workbook)
    "Brent_USD_per_bbl": ("RBRTE", "https://www.eia.gov/dnav/pet/hist_xls/RBRTEd.xls"),
    "WTI_USD_per_bbl": ("RWTC", "https://www.eia.gov/dnav/pet/hist_xls/RWTCd.xls"),
}
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
}
TIMEOUT = (10, 60)
SPREAD = "Brent_less_WTI"
STALE_AFTER_DAYS = 10
OVERLAP_DAYS = 28  # incremental runs re-read the last four weeks (EIA occasionally revises recent days)

DEFAULT_OUT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "output", "Data and Chart Outputs",
    "brent_wti_daily.xlsx"
)

FETCH_ATTEMPTS = 3
RETRY_BACKOFF_SECONDS = [5, 20]


def get(url, **kw):
    last_error = None
    for attempt in range(1, FETCH_ATTEMPTS + 1):
        try:
            r = requests.get(url, headers=HEADERS, timeout=TIMEOUT, **kw)
            r.raise_for_status()
            return r
        except requests.RequestException as e:
            last_error = e
            print(f"  attempt {attempt}/{FETCH_ATTEMPTS} failed: {type(e).__name__}: {e}", file=sys.stderr)
            if attempt < FETCH_ATTEMPTS:
                time.sleep(RETRY_BACKOFF_SECONDS[attempt - 1])
    raise last_error


def tidy(dates, values, col):
    df = pd.DataFrame({"date": pd.DatetimeIndex(pd.to_datetime(list(dates), errors="coerce")).date,
                       col: pd.to_numeric(list(values), errors="coerce")})
    df = df.dropna().set_index("date").sort_index()
    return df[~df.index.duplicated(keep="last")]


def fetch_api_one(col, series, start=None):
    key = os.environ.get("EIA_API_KEY")
    if not key:
        raise RuntimeError("EIA_API_KEY not set")
    rows, offset = [], 0
    while True:
        params = {"api_key": key, "frequency": "daily", "data[0]": "value", "facets[series][]": series,
                  "sort[0][column]": "period", "sort[0][direction]": "asc", "offset": offset, "length": 5000}
        if start:
            params["start"] = start.isoformat()
        resp = get(EIA_API, params=params).json()["response"]
        data = resp.get("data", [])
        rows += data
        offset += len(data)
        if not data or offset >= int(resp.get("total", 0)):
            break
    print(f"  EIA API {series}: {len(rows)} rows, units {ascii(sorted({str(r.get('units')) for r in rows}))}",
          file=sys.stderr)
    return tidy([r["period"] for r in rows], [r["value"] for r in rows], col)


def fetch_xls_one(col, url):
    """Keyless history workbook (sheet 'Data 1': Date, price, header on row 3)."""
    raw = pd.read_excel(io.BytesIO(get(url).content), sheet_name="Data 1", header=2)
    print(f"  EIA XLS columns: {list(raw.columns)}", file=sys.stderr)
    return tidy(raw.iloc[:, 0], raw.iloc[:, 1], col)


def fetch_eia_api(start=None):
    return pd.concat([fetch_api_one(c, s, start) for c, (s, _) in SERIES.items()], axis=1)


def fetch_eia_xls(start=None):
    return pd.concat([fetch_xls_one(c, u) for c, (_, u) in SERIES.items()], axis=1)


SOURCES = [("EIA API v2 (petroleum/pri/spt RBRTE, RWTC)", fetch_eia_api),
           ("EIA history XLS (RBRTEd.xls, RWTCd.xls)", fetch_eia_xls)]


def fetch_series(start=None):
    for name, fn in SOURCES:
        print(f"Fetching from {name} ...", file=sys.stderr)
        try:
            df = fn(start)
        except Exception as e:  # noqa: BLE001 - fall through to the next source
            print(f"  {name} failed: {type(e).__name__}: {e}", file=sys.stderr)
            continue
        if not df.empty:
            df.index.name = "date"
            return df.sort_index(), name
        print(f"  {name} returned no rows", file=sys.stderr)
    raise SystemExit("all Brent/WTI sources failed")


def load_archive(path):
    try:
        df = pd.read_excel(path, sheet_name="Data", index_col=0)
    except (FileNotFoundError, ValueError):
        return pd.DataFrame()
    df.index = pd.to_datetime(df.index).date
    df.index.name = "date"
    return df[list(SERIES)]


def upsert(existing, new_df):
    """Per-column upsert: a day's value from the new pull replaces the old one; other days are kept."""
    if existing.empty:
        combined = new_df.copy()
    else:
        combined = new_df.combine_first(existing)  # new values win, old fill the rest
    combined = combined[list(SERIES)].dropna(how="all")
    combined.index.name = "date"
    return combined.sort_index()


def with_spread(df):
    out = df.copy()
    out[SPREAD] = (out["Brent_USD_per_bbl"] - out["WTI_USD_per_bbl"]).round(4)
    return out


def monthly(df):
    m = df.copy()
    m.index = pd.to_datetime(m.index)
    m = m.resample("MS").mean().dropna(how="all").round(4)
    m.index = m.index.date
    m.index.name = "month"
    return m


def format_dates(path):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    for sheet, fmt in (("Data", "dd-mmm-yyyy"), ("Monthly", "mmm/yy")):
        if sheet in wb.sheetnames:
            ws = wb[sheet]
            for (cell,) in ws.iter_rows(min_row=2, max_col=1):
                cell.number_format = fmt
            ws.column_dimensions["A"].width = 14
    wb.save(path)


NOTES_LINES = [
    "UNITS",
    "Brent_USD_per_bbl: US dollars per barrel, Europe Brent Spot Price FOB (EIA series RBRTE).",
    "WTI_USD_per_bbl: US dollars per barrel, Cushing OK WTI Spot Price FOB (EIA series RWTC).",
    "Brent_less_WTI: Brent minus WTI, USD per barrel, only on days both are published.",
    "",
    "COVERAGE",
    "Daily; Brent from 1987-05-20, WTI from 1986-01-02. Business days only (no weekend/holiday rows) - gaps in "
    "the date index are expected. EIA publishes with a lag of a few business days. Missing observations are "
    "dropped, never filled. Negative prices are real and kept (WTI -36.98 on 2020-04-20). Sheet 'Monthly' holds "
    "the simple monthly mean of the published daily values.",
    "",
    "TIMESTAMPS",
    "The 'date' index is the trade date EIA assigns the observation to, as published.",
    "",
    "SOURCE",
    f"EIA (US Energy Information Administration) spot prices via EIA API v2 ({EIA_API}, needs an API key), "
    "falling back to EIA's keyless history workbooks "
    + ", ".join(u for _, u in SERIES.values()) + ". Only days a source returns are upserted; "
    "existing days are never deleted.",
]
NOTES_SECTION_TITLES = {"UNITS", "COVERAGE", "TIMESTAMPS", "SOURCE"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default=DEFAULT_OUT)
    args = parser.parse_args()

    existing = load_archive(args.out)
    start = max(existing.index) - timedelta(days=OVERLAP_DAYS) if not existing.empty else None
    new_df, source = fetch_series(start)
    print(f"  {source}: {len(new_df)} days, {new_df.index.min()} to {new_df.index.max()}", file=sys.stderr)

    before = set(existing.index) if not existing.empty else set()
    combined = upsert(existing, new_df)
    new_days = sorted(set(combined.index) - before)
    data = with_spread(combined)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_notes.write_workbook(args.out, {"Data": data, "Monthly": monthly(combined)},
                              NOTES_LINES + ["", f"Last run's source: {source}"], NOTES_SECTION_TITLES)
    format_dates(args.out)

    print(f"Source used: {source}")
    print(f"Added {len(new_days)} new day(s){' - ' + str(new_days[-5:]) if new_days else ''}")
    print(f"Archive now has {len(data)} days ({data.index.min()} to {data.index.max()}). Saved to {args.out}")
    print(data.tail())

    latest = max(combined.index)
    age = (date.today() - latest).days
    if age > STALE_AFTER_DAYS:
        print(f"STALE SOURCE: newest saved day is {latest} ({age} days old) - "
              f"{source} may have changed or stopped updating.", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
