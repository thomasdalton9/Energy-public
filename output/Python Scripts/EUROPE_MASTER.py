"""
Europe master workbook: one file with every European dataset this repo pulls, and Dashboard front pages carrying
all their charts - the same layout as the South & Central America, North America and Australia + NZ masters
(south_america/SOUTH_AMERICA_MASTER.py, whose table/chart/dashboard code this reuses).

  Dashboard                    - gas: EU gas storage (GIE AGSI+, water year, EU and the main countries) and LNG
                                 terminal send-out / inventory (GIE ALSI); Ireland gas demand, supply and sector use (GNI)
  Dashboard - Power            - Europe power generation by source and installed capacity (sum of the countries below),
                                 day-ahead prices, net imports/exports by country, then one generation chart and one
                                 capacity chart per country (ENTSO-E), Ireland (EirGrid), Turkey (EPIAS), Cyprus
                                 (TSOC), the Rhine at Kaub, and finally a supply/demand BALANCE chart per country
                                 (generation by fuel + net imports + pumped storage/batteries net, against load)
  Dashboard - Capacity factors - capacity factor by generation type, Europe and per country
  <CC> <chart> data            - the table each Dashboard chart plots
  <CC> <dataset> raw           - the full data sheet(s) from each source workbook
  Sources                      - where each dataset comes from, units and notes

Sources: ENTSO-E Transparency Platform (the TSOs' own statutory reporting) is the primary source for power across
Europe - in a settled-week check it matched SMARD (Germany) and RTE (France) to within 0.5% - labelled as such on
every chart's source line. National feeds replace it where they are better (Ireland EirGrid, Turkey EPIAS, Cyprus
TSOC are national). Gas storage and LNG are GIE's aggregation of the operators' own reports.
GB is not in the ENTSO-E generation set (no GB data after 2020); it needs its own feed (Elexon BMRS).

Reads (doesn't refetch) the workbooks the scheduled pulls write to "output/Data and Chart Outputs/". A missing input
is listed on the Dashboard and skipped rather than stopping the rest.

Usage: python3 EUROPE_MASTER.py [--out "output/Data and Chart Outputs/europe_master.xlsx"]
"""
import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
sys.path.insert(0, os.path.join(ROOT, "europe"))
import add_charts  # noqa: E402
import capacity_factors  # noqa: E402
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402
from europe_countries import COUNTRIES  # noqa: E402

DATA_DIR = sam.DATA_DIR
FUELS = ["Hydro", "Gas", "Wind", "Solar", "Coal", "Nuclear", "Other"]   # same order/colours as the other dashboards
START = "2021-01-01"
MIN_SHARE = 0.95   # a country joins the Europe total if it has this share of the months since START

ENTSOE = ("ENTSO-E Transparency Platform (TSOs' reporting): actual generation per production type",
          "https://transparency.entsoe.eu/")

# (code, country, workbook, raw sheet / tuple of raw sheets / "*", short dataset name)
DATASETS = [
    ("EU", "Europe", "eu_gas_storage_daily.xlsx", "Daily", "gas storage"),
    ("EU", "Europe", "eu_lng_terminals_daily.xlsx", "Daily", "LNG terminals"),
    ("IE", "Ireland", "ireland_gas_combined_daily.xlsx", "*", "gas"),
    ("IE", "Ireland", "ireland_gni_transparency_daily.xlsx", "*", "gas by sector"),
    ("EU", "Europe", "europe_gas_flows_daily.xlsx", "*", "ENTSOG"),
]
RAW_POWER_DATASETS = (
    [(code, name, f"{slug}_power_generation_daily.xlsx", "Daily", "power") for code, (name, slug, _) in COUNTRIES.items()]
    + [("IE-EG", "Ireland (EirGrid)", "ireland_smartgrid_15min.xlsx", (), "power"),
       ("TR", "Turkey", "turkey_generation_mix_dashboard_daily.xlsx", "*", "power"),
       ("CY", "Cyprus", "cyprus_generation_mix_daily.xlsx", "*", "power")])
CAPACITY_DATASETS = [(code, name, f"{slug}_power_capacity.xlsx", "Monthly", "capacity")
                     for code, (name, slug, _) in COUNTRIES.items()]
PRICE_DATASETS = [("EU", "Europe", "europe_power_prices_daily.xlsx", "Daily", "power prices"),
                  ("EU", "Europe", "europe_cross_border_flows_daily.xlsx", "*", "flows")]
FLOWS_FILE = "europe_cross_border_flows_daily.xlsx"
HYDRO_DATASETS = [("DE", "Germany (Rhine)", "rhine_kaub_level_daily.xlsx", "Data", "river level")]
HYDRO_EXTRA = {}
DASHBOARD_ONLY = {}
EMBER = set()
OPERATORS = {}


def rhine(path):
    """rhine_kaub_level_daily.xlsx draws its own chart (REGISTRY entry is None); the master needs a spec."""
    d = add_charts.by_date(add_charts.read(path, "Data"), "date")
    return [{"name": "Rhine", "water_year": d["level_cm"], "title": "Rhine water level at Kaub (PEGELONLINE)",
             "units": "cm (gauge datum)", "y_decimals": 0}]


MASTER_SPECS = {"rhine_kaub_level_daily.xlsx": rhine}

SOURCES = {
    "eu_gas_storage_daily.xlsx": ("Gas Infrastructure Europe, AGSI+ (storage operators' own reports)", "https://agsi.gie.eu/"),
    "eu_lng_terminals_daily.xlsx": ("Gas Infrastructure Europe, ALSI (LNG terminal operators' own reports)", "https://alsi.gie.eu/"),
    "europe_gas_flows_daily.xlsx": ("ENTSOG Transparency Platform: physical flows at interconnection points (TSOs' own reporting)",
                                    "https://transparency.entsog.eu/"),
    "ireland_gas_combined_daily.xlsx": ("Gas Networks Ireland (GNI) transparency and open data",
                                        "https://www.gasnetworks.ie/corporate/gas-regulation/transparency/"),
    "ireland_gni_transparency_daily.xlsx": ("Gas Networks Ireland (GNI) transparency pages",
                                            "https://www.gasnetworks.ie/corporate/gas-regulation/transparency/"),
    "ireland_smartgrid_15min.xlsx": ("EirGrid / SONI Smart Grid Dashboard", "https://www.smartgriddashboard.com/"),
    "turkey_generation_mix_dashboard_daily.xlsx": ("EPIAS Transparency Platform (Turkey)", "https://seffaflik.epias.com.tr/"),
    "cyprus_generation_mix_daily.xlsx": ("Transmission System Operator Cyprus (TSOC)", "https://tsoc.org.cy/"),
    "rhine_kaub_level_daily.xlsx": ("PEGELONLINE (German federal waterways administration), Kaub gauge",
                                    "https://www.pegelonline.wsv.de/"),
    "europe_power_prices_daily.xlsx": ("ENTSO-E Transparency Platform: day-ahead prices", "https://transparency.entsoe.eu/"),
    "europe_cross_border_flows_daily.xlsx": ("ENTSO-E Transparency Platform: cross-border physical flows",
                                             "https://transparency.entsoe.eu/"),
}
GAS_BALANCE_SRC = ("ENTSOG (production, pipeline flows, consumption), GIE ALSI (LNG send-out), GIE AGSI+ (storage)",
                   "https://transparency.entsog.eu/")
GAS_FLOWS_FILE = "europe_gas_flows_daily.xlsx"
BALANCE_SRC = ("ENTSO-E Transparency Platform: generation, load and cross-border physical flows",
               "https://transparency.entsoe.eu/")
for _code, (_name, _slug, _zones) in COUNTRIES.items():
    SOURCES[f"{_slug}_power_generation_daily.xlsx"] = ENTSOE
    SOURCES[f"{_slug}_power_capacity.xlsx"] = ("ENTSO-E Transparency Platform: installed capacity per production type",
                                               "https://transparency.entsoe.eu/")


def monthly_gwh(path):
    """Standard Daily sheet (MWh) -> complete months, GWh, in the dashboard fuel groups (storage/pumping left out)."""
    d = add_charts.by_date(add_charts.read(path, "Daily"), "date")
    d = d[[c for c in d.columns if str(c).endswith("_MWh")]].apply(pd.to_numeric, errors="coerce")
    m = d.resample("MS").sum(min_count=1) / 1000.0
    last = d.dropna(how="all").index.max()
    if last < last + pd.offsets.MonthEnd(0):   # drop the month in progress
        m = m[m.index < last.to_period("M").to_timestamp()]
    m = m[m.index >= START].rename(columns=lambda c: c.replace("_MWh", "_GWh"))
    m = m.rename(columns={"Oil_GWh": "Other Fossil_GWh", "Other_GWh": "Other Renewables_GWh"})
    return add_charts.power_mix(m)


def europe_generation(data_dir, frames_out=None):
    """Sum of the ENTSO-E countries with a near-complete record, GWh per month, over the months they all have."""
    frames, notes = {}, []
    for code, (name, slug, _) in COUNTRIES.items():
        fname = f"{slug}_power_generation_daily.xlsx"
        try:
            m = monthly_gwh(os.path.join(data_dir, fname))
            frames[name] = m[m.sum(axis=1) > 0]
        except Exception as e:  # noqa: BLE001
            notes.append(f"NOT INCLUDED: {name} ({type(e).__name__}: {e})")
    if not frames:
        return pd.DataFrame(), notes
    last = max(f.index.max() for f in frames.values())
    window = pd.date_range(START, last, freq="MS")
    keep = {}
    for name, f in frames.items():
        share = len(f.index.intersection(window)) / len(window)
        if share >= MIN_SHARE:
            keep[name] = f
        else:
            notes.append(f"NOT INCLUDED: {name} (only {share:.0%} of months since {pd.Timestamp(START):%b/%y}: "
                         f"{f.index.min():%b/%y}-{f.index.max():%b/%y})")
    if frames_out is not None:
        frames_out.update(frames)
    months = sorted(set.intersection(*(set(f.index) for f in keep.values()))) if keep else []
    if not months:
        return pd.DataFrame(), notes + ["no month common to every country yet"]
    total = sum(f.reindex(index=months, columns=FUELS).fillna(0) for f in keep.values())
    total.index.name = "date"
    notes.insert(0, f"{len(keep)} countries summed (ENTSO-E, {months[0]:%b/%y}-{months[-1]:%b/%y}): " + ", ".join(keep))
    return total, notes


def europe_capacity(data_dir):
    """Annual installed capacity, GW, summed over the countries with figures for the year (ENTSO-E, 1 January)."""
    frames, notes = {}, []
    for code, (name, slug, _) in COUNTRIES.items():
        path = os.path.join(data_dir, f"{slug}_power_capacity.xlsx")
        try:
            d = add_charts.by_date(add_charts.read(path, "Monthly"), "date")
            frames[name] = add_charts.capacity_groups(pd.DataFrame(
                {f: d.get(f"{f}_MW") for f in add_charts.CAPACITY_FUELS}, index=d.index)) / 1000.0
        except Exception as e:  # noqa: BLE001
            notes.append(f"NOT INCLUDED: {name} ({type(e).__name__}: {e})")
    if not frames:
        return pd.DataFrame(), notes
    first = min(f.index.min() for f in frames.values())
    last = max(f.index.max() for f in frames.values())
    rng = pd.date_range(first, last, freq="YS")
    core = {}
    for name, f in frames.items():
        share = len(f.index.intersection(rng)) / len(rng)
        if share >= MIN_SHARE:
            core[name] = f
        else:
            notes.append(f"NOT INCLUDED: {name} (figures for only {share:.0%} of {first:%Y}-{last:%Y}: "
                         f"{f.index.min():%Y}-{f.index.max():%Y})")
    years = sorted(set.intersection(*(set(f.index) for f in core.values()))) if core else []
    rows, used_years = [sum(core[n].loc[y] for n in core) for y in years], years
    frames = core
    if not rows:
        return pd.DataFrame(), notes + ["no year with figures for the countries counted"]
    total = pd.DataFrame(rows, index=pd.DatetimeIndex(used_years))
    total.index.name = "date"
    notes.insert(0, f"{len(frames)} countries (ENTSO-E annual installed capacity, 1 January; storage excluded)")
    return total, notes


def total_chart(wb, used, power, pos, df_total, notes, sheet, title, units, src, note_head, date_format="%Y-%m",
                label="Europe", line_cols=()):
    """Chart + data tab; pos=None appends to the dashboard lists, otherwise inserts at pos."""
    ws = wb.create_sheet(sam.sheet_name(sheet, used))
    df, n_bars = xlsx_charts.prepare(df_total, line_cols)
    xlsx_charts.write_table(ws, df, date_format)
    ws.cell(row=1, column=df.shape[1] + 4, value=note_head)
    for i, note in enumerate(notes, start=2):
        ws.cell(row=i, column=df.shape[1] + 4, value=note)
    chart = xlsx_charts.build_chart(ws, df, n_bars, title, units, "stacked_bar", date_format, width=sam.CHART_W,
                                    height=sam.CHART_H, gridlines=False, inner=xlsx_charts.DASHBOARD_INNER)
    missing = [n.split(":", 1)[1].split("(")[0].strip() for n in notes if n.startswith("NOT INCLUDED")]
    row = (label, title + (f" (not counted: {', '.join(missing)})" if missing else ""),
           df.index.max().strftime("%b/%y" if date_format != "%Y" else "%Y"), ws.title, *src)
    if pos is None:
        power[0].append((chart, (src[0], src[1])))
        power[1].append(row)
    else:
        power[0].insert(pos, (chart, (src[0], src[1])))
        power[1].insert(pos, row)


def monthly_cover(daily, min_cover=0.9):
    """Daily frame -> monthly totals scaled to the full month, blank where fewer than min_cover of the days exist."""
    g = daily.resample("MS")
    n = g.count()
    dim = pd.Series(n.index.days_in_month, index=n.index)
    out = g.sum(min_count=1).mul(dim, axis=0) / n.where(n > 0)
    return out.where(n.div(dim, axis=0) >= min_cover)


BALANCE_COLS = ["Hydro", "Gas", "Wind", "Solar", "Coal", "Nuclear", "Other", "Net imports",
                "Pumped & battery (net)", "Load"]


def country_balance(gen_path, net_imports):
    """Monthly GWh supply/demand balance for one country: generation by fuel group, net imports, pumped storage and
    batteries (discharge minus pumping/charging) and load - over the days that have load and flows."""
    d = add_charts.by_date(add_charts.read(gen_path, "Daily"), "date").apply(pd.to_numeric, errors="coerce")

    def col(name):
        return d[name] if name in d else pd.Series(float("nan"), index=d.index)

    day = pd.DataFrame({
        "Hydro": col("Hydro_MWh"), "Gas": col("Gas_MWh"), "Wind": col("Wind_MWh"), "Solar": col("Solar_MWh"),
        "Coal": col("Coal_MWh"), "Nuclear": col("Nuclear_MWh"),
        "Other": d[[c for c in ("Bioenergy_MWh", "Oil_MWh", "Other_MWh") if c in d]].sum(axis=1, min_count=1),
        "Pumped & battery (net)": (col("PumpedStorage_MWh").fillna(0) + col("Storage_MWh").fillna(0)
                                   - col("PumpedStorageConsumption_MWh").fillna(0) - col("StorageCharging_MWh").fillna(0)),
        "Load": col("Load_MWh")}) / 1000.0
    day["Net imports"] = net_imports
    day = day.dropna(subset=["Load", "Net imports", "Gas"]).fillna(0)
    day = day[day.index >= START]
    if day.empty:
        return pd.DataFrame()
    m = monthly_cover(day[BALANCE_COLS]).dropna(how="all")
    last = day.index.max()
    if last < last + pd.offsets.MonthEnd(0):   # drop the month in progress
        m = m[m.index < last.to_period("M").to_timestamp()]
    return m.dropna(how="any")


GAS_BAL_COLS = ["Production", "Pipeline imports", "LNG send-out", "Storage withdrawals", "Pipeline exports",
                "Storage injections", "Consumption"]
EU27_GAS = ["AT", "BE", "BG", "HR", "CZ", "DK", "EE", "FI", "FR", "DE", "GR", "HU", "IE", "IT", "LV", "LT", "LU", "NL", "PL",
            "PT", "RO", "SK", "SI", "ES", "SE"]
GAS_NAMES = {"AT": "Austria", "BE": "Belgium", "BG": "Bulgaria", "HR": "Croatia", "CZ": "Czechia", "DK": "Denmark",
             "EE": "Estonia", "FI": "Finland", "FR": "France", "DE": "Germany", "GR": "Greece", "HU": "Hungary",
             "IE": "Ireland", "IT": "Italy", "LV": "Latvia", "LT": "Lithuania", "LU": "Luxembourg", "NL": "Netherlands",
             "PL": "Poland", "PT": "Portugal", "RO": "Romania", "SK": "Slovakia", "SI": "Slovenia", "ES": "Spain",
             "SE": "Sweden", "UK": "United Kingdom"}


def _col(df, name):
    return df[name] if name in df else pd.Series(float("nan"), index=df.index)


def _monthly_twh(day, line_floor=12):
    """Daily GWh frame -> complete months in TWh (months with >= 90% of days), blank rows dropped."""
    m = monthly_cover(day).dropna(how="all") / 1000.0
    last = day.dropna(how="all").index.max()
    if len(day) and last < last + pd.offsets.MonthEnd(0):   # drop the month in progress
        m = m[m.index < last.to_period("M").to_timestamp()]
    return m.dropna(how="any") if len(m) >= line_floor else pd.DataFrame()


def gas_country_balance(bal, cc, storage, lng):
    """Monthly TWh gas balance for one ENTSOG country: production, pipeline imports, LNG send-out (ALSI) and storage
    withdrawals (AGSI+) as supply; pipeline exports and storage injections as negatives; consumption (distribution +
    final consumers) as a line. Supply less the negatives should land near the consumption line; the gap is the
    unreported or unclassified flow."""
    a = "GB" if cc == "UK" else cc
    day = pd.DataFrame(index=bal.index)
    day["Production"] = _col(bal, f"{cc}_production_GWhd").fillna(0)
    day["Pipeline imports"] = _col(bal, f"{cc}_imports_GWhd")
    lng_s = _col(lng, f"{a}_sendout_GWhd").reindex(bal.index)
    day["LNG send-out"] = lng_s.where(lng_s.notna(), _col(bal, f"{cc}_lng_GWhd")).fillna(0)
    out_s, in_s = _col(storage, f"{a}_withdrawal_GWhd").reindex(bal.index), _col(storage, f"{a}_injection_GWhd").reindex(bal.index)
    day["Storage withdrawals"] = out_s.where(out_s.notna(), _col(bal, f"{cc}_storage_out_GWhd")).fillna(0)
    day["Pipeline exports"] = -_col(bal, f"{cc}_exports_GWhd").fillna(0)
    day["Storage injections"] = -in_s.where(in_s.notna(), _col(bal, f"{cc}_storage_in_GWhd")).fillna(0)
    day["Consumption"] = pd.concat([_col(bal, f"{cc}_distribution_GWhd"), _col(bal, f"{cc}_final_consumers_GWhd")], axis=1).sum(axis=1, min_count=1)
    day = day.dropna(subset=["Pipeline imports", "Consumption"])
    day = day[day.index >= "2021-10-01"]
    return _monthly_twh(day[GAS_BAL_COLS]) if len(day) else pd.DataFrame()


def eu_gas_balance(bal, org, dst, storage, lng):
    """EU27 gas balance: production and consumption summed over the countries; extra-EU pipeline imports and exports
    from the origin / destination sheets; LNG and storage from GIE's EU aggregates."""
    cols = lambda cat: [f"{c}_{cat}_GWhd" for c in EU27_GAS if f"{c}_{cat}_GWhd" in bal]   # noqa: E731
    day = pd.DataFrame(index=bal.index)
    day["Production"] = bal[cols("production")].sum(axis=1, min_count=1).fillna(0)
    day["Pipeline imports"] = org.reindex(bal.index).sum(axis=1, min_count=1) if len(org) else float("nan")
    day["LNG send-out"] = _col(lng, "EU_sendout_GWhd").reindex(bal.index).fillna(0)
    day["Storage withdrawals"] = _col(storage, "EU_withdrawal_GWhd").reindex(bal.index).fillna(0)
    day["Pipeline exports"] = -(dst.reindex(bal.index).sum(axis=1, min_count=1) if len(dst) else 0.0)
    day["Storage injections"] = -_col(storage, "EU_injection_GWhd").reindex(bal.index).fillna(0)
    day["Consumption"] = bal[cols("distribution") + cols("final_consumers")].sum(axis=1, min_count=1)
    day = day.dropna(subset=["Pipeline imports", "Consumption"])
    day = day[day.index >= "2021-10-01"]
    return _monthly_twh(day[GAS_BAL_COLS]) if len(day) else pd.DataFrame()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "europe_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()
    cfg = sys.modules[__name__]

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    dash2 = wb.create_sheet("Dashboard - Power")
    used = {"Dashboard", "Dashboard - Power", "Dashboard - Capacity factors", "Sources"}
    sources = []

    gas = sam.collect(wb, DATASETS, args.data_dir, used, sources, cfg=cfg)
    parts = [sam.collect(wb, RAW_POWER_DATASETS, args.data_dir, used, sources, cfg=cfg),
             sam.collect(wb, PRICE_DATASETS, args.data_dir, used, sources, cfg=cfg),
             sam.collect(wb, CAPACITY_DATASETS, args.data_dir, used, sources, cfg=cfg),
             sam.collect(wb, HYDRO_DATASETS, args.data_dir, used, sources, cfg=cfg)]
    power = tuple(sum((p[i] for p in parts), []) for i in range(3))

    pos = 0
    gen_frames = {}
    total, notes = europe_generation(args.data_dir, gen_frames)
    if not total.empty:
        total_chart(wb, used, power, pos, total, notes, "Europe generation total data",
                    "Europe power generation by source (ENTSO-E countries)", "GWh per month", ENTSOE,
                    "Countries summed (only months all of them have):")
        pos += 1
        print("Europe generation total:", "; ".join(notes))
    else:
        power[2].append("Europe generation total: " + "; ".join(notes))
    cap, cap_notes = europe_capacity(args.data_dir)
    if not cap.empty:
        total_chart(wb, used, power, pos, cap, cap_notes, "Europe capacity total data",
                    "Europe installed generating capacity (ENTSO-E countries)", "GW installed",
                    ("ENTSO-E Transparency Platform: installed capacity per production type", "https://transparency.entsoe.eu/"),
                    "Countries summed (annual figures, 1 January; storage excluded):", date_format="%Y")
        pos += 1
        print("Europe capacity total:", "; ".join(cap_notes))

    cf_chart, cf_row, cf_missing, cf_countries = capacity_factors.add_capacity_factor_sheets(
        wb, used, sam.sheet_name, gen_frames,
        {name: os.path.join(args.data_dir, f"{slug}_power_capacity.xlsx") for _, (name, slug, _) in COUNTRIES.items()},
        sam.CHART_W, sam.CHART_H,
        "Generation (ENTSO-E) / installed capacity (ENTSO-E, annual)", "Europe")
    if cf_chart:
        power[0].insert(pos, cf_chart)
        power[1].insert(pos, cf_row)
    power[2].extend(cf_missing)
    # gas balance: EU27 first, then each country (ENTSOG flows + ALSI LNG + AGSI+ storage)
    try:
        gbal = add_charts.by_date(add_charts.read(os.path.join(args.data_dir, GAS_FLOWS_FILE), "Country balance"), "date")
        gorg = add_charts._sheet(os.path.join(args.data_dir, GAS_FLOWS_FILE), "Imports by origin", "date")
        gdst = add_charts._sheet(os.path.join(args.data_dir, GAS_FLOWS_FILE), "Exports by destination", "date")
        gsto = add_charts._sheet(os.path.join(args.data_dir, "eu_gas_storage_daily.xlsx"), "Daily", "date")
        glng = add_charts._sheet(os.path.join(args.data_dir, "eu_lng_terminals_daily.xlsx"), "Daily", "date")
    except Exception as e:  # noqa: BLE001
        gbal = pd.DataFrame()
        gas[2].append(f"gas balance charts need {GAS_FLOWS_FILE} ({type(e).__name__}: {e})")
    if len(gbal):
        eu = eu_gas_balance(gbal, gorg, gdst, gsto, glng)
        if not eu.empty:
            total_chart(wb, used, gas, 0, eu, ["EU27: production and consumption summed over the countries (ENTSOG); pipeline imports/exports "
                                               "from/to outside the EU; LNG and storage from GIE's EU aggregates"],
                        "EU gas balance data", "EU gas balance: supply and storage vs consumption (TWh per month)",
                        "TWh per month", GAS_BALANCE_SRC, "Notes:", label="Europe", line_cols=("Consumption",))
        for cc in [c for c in GAS_NAMES if any(col.startswith(f"{c}_") for col in gbal.columns)]:
            b = gas_country_balance(gbal, cc, gsto, glng)
            if b.empty:
                gas[2].append(f"{GAS_NAMES[cc]} gas balance: too little data")
                continue
            total_chart(wb, used, gas, None, b, ["Supply (production, pipeline imports, LNG, storage withdrawals) less exports and "
                                                 "storage injections should land near consumption (distribution + final consumers); "
                                                 "the gap is unreported or unclassified flow"],
                        f"{GAS_NAMES[cc]} gas balance data", f"{GAS_NAMES[cc]} gas balance: supply and storage vs consumption",
                        "TWh per month", GAS_BALANCE_SRC, "Notes:", label=GAS_NAMES[cc], line_cols=("Consumption",))
    # supply/demand balance per country (needs the flows workbook and each country's load)
    flows_path = os.path.join(args.data_dir, FLOWS_FILE)
    try:
        net_all = add_charts.by_date(add_charts.read(flows_path, "Net imports"), "date")
    except Exception as e:  # noqa: BLE001
        net_all = pd.DataFrame()
        power[2].append(f"balance charts need {FLOWS_FILE} ({type(e).__name__}: {e})")
    bal_frames = {}
    for code, (name, slug, _) in COUNTRIES.items():
        if name not in net_all:
            continue
        try:
            b = country_balance(os.path.join(args.data_dir, f"{slug}_power_generation_daily.xlsx"), net_all[name])
        except Exception as e:  # noqa: BLE001
            power[2].append(f"{name} balance ({type(e).__name__}: {e})")
            continue
        if len(b) >= 3:
            bal_frames[name] = b
        else:
            power[2].append(f"{name} balance: fewer than 3 months with load and flows")
    for name, b in bal_frames.items():
        total_chart(wb, used, power, None, b, [f"{name}: generation + net imports + pumped storage/batteries net vs load; "
                                               "months with >= 90% of days; ENTSO-E"],
                    f"{name} balance data", f"{name} power balance: supply by source and net imports vs load",
                    "GWh per month", BALANCE_SRC, "Notes:", label=name, line_cols=("Pumped & battery (net)", "Load"))
    if len(bal_frames) >= 10:
        window = pd.date_range(START, max(b.index.max() for b in bal_frames.values()), freq="MS")
        core = {n: b for n, b in bal_frames.items() if len(b.index.intersection(window)) / len(window) >= MIN_SHARE}
        skipped = [n for n in bal_frames if n not in core]
        months = sorted(set.intersection(*(set(b.index) for b in core.values()))) if core else []
        if months:
            eu = sum(b.loc[months] for b in core.values())
            eu.index.name = "date"
            total_chart(wb, used, power, pos, eu, [f"{len(core)} countries with load and flows: " + ", ".join(core)]
                        + [f"NOT INCLUDED: {n} (balance months since {pd.Timestamp(START):%b/%y}: too few)" for n in skipped],
                        "Europe balance data", "Europe power balance: supply by source and net imports vs load",
                        "GWh per month", BALANCE_SRC, "Countries summed (only months all of them have):",
                        line_cols=("Pumped & battery (net)", "Load"))
            pos += 1
    sam.draw_dashboard(dash, "Europe energy - gas dashboard", *gas)
    sam.draw_dashboard(dash2, "Europe energy - power generation, capacity and prices", *power)
    cf_dash = wb.create_sheet("Dashboard - Capacity factors", 2)
    sam.draw_dashboard(cf_dash, "Europe - capacity factor by generation type (generation / capacity x hours)",
                       [c for c, _ in cf_countries], [r for _, r in cf_countries], cf_missing)

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    for s in sources:
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    for col, w in (("A", 24), ("B", 16), ("C", 44), ("D", 60), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: {len(gas[0])} gas charts, {len(power[0])} power charts; {len(wb.sheetnames)} tabs")
    for label, m in (("gas", gas[2]), ("power", power[2])):
        if m:
            print(f"missing ({label}):", m)


if __name__ == "__main__":
    main()
