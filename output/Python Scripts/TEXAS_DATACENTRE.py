"""
Texas gas demand to Dec 2033 (Bcf/d): LNG feedgas build 2026-28 ('overhang') followed by growing data-centre power demand.
Called by TEXAS_PRODUCTION_FORECAST.py (nothing runs on its own); adds three tabs to texas_production_forecast.xlsx:

  'Assump - Data centres'  yellow input cells (editable, read back from the committed workbook on every run), a derived block
                           (live formulas), the gas conversion, cross-checks and a SOURCES block. EVERY input carries a status:
                           SOURCED (a published figure: publisher, document, as-of, URL in the Sources block), DERIVED (arithmetic on
                           sourced figures), PROXY (an observed number used for something it does not strictly measure) or
                           ASSUMPTION (no source - a pure scenario input). The count is printed on the charts.
  'Data centres'           year-end energised Texas data-centre load (GW) for LOW / BASE / HIGH, 2024-2033, and its gas burn in
                           Bcf/d - LIVE Excel formulas on the assumption cells; ERCOT's own data-centre forecast alongside.
  'Demand to 2033'         the script's own evaluation, month by month to Dec 2033 (used for the charts, PNGs and dashboard):
                           stacked demand (EIA sectors + LNG feedgas + Mexico + data centres), LOW/BASE/HIGH totals, dry
                           production and the implied net outflow to other states.

US LOAD  IEA 'Energy and AI' (Apr 2025) Annex A: US data centres 183 TWh in 2024 and 426 TWh in 2030 (Base Case, Table A.4); world
         totals by case for 2030 and 2035 (Table A.1: Headwinds / Base / Lift-Off; 2035 is exploratory in the IEA's words). The IEA
         publishes the US only for the Base Case, so LOW and HIGH scale the US Base by the world ratio of that case (derived, the
         US share of the world is held at its Base-Case value). 2025-29 and 2031-34 are LINEAR INTERPOLATIONS (assumption).
         Cross-checks, not anchors: LBNL 2024 report (176 TWh in 2023; 325-580 TWh in 2028), EIA AEO2026 'Data Center Servers'.
TEXAS    Texas share of US data-centre GROWTH = ERCOT's own data-centre forecast for summer 2030 (CDR Dec 2025, 22.2 GW, TSP-reported
         contracts + officer letters) x a realisation ratio, over the IEA's US installed-capacity addition 2024-30 (100 - 42 GW).
         Realisation: LOW = 0.498 (ERCOT's observed ratio of data-centre site peak to requested MW, 2025 LTLF report), HIGH = 1.0 (the
         ERCOT forecast at face value), BASE = midpoint (ASSUMPTION). Energised GW = TWh / 8.76 / load factor (IEA US load factor 49%).
GAS      Bcf/d = GW x 1000 x load factor x gas share x heat rate (MMBtu/MWh) x 24 h / (MMBtu per Mcf) / 1e6.
         Gas share = ERCOT's observed share of generation from natural gas (EIA-930, eia930_fuel_mix_daily.xlsx): last 12 months for
         BASE, the lowest / highest rolling 12 months since 2019 for LOW / HIGH. It is an AVERAGE-mix share used as a PROXY for the share
         of the extra load met by gas (the marginal share is not observed). Heat rate and MMBtu/Mcf = mean of the last 12 calibrated
         months of ercot_gas_burn_daily.xlsx (EIA-923). Only the burn ADDED after the last EIA actual month is stacked (EIA's
         electric-power history already contains today's data-centre load); the annual tab shows the total.
2033     LNG: the train table of texas_gas_monthly.xlsx ('Assump - LNG'; last train 2031) held at its steady utilisation by
       TEXAS_GAS_FORECAST's own logic run to Dec 2033. Other sectors: the same seasonal-trend base; electric power is held flat
       after 2028 (editable) because data centres are the explicit power-growth layer. Production: STEO to Dec 2027, 2028 as in
       the base workbook, then each region's growth DAMPED (editable factor per year, compounded month by month from the Dec 2028 level, no seasonal shape) - 'extension, not STEO' - capped by the
       takeaway table (capacity held after the last listed pipeline). Everything after Dec 2028 is a SCENARIO.
"""
import math
import os
import re

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

import TEXAS_GAS_FORECAST as gf

END = pd.Timestamp("2033-12-01")
CHART_START = pd.Timestamp("2024-01-01")
AS = "Assump - Data centres"
DC = "Data centres"
VIEW = "Demand to 2033"
CASES = ("LOW", "BASE", "HIGH")
YEARS = list(range(2024, 2034))
FILL_IN = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
FILL_HEAD = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
FILL_LIGHT = PatternFill("solid", start_color="EDEDED", end_color="EDEDED")
FILL_STATUS = {"SOURCED": "E2EFDA", "DERIVED": "DDEBF7", "PROXY": "FCE4D6", "ASSUMPTION": "F8CBAD"}
BOLD = Font(bold=True)
LAYOUT = "layout v2 (sourced anchors, Oct 2026)"
IEA_URL = "https://iea.blob.core.windows.net/assets/dd7c2387-2f60-4b60-8c5f-6563b6aa1e4c/EnergyandAI.pdf"
LBNL_URL = "https://escholarship.org/uc/item/32d6m0d1"
CDR_URL = "https://www.ercot.com/files/docs/2025/12/19/CapacityDemandandReservesReport_December2025.xlsx"
LTLF_URL = "https://www.ercot.com/files/docs/2025/04/08/2025_LTLF_Report.docx"
HEARING_URL = "https://www.ercot.com/files/docs/2026/04/09/ERCOTLargeLoadUpdate-April9HouseStateAffairsHearing.pdf"
BZ_URL = "https://www.ercot.com/files/docs/2026/08/19/ERCOTPanel1DataCenters.pdf"
AEO_URL = "https://www.eia.gov/outlooks/aeo/"
EIA930_URL = "https://www.eia.gov/opendata/ (Form EIA-930, hourly electric grid monitor)"
QUAD_TWH = 293.071      # 1 quad = 293.071 TWh

# Sources block: id -> (publisher, document / table, exact figure(s), unit, as-of, observed or forecast, URL)
SOURCES = [
    ("S1", "IEA", "Energy and AI, Annex A, Table A.4 (p260), Base Case", "US data centres: 154 (2023), 183 (2024), 426 (2030)", "TWh a year", "Apr 2025",
     "estimate 2023-24 / Base Case forecast 2030", IEA_URL),
    ("S2", "IEA", "Energy and AI, Annex A, Table A.1 (p258), WORLD data centres by case", "2030: Headwinds 669, High Efficiency 792, Base 946, Lift-Off 1,264; "
     "2035: Headwinds 707, High Efficiency 972, Base 1,193, Lift-Off 1,719 (2035 'exploratory scenarios')", "TWh a year", "Apr 2025", "forecast (scenarios)", IEA_URL),
    ("S3", "IEA", "Energy and AI, Annex A, Table A.2 (p259)", "US total installed capacity: 35 (2023), 42 (2024), 100 (2030, Base)", "GW", "Apr 2025",
     "estimate 2023-24 / Base Case forecast 2030", IEA_URL),
    ("S4", "IEA", "Energy and AI, Annex A, Table A.3 (p259)", "US load factor: 53 (2020), 51 (2023), 50 (2024), 49 (2030)", "% of installed capacity", "Apr 2025",
     "estimate / Base Case forecast", IEA_URL),
    ("S5", "Lawrence Berkeley National Laboratory (for US DOE)", "2024 United States Data Center Energy Usage Report (Shehabi et al.)",
     "176 TWh in 2023 (4.4% of US electricity); 325 to 580 TWh in 2028 (6.7-12.0%); 74-132 GW at 50% utilisation; Texas is the third-largest state after Virginia and California (no figure)",
     "TWh a year / GW", "Dec 2024", "estimate 2023 / scenario range 2028", LBNL_URL),
    ("S6", "ERCOT", "Capacity, Demand and Reserves report Dec 2025, tab LoadResourceScenarios, Scenario 1, 'Data Centers' (summer peak load hour)",
     "2026 2,433; 2027 6,660; 2028 13,903; 2029 17,995; 2030 22,175 (winter 2030/31: 24,150)", "MW", "19 Dec 2025",
     "FORECAST (TSP-reported contracts + officer letters), not observed load", CDR_URL),
    ("S7", "ERCOT (as read by texas_demand_regression.xlsx 'Load sources' S9; not re-read here)", "2025 Long-Term Demand and Energy Forecast report",
     "0.498 = average data-centre site peak consumption / requested MW, sites in service 2022-24", "ratio", "Apr 2025", "OBSERVED", LTLF_URL),
    ("S8", "ERCOT", "Large Load Update, House Committee on State Affairs hearing (Pablo Vegas), p1/p3",
     "approx. 410 GW of large loads seeking interconnection, ~87% data centres (about 357 GW)", "GW", "26 Mar 2026 (hearing 9 Apr 2026)", "QUEUE (requests, not load)", HEARING_URL),
    ("S9", "ERCOT", "Data centres panel, House Committee on State Affairs, p3 (Batch Zero eligibility)",
     "approx. 205 GW of large load eligible for Batch Zero based on existing studies", "GW", "28 Jul 2026 (hearing 19 Aug 2026)", "QUEUE (requests, not load)", BZ_URL),
    ("S10", "US Energy Information Administration", "Annual Energy Outlook 2026, API v2 route aeo/2026, series cnsm_NA_comm_NA_prc_datactrserv_usa_qbtu "
     "(Commercial: Purchased Electricity: Data Center Servers)", "see the cross-check block (quads x 293.071 = TWh)", "quads", "AEO2026",
     "FORECAST. 'Data Center Servers' end use only: far below the IEA/LBNL totals, so coverage differs; shown for comparison, NOT used as an anchor", AEO_URL),
    ("S11", "US Energy Information Administration, Form EIA-930", "eia930_fuel_mix_daily.xlsx sheet ERCOT (this repo): Natural_Gas_MWh / Total_MWh",
     "set by the script on every run (see the gas share row)", "share of generation", "to yesterday", "OBSERVED (average mix)", EIA930_URL),
    ("S12", "US EIA Form EIA-923 via this repo", "ercot_gas_burn_daily.xlsx sheet Monthly: Heat_rate_used_MMBtu_per_MWh, MMBtu_per_Mcf_used (calibrated months)",
     "set by the script on every run (see the heat-rate rows)", "MMBtu/MWh, MMBtu/Mcf", "last 12 calibrated months", "ESTIMATE calibrated on EIA-923", "https://www.eia.gov/electricity/data/eia923/"),
]

# (key, label, low, base, high, unit, status, source ids, note); rows start at ROW0
PARAMS = [
    ("us24", "US data-centre electricity, 2024", 183, 183, 183, "TWh", "SOURCED", "S1", "IEA Base Case estimate (the text rounds to 'around 180'); LBNL has 176 TWh for 2023, IEA 154"),
    ("us30b", "US data-centre electricity, 2030, IEA Base Case", 426, 426, 426, "TWh", "SOURCED", "S1", "the IEA publishes the US only for its Base Case"),
    ("w30", "IEA WORLD data-centre electricity 2030 by case (LOW = Headwinds, BASE = Base, HIGH = Lift-Off)", 669, 946, 1264, "TWh", "SOURCED", "S2",
     "used only as a ratio to the Base Case. High Efficiency (792) not used"),
    ("w35", "IEA WORLD data-centre electricity 2035 by case", 707, 1193, 1719, "TWh", "SOURCED", "S2", "IEA calls the 2035 numbers 'exploratory scenarios'"),
    ("cap24", "US data-centre installed capacity, 2024", 42, 42, 42, "GW", "SOURCED", "S3", "IEA convention: includes cooling and other overhead (PUE)"),
    ("cap30", "US data-centre installed capacity, 2030, IEA Base Case", 100, 100, 100, "GW", "SOURCED", "S3", ""),
    ("lf", "Load factor (average load / installed capacity)", 0.49, 0.49, 0.49, "ratio", "SOURCED", "S4", "IEA US value for 2030 (50% in 2024); LBNL uses 50% for its GW range"),
    ("cdr30", "ERCOT data-centre forecast, summer 2030 (TSP-reported contracts + officer letters)", 22.175, 22.175, 22.175, "GW", "SOURCED", "S6",
     "a FORECAST of requested MW, not observed load; ERCOT's own scaled variants are 2.4x higher and are not used"),
    ("real", "Realisation ratio applied to the ERCOT forecast", 0.498, 0.749, 1.0, "ratio", "ASSUMPTION", "S7",
     "LOW = ERCOT's observed site peak / requested MW (S7); HIGH = the ERCOT forecast at face value; BASE = midpoint of the two: a judgement, not a measurement"),
    ("sh24", "Texas share of the US 2024 data-centre stock", 0.10, 0.10, 0.10, "share", "ASSUMPTION", "S5",
     "only the rank is published (LBNL: third after Virginia and California); the number is a guess. It does NOT affect the charts (only the burn ADDED after the last EIA actual month is stacked)"),
    ("gas", "Gas share of the extra load (proxy: ERCOT gas share of generation)", None, None, None, "share", "PROXY", "S11",
     "set by the script: LOW / HIGH = lowest / highest rolling 12-month gas share of ERCOT generation since 2019, BASE = last 12 months. An AVERAGE-mix share standing in for the marginal one"),
    ("hr", "Effective heat rate", None, None, None, "MMBtu/MWh", "SOURCED", "S12", "set by the script: mean of the last 12 calibrated months (EIA-923 basis)"),
    ("mcf", "MMBtu per Mcf", None, None, None, "MMBtu/Mcf", "SOURCED", "S12", "set by the script: mean of the same months (EIA-923 fuel heat content)"),
]
EDITABLE = ("us24", "us30b", "w30", "w35", "cap24", "cap30", "lf", "cdr30", "real", "sh24")
ROW0 = 5
R = {k: ROW0 + i for i, (k, *_r) in enumerate(PARAMS)}
DERIVED = [   # (key, label, unit, excel formula template with {c} = column letter, note)
    ("us30", "US data-centre electricity, 2030, by case", "TWh", "={c}{us30b}*{c}{w30}/$C${w30}", "IEA US Base x (case world / Base world): the US share of the world is held at its Base value"),
    ("us35", "US data-centre electricity, 2035, by case", "TWh", "={c}{us30b}*{c}{w35}/$C${w30}", "IEA US Base 2030 x (case world 2035 / Base world 2030)"),
    ("shg", "Texas share of US data-centre GROWTH", "share", "={c}{cdr30}*{c}{real}/({c}{cap30}-{c}{cap24})", "ERCOT summer-2030 data-centre forecast x realisation / IEA US capacity addition 2024-30"),
]
RDER0 = ROW0 + len(PARAMS) + 2
RD = {k: RDER0 + i for i, (k, *_r) in enumerate(DERIVED)}
R_PERGW = RDER0 + len(DERIVED)                    # Bcf/d per GW (formula)
INTERP = "Linear interpolation of US TWh 2024->2030 and 2030->2035 (and so 2025-29, 2031-34)"
SCAL0 = R_PERGW + 3                                # scalar block
SCALARS = [
    ("gelec", "Electric-power sector underlying growth after Dec 2028", 0.0, "per year", "ASSUMPTION",
     "0 = flat: data centres are the explicit power-demand growth layer, so EIA's trend growth is not extended (avoids double counting)"),
    ("damp", "Production extension after 2028: growth multiplier per year", 0.5, "x", "ASSUMPTION",
     "STEO regional growth of 2027/2026 is damped by this factor each year after 2028 (0.5 = halves every year). 'Extension, not STEO'"),
]
RS = {k: SCAL0 + i for i, (k, *_r) in enumerate(SCALARS)}
YR0 = 7                                          # first data row (2024) of the 'Data centres' tab
CDR_GW = {2026: 2.433, 2027: 6.660, 2028: 13.903, 2029: 17.995, 2030: 22.175}       # S6, summer peak, GW
# anchors counted on the charts: every PARAMS row, every DERIVED row and the interpolation rule
STATUSES = [p[6] for p in PARAMS] + ["DERIVED"] * len(DERIVED) + ["ASSUMPTION"]


def anchor_count():
    n = len(STATUSES)
    c = {s: STATUSES.count(s) for s in ("SOURCED", "DERIVED", "PROXY", "ASSUMPTION")}
    return c, f"{c['SOURCED']} of {n} sourced, {c['DERIVED']} derived, {c['PROXY']} proxy, {c['ASSUMPTION']} assumption"


def ercot_heat_rate(path):
    """Mean heat rate (MMBtu/MWh) and MMBtu/Mcf of the last 12 calibrated complete months of ercot_gas_burn_daily.xlsx; (hr, mcf, text)."""
    try:
        m = pd.read_excel(path, sheet_name="Monthly")
        ok = m[(m["ERCOT_days"] >= 28) & ~m["Heat_rate_basis"].astype(str).str.startswith("estimated")].tail(12)
        return float(ok["Heat_rate_used_MMBtu_per_MWh"].mean()), float(ok["MMBtu_per_Mcf_used"].mean()), (
            f"{ok['month'].iloc[0]:%b/%y}-{ok['month'].iloc[-1]:%b/%y} ({len(ok)} months)")
    except Exception as e:  # noqa: BLE001
        print(f"  ERCOT heat rate unreadable ({type(e).__name__}: {e}) - 8.5 MMBtu/MWh, 1.015 MMBtu/Mcf used", flush=True)
        return 8.5, 1.015, "FALLBACK 8.5 / 1.015 (ercot_gas_burn_daily.xlsx unreadable) - treat as ASSUMPTION"


def ercot_gas_share(path):
    """ERCOT gas share of generation (EIA-930): (low, base, high, text) = lowest rolling 365-day share since 2019, last 365 days, highest."""
    try:
        d = pd.read_excel(path, sheet_name="ERCOT", usecols=["date", "Natural_Gas_MWh", "Total_MWh"])
        d["date"] = pd.to_datetime(d["date"], errors="coerce")
        d = d.dropna(subset=["date"]).set_index("date").sort_index()
        r = d.rolling(365, min_periods=360).sum()
        s = (r["Natural_Gas_MWh"] / r["Total_MWh"]).dropna()
        return float(s.min()), float(s.iloc[-1]), float(s.max()), (
            f"EIA-930 ERCOT, 365-day windows {d.index.min():%b/%y}-{d.index.max():%d %b %Y}: lowest {s.min():.1%} (to {s.idxmin():%b/%y}), last {s.iloc[-1]:.1%}, highest {s.max():.1%} (to {s.idxmax():%b/%y})")
    except Exception as e:  # noqa: BLE001
        print(f"  EIA-930 gas share unreadable ({type(e).__name__}: {e}) - 0.40 used", flush=True)
        return 0.40, 0.40, 0.40, "FALLBACK 0.40 (eia930_fuel_mix_daily.xlsx unreadable) - treat as ASSUMPTION"


def aeo_servers(key):
    """EIA AEO2026 Data Center Servers purchased electricity, TWh: {year: (min over scenarios, max over scenarios)}; fallback = the 5 Oct 2026 probe values."""
    fb = {2025: (104.8, 104.8), 2028: (147.3, 153.3), 2030: (176.8, 190.7), 2033: (221.9, 254.3)}
    if not key:
        return fb, "stored (EIA_API_KEY missing)"
    try:
        r = requests.get("https://api.eia.gov/v2/aeo/2026/data", params={"api_key": key, "frequency": "annual", "data[0]": "value", "length": "500",
                         "facets[seriesId][]": "cnsm_NA_comm_NA_prc_datactrserv_usa_qbtu"}, timeout=60)
        rows = r.json()["response"]["data"]
        out = {}
        for y in (2025, 2028, 2030, 2033):
            v = [float(x["value"]) * QUAD_TWH for x in rows if x["period"] == str(y) and x["value"] is not None]
            out[y] = (min(v), max(v))
        return out, f"live from the EIA API, {len({x['scenario'] for x in rows})} scenarios"
    except Exception as e:  # noqa: BLE001
        print(f"  AEO2026 unreadable ({type(e).__name__}: {e}) - stored values", flush=True)
        return fb, "stored (API failed)"


def read_prior(path):
    """Editable cells of the committed workbook (None when absent or from the older, unsourced layout). Called BEFORE the workbook is rewritten."""
    if not os.path.exists(path):
        return None
    try:
        wb = load_workbook(path, data_only=False)
        if AS not in wb.sheetnames:
            return None
        a = wb[AS]
        if a["A3"].value != LAYOUT:
            print("  prior data-centre tab is the old layout - sourced defaults used", flush=True)
            return None
        p = {k: [a.cell(R[k], c).value for c in (2, 3, 4)] for k in EDITABLE}
        s = {k: a.cell(RS[k], 2).value for k, *_x in SCALARS}
        return {"p": p, "s": s}
    except Exception as e:  # noqa: BLE001
        print(f"  prior data-centre assumptions unreadable ({type(e).__name__}: {e}) - defaults used", flush=True)
        return None


def assumptions(prior, hr, mcf, gas):
    par = {k: [lo, ba, hi] for k, _l, lo, ba, hi, *_x in PARAMS}
    par["hr"] = [hr, hr, hr]
    par["mcf"] = [mcf, mcf, mcf]
    par["gas"] = list(gas)
    sc = {k: v for k, _l, v, *_x in SCALARS}
    if prior:
        for k, v in prior["p"].items():
            if all(isinstance(x, (int, float)) for x in v):
                par[k] = [float(x) for x in v]
        for k, v in prior["s"].items():
            if isinstance(v, (int, float)):
                sc[k] = float(v)
    return par, sc


# --------------------------------------------------------------------------- python model of the Excel formulas
def derived(par, ci):
    """US TWh 2030 / 2035 and the Texas share of US growth for one case (the 'Derived' rows of the assumption tab)."""
    g = lambda k: par[k][ci]   # noqa: E731
    us30 = g("us30b") * g("w30") / par["w30"][1]
    us35 = g("us30b") * g("w35") / par["w30"][1]
    shg = g("cdr30") * g("real") / (g("cap30") - g("cap24"))
    return us30, us35, shg


def annual(par, ci):
    """Year-end energised load and gas burn for one case: {year: dict(us, tx_twh, gw, bcfd)}."""
    g = lambda k: par[k][ci]   # noqa: E731
    us30, us35, shg = derived(par, ci)
    per_gw = 1000 * g("lf") * g("gas") * g("hr") * 24 / g("mcf") / 1e6
    out = {}
    for y in YEARS:
        us = g("us24") + (us30 - g("us24")) * (y - 2024) / 6 if y <= 2030 else us30 + (us35 - us30) * (y - 2030) / 5
        tx = g("us24") * g("sh24") + shg * (us - g("us24"))
        gw = tx / 8.76 / g("lf")
        out[y] = {"us": us, "tx_twh": tx, "gw": gw, "bcfd": gw * per_gw}
    return out, per_gw


def monthly_path(ann, key):
    """Year-end annual values -> monthly path (linear within the year; month m of year y = year-end y-1 + m/12 of the step)."""
    out = {}
    for y in YEARS[1:]:
        for m in range(1, 13):
            out[pd.Timestamp(year=y, month=m, day=1)] = ann[y - 1][key] + (ann[y][key] - ann[y - 1][key]) * m / 12
    return pd.Series(out)


# --------------------------------------------------------------------------- production extension
def takeaway_caps(tk, local, delay, idx):
    add = tk[~tk["Pipeline"].str.contains("Existing", case=False) & tk["Status"].astype(str).str.lower().ne("in service")]
    base_ex = float(tk[tk["Pipeline"].str.contains("Existing", case=False)]["Capacity Bcf/d"].sum())
    caps = {}
    for nm, d in (("base", 0), ("delayed", delay)):
        c = pd.Series(base_ex + local, index=idx, dtype=float)
        for _, r in add.iterrows():
            c[c.index >= pd.Timestamp(r["In-service month"] + "-01") + pd.DateOffset(months=d)] += float(r["Capacity Bcf/d"])
        caps[nm] = c
    return caps


def extend_production(df, tk, shares, other, local, delay, damp, loss_share):
    """Marketed and dry Texas production to Dec 2033: months to Dec 2028 as the base workbook, then damped STEO-region growth
    capped by the takeaway table (capacity held after the last pipeline)."""
    idx = pd.date_range(df.index.min(), END, freq="MS")
    sp, se, sh = shares
    reg = {"Permian": "STEO Permian marketed (Bcf/d)", "Eagle Ford": "STEO Eagle Ford marketed (Bcf/d)",
           "Haynesville": "STEO Haynesville marketed (Bcf/d)"}
    s = {k: df[c].reindex(idx) for k, c in reg.items()}
    for k, ser in s.items():
        g = ser["2027-01-01":"2027-12-01"].mean() / ser["2026-01-01":"2026-12-01"].mean() - 1
        for m in idx[idx > pd.Timestamp("2028-12-01")]:
            # month-on-month from the Dec 2028 level (no seasonal shape): a same-month-last-year rule would step down at the join
            ser[m] = ser[m - pd.DateOffset(months=1)] * (1 + g * damp ** (m.year - 2028)) ** (1 / 12)
    caps = takeaway_caps(tk, local, delay, idx)
    out = {}
    for sc in ("base", "delayed"):
        perm = pd.concat([s["Permian"], caps[sc]], axis=1).min(axis=1)
        tx = sp * perm + se * s["Eagle Ford"] + sh * s["Haynesville"] + other
        tx = tx.where(idx > pd.Timestamp("2028-12-01"), df[f"Texas production forecast, {sc}"].reindex(idx))
        out[sc] = tx
    return idx, out, caps, s


# --------------------------------------------------------------------------- build
def build(prod_df, sdf, tk, params, texas_xlsx, ercot_xlsx, prior, loss_share, other, delay, last_hist, stor=None):
    """prod_df: 'Forecast' frame of TEXAS_PRODUCTION_FORECAST (Month index); sdf: 'Supply and demand' frame. Returns (view frame,
    context for the Excel tabs and notes)."""
    hr, mcf, hr_txt = ercot_heat_rate(ercot_xlsx)
    gl, gb, gh, gas_txt = ercot_gas_share(os.path.join(os.path.dirname(ercot_xlsx), "eia930_fuel_mix_daily.xlsx"))
    aeo, aeo_txt = aeo_servers(os.environ.get("EIA_API_KEY", ""))
    par, sc = assumptions(prior, hr, mcf, (gl, gb, gh))
    ann, per_gw = {}, {}
    for ci, c in enumerate(CASES):
        ann[c], per_gw[c] = annual(par, ci)
    # --- gas demand to 2033 from the existing LNG / sector model run to Dec 2033
    cons = pd.read_excel(texas_xlsx, sheet_name="Consumption by sector", index_col=0, parse_dates=True)
    exports = pd.read_excel(texas_xlsx, sheet_name="Exports", index_col=0, parse_dates=True)
    bal = pd.read_excel(texas_xlsx, sheet_name="Balance", index_col=0, parse_dates=True)
    old_last = gf.LAST
    gf.LAST = END
    try:
        f, ctx = gf.build(texas_xlsx, cons, exports, bal)
    finally:
        gf.LAST = old_last
    fv = gf.values_sheet(f)
    N = gf.NAME
    idx = pd.date_range(CHART_START, END, freq="MS")
    v = pd.DataFrame(index=idx)
    first_fc = ctx["fs"]
    last_act = first_fc - pd.offsets.MonthBegin(1)
    v["Type"] = ["Actual" if m < first_fc else ("Forecast" if m <= pd.Timestamp("2028-12-01") else "Scenario") for m in idx]
    for c in gf.SECT:
        v[c] = fv[c].reindex(idx).astype(float)
    ge = sc["gelec"]
    j = idx > pd.Timestamp("2028-12-01")
    v.loc[j, "Electric power"] = [v.at[m - pd.DateOffset(years=m.year - 2028), "Electric power"] * (1 + ge) ** (m.year - 2028) for m in idx[j]]
    v["LNG feedgas (base case)"] = fv[N["lng_b"]].reindex(idx).astype(float)
    v["Pipeline exports to Mexico"] = fv[N["mex"]].reindex(idx).astype(float)
    # join check against the shipped 2028 workbook (must be identical before 2029)
    old = sdf.reindex(idx)
    k = idx <= pd.Timestamp("2028-12-01")
    diff = float(np.nanmax(np.abs(v.loc[k, "LNG feedgas (base case)"] - old.loc[k, "LNG feedgas (base)"])))
    diff2 = float(np.nanmax(np.abs(v.loc[k, gf.SECT].sum(axis=1) - old.loc[k, gf.SECT].sum(axis=1))))
    # --- data centres: added burn since the last EIA actual month
    for ci, c in enumerate(CASES):
        g_ = monthly_path(ann[c], "bcfd")
        gw_ = monthly_path(ann[c], "gw")
        base_m = g_[last_act]
        v[f"Data centres GW (year-end path), {c}"] = gw_.reindex(idx)
        v[f"Data centres gas burn, total, {c}"] = g_.reindex(idx)
        v[f"Data centres added gas burn, {c}"] = (g_.reindex(idx) - base_m).where(idx > last_act, 0.0)
    stack = v[gf.SECT + ["LNG feedgas (base case)", "Pipeline exports to Mexico"]].sum(axis=1, min_count=7)
    for c in CASES:
        v[f"Total demand incl. data centres, {c}"] = stack + v[f"Data centres added gas burn, {c}"]
    # --- production
    shares = (params["Permian Texas share of STEO Permian"], params["Eagle Ford Texas share"],
              params["Haynesville Texas share of STEO Haynesville"])
    pidx, mkt, caps, steo = extend_production(prod_df, tk, shares, other, params["Local Permian demand (power, industrial, Mexico-bound not included)"],
                                              delay, sc["damp"], loss_share)
    k_ = 1 - loss_share
    hist_dry = sdf["Dry production, base"].where(sdf["Type"].eq("Actual")).reindex(idx)
    v["Dry production, base"] = hist_dry.fillna(mkt["base"].reindex(idx) * k_)
    v["Dry production, takeaway delayed"] = hist_dry.fillna(mkt["delayed"].reindex(idx) * k_)
    for c in CASES:
        v[f"Implied net outflow before storage, {c}"] = v["Dry production, base"] - v[f"Total demand incl. data centres, {c}"]
    if stor is not None:
        # storage (TEXAS_SUPPLY_DEMAND.storage_path): gross injection on the demand side, withdrawal on the supply side
        v["Storage injection (demand)"] = stor["inj"].reindex(idx)
        v["Storage withdrawal (supply)"] = stor["wd"].reindex(idx)
    else:
        v["Storage injection (demand)"] = 0.0
        v["Storage withdrawal (supply)"] = 0.0
    v["Dry production + storage withdrawal, base"] = v["Dry production, base"] + v["Storage withdrawal (supply)"]
    v["Dry production + storage withdrawal, takeaway delayed"] = v["Dry production, takeaway delayed"] + v["Storage withdrawal (supply)"]
    for c in CASES:
        v[f"Implied net outflow, {c}"] = (v["Dry production + storage withdrawal, base"] - v[f"Total demand incl. data centres, {c}"]
                                          - v["Storage injection (demand)"])
    v["Permian takeaway + local demand, base (Bcf/d)"] = caps["base"].reindex(idx)
    v["Permian STEO / extension marketed (Bcf/d)"] = steo["Permian"].reindex(idx)
    v["Data-centre inputs"] = anchor_count()[1]
    v.index.name = "Month"
    ctx2 = {"par": par, "sc": sc, "ann": ann, "per_gw": per_gw, "hr": hr, "hr_txt": hr_txt, "first_fc": first_fc, "last_act": last_act,
            "join_lng": diff, "join_sec": diff2, "loss_share": loss_share, "mcf": mcf, "gas_txt": gas_txt, "aeo": aeo, "aeo_txt": aeo_txt}
    return v, ctx2


def checks(v, c2):
    """Sanity flags; returns text lines (printed and put on the Units tab)."""
    out = []
    jan = pd.Timestamp("2029-01-01")
    dec = pd.Timestamp("2028-12-01")
    out.append(f"model run to 2033 reproduces the 2028 workbook: max |LNG diff| {c2['join_lng']:.1e}, sectors {c2['join_sec']:.1e} Bcf/d")
    for ci, c in enumerate(CASES):
        a28, a30, a33 = (c2["ann"][c][y] for y in (2028, 2030, 2033))
        out.append(f"data centres {c}: US {a28['us']:.0f} TWh in 2028 (LBNL 2028 range 325-580), Texas {a28['gw']:.1f} / {a30['gw']:.1f} / {a33['gw']:.1f} GW energised at Dec/28 / Dec/30 / Dec/33 (ERCOT CDR 2030: 22.2 GW requested)")
    for c in CASES:
        g = v[f"Data centres gas burn, total, {c}"]
        out.append(f"data-centre gas burn {c}: Dec/28 {g[dec]:.2f}, Dec/30 {g[pd.Timestamp('2030-12-01')]:.2f}, Dec/33 {g[END]:.2f} Bcf/d"
                   + ("  ** EXCEEDS 10 Bcf/d **" if g.max() > 10 else " (below the 10 Bcf/d flag)"))
    for c in CASES:
        o = v[f"Implied net outflow, {c}"]
        o_all, o = o, o[o.index >= c2["first_fc"]]
        neg = o[o < 0]
        out.append(f"implied outflow {c}: Dec/28 {o[dec]:.1f}, Dec/30 {o[pd.Timestamp('2030-12-01')]:.1f}, Dec/33 {o[END]:.1f}, min {o.min():.1f} ({o.idxmin():%b/%y})"
                   + (f"  ** NEGATIVE from {neg.index[0]:%b/%y} ({len(neg)} months): production cannot meet Texas demand, Texas would have to draw gas in from other states **" if len(neg) else ""))
    tot = v["Total demand incl. data centres, BASE"]
    typical = tot.diff().abs()[(tot.index > pd.Timestamp("2027-01-01")) & (tot.index <= dec)].median()
    jt, jp = tot[jan] - tot[dec], v["Dry production, base"][jan] - v["Dry production, base"][dec]
    jo = v["Implied net outflow, BASE"][jan] - v["Implied net outflow, BASE"][dec]
    ob = v["Implied net outflow, BASE"]
    jo_prev = ob[pd.Timestamp("2028-01-01")] - ob[pd.Timestamp("2027-12-01")]      # same seasonal step a year earlier
    out.append(f"join Dec/28 -> Jan/29 (BASE): total demand {jt:+.2f} (median monthly step 2027-28 {typical:.2f}), dry production {jp:+.2f}, outflow {jo:+.2f} Bcf/d "
               f"(the same Dec -> Jan step a year earlier was {jo_prev:+.2f}: winter demand) " + ("** JUMP **" if abs(jo - jo_prev) > 1.0 else "- no jump beyond the seasonal step"))
    return out


# --------------------------------------------------------------------------- Excel
def _hdr(ws, row, labels, c0=1):
    for j, t in enumerate(labels, start=c0):
        c = ws.cell(row, j, t)
        c.font = BOLD
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")


def _status(cell, s):
    cell.value = s
    cell.font = BOLD
    cell.fill = PatternFill("solid", start_color=FILL_STATUS[s], end_color=FILL_STATUS[s])


def write_sheets(path, v, c2, check_lines):
    par, sc = c2["par"], c2["sc"]
    cnt, cnt_txt = anchor_count()
    wb = load_workbook(path)
    for n in (AS, DC, VIEW):
        if n in wb.sheetnames:
            del wb[n]
    a = wb.create_sheet(AS)
    a["A1"] = "Texas data-centre assumptions (yellow cells are editable; 'Data centres' recalculates in Excel, 'Demand to 2033' and the charts follow on the next run)"
    a["A1"].font = Font(bold=True, size=13)
    a["A2"] = (f"Every input has a STATUS: SOURCED = a published figure (publisher, document, as-of, URL in the Sources block below); DERIVED = arithmetic on sourced figures; "
               f"PROXY = an observed number used for something it does not strictly measure; ASSUMPTION = no source, a pure scenario input. Count: {cnt_txt}. "
               "Forecasts are labelled as forecasts: none of the sourced US figures is observed Texas load.")
    a["A2"].alignment = Alignment(wrap_text=True)
    a["A3"] = LAYOUT
    a["A3"].font = Font(italic=True, color="808080")
    _hdr(a, 4, ["Input", "LOW", "BASE", "HIGH", "Unit", "Status", "Source", "Note"])
    fmt = {"us24": "0.0", "us30b": "0.0", "w30": "0", "w35": "0", "cap24": "0.0", "cap30": "0.0", "lf": "0.00", "cdr30": "0.000", "real": "0.000",
           "sh24": "0.00", "gas": "0.0%", "hr": "0.00", "mcf": "0.000"}
    for i, (k, lab, *_r, unit, status, sid, note) in enumerate(PARAMS):
        r = ROW0 + i
        a.cell(r, 1, lab)
        for j in range(3):
            c = a.cell(r, 2 + j, par[k][j])
            c.number_format = fmt[k]
            if k in EDITABLE:
                c.fill = FILL_IN
        a.cell(r, 5, unit)
        _status(a.cell(r, 6), status)
        a.cell(r, 7, sid)
        extra = {"hr": f"; {c2['hr_txt']}", "mcf": f"; {c2['hr_txt']}", "gas": f"; {c2['gas_txt']}"}.get(k, "")
        a.cell(r, 8, note + extra)
    a.cell(RDER0 - 1, 1, "Derived (live formulas, not editable)").font = BOLD
    rows = {**R, **RD}
    for i, (k, lab, unit, tpl, note) in enumerate(DERIVED):
        r = RDER0 + i
        a.cell(r, 1, lab)
        for j, col in enumerate("BCD"):
            c = a.cell(r, 2 + j, tpl.format(c=col, **R))
            c.number_format = "0.0%" if k == "shg" else "0.0"
        a.cell(r, 5, unit)
        _status(a.cell(r, 6), "DERIVED")
        a.cell(r, 7, "S1, S2, S3, S6, S7")
        a.cell(r, 8, note + ("; BASE uses the ASSUMED midpoint realisation" if k == "shg" else ""))
    a.cell(R_PERGW, 1, "Gas burn per GW of energised load (formula)").font = BOLD
    for j, col in enumerate("BCD"):
        c = a.cell(R_PERGW, 2 + j, f"=1000*{col}{R['lf']}*{col}{R['gas']}*{col}{R['hr']}*24/{col}{R['mcf']}/1000000")
        c.number_format = "0.0000"
    a.cell(R_PERGW, 5, "Bcf/d per GW")
    a.cell(R_PERGW, 8, "= 1000 MW/GW x load factor x gas share x heat rate x 24 h / (MMBtu per Mcf) / 1e6")
    a.cell(R_PERGW + 1, 1, INTERP)
    _status(a.cell(R_PERGW + 1, 6), "ASSUMPTION")
    a.cell(R_PERGW + 1, 8, "the IEA gives 2024, 2030 and (exploratory) 2035 only; real build-out is lumpy, so intermediate years are a convenience, not data")
    a.cell(SCAL0 - 1, 1, "Other inputs (not data-centre anchors, not counted)").font = BOLD
    for i, (k, lab, _v, unit, status, note) in enumerate(SCALARS):
        r = SCAL0 + i
        a.cell(r, 1, lab)
        c = a.cell(r, 2, sc[k])
        c.fill = FILL_IN
        c.number_format = "0.00"
        a.cell(r, 5, unit)
        _status(a.cell(r, 6), status)
        a.cell(r, 8, note)
    # ---- cross-checks
    r = SCAL0 + len(SCALARS) + 1
    a.cell(r, 1, "Cross-checks (published figures that are NOT anchors, set against this model)").font = BOLD
    _hdr(a, r + 1, ["Item", "LOW / low end", "BASE", "HIGH / high end", "Unit", "Type", "Source", "Note"])
    us = lambda y, ci: f"'{DC}'!{L(3 + 5 * ci)}{YR0 + y - 2024}"       # noqa: E731  US TWh of the year, case ci
    gwc = lambda y, ci: f"'{DC}'!{L(5 + 5 * ci)}{YR0 + y - 2024}"      # noqa: E731  Texas GW
    xr = r + 2

    def xrow(lab, vals, unit, typ, sid, note, fm="0"):
        nonlocal xr
        a.cell(xr, 1, lab)
        for j, x in enumerate(vals):
            if x is not None:
                c = a.cell(xr, 2 + j, x)
                c.number_format = fm
        a.cell(xr, 5, unit)
        a.cell(xr, 6, typ)
        a.cell(xr, 7, sid)
        a.cell(xr, 8, note)
        xr += 1
    xrow("US data centres 2028: LBNL range", [325, None, 580], "TWh", "forecast range", "S5", "LBNL scenarios for 2028 (6.7-12.0% of US electricity); 176 TWh in 2023")
    xrow("US data centres 2028: this model", [f"={us(2028, ci)}" for ci in range(3)], "TWh", "model", "", "BASE sits near the low end of LBNL's range: the IEA Base path is the more modest one; HIGH is below LBNL's 580", "0")
    xrow("US data centres 2030: IEA Base Case", [None, 426, None], "TWh", "forecast", "S1", "")
    xrow("US data centres 2030: this model", [f"={us(2030, ci)}" for ci in range(3)], "TWh", "model", "", "LOW / HIGH = IEA Headwinds / Lift-Off scaled to the US (derived)", "0")
    for y in (2025, 2028, 2030, 2033):
        lo, hi = c2["aeo"][y]
        xrow(f"EIA AEO2026 'Data Center Servers' electricity {y} (min / max of scenarios)", [lo, None, hi], "TWh", "forecast", "S10",
             "servers only, so well below IEA / LBNL totals (coverage differs): NOT used as an anchor; " + c2["aeo_txt"] if y == 2025 else "see 2025 row", "0")
    xrow("ERCOT data-centre forecast, summer 2030 (requested MW)", [None, CDR_GW[2030], None], "GW", "forecast", "S6", "yearly path on the 'Data centres' tab", "0.0")
    xrow("this model: Texas energised data-centre GW, Dec 2030", [f"={gwc(2030, ci)}" for ci in range(3)], "GW", "model", "",
         "BASE is built from the ERCOT forecast and the IEA US capacity addition, so agreement with S6 is by construction, not independent confirmation", "0.0")
    xrow("ERCOT large-load queue, data centres (410 GW x 87%)", [None, "=410*0.87", None], "GW", "queue (requests)", "S8",
         "requests under study, not load; ERCOT says ~410 GW of large loads seeking interconnection, ~87% data centres (26 Mar 2026)", "0")
    xrow("this model BASE Texas GW 2030 as a share of that queue", [None, f"={gwc(2030, 1)}/(410*0.87)", None], "share", "model", "", "queues are speculative; a small share is plausible", "0.0%")
    xrow("ERCOT large load eligible for Batch Zero (all types)", [None, 205, None], "GW", "queue (requests)", "S9", "Batch Zero eligibility on existing studies, 28 Jul 2026", "0")
    # ---- sensitivity to the Texas share
    r = xr + 1
    a.cell(r, 1, "Sensitivity to the Texas share of US growth (BASE case; the share is the least certain input)").font = BOLD
    _hdr(a, r + 1, ["Texas share of US growth", "Texas GW 2030", "Burn Bcf/d 2030", "Burn added since Dec 2025, Bcf/d", "", "", "", "Note"])
    for i, s_ in enumerate((0.10, 0.20, 0.30, 0.40)):
        rr = r + 2 + i
        c = a.cell(rr, 1, s_)
        c.number_format = "0%"
        c.alignment = Alignment(horizontal="left")
        tx30 = f"($C${R['us24']}*$C${R['sh24']}+$A{rr}*({us(2030, 1)}-$C${R['us24']}))"
        tx25 = f"($C${R['us24']}*$C${R['sh24']}+$A{rr}*({us(2025, 1)}-$C${R['us24']}))"
        a.cell(rr, 2, f"={tx30}/8.76/$C${R['lf']}").number_format = "0.0"
        a.cell(rr, 3, f"=B{rr}*$C${R_PERGW}").number_format = "0.00"
        a.cell(rr, 4, f"=C{rr}-{tx25}/8.76/$C${R['lf']}*$C${R_PERGW}").number_format = "0.00"
        a.cell(rr, 8, "BASE TWh path, BASE gas factors; the model's BASE share is in the Derived block" if i == 0 else "")
    # ---- sources block
    r = r + 8
    a.cell(r, 1, "SOURCES (exact figure, unit, as-of, publisher, URL, observed or forecast)").font = BOLD
    _hdr(a, r + 1, ["Source (id, publisher, document)", "", "", "", "Unit", "As-of", "Id", "Exact figure(s)", "URL", "Observed or forecast"])
    for i, (sid, pub, doc, fig, unit, asof, typ, url) in enumerate(SOURCES):
        rr = r + 2 + i
        a.cell(rr, 1, f"{sid}  {pub}: {doc}").alignment = Alignment(wrap_text=True, vertical="top")
        a.cell(rr, 5, unit)
        a.cell(rr, 6, asof).alignment = Alignment(wrap_text=True, vertical="top")
        a.cell(rr, 7, sid)
        a.cell(rr, 8, fig).alignment = Alignment(wrap_text=True, vertical="top")
        a.cell(rr, 9, url)
        a.cell(rr, 10, typ).alignment = Alignment(wrap_text=True, vertical="top")
        a.row_dimensions[rr].height = 48
    rr = r + 2 + len(SOURCES)
    a.cell(rr, 1, "Tried, not found: the IEA web pages (iea.org) return 403 from Actions - the report PDF on iea.blob.core.windows.net was readable; the LBNL web PDF "
                  "(eta-publications.lbl.gov) returns HTTP 202 with no body - the escholarship.org copy was readable; ERCOT publishes no observed data-centre GW "
                  "(only the large-load aggregate: see texas_demand_regression.xlsx 'Load sources'), no Texas data-centre share of the US total (LBNL gives only the "
                  "rank) and no state-level IEA split. Probes: discovery_archive/DATACENTRE_SOURCES_PROBE*.py (5 Oct 2026).").alignment = Alignment(wrap_text=True, vertical="top")
    a.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=10)
    a.row_dimensions[rr].height = 62
    for col, w in zip("ABCDEFGHIJ", (74, 12, 12, 12, 16, 15, 12, 120, 60, 36)):
        a.column_dimensions[col].width = w
    a.row_dimensions[2].height = 48
    a.merge_cells("A2:H2")
    a.freeze_panes = "B5"
    # ---- Data centres (live formulas)
    d = wb.create_sheet(DC)
    d["A1"] = "Texas data-centre load (year-end energised GW) and gas burn, LOW / BASE / HIGH - live formulas on 'Assump - Data centres'"
    d["A1"].font = Font(bold=True, size=13)
    d["A2"] = ("US TWh: IEA anchors 2024 / 2030 / 2035 (Base Case; LOW and HIGH scaled by the IEA's world Headwinds / Lift-Off ratio), linear in between. "
               "Texas TWh = US 2024 TWh x Texas stock share (ASSUMPTION) + Texas share of growth (DERIVED from ERCOT's CDR forecast) x (US TWh - US 2024 TWh); "
               "energised GW = TWh / 8.76 / load factor (IEA 49%); gas Bcf/d = GW x gas burn per GW (assumption tab).")
    d["A3"] = (f"Data-centre inputs: {cnt_txt}. EIA's electric-power history already contains the 2024-26 data-centre burn: the charts stack only the burn ADDED after the last "
               "EIA actual month (see 'Demand to 2033'); this tab shows the total. Last column = ERCOT's own data-centre forecast (requested MW, FORECAST) for comparison.")
    for rr in (2, 3):
        d[f"A{rr}"].alignment = Alignment(wrap_text=True)
        d.merge_cells(f"A{rr}:R{rr}")
        d.row_dimensions[rr].height = 46
    labels = ["Year (year-end)", "Basis"]
    for c in CASES:
        labels += [f"{c}: US TWh", f"{c}: Texas TWh", f"{c}: Texas energised GW", f"{c}: gas burn Bcf/d", f"{c}: added since Dec 2025, Bcf/d"]
    labels += ["ERCOT CDR Dec 2025: data centres, summer peak, GW (FORECAST, S6)"]
    _hdr(d, YR0 - 1, labels)
    d.row_dimensions[YR0 - 1].height = 62
    for i, y in enumerate(YEARS):
        r = YR0 + i
        d.cell(r, 1, y)
        d.cell(r, 2, "IEA US estimate (S1)" if y == 2024 else ("IEA anchor 2030 (S1, S2)" if y == 2030 else ("IEA anchor 2035 (S2, exploratory)" if y == 2035 else
                                                                  ("interpolated 2024-30 (ASSUMPTION)" if y < 2030 else "interpolated 2030-35 (ASSUMPTION)"))))
        if y > 2030:
            for cc in range(1, 19):
                d.cell(r, cc).fill = FILL_LIGHT
        for ci in range(3):
            col = "BCD"[ci]
            A = lambda k: f"'{AS}'!${col}${R[k]}"      # noqa: E731
            Dv = lambda k: f"'{AS}'!${col}${RD[k]}"    # noqa: E731
            c0 = 3 + 5 * ci
            us_, tx, gw, bc, inc = (L(c0 + n) for n in range(5))
            d[f"{us_}{r}"] = (f"=IF($A{r}<=2030,{A('us24')}+({Dv('us30')}-{A('us24')})*($A{r}-2024)/6,"
                              f"{Dv('us30')}+({Dv('us35')}-{Dv('us30')})*($A{r}-2030)/5)")
            d[f"{tx}{r}"] = f"={A('us24')}*{A('sh24')}+{Dv('shg')}*({us_}{r}-{A('us24')})"
            d[f"{gw}{r}"] = f"={tx}{r}/8.76/{A('lf')}"
            d[f"{bc}{r}"] = f"={gw}{r}*'{AS}'!${col}${R_PERGW}"
            d[f"{inc}{r}"] = f"={bc}{r}-{bc}${YR0 + 1}"
            for n, fm in zip((us_, tx, gw, bc, inc), ("0.0", "0.0", "0.00", "0.00", "0.00")):
                d[f"{n}{r}"].number_format = fm
        if y in CDR_GW:
            d.cell(r, 18, CDR_GW[y]).number_format = "0.0"
    d.column_dimensions["A"].width = 14
    d.column_dimensions["B"].width = 38
    for cc in range(3, 19):
        d.column_dimensions[L(cc)].width = 13
    d.freeze_panes = f"C{YR0}"
    # ---- values view
    w = wb.create_sheet(VIEW)
    cols = ["Month"] + list(v.columns)
    _hdr(w, 1, cols)
    w.row_dimensions[1].height = 62
    for i, (m, row) in enumerate(v.iterrows(), start=2):
        w.cell(i, 1, m.to_pydatetime()).number_format = "mmm/yy"
        for j, x in enumerate(row.values, start=2):
            if isinstance(x, str):
                w.cell(i, j, x)
            elif x is not None and not (isinstance(x, float) and math.isnan(x)):
                w.cell(i, j, float(x)).number_format = "0.00"
            if row["Type"] == "Scenario":
                w.cell(i, j).fill = FILL_LIGHT
    for j in range(1, len(cols) + 1):
        w.column_dimensions[L(j)].width = 13
    w.freeze_panes = "C2"
    base = [s for s in wb._sheets if s.title not in (AS, DC, VIEW)]
    # keep the new tabs together just before 'Supply and demand' (or at the end)
    names = [s.title for s in base]
    pos = names.index("Supply and demand") + 1 if "Supply and demand" in names else len(base)
    wb._sheets = base[:pos] + [wb[AS], wb[DC], wb[VIEW]] + base[pos:]
    root, ext = os.path.splitext(path)
    tmp = f"{root}.tmp{os.getpid()}{ext}"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def notes_lines(c2, check_lines):
    par = c2["par"]
    cnt, cnt_txt = anchor_count()
    return [
        "", "DATA CENTRES AND THE 2033 VIEW (americas/TEXAS_DATACENTRE.py; Bcf/d)",
        "Tabs 'Assump - Data centres' (yellow cells editable, survive reruns; every input has a status SOURCED / DERIVED / PROXY / ASSUMPTION and a Sources block with publisher, figure, as-of, URL), 'Data centres' (live formulas: year-end energised Texas data-centre GW and gas burn, LOW/BASE/HIGH, 2024-33) and 'Demand to 2033' (script values, month by month to Dec 2033, charted).",
        f"Data-centre inputs: {cnt_txt}. SOURCED: IEA 'Energy and AI' (Apr 2025) US data-centre electricity 183 TWh (2024) and 426 TWh (2030, Base Case), world totals by case (Headwinds / Base / Lift-Off) for 2030 and 2035, US installed capacity 42 / 100 GW and load factor 49%; ERCOT's data-centre forecast for summer 2030 (CDR Dec 2025, 22.2 GW requested, a forecast); heat rate and MMBtu/Mcf (EIA-923 calibration in ercot_gas_burn_daily.xlsx). DERIVED: US 2030 / 2035 for LOW and HIGH (IEA Base scaled by the world case ratio) and the Texas share of US growth (ERCOT forecast x realisation / IEA US capacity addition = {par['cdr30'][1] * par['real'][0] / (par['cap30'][1] - par['cap24'][1]):.0%} / {par['cdr30'][1] * par['real'][1] / (par['cap30'][1] - par['cap24'][1]):.0%} / {par['cdr30'][1] * par['real'][2] / (par['cap30'][1] - par['cap24'][1]):.0%}). PROXY: gas share = ERCOT's observed average gas share of generation (EIA-930), {par['gas'][0]:.1%} / {par['gas'][1]:.1%} / {par['gas'][2]:.1%}. ASSUMPTION: the realisation ratio's BASE midpoint (LOW 0.498 is ERCOT's observed site ratio), the Texas stock share (no effect on the charts) and the linear interpolation between IEA anchors.",
        "Cross-checks, not anchors: LBNL 2024 report (176 TWh in 2023; 325-580 TWh in 2028), EIA AEO2026 'Data Center Servers' (servers only, well below the totals), ERCOT's large-load queue (about 410 GW, ~87% data centres, 26 Mar 2026) and Batch Zero eligibility (about 205 GW, 28 Jul 2026). The IEA web pages return 403 from Actions: the report PDF on iea.blob.core.windows.net was used.",
        f"Gas = GW x 1000 x load factor {par['lf'][1]:.2f} x gas share x heat rate {c2['hr']:.2f} MMBtu/MWh x 24 / {c2['mcf']:.3f} MMBtu per Mcf / 1e6 = {c2['per_gw']['BASE']:.3f} Bcf/d per GW in BASE (heat rate and MMBtu/Mcf: mean of the last 12 calibrated ERCOT months, {c2['hr_txt']}). {c2['gas_txt']}.",
        "Only burn ADDED after the last EIA actual month is stacked on the sector demand (EIA's electric-power history already includes the present data-centre load); the electric-power sector is held flat after 2028 so the trend and the data-centre layer do not double count.",
        "2029-33 is a SCENARIO (lighter bars): LNG = the existing train table (last train 2031) at steady utilisation; other sectors = the seasonal-trend base extended; production = STEO to Dec 2027, 2028 as in the base forecast, then region growth damped by the editable factor (extension, not STEO) and capped by the takeaway table with capacity held after the last listed pipeline.",
        "Implied net outflow = dry production (base) + storage withdrawal - demand incl. LNG feedgas, Mexico and data centres - storage injection (storage: gross EIA flows, forecast = seasonal pattern of the last 3 years, see the production workbook's Assumptions; the old no-storage outflow is kept as 'Implied net outflow before storage' columns), for LOW / BASE / HIGH data-centre cases; a negative value means Texas production cannot meet Texas demand and gas must flow in from other states - it is shown, never clipped.",
        "Sanity checks of the last run: " + "; ".join(check_lines),
    ]
