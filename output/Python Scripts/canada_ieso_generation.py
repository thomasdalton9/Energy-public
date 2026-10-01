"""
Pull Ontario's (Canada) hourly generation by fuel type from IESO's
public Generator Output and Capability report and maintain a growing
daily archive.

Data source: https://reports-public.ieso.ca/public/GenOutputCapability/PUB_GenOutputCapability.xml
Found via CANADA_IESO_DISCOVERY.py/DISCOVERY2.py - the old reports.ieso.ca
domain is dead (404); the live one is reports-public.ieso.ca. No key
needed. Confirmed live via CANADA_IESO_XML_PARSE_INSPECT.py: 188
Generator elements, each with GeneratorName, FuelType (one of BIOFUEL,
GAS, HYDRO, NUCLEAR, OTHER, SOLAR, WIND - Ontario has no coal
generation left), and an Outputs block of hourly {Hour, EnergyMW}
readings (actual generation - Capabilities/Capacities are separate
forecast-style blocks, not used here).

This XML is IESO's rolling "today" document - it only ever covers the
current operating day (hours populate through the day as they're
realized), no historical range query. Like MISO_FUEL_MIX_DAILY.py, this
upserts a daily archive: each run reads whatever hours are populated so
far and saves a day once enough of it has come in, building up history
from whenever this script first started running (does not backfill the
past).

Coverage note: Ontario only, not all of Canada - IESO is the Ontario
grid operator.
"""

import argparse
import os
import sys
import xml.etree.ElementTree as ET
from datetime import date, datetime
from zoneinfo import ZoneInfo

import pandas as pd
import requests

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # repo root, for xlsx_notes

import xlsx_notes

URL = "https://reports-public.ieso.ca/public/GenOutputCapability/PUB_GenOutputCapability.xml"
HEADERS = {"User-Agent": "Mozilla/5.0"}
TIMEOUT = (10, 45)

ONTARIO_ZONE = ZoneInfo("America/Toronto")
CATEGORIES = ["NUCLEAR", "GAS", "HYDRO", "WIND", "SOLAR", "BIOFUEL", "OTHER"]

EXPECTED_HOURS_PER_DAY = 24
MIN_HOURS_PER_DAY = 20

RENEWABLE_CATEGORIES = ["HYDRO", "WIND", "SOLAR", "BIOFUEL"]

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_OUT = os.path.join(REPO_ROOT, "output", "canada_ieso_generation_daily.xlsx")


def strip_ns(tag):
    return tag.split("}")[-1] if "}" in tag else tag


def fetch_today():
    r = requests.get(URL, headers=HEADERS, timeout=TIMEOUT)
    r.raise_for_status()
    root = ET.fromstring(r.content)

    report_date = None
    for elem in root.iter():
        if strip_ns(elem.tag) == "Date":
            report_date = elem.text.strip()
            break

    rows = []
    for gen in root.iter():
        if strip_ns(gen.tag) != "Generator":
            continue
        fuel_elem = gen.find("./{*}FuelType")
        fuel = fuel_elem.text.strip() if fuel_elem is not None and fuel_elem.text else "OTHER"
        outputs = gen.find("./{*}Outputs")
        if outputs is None:
            continue
        for output in outputs:
            if strip_ns(output.tag) != "Output":
                continue
            hour_elem = output.find("./{*}Hour")
            energy_elem = output.find("./{*}EnergyMW")
            if hour_elem is None or energy_elem is None or energy_elem.text is None:
                continue
            try:
                hour = int(hour_elem.text.strip())
                energy = float(energy_elem.text.strip())
            except ValueError:
                continue
            rows.append({"fuel": fuel if fuel in CATEGORIES else "OTHER", "hour": hour, "energy_mw": energy})

    return report_date, rows


def to_daily_row(rows, report_date):
    if not rows:
        return None, 0
    df = pd.DataFrame(rows)
    # Sum across all generators of the same fuel type, per hour -> total
    # MW by fuel by hour, then mean across the hours actually present.
    by_hour_fuel = df.groupby(["hour", "fuel"])["energy_mw"].sum().reset_index()
    n_hours = by_hour_fuel["hour"].nunique()
    if n_hours < MIN_HOURS_PER_DAY:
        return None, n_hours

    means = by_hour_fuel.groupby("fuel")["energy_mw"].mean()
    row = {f"{cat}_MW": means.get(cat, 0.0) for cat in CATEGORIES}
    total = sum(row.values())
    row["Total_MW"] = total
    renewable = sum(row.get(f"{cat}_MW", 0.0) for cat in RENEWABLE_CATEGORIES)
    row["Renewables_Share"] = renewable / total if total else None
    row["Hours_Reported"] = n_hours
    day = datetime.strptime(report_date, "%Y-%m-%d").date()
    return pd.Series(row, name=day), n_hours


def load_archive(path):
    try:
        df = pd.read_excel(path, sheet_name="Data", index_col=0)
    except (FileNotFoundError, ValueError):
        return pd.DataFrame()
    df.index = pd.to_datetime(df.index).date
    df.index.name = "date"
    return df


def upsert(existing, new_row):
    if new_row is None:
        return existing
    new_df = new_row.to_frame().T
    new_df.index.name = "date"
    if existing.empty:
        combined = new_df
    else:
        combined = pd.concat([existing, new_df])
        combined = combined[~combined.index.duplicated(keep="last")]
    combined.index.name = "date"
    return combined.sort_index()


def format_date_column(path, sheet="Data"):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    if sheet not in wb.sheetnames:
        return
    ws = wb[sheet]
    for (cell,) in ws.iter_rows(min_row=2, max_col=1):
        cell.number_format = "dd-mmm-yyyy"
    ws.column_dimensions["A"].width = 14
    wb.save(path)


NOTES_LINES = [
    "UNITS",
    "All *_MW columns are MW, the daily mean of IESO's hourly generation readings for that fuel "
    "category, summed across every generator of that fuel type - 'average MW for that day', not "
    "total daily energy. Total_MW is the sum of the category means.",
    "",
    "SCOPE",
    "Ontario only, not all of Canada - IESO is Ontario's grid operator. Categories: NUCLEAR, GAS, "
    "HYDRO, WIND, SOLAR, BIOFUEL, OTHER (Ontario has no coal generation left).",
    "Renewables_Share: (HYDRO + WIND + SOLAR + BIOFUEL) daily mean, as a share of Total_MW.",
    "Hours_Reported: how many of the expected 24 hourly readings that day actually came back; a day "
    f"is only saved if at least {MIN_HOURS_PER_DAY} came back.",
    "",
    "TIMESTAMPS",
    "The 'date' index is the report's own date field (Ontario/Eastern operating day).",
    "",
    "COVERAGE",
    "This is IESO's rolling 'today' XML document, not a historical range query - hours populate "
    "through the day as they're realized. Like MISO's /FuelMix/Today, this upserts by date into the "
    "Data tab on every run, building up history from whenever this script first started running, not "
    "backfilling the past.",
    "",
    "SOURCE",
    f"IESO (Independent Electricity System Operator, Ontario) public report: {URL}",
]
NOTES_SECTION_TITLES = {"UNITS", "SCOPE", "TIMESTAMPS", "COVERAGE", "SOURCE"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default=DEFAULT_OUT)
    args = parser.parse_args()

    report_date, rows = fetch_today()
    print(f"Report date: {report_date}, {len(rows)} (generator, hour) readings", file=sys.stderr)
    new_row, n_hours = to_daily_row(rows, report_date)

    existing = load_archive(args.out)
    before_days = set(existing.index) if not existing.empty else set()
    combined = upsert(existing, new_row)
    is_new = bool(new_row is not None and new_row.name not in before_days)

    if new_row is None:
        print(f"Not enough hours to save a day ({n_hours}/{EXPECTED_HOURS_PER_DAY}) - "
              "keeping the archive as-is.", file=sys.stderr)
    else:
        print(f"Day {new_row.name}: {n_hours}/{EXPECTED_HOURS_PER_DAY} hours "
              f"({'new' if is_new else 'refreshed'})")

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    xlsx_notes.write_workbook(args.out, {"Data": combined}, NOTES_LINES, NOTES_SECTION_TITLES)
    format_date_column(args.out)
    print(f"Saved to {args.out}")


if __name__ == "__main__":
    main()
