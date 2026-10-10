"""
Taiwan master workbook: every Taiwan dataset this repo pulls in one file, with Dashboard front pages carrying all their
charts - the same layout and chart code as the other masters (south_america/SOUTH_AMERICA_MASTER.py, china/CHINA_MASTER.py).

  Dashboard - Power      generation by source (average GW) and annual, installed capacity (GW), renewable generation,
                         electricity consumption by sector, Taipower daily peak load, supply capacity and reserve margin
  Dashboard - Gas        natural gas supply (LNG imports + domestic) and use by sector (Bcf/d), LNG imports by origin (Mt),
                         LNG import price
  Dashboard - Fuels      crude oil imports by origin, refinery intake, coal imports by origin, crude/coal prices,
                         primary energy supply by fuel
  Dashboard - Water      reservoir storage as Oct-Sep water-year charts (Water Resources Agency; history starts with the
                         first daily run - WRA's open API keeps only the last day)
  Dashboard - Long-term  annual summary and history from 2000 (fundamentals.py)
  <chart> data           the table each Dashboard chart plots
  <dataset> raw          the full data sheet(s) from each source workbook
  Sources                where each dataset comes from, units and notes, and the sources checked but not built

Sources are Taiwan's own agencies: the Energy Administration (Ministry of Economic Affairs) monthly energy statistics
(E-STAT open API; nationwide generation including IPPs and self-generation), Taipower open data (daily peak load and
reserve margin; 10-minute unit output kept as a raw tab only - it covers Taipower-dispatched units, not the national total)
and the Water Resources Agency. Ember appears only on the Long-term page. The Energy Administration's API serves the last
~24 months and annual rows from 2007; monthly history grows with every release because the pull keeps what it has stored.
World Bank does not cover Taiwan, so the Long-term page has no World Bank GDP, industry share, urbanisation or access
series; GDP growth and population come from the IMF WEO, power from Ember, fuels from the Energy Institute.

Reads (doesn't refetch) the workbooks the scheduled pulls write to "output/Data and Chart Outputs/". A missing input
is listed on its Dashboard and skipped.

Usage: python3 taiwan/TAIWAN_MASTER.py [--out "output/Data and Chart Outputs/Master Outputs/taiwan_master.xlsx"]
"""
import argparse
import os
import sys
from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
import add_charts  # noqa: E402,F401
import fundamentals  # noqa: E402
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR
ESIST = "taiwan_esist_monthly.xlsx"
TAIPOWER = "taiwan_taipower.xlsx"
RESERVOIRS = "taiwan_reservoirs_daily.xlsx"
ES_RAW = tuple(f"{t} {k}" for t in ("GEN", "CAP", "CONS", "REN", "GAS", "LNG", "SUP", "CRUDE", "CRUDESRC", "COAL", "COALSRC",
                                    "IMPPRICE", "OILPRICE") for k in ("M", "A")) + ("Series", "Releases")

# (code, country, workbook, raw sheets, short name, charts to show (None = all))
POWER = [("TW", "Taiwan", ESIST, ES_RAW, "EA monthly statistics",
          {"Generation", "Generation annual", "Capacity", "Capacity annual", "Renewables", "Consumption"}),
         ("TW", "Taiwan", TAIPOWER, ("Peak", "Daily"), "Taipower", None)]
GAS_SETS = [("TW", "Taiwan", ESIST, (), "EA monthly statistics",
             {"Gas supply", "Gas use", "Gas use annual", "LNG imports", "LNG imports annual", "LNG price"})]
FUELS = [("TW", "Taiwan", ESIST, (), "EA monthly statistics",
          {"Crude imports", "Refinery intake", "Coal imports", "Oil price", "Coal price", "Energy supply annual"})]
WATER = [("TW", "Taiwan", RESERVOIRS, ("Daily storage", "Daily level", "Reservoirs"), "WRA reservoirs", None)]

SOURCES = {
    ESIST: ("Energy Administration, Ministry of Economic Affairs (Taiwan): Monthly Energy Statistics, E-STAT open API "
            "(nationwide generation and capacity - 'Other' on the charts = biomass, waste, geothermal and pumped-storage output - "
            "consumption, natural gas, LNG, crude oil, coal, import prices; origin charts show the six largest origins and 'Others')",
            "https://ea01.moeaea.gov.tw/a0303/02/newest/monthly/"),
    TAIPOWER: ("Taiwan Power Company (Taipower) open data: past electricity supply and demand (peak load, supply capacity, "
               "reserve margin; d006005) and 10-minute generation by unit (d006010; raw tab only)",
               "https://data.gov.tw/dataset/19995"),
    RESERVOIRS: ("Water Resources Agency, Ministry of Economic Affairs: reservoir water-regime open data (dataset 45501)",
                 "https://data.gov.tw/dataset/45501"),
}
MASTER_SPECS = {}
HYDRO_DATASETS = []
HYDRO_EXTRA = {}
EMBER = set()
OPERATORS = {}
DASHBOARD_ONLY = {}

# Sources checked and NOT built (probes: discovery_archive/taiwan/, results discovery_archive/results/taiwan/, 10 Oct 2026).
NOT_AVAILABLE = [
    ("Taiwan", "Taipower real-time generation by fuel and unit (today)", "service.taipower.com.tw/data/opendata d006001",
     "Reachable, not used", "https://service.taipower.com.tw/data/opendata/apply/file/d006001/001.json",
     "Today's 10-minute snapshot by unit with no history; the 10-minute past-generation file (d006010) is pulled instead."),
    ("Taiwan", "Taipower past generation by unit (10-minute)", "service.taipower.com.tw d006010 (data.gov.tw 37331)",
     "Reachable, raw tab only", "https://data.gov.tw/dataset/37331",
     "A rolling three-month window (about 200 MB), replaced monthly and ending about three months back, so history only builds "
     "from the first run (window on 10 Oct 2026: 1 Mar - 31 May 2026). It covers Taipower-dispatched (EMS) units only - about 16 GW "
     "average against about 33 GW nationwide, wind and solar far below the national totals - so it is kept as the 'Daily' raw "
     "tab and not charted; the Energy Administration series is the national total."),
    ("Taiwan", "Taipower website (taipower.com.tw pages and genary.json)", "www.taipower.com.tw", "Blocked (HTTP 403)",
     "https://www.taipower.com.tw/", "CloudFront 403 from GitHub Actions; service.taipower.com.tw/data/opendata answers."),
    ("Taiwan", "Water Resources Agency reservoir history and names", "fhy.wra.gov.tw, data.wra.gov.tw, opendata.wra.gov.tw",
     "Partly reachable", "https://opendata.wra.gov.tw/",
     "The keyless open-data API (dataset 45501) returns only the last ~24 hours, so the reservoir pull runs daily and history starts "
     "with its first run. fhy.wra.gov.tw/Api/v2 answers 'MISSING_API_KEY', fhy.wra.gov.tw/WraApi/v1 returns HTTP 503, data.wra.gov.tw "
     "(legacy open data, with reservoir names and capacities) does not resolve from Actions, and opendata.wra.gov.tw/api/v1 needs "
     "a login. The reservoir name table could therefore not be read: charts are labelled by WRA reservoir identifier, and the unit "
     "of the storage field is not stated by the API (10,000 m3 by WRA convention)."),
    ("Taiwan", "Customs trade statistics (LNG, crude, coal by origin)", "portal.sw.nat.gov.tw, cuswebo.trade.gov.tw, web02.mof.gov.tw",
     "Not reachable", "https://portal.sw.nat.gov.tw/APGA/GA30",
     "No connection from GitHub Actions (portal.sw.nat.gov.tw, cuswebo.trade.gov.tw, cus.trade.gov.tw); web02.mof.gov.tw answers with "
     "an empty page. The Energy Administration tables give LNG (by origin, tonnes), crude (by origin, barrels) and coal (by origin, "
     "tonnes) imports and import prices from the same flows, and are used instead."),
    ("Taiwan", "CPC Corporation LNG receipts, LNG terminal stocks / days of cover", "cpc.com.tw", "No data series found",
     "https://www.cpc.com.tw/en/", "The CPC site answers but publishes no LNG receipt or inventory series; none was found "
     "(no stock or days-of-cover figure is built or estimated)."),
    ("Taiwan", "Taipower tariff / spot price", "taipower.com.tw", "Not built",
     "https://www.taipower.com.tw/", "Taiwan has no wholesale spot market and the tariff pages are blocked from Actions."),
    ("Taiwan", "World Bank GDP, industry share, urbanisation, electricity access", "World Bank API", "Taiwan not covered",
     "https://api.worldbank.org/v2/country/TWN", "The World Bank API rejects the code TWN; the Long-term page therefore has no GDP "
     "level, GDP per head, industry share or urbanisation for Taiwan. IMF WEO GDP growth and population (code TWN) and "
     "degree days (Taipei, Kaohsiung, Taichung; approximate metro populations 7.0, 2.7 and 2.8 million) are shown."),
    ("Taiwan", "Energy Administration monthly history beyond ~24 months", "ea01.moeaea.gov.tw database query",
     "No API found", "https://ea01.moeaea.gov.tw/a0303/02/database/search/electric-generation/",
     "The database-query pages have no open endpoint in their scripts (probe 9); the monthly tables of the open API cover 24 months "
     "and annual rows from 2007, and the committed workbook accumulates the months from the first run."),
]


def collect(wb, datasets, data_dir, used, sources):
    """sam.collect per dataset so each can show a chosen subset of its charts (names = spec names; None = all)."""
    charts, rows, missing = [], [], []
    for code, country, fname, raw, short, only in datasets:
        cfg = SimpleNamespace(**{k: getattr(sys.modules[__name__], k) for k in (
            "SOURCES", "MASTER_SPECS", "HYDRO_DATASETS", "HYDRO_EXTRA", "EMBER", "OPERATORS")},
            DASHBOARD_ONLY={fname: only} if only else {}, GAS_BCFD=True)   # gas volumes in Bcf/d
        c, r, m = sam.collect(wb, [(code, country, fname, raw, short)], data_dir, used, sources, cfg=cfg)
        charts += c
        rows += r
        missing += m
    return charts, rows, missing


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "Master Outputs", "taiwan_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()

    wb = Workbook()
    dash_power = wb.active
    dash_power.title = "Dashboard - Power"
    names = ["Dashboard - Gas", "Dashboard - Fuels", "Dashboard - Water"]
    dashes = [wb.create_sheet(n) for n in names]
    used = {"Dashboard - Power", *names, "Dashboard - Long-term", "Sources"}
    sources = []
    sections = [(dash_power, "Taiwan - power generation (average GW), capacity, consumption and peak load", POWER),
                (dashes[0], "Taiwan - natural gas (Bcf/d) and LNG imports", GAS_SETS),
                (dashes[1], "Taiwan - crude oil, coal and primary energy", FUELS),
                (dashes[2], "Taiwan - reservoir storage, Oct-Sep water year (WRA)", WATER)]
    counts = []
    for dash, heading, datasets in sections:
        charts, rows, missing = collect(wb, datasets, args.data_dir, used, sources)
        sam.draw_dashboard(dash, heading, charts, rows, missing)
        counts.append(len(charts))
        if missing:
            print(f"missing ({heading}):", missing)

    fundamentals.add_long_term_dashboard(wb, used, sam.sheet_name, "Taiwan", args.data_dir, sam.CHART_W, sam.CHART_H,
                                         rows_per_chart=sam.ROWS_PER_CHART, degree_days=("CDD_18", "HDD_18"),
                                         index=sum(s.startswith("Dashboard") for s in wb.sheetnames))

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    seen = set()
    for s in sources:
        if s[2] in seen:   # a workbook read for two dashboards is listed once
            continue
        seen.add(s[2])
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    if NOT_AVAILABLE:
        src.append([])
        src.append(["Checked, NOT built", "Dataset", "Publisher / page", "Result from GitHub Actions", "Link", "Notes"])
        for cell in src[src.max_row]:
            cell.font = Font(bold=True)
        for row in NOT_AVAILABLE:
            src.append(list(row))
            if row[4]:
                src.cell(row=src.max_row, column=5).hyperlink = row[4]
    for col, w in (("A", 18), ("B", 22), ("C", 44), ("D", 60), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: charts per dashboard {dict(zip(['power', *names], counts))}; tabs {wb.sheetnames}")


if __name__ == "__main__":
    main()
