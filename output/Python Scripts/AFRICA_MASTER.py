"""
Africa master workbook: one file with every African dataset this repo pulls from a raw source (grid operator,
regulator, ministry, national oil company, statistics office - NO Ember), and Dashboard front pages carrying all
their charts. Same layout as the North America / South & Central America masters (south_america/SOUTH_AMERICA_MASTER.py,
whose table/chart/dashboard code this reuses).

  Dashboard          - gas and oil: Nigeria (NNPC Ltd monthly report)
  Dashboard - Power  - generation by source, then lake levels and capacity where a raw feed exists:
                       South Africa (Eskom), Ghana (Energy Commission: weekly WEM statistics, annual statistics with
                       Akosombo / Bui water-year level charts), Nigeria (NERC quarterly report)
  <CC> <chart> data  - the table each Dashboard chart plots (CC = ZA, GH, NG ... : one tab group per country)
  <CC> <dataset> raw - the full data sheet(s) from each source workbook
  Sources            - where each dataset comes from, units and notes

Reads (doesn't refetch) the workbooks the scheduled pulls write to "output/Data and Chart Outputs/". A missing
input is listed on the Dashboard and skipped rather than stopping the rest.

TO ADD A COUNTRY: append its workbook to the lists below -
  DATASETS            gas / oil / other non-power workbooks   (Dashboard)
  RAW_POWER_DATASETS  generation workbooks                    (Dashboard - Power)
  OTHER_POWER_DATASETS  demand, water levels, gas burn, ...   (Dashboard - Power)
  CAPACITY_DATASETS   installed capacity workbooks            (Dashboard - Power)
each as (country code, country, workbook, raw sheet / tuple of sheets / "*", short name), add a SOURCES entry
(publisher + link), and register the workbook's chart specs in add_charts.py REGISTRY (or MASTER_SPECS here).

Usage: python3 AFRICA_MASTER.py [--out "output/Data and Chart Outputs/Master Outputs/africa_master.xlsx"]
"""
import argparse
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR

# (country code, country, workbook, raw sheet / tuple of raw sheets / "*", short dataset name)
DATASETS = [
    ("NG", "Nigeria", "nigeria_nnpc_gas_monthly.xlsx", "Monthly", "NNPC gas and oil"),
]
RAW_POWER_DATASETS = [
    ("ZA", "South Africa", "south_africa_generation_mix_daily.xlsx",
     ("Data", "System", "Weekly EAF", "Pumped storage", "Load shedding"), "power"),
    ("GH", "Ghana", "ghana_wem_weekly_generation_daily.xlsx", ("Daily", "Plants_GWh"), "WEM weekly power"),
    ("NG", "Nigeria", "nigeria_power_generation_quarterly.xlsx", ("Quarterly", "Plants"), "power"),
]
OTHER_POWER_DATASETS = [
    # annual generation by type plus Akosombo / Bui water-year level charts
    ("GH", "Ghana", "ghana_energy_commission_power.xlsx", ("Annual", "Akosombo", "Bui"), "annual statistics and lake levels"),
]
CAPACITY_DATASETS = []
HYDRO_DATASETS = []
HYDRO_EXTRA = {}
DASHBOARD_ONLY = {}
MASTER_SPECS = {}
EMBER = set()          # no Ember fallback in this master
OPERATORS = {}
GAS_BCFD = True

SOURCES = {
    "south_africa_generation_mix_daily.xlsx": (
        "Eskom Data Portal, Station Build Up (hourly generation by station, rolling last 7 days; this repo archives it "
        "every 3 days, so history starts when the pull started)", "https://www.eskom.co.za/dataportal/"),
    "ghana_wem_weekly_generation_daily.xlsx": (
        "Energy Commission of Ghana, Weekly Wholesale Electricity Market (WEM) Statistics (market-operator dispatch data)",
        "https://www.energycom.gov.gh/index.php/planning/weekly-wholesale-electricity-market-wem-statistics"),
    "ghana_energy_commission_power.xlsx": (
        "Energy Commission of Ghana, National Energy Statistics (GRIDCo, VRA, ECG and IPP returns)",
        "https://www.energycom.gov.gh/index.php/planning/energy-statistics"),
    "nigeria_power_generation_quarterly.xlsx": (
        "NERC (Nigerian Electricity Regulatory Commission), Quarterly Reports: NISO-metered generation of grid-connected plants",
        "https://nerc.gov.ng/resource-category/nerc-reports/"),
    "nigeria_nnpc_gas_monthly.xlsx": (
        "NNPC Ltd, Monthly Report Summary (gas production and sales, crude and condensate production)",
        "https://nnpcgroup.com/insights"),
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "Master Outputs", "africa_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()
    cfg = sys.modules[__name__]

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    dash2 = wb.create_sheet("Dashboard - Power")
    used = {"Dashboard", "Dashboard - Power", "Sources"}
    sources = []

    gas = sam.collect(wb, DATASETS, args.data_dir, used, sources, cfg=cfg)
    # one block per country (dashboard order = order of first appearance in the lists), generation then levels/capacity
    pd_all = RAW_POWER_DATASETS + OTHER_POWER_DATASETS + CAPACITY_DATASETS
    order = {}
    for d in pd_all:
        order.setdefault(d[1], len(order))
    power_sets = sorted(pd_all, key=lambda d: order[d[1]])   # stable: keeps list order within a country
    power = sam.collect(wb, power_sets, args.data_dir, used, sources, cfg=cfg)

    sam.draw_dashboard(dash, "Africa energy - gas and oil dashboard", *gas)
    sam.draw_dashboard(dash2, "Africa energy - power generation (South Africa, Ghana, Nigeria)", *power)

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    for s in sources:
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    for col, w in (("A", 14), ("B", 24), ("C", 38), ("D", 40), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: {len(gas[0])} gas charts, {len(power[0])} power charts; tabs {wb.sheetnames}")
    for label, m in (("gas", gas[2]), ("power", power[2])):
        if m:
            print(f"missing ({label}):", m)


if __name__ == "__main__":
    main()
