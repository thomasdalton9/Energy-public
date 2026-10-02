"""
Add native Excel charts to a pulled workbook (see CLAUDE.md: every xlsx
we pull carries charts of its data). Run by each scheduled workflow
right after its pull:

    python3 add_charts.py "output/Data and Chart Outputs/colombia_gas_demand_by_sector.xlsx"

A registry keyed by file name says what to chart. Each entry returns a
list of chart specs; each spec becomes its own sheet ("Chart", "Chart -
<name>") holding the plotted table and a native Excel chart, rebuilt on
every run. Storage/level series use the AGSI-style water-year chart.
Files not in the registry get a generic line chart of their first data
sheet's numeric columns, so nothing is left without a chart.
"""
import os
import re
import sys
import warnings

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import water_year_chart  # noqa: E402
import xlsx_charts  # noqa: E402

warnings.filterwarnings("ignore")


def read(path, sheet):
    df = pd.read_excel(path, sheet_name=sheet)
    return df.drop(columns=[c for c in df.columns if str(c).startswith("Unnamed")])


def by_date(df, col):
    df = df.copy()
    df[col] = pd.to_datetime(df[col].astype(str), errors="coerce")
    return df.dropna(subset=[col]).set_index(col).sort_index()


def cols(df, *names):
    return [c for c in names if c in df.columns]


def spec(name, df, title, units, kind="line", date_format="%Y-%m", line_cols=()):
    return {"name": name, "df": df, "title": title, "units": units, "kind": kind, "date_format": date_format,
            "line_cols": line_cols}


def daily(df, start=None):
    """Daily series charted from `start` (keeps native charts responsive)."""
    return df[df.index >= start] if start else df


def monthly_mean(df, start=None):
    """Long daily series -> monthly averages (daily stacks over years are unreadable)."""
    d = daily(df, start)
    return d.resample("MS").mean().dropna(how="all")


def weekly_mean(df, start=None):
    d = daily(df, start)
    return d.resample("W-MON", label="left", closed="left").mean().dropna(how="all")


UNIT_SUFFIX = re.compile(r"_(Mt|Bcm|GWh|MWh|MW|TWh|10k_units|10k_kW|kt)$")


def label_and_unit(col):
    m = UNIT_SUFFIX.search(str(col))
    unit = {"10k_units": "10,000 units", "10k_kW": "10,000 kW"}.get(m.group(1), m.group(1)) if m else ""
    label = UNIT_SUFFIX.sub("", str(col)).replace("_", " ")
    return label, unit


# ------------------------------------------------------------------ registry

def per_day(df):
    """Monthly totals (million m3 per month) -> monthly averages in million m3/day."""
    return df.div(df.index.days_in_month, axis=0)


AR_DOMESTIC = "Net domestic supply (injected into pipelines + direct)"
AR_WITH_IMPORTS = "Net domestic supply + imports (LNG, Bolivia, Chile)"


def argentina(p):
    """Argentina gas (ARGENTINA_GAS.py), all charts in mcm/d (monthly averages): demand + exports vs measured net
    supply and imports; gross production vs net injection; imports; exports by destination; daily supply."""
    sheets = pd.ExcelFile(p).sheet_names
    d = by_date(read(p, "National"), "date")
    d = d[cols(d, "produccion_gas_natural", "centrales_electricas", "industria", "residencial", "comercial",
               "entes_oficiales", "gnc")].apply(pd.to_numeric, errors="coerce")
    out = pd.DataFrame({"Power": d.get("centrales_electricas"), "Industry": d.get("industria"),
                        "Residential": d.get("residencial"),
                        "Commercial & public": d.get("comercial") + d.get("entes_oficiales"),
                        "Vehicle CNG": d.get("gnc")}, index=d.index).dropna(how="all")
    exports = by_date(read(p, "Exports by destination"), "date") if "Exports by destination" in sheets else None
    net = by_date(read(p, "Supply net"), "date") if "Supply net" in sheets else None
    check = by_date(read(p, "Balance check"), "date") if "Balance check" in sheets else None
    deliv = by_date(read(p, "Deliveries (ENARGAS)"), "date") if "Deliveries (ENARGAS)" in sheets else None
    if deliv is not None and "Tra | RTP" in deliv:
        # ENARGAS deliveries the sector series leaves out: small local utilities (end users, mostly homes and small
        # business) and TGS's Cerri processing plant (NGL extraction + fuel, i.e. system use, not end demand)
        sub = deliv[cols(deliv, "Dis | Subdistribuidor", "Tra | Subdistribuidor")].sum(axis=1, min_count=1)
        out["Sub-distributors (local utilities)"] = sub.reindex(out.index)
        out["Cerri processing plant (system use)"] = deliv["Tra | RTP"].reindex(out.index)
    elif check is not None and "Other_deliveries_subdistributors_Cerri" in check:
        out["Sub-distributors & Cerri plant"] = check["Other_deliveries_subdistributors_Cerri"].reindex(out.index)
    if exports is not None and "Total_exports" in exports:
        out["Exports"] = exports["Total_exports"].reindex(out.index)
    lines = ()
    if net is not None and "Net_domestic_supply" in net:
        # supply lines stacked: domestic, then domestic + imports, so the top line compares with the bars
        out[AR_DOMESTIC] = net["Net_domestic_supply"].reindex(out.index)
        out[AR_WITH_IMPORTS] = net["Total_net_supply"].reindex(out.index)
        lines = (AR_DOMESTIC, AR_WITH_IMPORTS)
        out = out.dropna(subset=[AR_WITH_IMPORTS])      # months with both sides measured
        title = "Argentina gas demand + exports vs net supply and imports"
    else:
        title = "Argentina gas demand by sector + exports"
    specs = [spec("Demand", per_day(out), title, "million m3/day (monthly average)", "stacked_bar", line_cols=lines)]
    if net is not None and "produccion_gas_natural" in d:
        g = pd.DataFrame({"Gross production (wellhead)": d["produccion_gas_natural"],
                          "Net domestic supply (pipelines + direct)": net.get("Net_domestic_supply"),
                          "Injected into pipelines (domestic basins)": net.get("Domestic_injection")})
        g = g.dropna(subset=["Gross production (wellhead)", "Injected into pipelines (domestic basins)"])
        specs.append(spec("Production", per_day(g), "Argentina gross gas production vs net injection",
                          "million m3/day (monthly average)", "line"))
    if net is not None:
        imp = net[cols(net, "LNG_Escobar", "LNG_BahiaBlanca", "Imports_Bolivia", "Imports_Chile")].rename(
            columns={"LNG_Escobar": "LNG Escobar", "LNG_BahiaBlanca": "LNG Bahia Blanca",
                     "Imports_Bolivia": "Bolivia (pipeline)", "Imports_Chile": "Chile (pipeline)"})
        imp = imp.dropna(how="all")
        imp = imp.loc[:, imp.fillna(0).ne(0).any()]
        if not imp.empty:
            specs.append(spec("Imports", per_day(imp), "Argentina gas imports (LNG and pipeline)",
                              "million m3/day (monthly average)", "stacked_bar"))
    if exports is not None:
        dest = exports.drop(columns=[c for c in exports.columns if c in ("Total_exports", "country")])
        dest = dest.loc[:, dest.fillna(0).ne(0).any()]   # destinations with any flow since the start date
        if not dest.empty:
            specs.append(spec("Exports", per_day(dest.rename(columns=lambda c: str(c).replace("_", " "))),
                              "Argentina gas exports by destination", "million m3/day (monthly average)",
                              "stacked_bar"))
    if "Supply daily" in sheets:
        s = by_date(read(p, "Supply daily"), "date")
        if "Domestic_injection_mcmd" in s:
            start = pd.Timestamp(year=s.index.max().year - (1 if s.index.max().month >= 5 else 2), month=5, day=1)
            day = pd.DataFrame({"Domestic injection": s["Domestic_injection_mcmd"],
                                "LNG (Escobar, Bahia Blanca)": s[cols(s, "LNG_Escobar_mcmd",
                                                                      "LNG_BahiaBlanca_mcmd")].sum(axis=1, min_count=1),
                                "Pipeline imports (Bolivia, Chile)": s[cols(s, "Imports_Bolivia_mcmd",
                                                                            "Imports_Chile_mcmd")].sum(axis=1, min_count=1)})
            day = daily(day.dropna(subset=["Domestic injection"]), start)
            if len(day) > 30:
                specs.append(spec("Supply daily", day, "Argentina daily gas supply: domestic injection + imports",
                                  "million m3/day", "stacked_area", date_format="%Y-%m-%d"))
    return specs


def brazil(p):
    d = by_date(read(p, "Demand by segment"), "date")
    out = [spec("Demand", d[cols(d, "Power_Generation", "Industrial", "Residential", "Commercial", "Automotive",
                                 "Cogeneration", "Other_incl_CNG")].rename(
                    columns={"Power_Generation": "Power generation", "Automotive": "Vehicle CNG",
                             "Other_incl_CNG": "Other"}),
                "Brazil gas demand by segment (MME, nothing newer published)", "million m3/day", "stacked_bar")]
    sheets = pd.ExcelFile(p).sheet_names
    if "Grid demand (ANP)" in sheets:   # pipeline-grid deliveries by consumer type, current (ANP)
        g = by_date(read(p, "Grid demand (ANP)"), "date")
        g = g[cols(g, "Power_Generation", "Distributors", "Refineries", "Fertiliser", "Other")].copy()
        part = g.drop(columns="Other").isna().any(axis=1)
        g.loc[part] = float("nan")              # months ANP reported only in part: an empty bar, not a partial stack
        g["Other"] = g["Other"].fillna(0.0)     # (keeps those rows, so the gap shows on the time axis)
        out.append(spec("Grid demand", g
                        .rename(columns={"Power_Generation": "Power generation", "Distributors": "Distributors (city gates)"}),
                        "Brazil gas delivered from the transport pipelines, by consumer (ANP)", "million m3/day",
                        "stacked_bar"))
    if "Supply (ANP)" in sheets:
        s = by_date(read(p, "Supply (ANP)"), "date")
        s = s[cols(s, "Domestic_Available", "Bolivia_Pipeline", "Argentina_Pipeline", "LNG_Implied")].dropna(how="any")
        out.append(spec("Supply", s.rename(columns={"Domestic_Available": "Domestic (available)",
                                                    "Bolivia_Pipeline": "Bolivia (pipeline)",
                                                    "Argentina_Pipeline": "Argentina (pipeline)",
                                                    "LNG_Implied": "LNG (imports less pipeline)"}),
                        "Brazil gas supply by source (ANP)", "million m3/day", "stacked_bar"))
    return out


def bolivia(p):
    d = by_date(read(p, "Demand by sector"), "Month")
    c = [x for x in d.columns if x.endswith("_mcm_per_day") and not x.startswith("Total")]
    specs = [spec("Demand", d[c].rename(columns=lambda x: x.replace("_mcm_per_day", "").replace("_", " ")),
                  "Bolivia gas demand by sector", "million m3/day", "stacked_bar")]
    try:
        pe = by_date(read(p, "Production and exports"), "Month")
    except ValueError:   # older workbook without the production / exports sheet
        return specs
    names = {"Exports_Brazil_mcm_per_day": "Exports to Brazil", "Exports_Argentina_mcm_per_day": "Exports to Argentina",
             "Exports_other_mcm_per_day": "Exports, other", "Domestic_market_mcm_per_day": "Domestic market",
             "Production_mcm_per_day": "Production"}
    bal = pe[cols(pe, *names)].rename(columns=names)
    # months where production and every use are published (INE's production / domestic tables lag the customs data)
    bal = bal[bal.notna().all(axis=1)]
    specs.append(spec("Production", bal, "Bolivia gas production vs exports + domestic demand", "million m3/day",
                      "stacked_bar", line_cols=("Production",)))
    exp = pe[cols(pe, "Exports_Brazil_mcm_per_day", "Exports_Argentina_mcm_per_day", "Exports_other_mcm_per_day")]
    specs.append(spec("Exports", exp.rename(columns=names).dropna(how="all"),
                      "Bolivia gas exports by destination (INE customs)", "million m3/day", "stacked_bar"))
    return specs


def peru(p):
    d = by_date(read(p, "Demand by sector"), "Month")
    names = {"Power": "Power", "Industrial": "Industrial", "Vehicle_CNG": "Vehicle CNG",
             "Residential_commercial": "Residential & commercial"}
    out = [spec("Demand", d[cols(d, *[f"{k}_mcm_per_day" for k in names])].rename(
        columns=lambda x: names[x.replace("_mcm_per_day", "")]), "Peru gas demand by sector", "million m3/day",
        "stacked_bar")]
    sheets = pd.ExcelFile(p).sheet_names
    if "Production by lot" in sheets:
        s = by_date(read(p, "Production by lot"), "Month")
        lots = {"Lot_88": "Camisea Lot 88", "Lot_56": "Camisea Lot 56", "Lot_57": "Lot 57 (Repsol)",
                "Aguaytia_31C": "Aguaytia 31C", "Northwest": "Northwest", "Other": "Other lots"}
        g = s[cols(s, *[f"{k}_mcm_per_day" for k in lots])].rename(columns=lambda x: lots[x.replace("_mcm_per_day", "")])
        g = g.loc[:, g.abs().sum() > 0]
        # demand lines stacked: domestic, then domestic + LNG exports, so the top line compares with production
        g["Domestic demand"] = d.get("Total_mcm_per_day")
        lines = ("Domestic demand",)
        if "LNG exports" in sheets:
            lng = by_date(read(p, "LNG exports"), "Month").get("LNG_exports_mcm_per_day")
            if lng is not None:
                g["Domestic demand + LNG exports"] = g["Domestic demand"] + lng.reindex(g.index).fillna(0)
                lines += ("Domestic demand + LNG exports",)
        out.append(spec("Supply", g.dropna(subset=[g.columns[0]]),
                        "Peru gas supply vs demand + LNG", "million m3/day",
                        "stacked_bar", line_cols=lines))
    return out


def uruguay(p):
    d = by_date(read(p, "Demand by sector"), "Month")
    c = [f"{x}_mcm_per_day" for x in ("Residential", "Commercial", "Industrial", "Power", "Energy_own_use")]
    return [spec("Demand", d[cols(d, *c)].rename(columns=lambda x: x.replace("_mcm_per_day", "").replace("_", " ")),
                 "Uruguay gas demand by sector", "million m3/day", "stacked_bar")]


def chile_imports(p):
    """Chile gas (CHILE_GAS_IMPORTS.py): imports by use / terminal region (CNE import workbook), with the monthly
    report's total for months after the workbook ends; domestic production (ENAP + CEOP)."""
    sheets = pd.ExcelFile(p).sheet_names
    d = by_date(read(p, "Gas imports"), "Month")
    total = d["Imports_mcm_per_day_approx"]
    if "Imports by use" in sheets:
        u = by_date(read(p, "Imports by use"), "Month")
        names = {"LNG_V_region_Quintero_mcm_per_day": "LNG - Quintero (central)",
                 "LNG_II_region_Mejillones_mcm_per_day": "LNG - Mejillones (north)",
                 "Pipeline_energy_RM_V_region_mcm_per_day": "Argentina pipeline - central (energy use)",
                 "Pipeline_energy_II_region_mcm_per_day": "Argentina pipeline - north (energy use)",
                 "Pipeline_energy_VIII_region_mcm_per_day": "Argentina pipeline - Biobio (energy use)",
                 "Pipeline_petrochemical_Magallanes_mcm_per_day": "Argentina pipeline - Magallanes (methanol)"}
        imp = u[cols(u, *names)].rename(columns=names).reindex(d.index)
        no_split = imp.isna().all(axis=1)
        imp["Total, monthly report (no split)"] = total.where(no_split)
        title = "Chile natural gas imports by use and entry region"
    else:
        imp = total.to_frame("Gas imports")
        title = "Chile natural gas imports (LNG + Argentina)"
    specs = [spec("Imports", imp, title, "million m3/day (monthly average)", "stacked_bar")]
    if "Domestic production" in sheets:
        q = by_date(read(p, "Domestic production"), "Month")
        q = q[["ENAP_mcm_per_day", "CEOP_mcm_per_day"]].rename(
            columns={"ENAP_mcm_per_day": "ENAP", "CEOP_mcm_per_day": "CEOP (private operators)"})
        specs.append(spec("Production", q, "Chile domestic gas production (Magallanes)", "million m3/day",
                          "stacked_bar"))
    return specs


def power_mix(d):
    """Ember generation-by-type sheet (GWh columns) -> chart frame with a fixed fuel order."""
    g = pd.DataFrame({"Hydro": d.get("Hydro_GWh"), "Gas": d.get("Gas_GWh"), "Wind": d.get("Wind_GWh"),
                      "Solar": d.get("Solar_GWh"), "Coal": d.get("Coal_GWh"), "Nuclear": d.get("Nuclear_GWh"),
                      "Other": d[cols(d, "Bioenergy_GWh", "Other Fossil_GWh", "Other Renewables_GWh")].sum(axis=1)})
    # every chart carries the full fuel list (zeros where a country has none) so each fuel keeps its colour
    return g.fillna(0)


# Standard layout for raw grid-operator generation workbooks (one per country):
#   sheet "Daily": date, Hydro_MWh, Gas_MWh, Wind_MWh, Solar_MWh, Coal_MWh, Nuclear_MWh, Oil_MWh,
#                  Bioenergy_MWh, Other_MWh, Total_MWh   (MWh per day; absent fuels may be omitted)
# Charted as monthly GWh with the same fuel order/colours as power_mix.
def power_daily(title):
    def f(p):
        d = by_date(read(p, "Daily"), "date")
        m = d[[c for c in d.columns if str(c).endswith("_MWh") and c != "Total_MWh"]].resample("MS").sum(min_count=1) / 1000
        m = m[m.index >= "2021-01-01"].rename(columns=lambda c: c.replace("_MWh", "_GWh"))
        m = m.rename(columns={"Oil_GWh": "Other Fossil_GWh", "Other_GWh": "Other Renewables_GWh"})
        return [spec("Generation", power_mix(m), title, "GWh per month", "stacked_bar")]
    return f


CAPACITY_FUELS = ["Hydro", "Gas", "Wind", "Solar", "Coal", "Nuclear", "Oil", "Bioenergy", "Other"]


CAPACITY_GROUPS = ["Hydro", "Gas", "Wind", "Solar", "Oil", "Bioenergy", "Coal", "Nuclear & other"]


def capacity_groups(mw):
    """CAPACITY_FUELS columns -> the 8 fixed chart groups (same order and colours on every capacity chart, zero
    groups kept so colours never shift)."""
    mw = mw.apply(pd.to_numeric, errors="coerce").fillna(0)
    g = pd.DataFrame({k: mw[k] if k in mw else 0.0 for k in CAPACITY_GROUPS[:-1]}, index=mw.index)
    g["Nuclear & other"] = (mw["Nuclear"] if "Nuclear" in mw else 0.0) + (mw["Other"] if "Other" in mw else 0.0)
    return g


def power_capacity(title):
    """Standard capacity workbook: sheet "Monthly" with date (1st of month) and <Fuel>_MW columns for
    CAPACITY_FUELS plus Total_MW. Annual-only sources use one row per year dated 1 January."""
    def f(p):
        d = by_date(read(p, "Monthly"), "date")
        d = d[d.index >= "2021-01-01"]
        g = capacity_groups(pd.DataFrame({fuel: d.get(f"{fuel}_MW") for fuel in CAPACITY_FUELS}, index=d.index))
        g = g / 1000.0   # MW -> GW
        annual = len(d) > 1 and d.index.to_series().diff().median().days > 300
        return [spec("Capacity", g, title, "GW installed", "stacked_bar", "%Y" if annual else "%Y-%m")]
    return f


def honduras_power(p):
    """ODS daily mix (from 2026-06) plus ODS's monthly history by technology (thermal not split by fuel)."""
    out = power_daily("Honduras power generation by type (ODS)")(p)
    try:
        raw = pd.read_excel(p, sheet_name="Monthly_GWh")
        m = by_date(raw, raw.columns[0])   # month column (header may be blank)
    except Exception:  # noqa: BLE001
        return out
    g = pd.DataFrame({"Hydro": m.get("Hydro_GWh"), "Thermal (oil + coal)": m.get("Thermal_GWh"),
                      "Wind": m.get("Wind_GWh"), "Solar": m.get("Solar_GWh"), "Bioenergy": m.get("Bioenergy_GWh"),
                      "Geothermal": m.get("Other_GWh")}).fillna(0)
    out.append(spec("History", g[g.index >= "2021-01-01"], "Honduras monthly production by technology (ODS reports)",
                    "GWh per month", "stacked_bar"))
    return out


def chile_power(p):
    d = by_date(read(p, "Generation by type"), "Month")
    return [spec("Generation", power_mix(d), "Chile power generation by type", "GWh per month", "stacked_bar")]


def sa_power(p):
    """One chart per country sheet of the Ember South America workbook."""
    out = []
    for sheet in pd.ExcelFile(p).sheet_names:
        if sheet.lower() in ("units", "notes") or sheet.startswith("Chart"):
            continue
        d = by_date(read(p, sheet), "Month")
        out.append(spec(sheet, power_mix(d), f"{sheet} power generation by type", "GWh per month", "stacked_bar"))
    return out


def colombia(p):
    d = by_date(read(p, "Demand by sector"), "Month")
    z = lambda *c: d[cols(d, *c)].sum(axis=1, min_count=1)  # noqa: E731
    g = pd.DataFrame({"Power": z("Power"), "Industrial (incl. oil sector)": z("Industrial", "Oil_sector"),
                      "Residential & commercial": z("Residential", "Commercial"),
                      "Refinery & petrochemical": z("Refinery", "Petrochemical"),
                      "Vehicle CNG & compressors": z("Vehicle_CNG", "Compressors")})
    mcm = 0.0283   # GBTU -> million m3 at 1,000 Btu/cf (same factor as the Supply sheet's mcm columns)
    out = [spec("Demand", g * mcm, "Colombia gas demand by sector", "mcm/d (1,000 Btu/cf)", "stacked_bar")]
    if "Supply by source" in pd.ExcelFile(p).sheet_names:
        s = by_date(read(p, "Supply by source"), "Month")
        sup = pd.DataFrame({"Cusiana/Cupiagua (Piedemonte)": s["Piedemonte_Cusiana_Cupiagua"],
                            "Guajira (Chuchupa/Ballena)": s["Guajira_Chuchupa_Ballena"],
                            "Canacol (VIM-5, VIM-21, Esperanza)": s["Canacol_VIM5_VIM21_Esperanza"],
                            "Other domestic fields": s["Other_fields"],
                            "LNG imports (SPEC Cartagena)": s["LNG_imports_SPEC"]})
        if s["Venezuela_imports"].notna().any():
            sup["Imports from Venezuela"] = s["Venezuela_imports"]
        sup["Total demand"] = d["Reported_total"].reindex(sup.index)
        out.append(spec("Supply", sup * mcm, "Colombia gas supply by source vs demand", "mcm/d (1,000 Btu/cf)",
                        "stacked_bar", line_cols=("Total demand",)))
    return out


MMBTU_PER_MCM = 36374.0   # MMBtu per million m3 at 1,030 Btu/cf (gross); used to show gas in mcm/d


def panama_gas(p):
    d = by_date(read(p, "Gas use"), "Month")
    g = pd.DataFrame({"Gas use, estimate (CND gas-fired MWh x 7.0 MMBtu/MWh)": d.get("Gas_use_MMBtu_per_day_est")})
    return [spec("Use", g / MMBTU_PER_MCM, "Panama LNG use for power (estimate from CND gas generation)",
                 "mcm/d (1,030 Btu/cf)", "stacked_bar")]


def ecuador(p):
    d = by_date(read(p, "Gas by use"), "Month")
    g = pd.DataFrame({"Power (Machala)": d.get("Power_MMBtu_per_day"),
                      "Industry (Bajo Alto LNG)": d.get("Industrial_LNG_MMBtu_per_day")})
    out = [spec("Use", g / MMBTU_PER_MCM, "Ecuador domestic gas (Amistad) by use", "mcm/d (1,030 Btu/cf)",
                "stacked_bar")]
    try:
        t = by_date(read(p, "Total demand"), "Month")
    except ValueError:  # workbook from before the total-demand sheet
        return out
    td = pd.DataFrame({"Domestic - power (Machala)": t.get("Domestic_power_MMBtu_per_day"),
                       "Domestic - industry (Bajo Alto LNG)": t.get("Domestic_industry_MMBtu_per_day"),
                       "Imports - power": t.get("Imports_power_MMBtu_per_day"),
                       "Imports - industry (LNG, customs)": t.get("Imports_industry_MMBtu_per_day")})
    td = td.loc[:, td.fillna(0).ne(0).any()]  # no imports for power so far: leave out an all-zero series
    out.append(spec("Total demand", td / MMBTU_PER_MCM, "Ecuador gas demand: domestic + imports",
                    "mcm/d (1,030 Btu/cf)", "stacked_bar"))
    return out


def el_salvador_gas(p):
    d = by_date(read(p, "Gas use"), "Month")
    g = pd.DataFrame({"Power (Energia del Pacifico, estimate)": d.get("Gas_use_MMBtu_per_day")})
    return [spec("Use", g / MMBTU_PER_MCM, "El Salvador gas use for power (estimate)", "mcm/d (1,030 Btu/cf)",
                 "stacked_bar")]


def trinidad(p):
    out = []
    u = read(p, "Utilization by sector")
    u = u[u["sector"].astype(str).str.upper() != "TOTAL"]
    w = u.pivot_table(index=pd.to_datetime(u["date"]), columns="sector", values="mmscfd", aggfunc="sum")
    pick = lambda pat: w[[c for c in w.columns if re.search(pat, c, re.I)]].sum(axis=1, min_count=1)  # noqa: E731
    g = pd.DataFrame({"LNG": pick(r"^LNG"), "Methanol": pick("methanol"), "Ammonia & derivatives": pick("ammonia|urea"),
                      "Power": pick("power")})
    g["Other"] = w.sum(axis=1) - g.sum(axis=1)
    out.append(spec("Use", g, "Trinidad & Tobago gas use by sector", "MMscf/d", "stacked_bar"))
    pr = read(p, "Production by company")
    pr = pr[pr["company"].astype(str).str.upper() != "TOTAL"].copy()
    pr["company"] = pr["company"].astype(str).str.strip().str.upper().replace(
        {"BHP": "PERENCO (ex-BHP/Woodside)", "WOODSIDE": "PERENCO (ex-BHP/Woodside)", "PEREN": "PERENCO (ex-BHP/Woodside)",
         "PERENCO": "PERENCO (ex-BHP/Woodside)"})
    pw = pr.pivot_table(index=pd.to_datetime(pr["date"]), columns="company", values="mmscfd", aggfunc="sum")
    out.append(spec("Production", pw, "Trinidad & Tobago gas production by company", "MMscf/d", "stacked_bar"))
    return out


def guyana_gas(p):
    """GUYANA_GAS.py (Ministry data centre): monthly gross production, reinjection as its own line, and production net
    of reinjection; then the disposition of the gas (stacked) with gross production as a line."""
    d = by_date(read(p, "Monthly"), "date")
    g = lambda c: d.get(f"{c}_mcm_per_day")  # noqa: E731
    prod = pd.DataFrame({"Gross production": g("Produced"), "Reinjected": g("Reinjected"),
                         "Net of reinjection": g("Produced") - g("Reinjected")})
    disp = pd.DataFrame({"Reinjected": g("Reinjected"), "Used as fuel": g("Used_as_fuel"), "Flared": g("Flared"),
                         "Other / onshore": g("Other_or_unreported"), "Gross production": g("Produced")})
    if disp["Other / onshore"].abs().max() < 0.01:   # nothing sent onshore yet: keep the legend clean
        disp = disp.drop(columns="Other / onshore")
    return [spec("Production", prod, "Guyana gas production and reinjection (Ministry of Natural Resources)",
                 "million m3/day"),
            spec("Disposition", disp, "Guyana associated gas: where it goes (Ministry of Natural Resources)",
                 "million m3/day", "stacked_bar", line_cols=("Gross production",))]


def power_annual(title):
    """Standard 'Daily' layout holding one row per YEAR (Jamaica): annual GWh bars labelled by year."""
    def f(p):
        d = by_date(read(p, "Daily"), "date")
        y = d[[c for c in d.columns if str(c).endswith("_MWh") and c != "Total_MWh"]] / 1000
        y = y[y.index >= "2021-01-01"].rename(columns=lambda c: c.replace("_MWh", "_GWh"))
        y = y.rename(columns={"Oil_GWh": "Other Fossil_GWh", "Other_GWh": "Other Renewables_GWh"})
        return [spec("Generation", power_mix(y), title, "GWh per year", "stacked_bar", "%Y")]
    return f


def paraguay_power(p):
    """PARAGUAY_POWER.py: Paraguay's 50% share of Itaipu + Yacyreta by month (standard 'Daily' layout, monthly
    rows), then where that energy went: Paraguay's own use vs the shares ceded to Brazil and Argentina."""
    out = power_daily("Paraguay power generation by type (ANDE / Itaipu / Yacyreta)")(p)
    try:
        e = by_date(read(p, "Exports"), "date")
    except Exception:  # noqa: BLE001
        return out
    e = e[e.index >= "2021-01-01"]
    g = pd.DataFrame({"Paraguay's own use (ANDE)": e.get("Paraguay_consumption_GWh"),
                      "Itaipu ceded to Brazil": e.get("Itaipu_to_Brazil_GWh"),
                      "Yacyreta ceded to Argentina": e.get("Yacyreta_to_Argentina_GWh"),
                      "Paraguay's 50% share of generation": e.get("Total_generation_share_GWh")})
    out.append(spec("Domestic vs exports", g, "Paraguay hydro: domestic use vs exports", "GWh per month",
                    "stacked_bar", line_cols=("Paraguay's 50% share of generation",)))
    return out


def jamaica_gas(p):
    d = by_date(read(p, "Gas use"), "Month")
    return [spec("Use", d[["Total_mcm_per_day"]].rename(columns={"Total_mcm_per_day": "Natural gas, all uses"}),
                 "Jamaica gas use (annual)", "million m3/day, annual average", "stacked_bar", "%Y")]


def gas_use(title, skip=r"^(Total|.*_terminal)_"):
    """Caribbean 'Gas use' sheets (Month + *_mcm_per_day columns): stack the parts, not the totals/groupings."""
    def f(p):
        d = by_date(read(p, "Gas use"), "Month")
        c = [x for x in d.columns if str(x).endswith("_mcm_per_day") and not re.match(skip, str(x))]
        return [spec("Use", d[c].rename(columns=lambda x: x.replace("_mcm_per_day", "").replace("_", " ")),
                     title, "million m3/day", "stacked_bar")]
    return f


def ireland_demand(p):
    d = by_date(read(p, "Data"), "date")
    s = d[cols(d, "PowerGen_GWh", "LDM_ex_PowerGen_GWh", "DM_GWh", "NDM_GWh")]
    names = {"PowerGen": "Power", "LDM_ex_PowerGen": "Large industry (ex power)", "DM": "Daily metered",
             "NDM": "Non-daily metered (homes, small business)"}
    return [spec("Demand", monthly_mean(s, "2021-01-01").rename(columns=lambda c: names.get(c.replace("_GWh", ""), c)),
                 "Ireland gas demand by sector (monthly average)", "GWh/day", "stacked_bar"),
            {"name": "Seasonal", "water_year": d["Total_ROI_GWh"], "title": "Ireland total gas demand", "units": "GWh/day"}]


def ireland_combined(p):
    names = {"PowerGen": "Power", "LDM_ex_PowerGen": "Large industry (ex power)", "DM": "Daily metered",
             "NDM": "Non-daily metered (homes, small business)"}
    d = by_date(read(p, "Demand"), "date")
    dem = d[cols(d, "PowerGen_GWh", "LDM_ex_PowerGen_GWh", "DM_GWh", "NDM_GWh")]
    s = by_date(read(p, "Supply"), "date")
    sup = s[cols(s, "Corrib_Production_GWh", "Inch_Production_GWh", "Moffat_Imports_GWh")]
    return [spec("Demand", monthly_mean(dem, "2021-01-01").rename(columns=lambda c: names.get(c.replace("_GWh", ""), c)),
                 "Ireland gas demand by sector (monthly average)", "GWh/day", "stacked_bar"),
            spec("Supply", monthly_mean(sup, "2021-01-01").rename(
                     columns={"Corrib_Production_GWh": "Corrib", "Inch_Production_GWh": "Inch",
                              "Moffat_Imports_GWh": "Moffat imports (UK)"}),
                 "Ireland gas supply by source (monthly average)", "GWh/day", "stacked_bar"),
            {"name": "Seasonal", "water_year": d["Total_ROI_GWh"], "title": "Ireland total gas demand", "units": "GWh/day"}]


def ireland_supply(p):
    d = by_date(read(p, "Data"), "date")
    s = d[cols(d, "Corrib_Production_GWh", "Inch_Production_GWh", "Moffat_Imports_GWh")]
    return [spec("Supply", monthly_mean(s, "2021-01-01").rename(columns=lambda c: c.replace("_GWh", "").replace("_", " ")),
                 "Ireland gas supply by source (monthly average)", "GWh/day", "stacked_bar")]


def ireland_gni(p):
    e = by_date(read(p, "Entry flows"), "date")
    c = by_date(read(p, "Consumption by sector"), "date")
    tidy = lambda c: c.replace("_GWh", "").replace("ROI_Power_Gen", "Power").replace("_", " ")  # noqa: E731
    return [spec("Entry flows", e[cols(e, "Bellanaboy_GWh", "Gormanston_GWh", "Moffat_GWh")].rename(columns=tidy),
                 "Ireland gas entry flows (GNI)", "GWh/day", "stacked_area", "%Y-%m-%d"),
            spec("Consumption", c[cols(c, "ROI_Power_Gen_GWh", "DM_GWh", "NDM_GWh")].rename(columns=tidy),
                 "Ireland gas consumption by market sector (GNI)", "GWh/day", "stacked_area", "%Y-%m-%d")]


def ireland_smartgrid(p):
    d = by_date(read(p, "Data"), "datetime")
    dd = weekly_mean(d, "2021-01-01")
    return [spec("Demand and wind", dd[cols(dd, "Demand_Actual_MW", "Wind_Actual_MW")].rename(
                     columns={"Demand_Actual_MW": "Demand", "Wind_Actual_MW": "Wind generation"}),
                 "Ireland electricity demand and wind (weekly average)", "MW", "line", "%Y-%m-%d"),
            spec("CO2 intensity", dd[cols(dd, "CO2_Intensity_gCO2_per_kWh")].rename(
                     columns={"CO2_Intensity_gCO2_per_kWh": "CO2 intensity"}),
                 "Ireland grid CO2 intensity (weekly average)", "gCO2/kWh", "line", "%Y-%m-%d")]


def mexico(p):
    d = by_date(read(p, "Data"), "date")
    w = weekly_mean(d[["Total_MWh"]]).rename(columns={"Total_MWh": "National demand"})
    return [spec("Demand", w, "Mexico national electricity demand (weekly average)", "MWh/day", "line", "%Y-%m-%d")]


def fuel_mix(title, units, drop=("Total", "Imports", "Renewables_Share", "share", "unknown")):
    def f(p):
        d = by_date(read(p, "Data"), "date")
        c = [x for x in d.columns if not any(k.lower() in str(x).lower() for k in drop) and pd.api.types.is_numeric_dtype(d[x])]
        m = d[c].rename(columns=lambda x: str(x).replace("_MW", "").replace("_", " "))
        if (m.index.max() - m.index.min()).days > 400:
            return [spec("Mix", monthly_mean(m), f"{title} (monthly average)", units, "stacked_bar")]
        return [spec("Mix", m, title, units, "stacked_area", "%Y-%m-%d")]
    return f


def puertorico(p):
    d = by_date(read(p, "By fuel type"), "period")
    names = {"NG": "Natural gas", "RFO": "Residual fuel oil", "DFO": "Distillate oil", "COL": "Coal", "SUN": "Solar",
             "HYC": "Hydro", "WND": "Wind", "LFG": "Landfill gas", "BIO": "Biomass"}
    c = [x for x in d.columns if pd.api.types.is_numeric_dtype(d[x]) and not re.search(r"total|all", str(x), re.I)]
    return [spec("By fuel", d[c].rename(columns=lambda x: names.get(x, x)), "Puerto Rico generation by fuel",
                 "GWh per month", "stacked_bar")]


def china_nbs(p):
    d = read(p, "Data")
    dc = next(c for c in d.columns if "month" in str(c).lower() or "date" in str(c).lower())
    d = by_date(d, dc)
    gen = [c for c in d.columns if str(c).endswith("_Generation_GWh") and not str(c).startswith("Total")]
    out = []
    if gen:
        out.append(spec("Generation", d[gen].rename(columns=lambda c: c.replace("_Generation_GWh", "")),
                        "China power generation by source (NBS)", "GWh per month", "stacked_bar"))
    for c in [c for c in d.columns if c not in gen and pd.api.types.is_numeric_dtype(d[c])]:
        label, unit = label_and_unit(c)
        out.append(spec(label, d[[c]].rename(columns={c: label}), f"China {label.lower()} (NBS)",
                        f"{unit} per month" if unit else "per month", "line"))
    return out


def china_nbs_series(p):
    """China NBS workbooks: the pull writes a 'Series' sheet (column, label, unit, chart, kind); one chart per group."""
    d = read(p, "Data")
    d = by_date(d, d.columns[0])
    s = pd.read_excel(p, sheet_name="Series", index_col=0)
    out, used = [], set()
    for group in dict.fromkeys(s["chart"].dropna()):
        rows = s[s["chart"] == group]
        c = [x for x in rows.index if x in d.columns and d[x].notna().any()]
        if str(group).strip() and c:
            name = re.sub(r"[\\/*?:\[\]]", "-", str(group))[:23]   # sheet = "Chart - " + name, max 31 chars
            name = name if name not in used else f"{name[:20]} {len(used)}"
            used.add(name)
            out.append(spec(name, d[c].rename(columns=rows["label"].to_dict()),
                            f"China {group} (NBS)", rows["unit"].iloc[0], rows["kind"].iloc[0]))
    return out


def giignl(p):
    d = read(p, "Data")
    d = d.set_index(pd.to_datetime(d["report_year"].astype(int).astype(str) + "-01-01"))
    exporters = [c for c in d.columns if c.endswith("_Total_MT") and "Basin" not in c and "Middle East" not in c]
    top = d[exporters].mean().sort_values(ascending=False).index[:7]
    t = d[top].rename(columns=lambda c: c.replace("_Total_MT", ""))
    return [spec("Exports", t, "LNG exports by country (GIIGNL)", "million tonnes per year", "stacked_bar", "%Y")]


def south_africa(p):
    return fuel_mix("South Africa generation mix (Eskom)", "MW (daily mean)")(p)


def henry_hub(p):
    d = by_date(read(p, "Data"), "date")[["Henry_Hub_USD_per_MMBtu"]].rename(
        columns={"Henry_Hub_USD_per_MMBtu": "Henry Hub spot"})
    return [spec("Daily", daily(d, "2021-01-01"), "Henry Hub natural gas spot price (daily)", "USD/MMBtu"),
            spec("Monthly", monthly_mean(d), "Henry Hub natural gas spot price (monthly average)", "USD/MMBtu")]


def _sheet(p, name, date_col):
    try:
        return by_date(read(p, name), date_col)
    except ValueError:   # sheet not in this workbook
        return pd.DataFrame()


def _full_years(d):
    return d[d["Coverage"].astype(str).eq("Full year")] if "Coverage" in d else d


def singapore_power(p):
    out = []
    d = _sheet(p, "Daily demand", "Date")
    if not d.empty:
        dm = monthly_mean(d[cols(d, "System_Demand_Avg_MW", "System_Demand_Peak_MW")], "2021-01-01")
        out.append(spec("Demand", dm.rename(columns={"System_Demand_Avg_MW": "Average system demand",
                                                     "System_Demand_Peak_MW": "Daily peak"}),
                        "Singapore electricity system demand (monthly average of daily values)", "MW"))
    g = _sheet(p, "Daily generation by type", "Date")
    if not g.empty:
        gc = [c for c in g.columns if str(c).endswith("_GWh") and not str(c).startswith("Total")]
        out.append(spec("Generation", monthly_mean(g[gc].clip(lower=0), "2021-01-01").rename(
                            columns=lambda c: c.replace("_GWh", "").replace("_", " ")),
                        "Singapore metered generation by plant type (monthly average)", "GWh/day", "stacked_bar"))
    m = _sheet(p, "Monthly generation", "Month")
    if not m.empty:
        out.append(spec("Monthly generation", m.loc[m.index >= "2015-01-01", ["Electricity_Generation_GWh"]].rename(
                            columns={"Electricity_Generation_GWh": "Electricity generation"}),
                        "Singapore electricity generation (SingStat/EMA)", "GWh per month"))
    c = _full_years(_sheet(p, "Annual consumption", "Year"))
    if not c.empty:
        cc = [x for x in c.columns if str(x).endswith("_GWh") and not str(x).startswith("Total")]
        out.append(spec("Consumption", c[cc].rename(columns=lambda x: x.replace("_GWh", "").replace("_", " ")),
                        "Singapore electricity consumption by sector", "GWh per year", "stacked_bar", "%Y"))
    f = _full_years(_sheet(p, "Annual fuel mix", "Year"))
    if not f.empty:
        fc = [x for x in f.columns if str(x).endswith("_pct")]
        out.append(spec("Fuel mix", f[fc].rename(columns=lambda x: x.replace("_pct", "").replace("_", " ")),
                        "Singapore fuel mix for electricity generation", "% of generation", "stacked_bar", "%Y"))
    return out


def singapore_gas(p):
    out = []
    d = _full_years(_sheet(p, "Annual demand by sector", "Year"))
    if not d.empty:
        dc = cols(d, "Power_generation_TJ", "Industrial_TJ", "Commerce_Services_TJ", "Households_TJ", "Transport_TJ",
                  "Others_TJ")
        out.append(spec("Demand", (d[dc] / 1000).rename(columns=lambda x: x.replace("_TJ", "").replace("_", " ")),
                        "Singapore natural gas demand by sector", "PJ per year", "stacked_bar", "%Y"))
    i = _sheet(p, "Annual imports", "Year")
    if not i.empty:
        out.append(spec("Imports", (i[cols(i, "Pipeline_TJ", "LNG_TJ")] / 1000).rename(
                            columns={"Pipeline_TJ": "Pipeline gas", "LNG_TJ": "LNG"}),
                        "Singapore natural gas imports, pipeline vs LNG", "PJ per year", "stacked_bar", "%Y"))
    e = _sheet(p, "Power burn monthly (est)", "Month")
    if not e.empty:
        out.append(spec("Power burn", e[["Gas_for_power_mcm_per_day_est"]].rename(
                            columns={"Gas_for_power_mcm_per_day_est": "Gas for power (estimate)"}),
                        "Singapore gas burn for power, estimated from metered CCGT generation",
                        "mcm/day (approx)"))
    t = _sheet(p, "Town gas quarterly", "Quarter_start")
    if not t.empty:
        out.append(spec("Town gas", t.loc[t.index >= "2015-01-01", cols(t, "Domestic_GWh", "Non_domestic_GWh")].rename(
                            columns={"Domestic_GWh": "Domestic", "Non_domestic_GWh": "Non-domestic"}),
                        "Singapore town gas sales", "GWh per quarter", "stacked_bar"))
    return out


def brazil_hydro(p):
    d = by_date(read(p, "Daily"), "date")
    areas = [("SIN", "Brazil reservoirs, SIN (national)"),
             ("SE_CO", "Brazil reservoirs, SE/CO"),
             ("S", "Brazil reservoirs, South"),
             ("NE", "Brazil reservoirs, Northeast"),
             ("N", "Brazil reservoirs, North")]
    return [{"name": a, "water_year": d[f"{a}_pct"], "title": t, "units": "% full",
             "sheet": f"Water year - {a.replace('_', '-')}"} for a, t in areas if f"{a}_pct" in d]


def colombia_hydro(p):
    d = by_date(read(p, "Daily"), "date")
    return [{"name": "Storage", "water_year": d["Storage_pct"], "title": "Colombia reservoirs (national)",
             "units": "% full"}]


def _join_short_gaps(s, max_gap=8):
    """Daily series with gaps of up to max_gap days filled by straight lines (longer gaps stay empty)."""
    full = s.dropna().resample("D").mean()
    gap = full.isna().groupby(full.notna().cumsum()).transform("sum")
    return full.interpolate(limit_area="inside").where(full.notna() | (gap <= max_gap)).dropna()


def argentina_hydro(p):
    """Lake levels (before 2023 one CAMMESA weekly-programme value per week, joined up across gaps of up to 15 days) and the
    Parana / Uruguay flows at the run-of-river binational plants. El Chocon first: no volume-based Comahue %."""
    d = by_date(read(p, "Daily"), "date")
    charts = [("ChoconLevel_m", "Chocon", "Argentina, El Chocón reservoir level (Comahue)", "m above sea level"),
              ("PiedraAguilaLevel_m", "PiedraAguila", "Argentina, Piedra del Águila reservoir level (Comahue)", "m above sea level"),
              ("AlicuraLevel_m", "Alicura", "Argentina, Alicurá reservoir level (Comahue)", "m above sea level"),
              ("CerrosColoradosLevel_m", "CerrosColorados", "Argentina, Cerros Colorados (Los Barreales) level (Comahue)",
               "m above sea level"),
              ("PichiPicunLevel_m", "PichiPicun", "Argentina, Pichi Picún Leufú reservoir level (Comahue)", "m above sea level"),
              ("FutaleufuLevel_m", "Futaleufu", "Argentina, Futaleufú reservoir level", "m above sea level"),
              ("Parana_Yacyreta_m3s", "Yacyreta", "Paraná river flow into Yacyretá", "m3/s (daily mean)"),
              ("Uruguay_SaltoGrande_m3s", "SaltoGrande", "Uruguay river flow at Salto Grande", "m3/s (daily mean)")]
    return [{"name": n, "water_year": _join_short_gaps(d[c], 15), "title": t, "units": u, "sheet": f"Water year - {n}"}
            for c, n, t, u in charts if c in d and d[c].notna().any()]


def chile_hydro(p):
    """DGA reservoir volumes: daily from mid-2024, month-end points before that (joined up for the chart).
    The six-reservoir total as % of capacity first, then the big generation reservoirs in Mm3."""
    d = by_date(read(p, "Daily"), "date")
    charts = [("Total_pct", "Total", "Chile hydro reservoirs, total of 6 generation reservoirs (DGA)", "% full"),
              ("Colbun_Mm3", "Colbun", "Chile, Colbún reservoir volume (DGA)", "million m3"),
              ("LagoLaja_Mm3", "LagoLaja", "Chile, Lago Laja volume (DGA)", "million m3"),
              ("Ralco_Mm3", "Ralco", "Chile, Ralco reservoir volume (DGA)", "million m3"),
              ("LagunaMaule_Mm3", "LagunaMaule", "Chile, Laguna del Maule volume (DGA)", "million m3"),
              ("Rapel_Mm3", "Rapel", "Chile, Rapel reservoir volume (DGA)", "million m3")]
    return [{"name": n, "water_year": _join_short_gaps(d[c], max_gap=35), "title": t, "units": u,
             "sheet": f"Water year - {n}"} for c, n, t, u in charts if c in d and d[c].notna().any()]


IGU_SECTORS = ["LNG liquefaction", "Ammonia/urea", "Methanol", "Steel DRI", "Alumina", "Cement", "Glass", "Ceramics",
               "Other"]


def _category_stacked_bar(path, sheet, t, title, y_title):
    """Native stacked column chart over a category (non-date) axis: rows of t are the categories, columns the series."""
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    wb = load_workbook(path)
    if sheet in wb.sheetnames:
        del wb[sheet]
    ws = wb.create_sheet(sheet, 1)
    ws.cell(row=1, column=1, value="Country")
    for j, c in enumerate(t.columns, start=2):
        ws.cell(row=1, column=j, value=str(c))
    for i, (k, r) in enumerate(t.iterrows(), start=2):
        ws.cell(row=i, column=1, value=str(k))
        for j, v in enumerate(r.values, start=2):
            ws.cell(row=i, column=j, value=None if pd.isna(v) else round(float(v), 3))
    ws.column_dimensions["A"].width = 20
    n = len(t) + 1
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "col", "stacked", 100, 50
    ch.add_data(Reference(ws, min_col=2, max_col=1 + t.shape[1], min_row=1, max_row=n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    for i, s in enumerate(ch.series):
        grey = str(t.columns[i]).startswith("Other")
        s.graphicalProperties = GraphicalProperties(
            solidFill=xlsx_charts.OTHER_GREY if grey else xlsx_charts.PALETTE[i % len(xlsx_charts.PALETTE)],
            ln=LineProperties(noFill=True))
    ch.title = title
    ch.y_axis.title = y_title
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))     # no borders
    ch.plot_area.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.width, ch.height = 28, 13
    xlsx_charts.tidy_layout(ch)
    ws.add_chart(ch, f"{chr(ord('A') + min(t.shape[1] + 2, 20))}2")
    xlsx_charts.save_atomic(wb, path)


def industrial_gas_users(p):
    """Plant register (south_america/INDUSTRIAL_GAS_USERS.py): static data, so a stacked bar over countries of the
    number of OPERATING plants, by sector (Petrochemical folds into 'Other': 8 palette slots)."""
    d = read(p, "Plants")
    op = d[d["counted_as"] == "Operating"].copy()
    op["sector"] = op["sector"].where(op["sector"] != "Petrochemical", "Other")
    t = op.pivot_table(index="country", columns="sector", values="plant", aggfunc="count").fillna(0)
    t = t[[s for s in IGU_SECTORS if s in t.columns]]
    t = t.loc[t.sum(axis=1).sort_values(ascending=False).index]
    t.index = [str(c).replace("Trinidad and Tobago", "Trinidad & Tobago") for c in t.index]
    t = t.rename(columns={"Other": "Other (petrochemical)"})
    return [{"name": "Plants by country", "custom": lambda path, sheet: _category_stacked_bar(
        path, sheet, t, "Large industrial gas users: operating plants by country and sector (source: plant register - "
        "USGS, Trinidad MEEI, GEM, company reports)", "number of operating plants")}]


CO_COAL_DEPTS = ["La_Guajira", "Cesar", "Boyaca", "Cundinamarca", "Norte_de_Santander", "Cordoba"]


def sa_coal(p):
    """South America coal (SOUTH_AMERICA_COAL.py): Colombia quarterly production by department (ANM, stacked bars)
    with DANE coal exports as a line; annual production by country (stacked bars, Mt/yr)."""
    out = []
    q = _sheet(p, "Colombia", "Quarter")
    if not q.empty:
        q = q.apply(pd.to_numeric, errors="coerce")
        dep = [f"{d}_Mt" for d in CO_COAL_DEPTS if f"{d}_Mt" in q]
        rest = [c for c in q.columns if c.endswith("_Mt") and c not in dep + ["Total_production_Mt", "Exports_DANE_Mt"]]
        df = q[dep].rename(columns=lambda c: c[:-3].replace("_", " ").replace("Boyaca", "Boyacá")
                           .replace("Cordoba", "Córdoba"))
        df["Other departments"] = q[rest].sum(axis=1, min_count=1)
        lines = ()
        if "Exports_DANE_Mt" in q and q["Exports_DANE_Mt"].notna().any():
            df["Exports (DANE)"] = q["Exports_DANE_Mt"]
            lines = ("Exports (DANE)",)
        out.append(spec("Colombia", df, "Colombia coal production by department (ANM) and exports (DANE)",
                        "Mt per quarter", "stacked_bar", line_cols=lines))
    a = _sheet(p, "Annual by country", "Year")
    if not a.empty:
        mt = [c for c in a.columns if str(c).endswith("_Mt") and c != "Total_Mt"]
        out.append(spec("Annual", a[mt].apply(pd.to_numeric, errors="coerce").rename(columns=lambda c: c[:-3]),
                        "South America coal production by country (Venezuela: EI estimate)", "Mt per year",
                        "stacked_bar", "%Y"))
    return out


# South America daily wholesale power prices (south_america/SA_POWER_PRICES_DAILY.py) - a separate workbook,
# not part of the South & Central America master.
PRICE_CHARTS = [  # sheet, title (with source)
    ("Brazil", "Brazil CMO by subsystem (ONS, daily average)"),
    ("Colombia", "Colombia Precio de Bolsa and Precio de Escasez (XM, daily average)"),
    ("Peru", "Peru marginal cost, Santa Rosa 220 kV (COES, daily average)"),
    ("Argentina", "Argentina CMO and sanctioned spot price (CAMMESA, daily average)"),
    ("Uruguay", "Uruguay spot sancionado (ADME, daily average)"),
    ("Bolivia", "Bolivia marginal cost (CNDC, daily)"),
]


def sa_power_prices(p):
    """Everything in US$/MWh: all markets as monthly averages of 'USD daily' (Chile's monthly PMM is on every day
    of its month there; CNDC's monthly Bolivian energy price is added since the daily series is short), then one
    chart per country (daily; Chile monthly). Local-currency prices stay in the country sheets only."""
    sheets = pd.ExcelFile(p).sheet_names
    u = by_date(read(p, "USD daily"), "date")
    m = monthly_mean(u, "2021-01-01")
    if "Bolivia monthly" in sheets:
        b = read(p, "Bolivia monthly")
        b = by_date(b, b.columns[0])
        if "Precio de energia (USD/MWh)" in b:
            m = m.join(b[["Precio de energia (USD/MWh)"]].rename(
                columns={"Precio de energia (USD/MWh)": "Bolivia (CNDC energy price, monthly)"}), how="outer")
    m = m.drop(columns=[c for c in m if c.startswith("Bolivia (CNDC marginal")], errors="ignore")
    # one price per market: averaging is fine, but a fold into "Other" would ADD prices, so the extra Brazil
    # subsystems and Argentina's capped spot price stay in their country charts
    m = m.drop(columns=[c for c in m if c.startswith(("Brazil S ", "Brazil NE", "Brazil N ", "Argentina (spot"))],
               errors="ignore")
    out = [spec("USD monthly", m[m.index >= "2021-01-01"],
                "South America wholesale power prices (monthly average)", "US$/MWh")]
    for sheet, title in PRICE_CHARTS:
        if sheet not in sheets:
            continue
        d = by_date(read(p, sheet), "date")
        usd = [c for c in d.columns if "(USD/MWh)" in str(c)]
        if usd:
            g = d[usd].rename(columns=lambda c: re.sub(r"\s*\([^)]*\)$", "", str(c)))
            g = daily(g, "2021-01-01").dropna(how="all")
            if len(g) and (g.index.max() - g.index.min()).days > 400:   # years of days: weekly averages read better
                g, title = weekly_mean(g), title.replace("daily average", "weekly average of daily prices")
            out.append(spec(sheet, g, title, "US$/MWh"))
    if "Chile" in sheets:
        c = read(p, "Chile")
        c = by_date(c, c.columns[0])
        if "PMM SEN (USD/MWh)" in c:
            out.append(spec("Chile", c[["PMM SEN (USD/MWh)"]].rename(columns={"PMM SEN (USD/MWh)": "PMM SEN"}),
                            "Chile Precio Medio de Mercado, SEN (CNE, monthly; FX dolar observado)", "US$/MWh"))
    return out


def peru_hydro(p):
    """COES weekly table: two readings a week (start and end of each COES week), joined up for the chart.
    The national useful-storage % first, then the main seasonal systems in hm3."""
    d = by_date(read(p, "Daily"), "date")
    charts = [("Total_pct", "Total", "Peru reservoirs and lagoons, total useful volume (COES)", "% of useful capacity"),
              ("Junin_hm3", "Junin", "Peru, Lake Junín useful volume (COES)", "million m3 (useful)"),
              ("Mantaro_lagoons_hm3", "Mantaro", "Peru, Mantaro-basin lagoons useful volume (COES)", "million m3 (useful)"),
              ("Rimac_hm3", "Rimac", "Peru, Rímac system useful volume (COES)", "million m3 (useful)"),
              ("Chili_hm3", "Chili", "Peru, Chili system (Arequipa) useful volume (COES)", "million m3 (useful)"),
              ("Aricota_hm3", "Aricota", "Peru, Aricota useful volume (COES)", "million m3 (useful)"),
              ("Sibinacocha_hm3", "Sibinacocha", "Peru, Sibinacocha useful volume (COES)", "million m3 (useful)")]
    return [{"name": n, "water_year": _join_short_gaps(d[c], max_gap=8), "title": t, "units": u,
             "sheet": f"Water year - {n}"} for c, n, t, u in charts if c in d and d[c].notna().any()]


def ecuador_hydro(p):
    """CELEC SUR daily levels: Mazar (the seasonal store behind the 2023-24 blackouts) first, then Amaluza."""
    d = by_date(read(p, "Daily"), "date")
    charts = [("MazarLevel_m", "Mazar", "Ecuador, Mazar reservoir level (CELEC SUR)", "m above sea level"),
              ("AmaluzaLevel_m", "Amaluza", "Ecuador, Amaluza (Paute) reservoir level (CELEC SUR)", "m above sea level")]
    return [{"name": n, "water_year": _join_short_gaps(d[c], max_gap=8), "title": t, "units": u,
             "sheet": f"Water year - {n}"} for c, n, t, u in charts if c in d and d[c].notna().any()]


def uruguay_hydro(p):
    """Rincón del Bonete (Uruguay's storage lake) first, then Salto Grande and the two lower Río Negro lakes."""
    d = by_date(read(p, "Daily"), "date")
    charts = [("BoneteLevel_m", "Bonete", "Uruguay, Rincón del Bonete lake level (ADME)", "m above sea level"),
              ("SaltoGrandeLevel_m", "SaltoGrande", "Uruguay/Argentina, Salto Grande lake level (INA)", "m"),
              ("PalmarLevel_m", "Palmar", "Uruguay, Palmar lake level (ADME)", "m above sea level"),
              ("BaygorriaLevel_m", "Baygorria", "Uruguay, Baygorria lake level (ADME)", "m above sea level")]
    return [{"name": n, "water_year": _join_short_gaps(d[c], max_gap=8), "title": t, "units": u,
             "sheet": f"Water year - {n}"} for c, n, t, u in charts if c in d and d[c].notna().any()]


def complete_through(d, active_months=12):
    """Rows up to the last month every 'active' column (one with data in the final active_months) has a value, so a
    stacked total does not drop because one country publishes later than the others."""
    d = d.dropna(how="all")
    if d.empty:
        return d
    last = d.index.max()
    active = [c for c in d.columns if d[c].last_valid_index() is not None
              and d[c].last_valid_index() >= last - pd.DateOffset(months=active_months)]
    end = min(d[c].last_valid_index() for c in active) if active else last
    return d[d.index <= end]


def sa_gas_balance(p):
    """SA_GAS_BALANCE.py: production, imports and exports by country (million m3/day), stacked."""
    out = []
    for sheet, title in [("Production", "South America gas production by country"),
                         ("Imports", "South America gas imports by country (pipeline + LNG)"),
                         ("Exports", "South America gas exports by country (pipeline + LNG)")]:
        d = by_date(read(p, sheet), "date")
        out.append(spec(sheet, complete_through(d), title, "million m3/day", "stacked_bar"))
    return out


# North America (americas/US_GAS_EIA.py, americas/CANADA_STATCAN.py) - gas in Bcf/d, the North American convention
BCF_TO_MCM = 28.3168   # 1 Bcf = 28.3168 million m3


def us_gas(p):
    """EIA: demand by sector, dry production, trade (exports stacked, Canadian pipeline imports as a line), then a
    water-year chart of Lower 48 working gas and one per storage region."""
    out = []
    d = _sheet(p, "Demand by sector", "Month")
    if not d.empty:
        z = lambda *c: d[[f"{x}_Bcf_per_day" for x in c if f"{x}_Bcf_per_day" in d]].sum(axis=1, min_count=1)  # noqa: E731
        g = pd.DataFrame({"Power": z("Electric_power"), "Industrial": z("Industrial"), "Residential": z("Residential"),
                          "Commercial": z("Commercial"), "Lease & plant fuel": z("Lease_and_plant_fuel", "Lease_fuel", "Plant_fuel"),
                          "Pipeline, distribution & vehicle": z("Pipeline_and_distribution", "Vehicle_fuel")})
        out.append(spec("Demand", g[g.index >= "2021-01-01"], "US natural gas consumption by sector (EIA)",
                        "Bcf/d", "stacked_bar"))
    t = _sheet(p, "Supply and trade", "Month")
    if not t.empty:
        t = t[t.index >= "2021-01-01"]
        if "Dry_production_Bcf_per_day" in t:
            out.append(spec("Production", t[["Dry_production_Bcf_per_day"]].rename(
                columns={"Dry_production_Bcf_per_day": "Dry gas production"}), "US dry natural gas production (EIA)",
                "Bcf/d"))
        names = {"LNG_exports_Bcf_per_day": "LNG exports", "Pipeline_exports_to_Mexico_Bcf_per_day": "Pipeline exports to Mexico",
                 "Pipeline_exports_to_Canada_Bcf_per_day": "Pipeline exports to Canada",
                 "Pipeline_imports_from_Canada_Bcf_per_day": "Pipeline imports from Canada"}
        tr = t[cols(t, *names)].rename(columns=names)
        if not tr.empty:
            out.append(spec("Trade", tr, "US natural gas exports and imports from Canada (EIA)", "Bcf/d", "stacked_bar",
                            line_cols=tuple(c for c in ("Pipeline imports from Canada",) if c in tr)))
    w = _sheet(p, "Storage weekly", "date")
    regions = [("Lower_48", "US Lower 48"), ("East", "US East"), ("Midwest", "US Midwest"), ("Mountain", "US Mountain"),
               ("Pacific", "US Pacific"), ("South_Central", "US South Central")]
    for col, label in regions:
        if f"{col}_Bcf" in w and w[f"{col}_Bcf"].notna().any():
            daily_s = w[f"{col}_Bcf"].dropna().resample("D").interpolate()   # weekly -> daily so the lines join up
            out.append({"name": "Storage" if col == "Lower_48" else f"Storage {label[3:]}", "water_year": daily_s,
                        "title": f"{label} working gas in storage (EIA)", "units": "Bcf",
                        "sheet": "Water year" if col == "Lower_48" else f"Water year - {col.replace('_', ' ')}"[:31],
                        "y_decimals": 0})
    return out


def mexico_gas(p):
    """Pipeline imports only: EIA's LNG-to-Mexico series (N9133MX2) has implausible months of 2-3.5 Bcf/d among
    months of ~0.1 (as published by EIA, one row per month), so it stays in the sheet but is not charted."""
    d = _sheet(p, "Imports from US", "Month")
    names = {"Pipeline_imports_from_US_Bcf_per_day": "Pipeline from the US"}
    return [spec("Imports", d[cols(d, *names)].rename(columns=names)[lambda x: x.index >= "2021-01-01"],
                 "Mexico pipeline gas imports from the US (EIA, US export data)", "Bcf/d", "stacked_bar")]


def us_mexico_pipeline_capacity(p):
    """Export capacity by line (stacked, Bcf/d, annual) and monthly US pipeline exports to Mexico against total capacity."""
    out = []
    d = _sheet(p, "Export capacity by line", "Year")
    if d.empty:
        return out
    d = d.apply(pd.to_numeric, errors="coerce") / 1000.0
    lines = d.drop(columns=["Total"], errors="ignore")
    out.append(spec("Pipeline capacity", lines[lines.index >= "2010-01-01"], "US to Mexico gas pipeline capacity by line (EIA)",
                    "Bcf/d", "stacked_bar", "%Y"))
    mx = os.path.join(os.path.dirname(p), "mexico_gas.xlsx")
    if os.path.exists(mx) and "Total" in d:
        e = by_date(read(mx, "Imports from US"), "Month")
        e = e[["Pipeline_imports_from_US_Bcf_per_day"]].rename(
            columns={"Pipeline_imports_from_US_Bcf_per_day": "US pipeline exports to Mexico"})
        e = e[e.index >= "2015-01-01"].copy()
        cap = d["Total"].reindex(pd.date_range(d.index.min(), e.index.max(), freq="MS")).ffill()
        e["Pipeline capacity (annual, EIA)"] = cap.reindex(e.index)
        out.append(spec("Exports vs capacity", e, "US pipeline exports to Mexico vs pipeline capacity (EIA)", "Bcf/d"))
    return out


def canada_gas(p):
    """StatCan supply and disposition: consumption by sector (stacked), production and trade (lines), in Bcf/d;
    then closing inventory (storage) as a water-year chart in Bcf. Items are matched by name; StatCan's
    ', supply' / ', disposition' / ', storage' tags are dropped from the labels."""
    d = _sheet(p, "Supply and disposition", "Month").apply(pd.to_numeric, errors="coerce") / BCF_TO_MCM
    label = lambda c: re.sub(r",\s*(supply|disposition|storage)$", "", re.sub(r"\s*\((mcm/d|mcm)\)$", "",  # noqa: E731
                                                                              re.sub(r"\s+", " ", str(c))))
    flow = {label(c): c for c in d.columns if str(c).endswith("(mcm/d)")}
    pick = lambda pat: {k: v for k, v in flow.items() if re.search(pat, k, re.I)}  # noqa: E731
    out = []
    recent = d[d.index >= "2021-01-01"]
    use = pick(r"^(residential|commercial|industrial) consumption$|^pipeline fuel$|^deliveries to natural gas processing")
    if use:
        out.append(spec("Demand", recent[list(use.values())].rename(columns={v: k for k, v in use.items()}),
                        "Canada natural gas consumption by sector (StatCan)", "Bcf/d", "stacked_bar"))
    flows = pick(r"^marketable production$|^total imports$|^total exports$|^exports to the united states$")
    if flows:
        out.append(spec("Supply", recent[list(flows.values())].rename(columns={v: k for k, v in flows.items()}),
                        "Canada natural gas production and trade (StatCan)", "Bcf/d"))
    inv = next((c for c in d.columns if re.match(r"closing\s+inventory", str(c), re.I)), None)
    if inv is not None and d[inv].notna().any():
        st = d[inv].dropna()
        st.index = st.index + pd.offsets.MonthEnd(0)   # closing inventory = end of the month
        out.append({"name": "Storage", "water_year": st.resample("D").interpolate(), "y_decimals": 0,
                    "title": "Canada natural gas in storage, closing inventory (StatCan)", "units": "Bcf"})
    return out or generic(p)


def canada_power(p):
    """StatCan monthly generation: non-combustible plants by type, fuel-burning plants as fossil and biomass."""
    d = by_date(read(p, "Daily"), "date")
    d = d[d.index >= "2021-01-01"]
    names = {"Hydro_MWh": "Hydro", "Nuclear_MWh": "Nuclear", "Wind_MWh": "Wind", "Solar_MWh": "Solar",
             "Fossil_MWh": "Fossil fuels (coal, gas, oil)", "Bioenergy_MWh": "Bioenergy", "Other_MWh": "Other"}
    out = [spec("Generation", d[cols(d, *names)].rename(columns=names) / 1000,
                "Canada power generation by source (StatCan)", "GWh per month", "stacked_bar")]
    pv = _sheet(p, "Provinces", "date")
    if not pv.empty:
        pv = pv[pv.index >= "2021-01-01"].rename(columns=lambda c: str(c).replace("_MWh", "")) / 1000
        top = list(pv.sum().sort_values(ascending=False).index[:6])
        g = pv[top].copy()
        g["Other provinces & territories"] = pv.drop(columns=top).sum(axis=1)
        out.append(spec("Provinces", g, "Canada power generation by province (StatCan)", "GWh per month",
                        "stacked_bar"))
    return out


def us_capacity(p):
    """EIA-860M monthly capacity in the standard groups, plus battery storage (kept out of the fuel columns)."""
    out = power_capacity("US installed generating capacity (EIA-860M, net summer)")(p)
    d = by_date(read(p, "Monthly"), "date")
    d = d[d.index >= "2021-01-01"]
    for col, label in (("Battery_storage_MW", "Battery storage"), ("Pumped_storage_MW", "Pumped storage")):
        if col in d:
            out[0]["df"][label] = pd.to_numeric(d[col], errors="coerce").fillna(0) / 1000.0
    return out



def capacity_with_storage(title):
    """Standard capacity chart plus battery / pumped storage columns (kept out of the fuel groups)."""
    def f(p):
        out = power_capacity(title)(p)
        d = by_date(read(p, "Monthly"), "date")
        d = d[d.index >= "2021-01-01"]
        for col, label in (("Battery_storage_MW", "Battery storage"), ("Pumped_storage_MW", "Pumped storage")):
            if col in d:
                out[0]["df"][label] = pd.to_numeric(d[col], errors="coerce").fillna(0) / 1000.0
        return out
    return f


def au_nem_power(p):
    """NEM generation by fuel (monthly GWh) and by state."""
    out = power_daily("Australia NEM power generation by source (AEMO unit SCADA)")(p)
    st = _sheet(p, "States", "date")
    if not st.empty:
        st = st[st.index >= "2021-01-01"].resample("MS").sum(min_count=1) / 1000
        out.append(spec("States", st.rename(columns=lambda c: str(c).replace("_MWh", "")),
                        "Australia NEM power generation by state (AEMO)", "GWh per month", "stacked_bar"))
    d = by_date(read(p, "Daily"), "date")
    if "Battery_charge_MWh" in d:
        last = d.dropna(how="all").index.max()
        m = d[d.index >= "2021-01-01"].resample("MS").sum(min_count=1) / 1000
        if last < last + pd.offsets.MonthEnd(0):
            m = m[m.index < last.to_period("M").to_timestamp()]
        bt = pd.DataFrame({"Battery discharge": m.get("Battery_discharge_MWh"),
                           "Battery charging (-)": -m["Battery_charge_MWh"]})
        out.append(spec("Batteries", bt, "Australia NEM grid batteries: charging vs discharging (AEMO)",
                        "GWh per month", "stacked_bar"))
    out += nem_state_balance(p)
    return out


def au_nem_prices(p):
    """NEM wholesale prices: monthly average by state (A$/MWh) and hours of negative prices per month."""
    out = []
    a = _sheet(p, "Daily average", "date")
    if not a.empty:
        out.append(spec("Prices", monthly_mean(a, "2021-01-01"), "Australia NEM wholesale power price by state (AEMO)",
                        "A$/MWh, monthly average"))
    n = _sheet(p, "Negative hours", "date")
    if not n.empty:
        last = n.dropna(how="all").index.max()
        m = n[n.index >= "2021-01-01"].resample("MS").sum(min_count=1)
        if last < last + pd.offsets.MonthEnd(0):
            m = m[m.index < last.to_period("M").to_timestamp()]
        out.append(spec("Negative prices", m, "Australia NEM hours with negative power prices by state (AEMO)",
                        "hours per month"))
    return out


NEM_STATE_NAMES = {"NSW": "New South Wales", "QLD": "Queensland", "VIC": "Victoria", "SA": "South Australia",
                   "TAS": "Tasmania"}


def nem_state_balance(p):
    """One chart per NEM state: monthly GWh of generation by fuel, rooftop solar, storage output and net
    interconnector imports (negative = net exports) as stacked bars, storage charging below zero, demand
    (operational + rooftop) as a line."""
    b = _sheet(p, "State balance", "date")
    if b.empty:
        return []
    b = b[b.index >= "2021-01-01"].apply(pd.to_numeric, errors="coerce")
    last = b.dropna(how="all").index.max()
    m = b.resample("MS").sum(min_count=1) / 1000.0
    if last < last + pd.offsets.MonthEnd(0):   # drop the month in progress
        m = m[m.index < last.to_period("M").to_timestamp()]
    out = []
    for st, name in NEM_STATE_NAMES.items():
        z = lambda *items: m[[f"{st}_{i}_MWh" for i in items if f"{st}_{i}_MWh" in m]].sum(axis=1, min_count=1)  # noqa: E731
        if not any(c.startswith(f"{st}_") for c in m.columns):
            continue
        g = pd.DataFrame({"Hydro": z("Hydro"), "Gas": z("Gas"), "Wind": z("Wind"), "Solar (utility)": z("Solar"),
                          "Rooftop solar": z("Rooftop_solar"), "Coal": z("Coal"), "Other": z("Oil", "Bioenergy", "Other"),
                          "Storage (battery, pumped hydro)": z("Battery_discharge", "Pumped_hydro"),
                          "Net imports (- = exports)": z("Net_imports"),
                          "Storage charging (-)": -z("Battery_charge", "Pumped_hydro_pumping")})
        g["Demand (incl. rooftop solar)"] = z("Operational_demand") + g["Rooftop solar"].fillna(0)
        g = g.dropna(axis=1, how="all")
        g = g.loc[:, (g.fillna(0) != 0).any()]
        out.append(spec(f"Balance {st}", g, f"{name} power balance (AEMO)", "GWh per month", "stacked_bar",
                        line_cols=tuple(c for c in ("Demand (incl. rooftop solar)",) if c in g)))
    return out


# Energy -> volume for the ANZ gas charts (sources publish TJ / PJ): 1 TJ = 1e12 J / (1,037 Btu/cf x 1,055.06 J/Btu)
# = 0.914 MMcf, EIA's average US heat content - the same basis as the US workbooks
MMCF_PER_TJ = 1e12 / (1037 * 1055.06) / 1e6
MMCFD = "MMcf/d (1 TJ = 0.914 MMcf)"


def _per_day(df, months=1):
    """Volume per month (or quarter) dated at its start/end -> per day over that period."""
    days = df.index.to_period("Q" if months == 3 else "M").days_in_month if months == 1 else \
        (df.index.to_period("Q").end_time - df.index.to_period("Q").start_time).days + 1
    return df.div(pd.Index(days, dtype=float).values, axis=0)


def au_gas(p):
    """GBB: east coast demand by sector and production by state (monthly average MMcf/d), storage (water year, Bcf),
    LNG cargoes (MMcf/d averaged over the month). Sheets hold AEMO's TJ / PJ; charts convert at MMCF_PER_TJ."""
    out = []
    d = _sheet(p, "Demand by sector", "date")
    if not d.empty:
        names = {"Gas_power_generation": "Power generation", "Large_industrial": "Large industrial",
                 "LNG_export_plants": "LNG export plants"}
        out.append(spec("Demand", monthly_mean(d[cols(d, *names)].rename(columns=names)) * MMCF_PER_TJ,
                        "Australia east coast gas demand: power, large industry, LNG (AEMO GBB)",
                        f"{MMCFD}, monthly average", "stacked_bar"))
    q = _sheet(p, "Production", "date")
    if not q.empty:
        q = q.drop(columns=["Total"], errors="ignore")
        out.append(spec("Production", monthly_mean(q, "2021-01-01") * MMCF_PER_TJ,
                        "Australia east coast gas production by state (AEMO GBB)", f"{MMCFD}, monthly average",
                        "stacked_bar"))
    st = _sheet(p, "Storage", "date")
    if "Total" in st and st["Total"].notna().any():
        tot = st["Total"].dropna() * MMCF_PER_TJ / 1000.0   # TJ -> Bcf
        med = tot.rolling(15, center=True, min_periods=5).median()
        tot = tot[(tot - med).abs() <= 0.15 * med]   # one-day reporting glitches (a facility missing or doubled)
        out.append({"name": "Storage", "water_year": _join_short_gaps(tot), "y_decimals": 1,
                    "title": "Australia east coast gas in storage (AEMO GBB)", "units": "Bcf (1 PJ = 0.914 Bcf)"})
    s = _sheet(p, "LNG shipments", "Month")
    if not s.empty:
        s = s.drop(columns=["Total", "Cargoes"], errors="ignore")
        s = _per_day(s[s.index >= "2021-01-01"] * 1000 * MMCF_PER_TJ)   # PJ per month -> MMcf/d
        out.append(spec("LNG", s, "Australia east coast LNG exports by plant (AEMO GBB cargoes)",
                        f"{MMCFD}, monthly average", "stacked_bar"))
    return out


def au_gas_prices(p):
    d = _sheet(p, "Daily", "date")
    return [spec("Prices", d, "Australia east coast gas hub prices, STTM ex-ante (AEMO)", "A$/GJ")]


def au_gas_hub_prices(p):
    d = _sheet(p, "Daily", "date")
    names = {"DWGM_BOD": "Victoria DWGM (beginning of day)", "Wallumbilla_benchmark": "Wallumbilla (QLD)"}
    return [spec("Hub prices", d[cols(d, *names)].rename(columns=names),
                 "Australia gas hub prices: Victoria DWGM and Wallumbilla (AEMO)", "A$/GJ")]


def au_wa_gas(p):
    """WA GBB: production by facility (top 6 + other) and consumption by user type, monthly average MMcf/d."""
    out = []
    pr = _sheet(p, "Production", "date")
    if not pr.empty:
        pr = pr.drop(columns=["Total"], errors="ignore")
        top = list(pr.sum().sort_values(ascending=False).index[:6])
        g = pr[top].copy()
        g["Other facilities"] = pr.drop(columns=top).sum(axis=1, min_count=1)
        out.append(spec("Production", monthly_mean(g, "2021-01-01") * MMCF_PER_TJ,
                        "Western Australia domestic gas production by facility (AEMO WA GBB)",
                        f"{MMCFD}, monthly average", "stacked_bar"))
    c = _sheet(p, "Consumption", "date")
    if not c.empty:
        c = c.drop(columns=["Total"], errors="ignore")
        out.append(spec("Demand", monthly_mean(c, "2021-01-01") * MMCF_PER_TJ,
                        "Western Australia gas consumption by user type (AEMO WA GBB)",
                        f"{MMCFD}, monthly average", "stacked_bar"))
    return out


def au_rooftop_solar(p):
    """CER small-scale solar: MW installed per month by state, and home battery installs."""
    out = []
    c = _sheet(p, "Solar capacity", "Month")
    if not c.empty:
        out.append(spec("Rooftop solar", c[c.index >= "2021-01-01"].drop(columns=["Total"], errors="ignore"),
                        "Australia rooftop solar installed per month by state (Clean Energy Regulator)",
                        "MW per month (latest 12 months still rising)", "stacked_bar"))
    b = _sheet(p, "Battery installs", "Month")
    if not b.empty:
        out.append(spec("Home batteries", b.drop(columns=["Total"], errors="ignore"),
                        "Australia home batteries installed per month by state (Clean Energy Regulator)",
                        "installations per month", "stacked_bar"))
    return out


def future_workbook(p):
    """One-off 'future' workbooks (future/*.py): year-indexed pipeline sheets as stacked bars (MW by fuel), the EIA
    STEO outlook as one line chart per series, other time-indexed outlook sheets as lines. Status / region / list
    sheets stay as tables."""
    region = {"north_america_future.xlsx": "North America", "south_america_future.xlsx": "South America",
              "australia_nz_future.xlsx": "Australia + NZ"}.get(os.path.basename(p), "")
    out = []
    for sh in pd.ExcelFile(p).sheet_names:
        if sh.lower() in ("units", "notes") or sh.startswith("Chart"):
            continue
        d = read(p, sh)
        first = d.columns[0]
        idx = pd.to_datetime(d[first], errors="coerce") if str(first) in ("Year", "Month", "date") else None
        if idx is None or idx.isna().all():
            continue
        d = d.set_index(idx).drop(columns=[first]).apply(pd.to_numeric, errors="coerce").dropna(how="all", axis=1)
        if d.empty:
            continue
        if str(first) == "Year":
            if sh == "NEM capacity outlook":
                out.append(spec(sh, d / 1000.0, "Australia NEM capacity outlook: in service + committed/anticipated "
                                "projects - announced closures (AEMO Generation Information)", "GW", "stacked_bar", "%Y"))
                continue
            unit = "GWh" if "generation" in sh.lower() else "MW"
            d = d.drop(columns=[c for c in d.columns if str(c).lower() == "total"])
            out.append(spec(sh, d, f"{region}: {sh}", unit, "stacked_bar", "%Y"))
        elif sh == "US STEO outlook":
            for c in d.columns:
                out.append(spec(f"STEO {c}"[:28], d[[c]], f"US Short-Term Energy Outlook: {c} (EIA)", c))
        else:
            out.append(spec(sh, d, f"{region}: {sh}", "as published"))
    return out


def au_hydro_storage(p):
    d = _sheet(p, "Weekly", "date")
    return [{"name": "Storage", "water_year": d["Total_GWh"].dropna().resample("D").interpolate(), "y_decimals": 0,
             "title": "Tasmania hydro energy in storage (Hydro Tasmania)", "units": "GWh"}]


def nz_gas(p):
    """MBIE: monthly gross/net production and stock change, quarterly consumption by sector - sheets in PJ as
    published, charts in MMcf/d (average over the month / quarter)."""
    out = []
    d = _sheet(p, "Monthly", "date")
    if not d.empty:
        d = _per_day(d[d.index >= "2021-01-01"].apply(pd.to_numeric, errors="coerce") * 1000 * MMCF_PER_TJ)
        prod = [c for c in d.columns if re.match(r"(gross|net) production", c, re.I)]
        if prod:
            out.append(spec("Production", d[prod], "New Zealand gas production (MBIE)", f"{MMCFD}, monthly average"))
        stock = [c for c in d.columns if re.match(r"stock change", c, re.I)]
        if stock:
            out.append(spec("Storage", d[stock], "New Zealand gas stock change, Ahuroa storage (MBIE)",
                            f"{MMCFD}, monthly average", "stacked_bar"))
    q = _sheet(p, "Quarterly consumption", "date")
    if not q.empty:
        q = q[q.index >= "2021-01-01"].apply(pd.to_numeric, errors="coerce")
        z = lambda *c: q[cols(q, *c)].sum(axis=1, min_count=1)  # noqa: E731
        g = pd.DataFrame({"Power (incl. cogeneration)": z("Electricity Generation", "Cogeneration"),
                          "Petrochemicals (non-energy: methanol, urea)": z("Non-Energy Use"),
                          "Industry": z("Agriculture/ Forestry/ Fishing", "Food Processing",
                                        "Wood, Pulp, Paper, and Printing", "Chemicals", "Basic Metals", "Other"),
                          "Commercial & residential": z("Commercial", "Residential"),
                          "Transport & other transformation": z("Transport", "Other Transformation")})
        g = _per_day(g.dropna(how="all") * 1000 * MMCF_PER_TJ, months=3)
        out.append(spec("Demand", g, "New Zealand gas consumption by sector (MBIE, quarterly)",
                        f"{MMCFD}, quarterly average", "stacked_bar"))
    return out


def generic(p):
    xl = pd.ExcelFile(p)
    for s in xl.sheet_names:
        if s.lower() in ("units", "notes") or s.lower().startswith("chart") or s == water_year_chart.SHEET:
            continue
        d = read(p, s)
        dc = next((c for c in d.columns if re.search(r"date|month|period|day|time", str(c), re.I)), None)
        if dc is None:
            continue
        d = by_date(d, dc)
        num = [c for c in d.columns if pd.api.types.is_numeric_dtype(d[c])][:8]
        if num and len(d) > 1:
            return [spec(s, d[num], f"{os.path.splitext(os.path.basename(p))[0]} - {s}", "", "line")]
    return []


REGISTRY = {
    "argentina_gas_monthly.xlsx": argentina,
    "brazil_gas_monthly.xlsx": brazil,
    "bolivia_gas_demand_by_sector.xlsx": bolivia,
    "uruguay_gas_demand_by_sector.xlsx": uruguay,
    "peru_gas_demand_by_sector.xlsx": peru,
    "chile_gas_imports.xlsx": chile_imports,
    "chile_power_by_type.xlsx": chile_power,
    "brazil_power_generation_daily.xlsx": power_daily("Brazil power generation by type (ONS)"),
    "colombia_power_generation_daily.xlsx": power_daily("Colombia power generation by type (XM)"),
    "panama_power_generation_daily.xlsx": power_daily("Panama power generation by type (CND)"),
    "costa_rica_power_generation_daily.xlsx": power_daily("Costa Rica power generation by type (CENCE)"),
    "nicaragua_power_generation_daily.xlsx": power_daily("Nicaragua power generation by type (CNDC)"),
    "south_america_power_by_type.xlsx": sa_power,
    "argentina_power_generation_daily.xlsx": power_daily("Argentina power generation by type (CAMMESA)"),
    "uruguay_power_generation_daily.xlsx": power_daily("Uruguay power generation by type (ADME)"),
    "bolivia_power_generation_daily.xlsx": power_daily("Bolivia power generation by type (CNDC)"),
    "ecuador_power_generation_daily.xlsx": power_daily("Ecuador power generation by type (CENACE)"),
    "peru_power_generation_daily.xlsx": power_daily("Peru power generation by type (COES)"),
    "chile_power_generation_daily.xlsx": power_daily("Chile power generation by type (CEN)"),
    "central_america_power_by_type.xlsx": sa_power,
    "guatemala_power_generation_daily.xlsx": power_daily("Guatemala power generation by type (AMM)"),
    "honduras_power_generation_daily.xlsx": honduras_power,
    "el_salvador_power_generation_daily.xlsx": power_daily("El Salvador power generation by type (SIGET, monthly net)"),
    # Caribbean
    "puerto_rico_power_generation_daily.xlsx": power_daily("Puerto Rico power generation by type (EIA-923)"),
    "dominican_republic_power_generation_daily.xlsx": power_daily("Dominican Republic power generation by type (OC-SENI)"),
    "puerto_rico_gas.xlsx": gas_use("Puerto Rico gas use"),
    "jamaica_power_generation_daily.xlsx": power_annual("Jamaica power generation by type (MSET / JPS, annual)"),
    # Paraguay: monthly rows (first of month); second chart = own use vs energy ceded to Brazil / Argentina
    "paraguay_power_generation_daily.xlsx": paraguay_power,
    "jamaica_gas.xlsx": jamaica_gas,
    "dominican_republic_gas.xlsx": gas_use("Dominican Republic gas use (gas burned for power, SIE)"),
    "colombia_gas_demand_by_sector.xlsx": colombia,
    "ecuador_gas.xlsx": ecuador,
    "panama_gas.xlsx": panama_gas,
    "trinidad_gas.xlsx": trinidad,
    "guyana_gas.xlsx": guyana_gas,
    "el_salvador_gas.xlsx": el_salvador_gas,
    "ireland_gas_demand_daily.xlsx": ireland_demand,
    "ireland_gas_supply_daily.xlsx": ireland_supply,
    "ireland_gas_combined_daily.xlsx": ireland_combined,
    "ireland_gni_transparency_daily.xlsx": ireland_gni,
    "ireland_smartgrid_15min.xlsx": ireland_smartgrid,
    "mexico_demanda_nacional_daily.xlsx": mexico,
    "miso_fuel_mix_daily.xlsx": fuel_mix("MISO fuel mix", "MW (daily mean)"),
    "turkey_generation_mix_dashboard_daily.xlsx": fuel_mix("Turkey generation mix", "MW"),
    "south_africa_generation_mix_daily.xlsx": south_africa,
    "puertorico_generation_monthly.xlsx": puertorico,
    "china_nbs_clean_energy_products_monthly.xlsx": china_nbs_series,
    "china_nbs_energy_production_monthly.xlsx": china_nbs_series,
    "china_nbs_industrial_output_monthly.xlsx": china_nbs_series,
    "china_nbs_market_prices_10day.xlsx": china_nbs_series,
    "china_nbs_capacity_utilization_quarterly.xlsx": china_nbs_series,
    "china_nbs_ppi_monthly.xlsx": china_nbs_series,
    "giignl_contracted_vs_spot_annual.xlsx": giignl,
    "singapore_power.xlsx": singapore_power,
    "singapore_gas.xlsx": singapore_gas,
    "henry_hub_daily.xlsx": henry_hub,
    "us_gas.xlsx": us_gas,
    "mexico_gas.xlsx": mexico_gas,
    "us_mexico_pipeline_capacity.xlsx": us_mexico_pipeline_capacity,
    "canada_gas.xlsx": canada_gas,
    "canada_power_generation_daily.xlsx": canada_power,
    # Australia and New Zealand
    "nz_power_generation_daily.xlsx": power_daily("New Zealand power generation by source (Electricity Authority EMI)"),
    "au_nem_power_generation_daily.xlsx": au_nem_power,
    "au_nem_prices.xlsx": au_nem_prices,
    "au_wem_power_generation_daily.xlsx": power_daily("Western Australia (WEM) power generation by source (AEMO)"),
    "au_power_capacity.xlsx": capacity_with_storage("Australia registered generating capacity (AEMO, NEM + WEM)"),
    "nz_power_capacity.xlsx": capacity_with_storage("New Zealand installed generating capacity (MBIE)"),
    "au_gas.xlsx": au_gas,
    "au_gas_prices.xlsx": au_gas_prices,
    # one-off future workbooks (future/*.py)
    "north_america_future.xlsx": future_workbook,
    "south_america_future.xlsx": future_workbook,
    "australia_nz_future.xlsx": future_workbook,
    "au_gas_hub_prices.xlsx": au_gas_hub_prices,
    "au_wa_gas.xlsx": au_wa_gas,
    "au_rooftop_solar.xlsx": au_rooftop_solar,
    "au_hydro_storage.xlsx": au_hydro_storage,
    "nz_gas.xlsx": nz_gas,
    "north_america_power_by_type.xlsx": sa_power,
    "us_power_capacity.xlsx": us_capacity,
    "canada_power_capacity.xlsx": power_capacity("Canada installed generating capacity (StatCan, annual)"),
    "mexico_power_capacity.xlsx": power_capacity("Mexico installed generating capacity (Ember - no raw feed, annual)"),   # Ember fallback for Canada and Mexico (same layout)
    "latin_america_industrial_gas_users.xlsx": industrial_gas_users,   # static plant register, category axis
    "south_america_power_prices_daily.xlsx": sa_power_prices,
    "south_america_gas_balance.xlsx": sa_gas_balance,
    "south_america_prices_vs_hydro.xlsx": None,   # two-panel charts drawn by SA_PRICES_VS_HYDRO.py itself
    "brazil_hydro_reservoirs.xlsx": brazil_hydro,
    "colombia_hydro_reservoirs.xlsx": colombia_hydro,
    "argentina_hydro_reservoirs.xlsx": argentina_hydro,
    "chile_hydro_reservoirs.xlsx": chile_hydro,
    "south_america_coal_production.xlsx": sa_coal,
    "peru_hydro_reservoirs.xlsx": peru_hydro,
    "ecuador_hydro_reservoirs.xlsx": ecuador_hydro,
    "uruguay_hydro_reservoirs.xlsx": uruguay_hydro,
    # installed generation capacity by technology (standard sheet "Monthly")
    "brazil_power_capacity.xlsx": power_capacity("Brazil installed generation capacity (ANEEL)"),
    "argentina_power_capacity.xlsx": power_capacity("Argentina installed generation capacity (CAMMESA)"),
    "chile_power_capacity.xlsx": power_capacity("Chile installed generation capacity (CNE)"),
    "colombia_power_capacity.xlsx": power_capacity("Colombia installed generation capacity (XM)"),
    "uruguay_power_capacity.xlsx": power_capacity("Uruguay installed generation capacity (MIEM / DNE)"),
    "bolivia_power_capacity.xlsx": power_capacity("Bolivia installed generation capacity (CNDC)"),
    "ecuador_power_capacity.xlsx": power_capacity("Ecuador installed generation capacity (ARCONEL)"),
    "peru_power_capacity.xlsx": power_capacity("Peru installed generation capacity (COES)"),
    "paraguay_power_capacity.xlsx": power_capacity("Paraguay installed generation capacity (ANDE)"),
    # these build their own charts in their pull scripts:
    "rhine_kaub_level_daily.xlsx": None,
    "gatun_lake_level.xlsx": None,
    "eia930_fuel_mix_daily.xlsx": None,
    "lng_feedgas_daily.xlsx": None,
}


def _drop_old_chart_sheets(path):
    """Charts are rebuilt from scratch each run; remove last run's chart sheets so renamed/retired ones don't linger."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    old = [n for n in wb.sheetnames if n == "Chart" or n.startswith("Chart - ")]
    if old:
        for n in old:
            del wb[n]
        root, ext = os.path.splitext(path)
        tmp = f"{root}.tmp{os.getpid()}{ext}"
        try:
            wb.save(tmp)
            os.replace(tmp, path)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)


def _order_chart_sheets(path, names):
    """Chart sheets straight after the Units tab, in registry order (each add inserts at position 1)."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    present = [n for n in names if n in wb.sheetnames]
    if len(present) < 2:
        return
    sheets = {ws.title: ws for ws in wb._sheets}
    rest = [ws for ws in wb._sheets if ws.title not in present]
    head = rest[:1] if rest and rest[0].title.lower() in ("units", "notes") else []
    wb._sheets = head + [sheets[n] for n in present] + [ws for ws in rest if ws not in head]
    root, ext = os.path.splitext(path)
    tmp = f"{root}.tmp{os.getpid()}{ext}"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def _add_in_one_pass(path, specs):
    """All chart sheets in one load/save. openpyxl drops the formatting of charts it reads back in (axis
    min/max and step - e.g. the 0-100% scale of storage charts - axis titles, date formats), so adding charts
    one save at a time left every chart but the last unformatted in multi-chart workbooks."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for n in [n for n in wb.sheetnames if n == "Chart" or n.startswith("Chart - ")]:
        del wb[n]
    names = []
    for i, s in enumerate(specs):
        sheet = "Chart" if i == 0 else f"Chart - {s['name']}"[:31]
        if "water_year" in s:
            sheet = s.get("sheet", water_year_chart.SHEET)
            water_year_chart.add_water_year_chart(path, s["water_year"], s["title"], s["units"], sheet_name=sheet,
                                                  y_decimals=s.get("y_decimals"), wb=wb)
        else:
            df = s["df"].dropna(how="all")
            if df.empty:
                continue
            xlsx_charts.add_chart_sheet(path, df, s["title"], s["units"], kind=s["kind"], sheet_name=sheet,
                                        date_format=s["date_format"], line_cols=s.get("line_cols", ()), wb=wb)
        names.append(sheet)
    # chart sheets straight after the Units tab, in registry order
    present = [n for n in names if n in wb.sheetnames]
    sheets = {ws.title: ws for ws in wb._sheets}
    rest = [ws for ws in wb._sheets if ws.title not in present]
    head = rest[:1] if rest and rest[0].title.lower() in ("units", "notes") else []
    wb._sheets = head + [sheets[n] for n in present] + [ws for ws in rest if ws not in head]
    xlsx_charts.save_atomic(wb, path)


def add_charts(path):
    name = os.path.basename(path)
    fn = REGISTRY.get(name, generic)
    if fn is None:
        print(f"{name}: charts are built by its own pull script")
        return 0
    specs = fn(path)
    if not any("custom" in s for s in specs):
        _add_in_one_pass(path, specs)
        print(f"{name}: {len(specs)} chart(s)")
        return len(specs)
    _drop_old_chart_sheets(path)
    for i, s in enumerate(specs):
        sheet = "Chart" if i == 0 else f"Chart - {s['name']}"[:31]
        if "custom" in s:   # chart that is not a date-indexed series (e.g. categories on the x-axis)
            s["custom"](path, sheet)
            continue
        if "water_year" in s:
            water_year_chart.add_water_year_chart(path, s["water_year"], s["title"], s["units"],
                                                  sheet_name=s.get("sheet", water_year_chart.SHEET),
                                                  y_decimals=s.get("y_decimals"))
            continue
        df = s["df"].dropna(how="all")
        if df.empty:
            continue
        xlsx_charts.add_chart_sheet(path, df, s["title"], s["units"], kind=s["kind"], sheet_name=sheet,
                                    date_format=s["date_format"], line_cols=s.get("line_cols", ()))
    # one name per spec, in registry order; an optional "sheet" key names a water-year sheet
    # (several water-year charts in one workbook)
    _order_chart_sheets(path, [sp.get("sheet", water_year_chart.SHEET) if "water_year" in sp
                               else ("Chart" if i == 0 else f"Chart - {sp['name']}"[:31])
                               for i, sp in enumerate(specs)])
    print(f"{name}: {len(specs)} chart(s)")
    return len(specs)


if __name__ == "__main__":
    for p in sys.argv[1:]:
        if os.path.exists(p):
            add_charts(p)
        else:
            print(f"{p}: not found, skipped")
