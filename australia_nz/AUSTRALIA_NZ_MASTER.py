"""
Australia + New Zealand master workbook: one file with every Australian and
New Zealand dataset this repo pulls, and Dashboard front pages carrying all
their charts - the same layout as the South & Central America and North
America masters (south_america/SOUTH_AMERICA_MASTER.py, whose
table/chart/dashboard code this reuses).

  Dashboard                  - gas: Australia east coast demand by sector, production, storage (water year), LNG
                               cargoes, STTM hub prices (AEMO); New Zealand production, stock change and
                               consumption (MBIE)
  Dashboard - Power & Hydro  - Australia + NZ generation by source (sum of NEM, WEM and NZ) and installed
                               capacity; then NEM by fuel and by state, WA (WEM), New Zealand (Electricity
                               Authority), capacity per country, Tasmania hydro storage (water year)
  <CC> <chart> data          - the table each Dashboard chart plots
  <CC> <dataset> raw         - the full data sheet(s) from each source workbook
  Sources                    - where each dataset comes from, units and notes

All sources are raw (grid operators, market operator, ministry); no Ember.
Reads (doesn't refetch) the workbooks the scheduled pulls write to
"output/Data and Chart Outputs/". A missing input is listed on the Dashboard
and skipped rather than stopping the rest.

Usage: python3 AUSTRALIA_NZ_MASTER.py [--out "output/Data and Chart Outputs/australia_nz_master.xlsx"]
"""
import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "south_america"))
import add_charts  # noqa: E402
import xlsx_charts  # noqa: E402
import SOUTH_AMERICA_MASTER as sam  # noqa: E402

DATA_DIR = sam.DATA_DIR
FUELS = ["Hydro", "Gas", "Wind", "Solar", "Coal", "Nuclear", "Other"]   # same order/colours as the other dashboards

# (country code, country, workbook, raw sheet / tuple of raw sheets / "*", short dataset name)
DATASETS = [
    ("AU", "Australia", "au_gas.xlsx", ("Demand by sector", "Production", "Storage", "LNG shipments"), "gas"),
    ("AU", "Australia", "au_gas_prices.xlsx", "Daily", "gas prices"),
    ("AU", "Australia", "au_gas_hub_prices.xlsx", "Daily", "gas hub prices"),
    ("WA", "Western Australia", "au_wa_gas.xlsx", ("Production", "Consumption", "Storage", "By zone"), "gas"),
    ("NZ", "New Zealand", "nz_gas.xlsx", ("Monthly", "Quarterly consumption"), "gas"),
]
RAW_POWER_DATASETS = [
    ("AU", "Australia NEM", "au_nem_power_generation_daily.xlsx", ("Daily", "States"), "power"),
    ("WA", "Western Australia", "au_wem_power_generation_daily.xlsx", "Daily", "power"),
    ("NZ", "New Zealand", "nz_power_generation_daily.xlsx", ("Daily", "By fuel code"), "power"),
]
CAPACITY_DATASETS = [
    ("AU", "Australia", "au_power_capacity.xlsx", ("Monthly", "By region"), "capacity"),
    ("NZ", "New Zealand", "nz_power_capacity.xlsx", ("Monthly", "By plant type"), "capacity"),
]
PRICE_DATASETS = [
    ("AU", "Australia NEM", "au_nem_prices.xlsx", ("Daily average", "Negative hours"), "power prices"),
    ("AU", "Australia", "au_rooftop_solar.xlsx", ("Solar capacity", "Solar installs", "Battery installs"), "rooftop solar"),
]
HYDRO_DATASETS = [
    ("TAS", "Tasmania", "au_hydro_storage.xlsx", "Weekly", "hydro storage"),
]
HYDRO_EXTRA = {}
DASHBOARD_ONLY = {}
MASTER_SPECS = {}
EMBER = set()
OPERATORS = {}

SOURCES = {
    "au_gas.xlsx": ("AEMO Gas Bulletin Board (actual flows and storage, LNG shipments)",
                    "https://aemo.com.au/energy-systems/gas/gas-bulletin-board-gbb"),
    "au_gas_prices.xlsx": ("AEMO Short Term Trading Market, INT651 ex-ante market price",
                           "https://aemo.com.au/energy-systems/gas/short-term-trading-market-sttm"),
    "nz_gas.xlsx": ("MBIE (NZ Ministry of Business, Innovation and Employment), gas statistics webtables",
                    "https://www.mbie.govt.nz/building-and-energy/energy-and-natural-resources/energy-statistics-and-modelling/energy-statistics/gas-statistics/"),
    "au_nem_power_generation_daily.xlsx": ("AEMO NEMWEB / MMSDM 5-minute unit SCADA (DISPATCH_UNIT_SCADA)",
                                           "https://nemweb.com.au/"),
    "au_wem_power_generation_daily.xlsx": ("AEMO WA facility SCADA (WEM)", "https://data.wa.aemo.com.au/"),
    "nz_power_generation_daily.xlsx": ("Electricity Authority (Te Mana Hiko), EMI Generation_MD metered data",
                                       "https://www.emi.ea.govt.nz/Wholesale/Datasets/Generation/Generation_MD"),
    "au_power_capacity.xlsx": ("AEMO NEM Registration and Exemption List + WEM facilities list (registered capacity)",
                               "https://aemo.com.au/en/energy-systems/electricity/national-electricity-market-nem/participate-in-the-market/registration"),
    "nz_power_capacity.xlsx": ("MBIE electricity statistics, installed capacity by plant type",
                               "https://www.mbie.govt.nz/building-and-energy/energy-and-natural-resources/energy-statistics-and-modelling/energy-statistics/electricity-statistics/"),
    "au_nem_prices.xlsx": ("AEMO NEMWEB / MMSDM DISPATCHPRICE (5-minute regional reference price)",
                           "https://nemweb.com.au/"),
    "au_gas_hub_prices.xlsx": ("AEMO: Victorian DWGM market prices (INT041) and Wallumbilla gas supply hub benchmark "
                               "price", "https://aemo.com.au/energy-systems/gas/declared-wholesale-gas-market-dwgm"),
    "au_wa_gas.xlsx": ("AEMO WA Gas Bulletin Board (actual flows, end-user consumption)", "https://gbbwa.aemo.com.au/"),
    "au_rooftop_solar.xlsx": ("Clean Energy Regulator, small-scale installation postcode data",
                              "https://cer.gov.au/markets/reports-and-data/small-scale-installation-postcode-data"),
    "au_hydro_storage.xlsx": ("Hydro Tasmania, Energy in Storage", "https://www.hydro.com.au/water/energy-in-storage"),
}

POWER_PARTS = [("Australia NEM", "au_nem_power_generation_daily.xlsx"),
               ("Western Australia", "au_wem_power_generation_daily.xlsx"),
               ("New Zealand", "nz_power_generation_daily.xlsx")]


def monthly_gwh(path):
    """Standard Daily sheet (MWh) -> complete months, GWh, in the dashboard fuel groups (storage left out)."""
    d = add_charts.by_date(add_charts.read(path, "Daily"), "date")
    d = d[[c for c in d.columns if str(c).endswith("_MWh")]].apply(pd.to_numeric, errors="coerce")
    m = d.resample("MS").sum(min_count=1) / 1000.0
    last = d.dropna(how="all").index.max()
    if last < last + pd.offsets.MonthEnd(0):   # drop the month in progress
        m = m[m.index < last.to_period("M").to_timestamp()]
    m = m[m.index >= "2021-01-01"].rename(columns=lambda c: c.replace("_MWh", "_GWh"))
    m = m.rename(columns={"Oil_GWh": "Other Fossil_GWh", "Other_GWh": "Other Renewables_GWh"})
    return add_charts.power_mix(m)


def anz_generation(data_dir):
    """Sum of NEM + WEM + NZ, GWh per month, over the months all three have."""
    frames, notes = {}, []
    for name, fname in POWER_PARTS:
        try:
            m = monthly_gwh(os.path.join(data_dir, fname))
            m = m[m.sum(axis=1) > 0]
            frames[name] = m
            notes.append(f"{name}: {SOURCES[fname][0]}, {m.index.min():%b/%y}-{m.index.max():%b/%y}")
        except Exception as e:  # noqa: BLE001
            notes.append(f"NOT INCLUDED: {name} ({type(e).__name__}: {e})")
    if not frames:
        return pd.DataFrame(), notes
    months = sorted(set.intersection(*(set(f.index) for f in frames.values())))
    if not months:
        return pd.DataFrame(), notes + ["no month common to every feed yet"]
    total = sum(f.reindex(index=months, columns=FUELS).fillna(0) for f in frames.values())
    total.index.name = "date"
    return total, notes


def anz_capacity(data_dir):
    """Australia (monthly snapshots from the first run) + NZ (annual, carried forward), GW, from the first month
    Australia has - earlier months are not estimated."""
    frames, notes = {}, []
    for _, country, fname, _, _ in CAPACITY_DATASETS:
        path = os.path.join(data_dir, fname)
        try:
            d = add_charts.by_date(add_charts.read(path, "Monthly"), "date")
            g = add_charts.capacity_groups(pd.DataFrame({f: d.get(f"{f}_MW") for f in add_charts.CAPACITY_FUELS},
                                                        index=d.index)) / 1000.0
            annual = len(d) > 1 and d.index.to_series().diff().median().days > 300
            if annual:   # annual rows hold the year-end value, dated 1 January
                g.index = g.index + pd.DateOffset(months=11)
            frames[country] = g
            notes.append(f"{country}: {SOURCES[fname][0]}, {'annual' if annual else 'monthly'} "
                         f"{g.index.min():%b/%y}-{g.index.max():%b/%y}")
        except Exception as e:  # noqa: BLE001
            notes.append(f"NOT INCLUDED: {country} ({type(e).__name__}: {e})")
    if "Australia" not in frames:
        return pd.DataFrame(), notes
    au = frames["Australia"]
    months = pd.date_range(au.index.min(), au.index.max(), freq="MS")
    total = au.reindex(months).ffill()
    if "New Zealand" in frames:
        nz = frames["New Zealand"]
        total = total + nz.reindex(nz.index.union(months)).sort_index().ffill().reindex(months).fillna(0)
    total.index.name = "date"
    return total, notes


def total_chart(wb, used, power, pos, df_total, notes, sheet, title, units, src_text, note_head):
    ws = wb.create_sheet(sam.sheet_name(sheet, used))
    df, n_bars = xlsx_charts.prepare(df_total)
    xlsx_charts.write_table(ws, df)
    ws.cell(row=1, column=df.shape[1] + 4, value=note_head)
    for i, note in enumerate(notes, start=2):
        ws.cell(row=i, column=df.shape[1] + 4, value=note)
    src = (src_text, None)
    power[0].insert(pos, (xlsx_charts.build_chart(ws, df, n_bars, title, units, "stacked_bar", width=sam.CHART_W,
                                                  height=sam.CHART_H, gridlines=False,
                                                  inner=xlsx_charts.DASHBOARD_INNER), src))
    missing = [n.split(":", 1)[1].split("(")[0].strip() for n in notes if n.startswith("NOT INCLUDED")]
    power[1].insert(pos, ("Australia + NZ", title + (f" (missing: {', '.join(missing)})" if missing else ""),
                          df.index.max().strftime("%b/%y"), ws.title, *src))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(DATA_DIR, "australia_nz_master.xlsx"))
    ap.add_argument("--data-dir", default=DATA_DIR)
    args = ap.parse_args()
    cfg = sys.modules[__name__]

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    dash2 = wb.create_sheet("Dashboard - Power & Hydro")
    used = {"Dashboard", "Dashboard - Power & Hydro", "Sources"}
    sources = []

    gas = sam.collect(wb, DATASETS, args.data_dir, used, sources, cfg=cfg)
    parts = [sam.collect(wb, RAW_POWER_DATASETS, args.data_dir, used, sources, cfg=cfg),
             sam.collect(wb, PRICE_DATASETS, args.data_dir, used, sources, cfg=cfg),
             sam.collect(wb, CAPACITY_DATASETS, args.data_dir, used, sources, cfg=cfg),
             sam.collect(wb, HYDRO_DATASETS, args.data_dir, used, sources, cfg=cfg)]
    power = tuple(sum((p[i] for p in parts), []) for i in range(3))

    pos = 0
    total, notes = anz_generation(args.data_dir)
    if not total.empty:
        total_chart(wb, used, power, pos, total, notes, "ANZ generation total data",
                    "Australia + New Zealand power generation by source", "GWh per month",
                    "Sum of NEM (AEMO SCADA), WEM (AEMO WA SCADA) and NZ (Electricity Authority EMI)",
                    "Feeds summed (only months all of them have):")
        pos += 1
        print("ANZ generation total:", "; ".join(notes))
    else:
        power[2].append("ANZ generation total: " + "; ".join(notes))
    cap, cap_notes = anz_capacity(args.data_dir)
    if not cap.empty:
        total_chart(wb, used, power, pos, cap, cap_notes, "ANZ capacity total data",
                    "Australia + New Zealand installed generating capacity", "GW installed",
                    "Sum of AEMO registered capacity (NEM + WEM) and MBIE (NZ)",
                    "Countries summed (NZ annual figures carried forward; storage excluded; from Australia's "
                    "first monthly snapshot):")
        print("ANZ capacity total:", "; ".join(cap_notes))

    sam.draw_dashboard(dash, "Australia and New Zealand energy - gas dashboard", *gas)
    sam.draw_dashboard(dash2, "Australia and New Zealand energy - power generation and hydro", *power)

    src = wb.create_sheet("Sources")
    src.append(["Country", "Dataset", "Workbook", "Publisher", "Link", "Units and notes (from the source workbook)"])
    for cell in src[1]:
        cell.font = Font(bold=True)
    for s in sources:
        src.append(list(s))
        if s[4]:
            src.cell(row=src.max_row, column=5).hyperlink = s[4]
    for col, w in (("A", 14), ("B", 14), ("C", 38), ("D", 40), ("E", 50), ("F", 160)):
        src.column_dimensions[col].width = w

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    xlsx_charts.save_atomic(wb, args.out)
    print(f"Saved {args.out}: {len(gas[0])} gas charts, {len(power[0])} power charts; tabs {wb.sheetnames}")
    for label, m in (("gas", gas[2]), ("power", power[2])):
        if m:
            print(f"missing ({label}):", m)


if __name__ == "__main__":
    main()
